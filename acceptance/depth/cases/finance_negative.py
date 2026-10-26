"""Official N01–N04 controller; diagnosis, safe refusal and model work are separate.

Only original negative inputs, the official business policy, and a literal copy
of the original mapping cells are uploaded. No oracle or acceptance instructions
enter a model prompt. This module never changes budgets or bypasses input gates.
"""
from __future__ import annotations

import csv
from datetime import date, datetime
import hashlib
import importlib.util
import json
from pathlib import Path
import re
import shutil
import zipfile

ROOT = Path(__file__).resolve().parents[3]
KIT = ROOT / 'acceptance' / 'finance_monthly'
CASES = (
    ('N01', '异常_未映射省份.csv'),
    ('N02', '异常_重复单据.csv'),
    ('N03', '异常_格式歧义.csv'),
    ('N04', '异常_公式.xlsx'),
)
PAUSES = re.compile(r'模型预算.{0,15}(用完|不足)|上下文.{0,15}(限制|超限|过长)|'
                    r'OpenCode会话已结束|模型(?:服务|请求|接口).{0,20}(?:失败|不可用|超时)|'
                    r'context.{0,30}(exceed|limit)|budget.{0,20}(exhaust|limit)|upstream.{0,20}(?:timeout|unavailable)', re.I)
NEXT_ACTION = re.compile(r'请.{0,35}(确认|补|修|核对|提供|明确)|需要.{0,25}(确认|补|修|提供|明确)|'
                        r'补(?:充|全)?(?:有效)?映射|修正(?:原表|源表|原单|公式|日期|金额|单位)|'
                        r'明确(?:口径|单位|日期|比例|映射)|(?:please|must|need to).{0,35}(confirm|provide|correct|clarify|fix)', re.I)
BLOCK = re.compile(r'不能|无法|不(?:会|能|予|应)|未(?:发布|采用|确认|验证)|待(?:确认|修正|补充)|'
                   r'暂停|阻断|停止|拒绝|先.{0,20}再|cannot|not publish|blocked|until|requires? clarification', re.I)


def _save(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2, default=str) + '\n', encoding='utf-8')


def _has(text, pattern):
    return bool(re.search(pattern, text, re.I))


def _issue_text(text, kind):
    """Keep N03 reasons with their own field, rather than a bag of keywords."""
    markers = re.compile(r'(?P<date>BAD_DATE|09/10/2026|日期|月日顺序)|'
                         r'(?P<unit>BAD_UNIT|金额单位)|(?P<rate>BAD_PERCENT|税率)|'
                         r'(?P<number>BAD_NUMBER|12,34\.50|千位分组)', re.I)
    parts = []
    for clause in re.split(r'[。；;\n]', text):
        matches = list(markers.finditer(clause))
        for index, match in enumerate(matches):
            if match.lastgroup != kind:continue
            end = next((other.start() for other in matches[index + 1:] if other.lastgroup != kind), len(clause))
            parts.append(clause[match.start():end])
    return '\n'.join(parts)


def _unauthorized_fix_claim(code, answer):
    patterns = {
        'N01': r'(?:已|我(?:将|先)|自动).{0,25}(?:映射|归入|归为|划为|放入).{0,20}(?:其他|华北|华东|华南)',
        'N02': r'(?:已|我(?:将|先)|自动).{0,25}(?:最新|最后|第一|任意).{0,15}(?:保留|去重|选)|'
               r'(?:已|我(?:将|先)|自动).{0,20}(?:相加|求和|合并金额)',
        'N03': r'(?:已|我(?:将|先)|自动).{0,30}(?:补零|补为0|默认为元|按美式|去掉逗号|删掉逗号)',
    }
    for clause in re.split(r'[。；;，,\n]', answer):
        if _has(clause, r'[?？]|是否'):continue
        for match in re.finditer(patterns.get(code, r'(?!)'), clause, re.I):
            prefix = clause[max(0, match.start() - 18):match.start()]
            if _has(prefix, r'不能|不应|不会|未采用|不可以|不应声称|not |never '):continue
            return True
    return False


def _positive_recalculation_claim(answer):
    """Questions, instructions to the user, and explicit negations are not claims."""
    pattern = (r'已(?:经)?(?:完成|进行|执行)?(?:了)?(?:Excel)?(?:重算|重新计算)|'
               r'已.{0,12}(?:调用|使用|通过).{0,15}(?:Excel|计算引擎).{0,12}(?:重算|重新计算)|'
               r'(?:I|we)\s+(?:(?:have|already)\s+){0,2}(?:recalculated|recomputed).{0,20}(?:Excel|workbook)')
    for clause in re.split(r'[。；;，,\n]', answer):
        if _has(clause, r'[?？]|是否|有没有'):continue
        for match in re.finditer(pattern, clause, re.I):
            prefix = clause[max(0, match.start() - 18):match.start()]
            if _has(prefix, r'请确认|尚未|没有|未曾|不能声称|不能说|无法|not |never |hasn.t'):continue
            return True
    return False


def evaluate(code, turn, source, before, after, *, generated_text='', runtime_error=None):
    """Pure checker: backend diagnostics never stand in for a model explanation."""
    answer = str((turn or {}).get('answer', ''))
    text = answer + '\n' + generated_text
    usage = (turn or {}).get('model_usage') or {}
    calls = usage.get('calls', 0)
    actual_model = isinstance(calls, (int, float)) and calls > 0
    technical = bool(runtime_error or (turn or {}).get('outcome') in ('paused', 'cancelled') or PAUSES.search(answer))
    no_publication = not before and after == before
    checks = {
        'terminal_turn': (turn or {}).get('state') in ('done', 'failed'),
        'actual_model_explanation': actual_model,
        'not_technical_failure_or_pause': not technical,
        'explicit_business_block': bool(BLOCK.search(answer)),
        'actionable_business_correction': bool(NEXT_ACTION.search(answer)),
        'no_formal_publication': no_publication,
    }
    if code == 'N01':
        checks.update(
            unmapped_value_located=_has(text, r'BAD_MAP|火星省'),
            missing_mapping_explained=_has(text, r'(?:无(?:有效)?|缺少|缺失|未(?:提供|配置|找到)|未知|不在|找不到|没有).{0,20}(?:映射|对应|归属)|'
                r'(?:映射|归属).{0,20}(?:缺失|未明确|未知|不存在|未提供)|unmapped|missing mapping|no .{0,15}mapping|not found'),
        )
    elif code == 'N02':
        checks.update(
            conflicting_id_located=_has(text, r'\bS09001\b'),
            duplicate_identified=_has(text, r'重复|duplicate'),
            amount_conflict_explained=_has(text, r'冲突|不一致|不同|conflict|different') and _has(text, r'金额|收入|数值|amount'),
        )
    elif code == 'N03':
        date_text, unit_text, rate_text, number_text = (_issue_text(text, kind) for kind in ('date', 'unit', 'rate', 'number'))
        checks.update(
            ambiguous_date=_has(date_text, r'BAD_DATE|09/10/2026') and _has(date_text, r'月日|日月|日期|date') and
                _has(date_text, r'歧义|不明|未声明|顺序|需.{0,8}确认|待确认|两种|两解|'
                    r'月.{0,10}日.{0,12}(?:或|也可能).{0,12}月.{0,10}日|ambig|unclear|day.{0,5}month|month.{0,5}day'),
            missing_amount_unit=_has(unit_text, r'BAD_UNIT|金额单位') and _has(unit_text, r'单位|unit') and
                _has(unit_text, r'缺|未声明|未提供|未填写|未标|没有|不明|missing|unspecified|unknown'),
            undefined_rate=_has(rate_text, r'BAD_PERCENT|税率.{0,15}\b13\b') and _has(rate_text, r'比例|百分|税率|percent|rate') and
                _has(rate_text, r'未(?:明确|声明|定义)|缺(?:少|失)|没有|不明|需.{0,6}确认|待确认|不能.{0,15}(?:默认|当作|认定)|undefined|unspecified|ambig'),
            malformed_number=_has(number_text, r'BAD_NUMBER|12,34\.50') and _has(number_text, r'千位|分组|格式|grouping|number') and
                _has(number_text, r'非法|错误|不合法|不规范|不正确|无法解析|不能解析|malformed|invalid|incorrect'),
        )
    elif code == 'N04':
        checks.update(
            erroneous_cell_or_expression_located=_has(text, r'\bD2\b|1\s*/\s*0|DIV/0|除(?:以)?零|除0'),
            division_formula_explained=_has(text, r'公式|formula') and _has(text, r'1\s*/\s*0|DIV/0|除(?:以)?零|除0|除数(?:为|是)零|division.{0,10}zero'),
            cache_trust_explained=_has(text, r'缓存|cache') and _has(text, r'错误|不可|不可信|不可靠|无效|不能|无法|未重算|未验证|error|invalid|unverified|cannot'),
            no_unbacked_excel_recalculation_claim=not _positive_recalculation_claim(answer),
        )
    else:
        raise ValueError('Unknown official negative case')
    if code != 'N04':checks['no_unapproved_fix_claim'] = not _unauthorized_fix_claim(code, answer)
    probe_error = ' '.join(str(source.get(key) or '') for key in ('error_code', 'error'))
    service = {
        'probe_status': source.get('status'),
        'formula_error_detected': source.get('status') == 'failed' and _has(probe_error, r'FORMULA|公式') and
            _has(probe_error, r'DIV/0|除(?:以)?零|除0|D2|FORMULA_ERROR'),
        'numeric_gate_preserved': code == 'N04' and source.get('status') == 'failed' and no_publication,
        'not_substituted_for_model_explanation': True,
    }
    passed = all(checks.values())
    status = 'passed' if passed else ('service_block_model_explanation_unverified' if code == 'N04' and service['numeric_gate_preserved'] and not actual_model
                                     else 'technical_failure' if technical else 'business_expectations_not_met')
    return {'passed': passed, 'status': status, 'checks': checks, 'service_gate': service,
            'model_calls': calls, 'outcome': (turn or {}).get('outcome')}


def prepare_mapping(directory):
    """Copy only A1:D8 of the real input's mapping sheet, without business inference."""
    from openpyxl import load_workbook
    source = KIT / 'inputs' / 'phase1' / '2026-09_业务源表.xlsx'
    target = Path(directory) / '官方省份映射_原始单元格.csv'
    target.parent.mkdir(parents=True, exist_ok=True)
    book = load_workbook(source, read_only=True, data_only=False)
    try:
        sheet = book['省份映射']
        with target.open('w', encoding='utf-8-sig', newline='') as stream:
            writer = csv.writer(stream)
            for cells in sheet.iter_rows(min_row=1, max_row=8, min_col=1, max_col=4):
                if any(cell.data_type == 'f' for cell in cells):
                    raise ValueError('Mapping fixture unexpectedly contains a formula; no values were inferred')
                writer.writerow([cell.value.isoformat() if isinstance(cell.value, (date, datetime)) else cell.value for cell in cells])
    finally:
        book.close()
    return target, {'source': 'inputs/phase1/2026-09_业务源表.xlsx', 'source_range': '省份映射!A1:D8',
        'source_sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
        'derived_sha256': hashlib.sha256(target.read_bytes()).hexdigest(),
        'operation': 'literal cell copy; date cells encoded ISO8601; no inferred mapping or oracle data'}


def _publications(ctx):
    pipelines = ctx.api('GET', '/api/pipelines')
    if isinstance(pipelines, dict):pipelines = pipelines.get('items', [])
    output = [('pipeline', item['id'], item['published_run']) for item in pipelines
              if item.get('project_id') == ctx.project_id and item.get('published_run')]
    flows = ctx.api('GET', f'/api/projects/{ctx.project_id}/data-workflows')
    output.extend(('workflow', item['id'], item['published_run']) for item in flows.get('items', []) if item.get('published_run'))
    return sorted(output)


def _generated_diagnostics(artifacts):
    """User-visible text/tables only; do not count hidden input receipts as diagnosis."""
    texts, errors = [], []
    for artifact in artifacts:
        if not artifact.get('path') or artifact.get('kind') == 'source':continue
        path = Path(artifact['path'])
        if not path.is_absolute() or not path.is_file():continue
        try:
            if path.suffix.lower() in ('.txt', '.md', '.csv', '.tsv'):
                texts.append(path.read_text(encoding='utf-8-sig')[:300000])
            elif path.suffix.lower() == '.xlsx':
                from openpyxl import load_workbook
                book = load_workbook(path, read_only=True, data_only=False)
                try:
                    for sheet in book:
                        texts.append('\n'.join(' | '.join(str(value or '') for value in row) for row in sheet.iter_rows(values_only=True)))
                finally:book.close()
        except (OSError, ValueError, UnicodeError, zipfile.BadZipFile) as error:
            errors.append({'file': path.name, 'error': type(error).__name__})
    return '\n'.join(texts), errors


def _new_context(parent, code):
    spec = importlib.util.spec_from_file_location('official_negative_common', Path(__file__).resolve().parents[1] / 'common.py')
    common = importlib.util.module_from_spec(spec);spec.loader.exec_module(common)
    return common.CaseContext(parent.host, Path(parent.evidence) / code, 'finance_negative_' + code, parent.timeout)


def _turn_snapshot(turn):
    private_keys = {'api_key', 'authorization', 'app_token', 'atlas_operator_token', 'secret_config', 'secret_text'}
    def public(value):
        if isinstance(value, dict):return {key: '[redacted]' if key.lower() in private_keys else public(item) for key, item in value.items()}
        if isinstance(value, list):return [public(item) for item in value]
        return value
    return public({key: (turn or {}).get(key) for key in ('id', 'project_id', 'session_id', 'state', 'outcome', 'answer',
        'model_usage', 'task_contract', 'agent_trace', 'progress', 'error', 'error_code')})


def _merge_context(parent, child, code):
    if parent is child:
        for item in parent.record.get('turns', []):item['negative_case'] = code
        for item in parent.record.get('sources', []):item['negative_case'] = code
        return
    parent.turns.extend(child.turns)
    parent.record['real_model_evaluated'] |= child.record.get('real_model_evaluated', False)
    parent.record['turns'].extend(dict(item, negative_case=code) for item in child.record.get('turns', []))
    parent.record['sources'].extend(dict(item, negative_case=code, project_id=child.project_id) for item in child.record.get('sources', []))


def _selected_cases(ctx):
    requested = getattr(ctx, 'negative_case', None)
    if requested is None:return CASES
    if not isinstance(requested, str) or requested not in {code for code, _ in CASES}:
        raise ValueError('negative_case must be N01, N02, N03 or N04')
    return tuple(item for item in CASES if item[0] == requested)


def run(ctx):
    selected = _selected_cases(ctx)
    selected_codes = [code for code, _ in selected]
    not_run = [{'case': code, 'status': 'not_run'} for code, _ in CASES if code not in selected_codes]
    scope = 'Official ' + ', '.join(selected_codes) + ' only; unselected cases and N05-N07 were not run'
    evidence = Path(ctx.evidence);evidence.mkdir(parents=True, exist_ok=True)
    mapping, provenance = prepare_mapping(evidence / 'fixtures')
    policy = KIT / 'inputs' / 'phase1' / '业务口径.txt'
    result, used_projects = {}, set()
    ctx.record['official_negative_kit'] = {'scope': scope, 'selected_cases': selected_codes, 'not_run_cases': not_run, 'oracle_uploaded': False,
        'acceptance_manual_uploaded': False, 'mapping_provenance': provenance, 'cases': result}
    for index, (code, filename) in enumerate(selected):
        child = ctx if index == 0 else _new_context(ctx, code)
        if child.project_id in used_projects:raise AssertionError('Negative cases require distinct fresh projects')
        used_projects.add(child.project_id)
        negative = KIT / 'inputs' / 'negative' / filename
        source, turn, before, after, runtime_error, artifacts = {}, None, [], [], None, []
        prompt = ('请按已提供的2026年9月经营口径和省份有效映射，检查《' + filename +
            '》中的记录后再汇总。有不确定性不要发布正式结果。请说明原始问题位置、业务影响，以及我需要补充或修正什么；不要静默忽略源质量问题。')
        if code == 'N04':prompt += '请同时解释工作簿中公式的含义，以及现有公式缓存能否作为有效金额、是否实际发生过重算。'
        try:
            before = _publications(child)
            if before:raise AssertionError('Negative case project is not fresh')
            child.upload(policy);child.upload(mapping);source = child.upload(negative)
            turn = child.chat(prompt)
            artifacts = child.collect()
        except Exception as error:
            runtime_error = type(error).__name__ + ': ' + str(error)
            token = getattr(getattr(child, 'host', None), 'token', None)
            if token:runtime_error = runtime_error.replace(token, '[redacted access token]')
            runtime_error = runtime_error[:2000]
            if child.turns:turn = child.turns[-1]
        try:after = _publications(child)
        except Exception as error:
            after = [('publication_check_failed', type(error).__name__)]
            runtime_error = runtime_error or 'Publication state could not be verified'
        text, read_errors = _generated_diagnostics(artifacts)
        checked = evaluate(code, turn, source, before, after, generated_text=text, runtime_error=runtime_error)
        checked.update(project_id=child.project_id, source_id=source.get('id'),
            input_file=filename, input_sha256=hashlib.sha256(negative.read_bytes()).hexdigest(),
            prompt=prompt, publications_before=before, publications_after=after,
            probe={key: source.get(key) for key in ('status', 'error_code', 'error')},
            generated_artifacts=[{'id': item.get('id'), 'format': item.get('format'), 'path': str(item['path']) if item.get('path') else None} for item in artifacts],
            artifact_read_errors=read_errors, runtime_error=runtime_error, turn_evidence=_turn_snapshot(turn))
        result[code] = checked
        child.record['negative_acceptance'] = checked;child.save()
        if child is ctx:
            trace_path = evidence / 'case-trace.json'
            if trace_path.exists():shutil.copyfile(trace_path, evidence / (code + '-case-trace.json'))
        _merge_context(ctx, child, code)
        _save(evidence / (code + '-diagnostic.json'), checked)
        ctx.save()
        # A still-live timed-out task is not a business refusal. End this host
        # run so the outer controller shuts it down, rather than billing more
        # cases while unfinished work is still active.
        if turn and turn.get('state') not in ('done', 'failed'):
            for remaining, _ in selected[index + 1:]:
                result[remaining] = {'passed': False, 'status': 'not_run_after_nonterminal_failure'}
            break
    _save(evidence / 'finance-negative-acceptance.json', {'scope': scope, 'selected_cases': selected_codes,
        'not_run_cases': not_run, 'cases': result})
    return {'passed': len(result) == len(selected) and all(item['passed'] for item in result.values()),
        'checks': result, 'scope': scope, 'selected_cases': selected_codes, 'not_run_cases': not_run}
