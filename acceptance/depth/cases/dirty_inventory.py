"""Two-turn inventory acceptance; the oracle never enters the model context."""
from __future__ import annotations

import csv
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import hashlib
import json
from pathlib import Path
import re


DETAIL = ('record_id', 'sku', 'warehouse', 'stock_date', 'updated_at', 'quantity', 'inventory_cost', 'unit_cost')
SUMMARY = ('warehouse', 'record_count', 'quantity_total', 'inventory_cost_total', 'weighted_unit_cost')
EXCEPTIONS = ('record_id', 'updated_at', 'quantity', 'inventory_cost', 'reason')
FOUR = Decimal('0.0001')


def fixture(folder):
    folder = Path(folder);folder.mkdir(parents=True, exist_ok=True)
    warehouses = ('东仓', '南仓', '西仓', '北仓')
    rows = []
    for number in range(1, 1001):
        quantity = Decimal((number * 7) % 37 + 1)
        cost = Decimal('5.25') + Decimal(number % 91) * Decimal('0.13')
        row = dict(record_id=f'INV{number:04}', sku=f'SKU{number % 137:03}',
            warehouse=warehouses[number % 4], stock_date='2026-10-31',
            updated_at='2026-10-31 18:00:00', quantity=str(quantity),
            inventory_cost=format(quantity * cost, '.2f'))
        if number % 53 == 0:
            row['quantity'] = ''
        elif number % 67 == 0:
            row['quantity'] = '0'
            row['inventory_cost'] = '0.00' if number % 134 == 0 else '17.25'
        rows.append(row)
    for number in range(50, 1001, 50):
        row = dict(rows[number - 1])
        quantity = Decimal(row['quantity']) + 2
        cost = Decimal('5.25') + Decimal(number % 91) * Decimal('0.13')
        row.update(quantity=str(quantity), inventory_cost=format(quantity * cost, '.2f'),
                   warehouse=warehouses[(number + 1) % 4], updated_at='2026-11-01 09:00:00')
        rows.append(row)
    source = folder / '虚构库存_2026-10-31.csv'
    with source.open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=DETAIL[:-1]);writer.writeheader();writer.writerows(rows)
    # Independent row-wise reference: resolve versions first, then validate.
    latest = {}
    for row in rows:
        key = row['record_id']
        if key not in latest or row['updated_at'] > latest[key]['updated_at']:
            latest[key] = row
    clean, rejected, totals = {}, {}, {}
    for row in rows:
        key = (row['record_id'], row['updated_at'])
        if row is not latest[row['record_id']]:
            reason = 'duplicate_older'
        elif row['quantity'] == '':
            reason = 'missing_quantity'
        elif Decimal(row['quantity']) == 0 and Decimal(row['inventory_cost']) != 0:
            reason = 'zero_qty_nonzero_cost'
        else:
            quantity, cost = Decimal(row['quantity']), Decimal(row['inventory_cost'])
            clean[row['record_id']] = dict(row, unit_cost='' if quantity == 0 else str((cost / quantity).quantize(FOUR, rounding=ROUND_HALF_UP)))
            total = totals.setdefault(row['warehouse'], {'record_count': 0, 'quantity_total': Decimal(0), 'inventory_cost_total': Decimal(0)})
            total['record_count'] += 1;total['quantity_total'] += quantity;total['inventory_cost_total'] += cost
            continue
        rejected[key] = dict(row, reason=reason)
    summary = {warehouse: dict(total, weighted_unit_cost=(total['inventory_cost_total'] / total['quantity_total']).quantize(FOUR, rounding=ROUND_HALF_UP))
               for warehouse, total in totals.items()}
    assert len(rows) == len(clean) + len(rejected)
    return source, {'input_rows': len(rows), 'clean': clean, 'exceptions': rejected, 'summary': summary}


def number(value):
    try:
        result = Decimal(str(value).strip().replace(',', ''))
    except (InvalidOperation, ValueError):
        raise ValueError('not a decimal value') from None
    if not result.is_finite():
        raise ValueError('nonfinite decimal value')
    return result


def timestamp(value):
    if isinstance(value, datetime):
        return value.isoformat(timespec='seconds')
    return datetime.fromisoformat(str(value).strip().replace('Z', '+00:00')).isoformat(timespec='seconds')


def tables(artifacts):
    """Read real generated tables; never treat model prose as numeric evidence."""
    values, errors = [], []
    for artifact in artifacts:
        path = artifact if isinstance(artifact, (str, Path)) else artifact.get('path')
        if not path:
            continue
        path = Path(path)
        if not path.is_absolute() or not path.is_file():
            continue
        kind = path.suffix.lower().lstrip('.')
        if kind not in ('csv', 'xlsx', 'tsv'):
            continue
        try:
            if kind in ('csv', 'tsv'):
                with path.open(encoding='utf-8-sig', newline='') as stream:
                    reader = csv.DictReader(stream, delimiter='\t' if kind == 'tsv' else ',')
                    columns = [str(value or '').strip() for value in reader.fieldnames or []]
                    if len(columns) != len(set(columns)):
                        raise ValueError('duplicate table headers')
                    rows = [{str(key or '').strip(): value for key, value in row.items()} for row in reader
                            if any(value not in ('', None) for value in row.values())]
                values.append({'path': str(path), 'sheet': None, 'format': kind, 'columns': columns, 'rows': rows})
            else:
                from openpyxl import load_workbook
                workbook = load_workbook(path, read_only=True, data_only=False)
                try:
                    for sheet in workbook:
                        iterator = sheet.iter_rows(values_only=True)
                        columns = [str(value or '').strip() for value in next(iterator, ())]
                        if len(columns) != len(set(columns)):
                            raise ValueError('duplicate table headers')
                        rows = []
                        for row in iterator:
                            if not any(value not in ('', None) for value in row):
                                continue
                            if any(isinstance(value, str) and value.startswith('=') for value in row):
                                raise ValueError('unrecalculated formulas are not result values')
                            rows.append(dict(zip(columns, row)))
                        values.append({'path': str(path), 'sheet': sheet.title, 'format': kind, 'columns': columns, 'rows': rows})
                finally:
                    workbook.close()
        except (OSError, ValueError, UnicodeError, csv.Error) as error:
            errors.append({'file': path.name, 'error': str(error)})
    return values, errors


def validate_detail(table, golden):
    errors, seen = [], set()
    for row in table['rows']:
        key = str(row.get('record_id', '')).strip()
        expected = golden['clean'].get(key)
        if key in seen or expected is None:
            errors.append('duplicate or unexpected record_id: ' + key);continue
        seen.add(key)
        try:
            for field in ('sku', 'warehouse', 'stock_date'):
                actual = row.get(field)
                if isinstance(actual, datetime):actual = actual.date().isoformat()
                if str(actual).strip() != expected[field]:raise ValueError(field + ' differs')
            if timestamp(row.get('updated_at')) != timestamp(expected['updated_at']):raise ValueError('wrong retained version')
            for field in ('quantity', 'inventory_cost'):
                if number(row.get(field)) != number(expected[field]):raise ValueError(field + ' differs')
            if expected['unit_cost'] == '':
                if row.get('unit_cost') not in ('', None):raise ValueError('zero quantity unit_cost must be blank')
            elif number(row.get('unit_cost')) != number(expected['unit_cost']):raise ValueError('unit_cost differs')
        except (ValueError, TypeError) as error:
            errors.append(key + ': ' + str(error))
    if seen != set(golden['clean']):errors.append('clean record coverage differs')
    return errors


def validate_summary(table, golden):
    errors, seen = [], set()
    for row in table['rows']:
        key = str(row.get('warehouse', '')).strip();expected = golden['summary'].get(key)
        if key in seen or expected is None:
            errors.append('duplicate or unexpected warehouse: ' + key);continue
        seen.add(key)
        try:
            for field, value in expected.items():
                if number(row.get(field)) != Decimal(value):raise ValueError(field + ' differs')
        except (ValueError, TypeError) as error:
            errors.append(key + ': ' + str(error))
    if seen != set(golden['summary']):errors.append('warehouse coverage differs')
    return errors


def validate_exceptions(table, golden):
    errors, seen = [], set()
    for row in table['rows']:
        try:key = (str(row.get('record_id', '')).strip(), timestamp(row.get('updated_at')))
        except (ValueError, TypeError):errors.append('exception has invalid identity');continue
        expected = golden['exceptions'].get((key[0], key[1].replace('T', ' ')))
        if key in seen or expected is None:
            errors.append('duplicate or unexpected exception: ' + key[0]);continue
        seen.add(key)
        try:
            for field in ('quantity', 'inventory_cost'):
                if expected[field] == '':
                    if row.get(field) not in ('', None):raise ValueError('missing value was silently imputed')
                elif number(row.get(field)) != number(expected[field]):raise ValueError('original ' + field + ' differs')
            reason = str(row.get('reason', '')).lower()
            patterns = {'duplicate_older': r'duplicate|重复|旧版本|较旧',
                        'missing_quantity': r'missing|缺失|缺少|为空|空值',
                        'zero_qty_nonzero_cost': r'zero|零|0'}
            if not re.search(patterns[expected['reason']], reason):raise ValueError('exception reason missing or incorrect')
        except (ValueError, TypeError) as error:errors.append(key[0] + ': ' + str(error))
    expected_keys = {(key, timestamp(stamp)) for key, stamp in golden['exceptions']}
    if seen != expected_keys:errors.append('exception coverage differs')
    return errors


def find_table(values, columns, validator, golden):
    candidates = [value for value in values if set(columns).issubset(value['columns'])]
    failures = []
    for table in candidates:
        errors = validator(table, golden)
        if not errors:return {'passed': True, 'file': Path(table['path']).name, 'sheet': table['sheet'], 'rows': len(table['rows'])}
        failures.append({'file': Path(table['path']).name, 'sheet': table['sheet'], 'errors': errors[:8]})
    return {'passed': False, 'candidates': len(candidates), 'failures': failures}


def publications(ctx):
    values = ctx.api('GET', '/api/pipelines')
    if isinstance(values, tuple):values = values[1]
    if isinstance(values, dict):values = values.get('items', [])
    pipelines = [(value['id'], value.get('published_run')) for value in values
                 if value.get('project_id') == ctx.project_id and value.get('published_run')]
    workflows = ctx.api('GET', f'/api/projects/{ctx.project_id}/data-workflows')
    if isinstance(workflows, tuple):workflows = workflows[1]
    return sorted(pipelines + [(value['id'], value.get('published_run')) for value in workflows.get('items', []) if value.get('published_run')])


def run(ctx):
    evidence = Path(ctx.evidence);evidence.mkdir(parents=True, exist_ok=True)
    source, golden = fixture(evidence / 'fixtures')
    (evidence / 'golden-inventory.json').write_text(json.dumps({
        'input_rows': golden['input_rows'], 'clean': golden['clean'],
        'exceptions': list(golden['exceptions'].values()), 'summary': golden['summary']},
        ensure_ascii=False, indent=2, default=str), encoding='utf-8')
    ctx.upload(source)
    before = publications(ctx)
    first = ctx.chat('帮我把这份库存做成可用的清理明细和各仓库汇总，算数量、库存成本和单位成本。'
                     '有空数量、重复记录和零数量的情况，别直接糊过去，也别擅自把异常当成零。先确认影响结果的处理口径。')
    answer = str(first.get('answer', ''))
    checks = {'initial_needs_input': first.get('outcome') == 'needs_input',
        'asks_missing_quantity_policy': bool(re.search(r'缺|空|missing', answer, re.I)),
        'asks_duplicate_policy': bool(re.search(r'重复|duplicate|旧版本', answer, re.I)),
        'asks_zero_quantity_policy': bool(re.search(r'零|zero|\b0\b', answer, re.I)),
        'initial_no_publication': publications(ctx) == before}
    # Policies and output schema are requirements; no expected counts/totals are sent.
    final = ctx.chat('口径确认：库存日期为2026-10-31，record_id重复时只留updated_at最新的一行，较旧版本进异常清单。'
        '最新记录quantity缺失就整行排除计算并留异常，不补零；quantity为0但inventory_cost非0也整行排除并留异常。'
        'quantity和inventory_cost都为0的行保留在清理明细，unit_cost留空，计入记录数。其余行unit_cost=inventory_cost/quantity。'
        '按warehouse汇总记录数、数量和库存成本，加权单位成本=库存成本合计/数量合计。成本保留分，单位成本四位小数，ROUND_HALF_UP。'
        '请交三张可下载表，CSV或同一个Excel的三个页签均可：清理明细列record_id,sku,warehouse,stock_date,updated_at,quantity,inventory_cost,unit_cost；'
        '仓库汇总列warehouse,record_count,quantity_total,inventory_cost_total,weighted_unit_cost；'
        '异常表列record_id,updated_at,quantity,inventory_cost,reason，保留异常原值并解释排除原因。'
        '逐行处理完整原文件，不要拿预览或前20行当全量。先交文件供我核对，不要发布应用或设置自动更新。')
    generated, read_errors = tables(ctx.collect())
    results = {'clean_detail': find_table(generated, DETAIL, validate_detail, golden),
               'warehouse_summary': find_table(generated, SUMMARY, validate_summary, golden),
               'all_exceptions': find_table(generated, EXCEPTIONS, validate_exceptions, golden)}
    checks.update({name: result['passed'] for name, result in results.items()})
    checks['followup_terminal'] = final.get('state') == 'done'
    checks['no_automatic_publication'] = publications(ctx) == before
    ctx.record.update(fixture={'rows': golden['input_rows'], 'sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
        'synthetic_only': True, 'oracle_uploaded': False}, table_validation=results, table_read_errors=read_errors)
    return {'passed': all(checks.values()), 'checks': checks}
