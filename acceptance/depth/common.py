"""Real-model acceptance controller; never substitutes for model tool work."""
from __future__ import annotations
import csv
import hashlib
import importlib.util
import io
import json
from pathlib import Path
import re
import time
from urllib.parse import urlencode
import uuid
import zipfile
from urllib.error import HTTPError
from urllib.request import Request, ProxyHandler, build_opener

SECRET_KEYS={'api_key','authorization','app_token','atlas_operator_token','secret_config'}

def public(value):
    if isinstance(value,dict):return {k:('[redacted]' if k.lower() in SECRET_KEYS else public(v)) for k,v in value.items()}
    if isinstance(value,list):return [public(v) for v in value]
    return value

def table_rows(path,format=None):
    """Read actual exported data, without recalculating or executing formulas."""
    path=Path(path);format=(format or path.suffix.lstrip('.')).lower()
    def csv_table(raw,name):
        text=raw.decode('utf-8-sig');reader=csv.DictReader(io.StringIO(text))
        return {'name':name,'columns':reader.fieldnames or [],'rows':list(reader)}
    if format in ('csv','tsv'):
        return [csv_table(path.read_bytes(),path.name)]
    if format in ('json','jsonl'):
        raw=path.read_text(encoding='utf-8-sig')
        values=[json.loads(line) for line in raw.splitlines() if line.strip()] if format=='jsonl' else json.loads(raw)
        if isinstance(values,dict):values=values.get('rows',values.get('data',[]))
        if isinstance(values,list) and values and all(isinstance(v,dict) for v in values):
            return [{'name':path.name,'columns':list(values[0]),'rows':values}]
        return []
    if format=='xlsx':
        from openpyxl import load_workbook
        book=load_workbook(path,read_only=True,data_only=True);tables=[]
        try:
            for sheet in book:
                rows=list(sheet.iter_rows(values_only=True))
                nonempty=[r for r in rows if any(v is not None for v in r)]
                if not nonempty:continue
                header=[str(v) if v is not None else '' for v in nonempty[0]]
                tables.append({'name':sheet.title,'columns':header,
                    'rows':[dict(zip(header,row)) for row in nonempty[1:]]})
            return tables
        finally:book.close()
    if format in ('handoff','zip') and zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as archive:
            return [csv_table(archive.read(n),n) for n in archive.namelist() if n.lower().endswith('.csv')]
    return []

class CaseContext:
    def __init__(self,host,evidence,case_name,timeout=600):
        self.host,self.evidence,self.case_name,self.timeout=host,Path(evidence),case_name,timeout
        self.evidence.mkdir(parents=True,exist_ok=True)
        self.sources=[];self.source_records=[];self.turns=[];self.baseline_ids=None
        self.model_id=host.json('GET','/api/models')['default_model_id']
        self.project_id=host.json('POST','/api/projects',{'name':'能力纵深验收 · '+case_name,'request_id':uuid.uuid4().hex})['id']
        self.record={'case':case_name,'project_id':self.project_id,'synthetic_data_only':True,
            'real_model_evaluated':False,'model_configuration_contents_logged':False,
            'default_budgets_used':True,'sources':[],'turns':[],'collections':[]}

    def api(self,method,path,value=None):
        if not path.startswith('/api/'):path='/api'+('/' if not path.startswith('/') else '')+path
        return self.host.json(method,path,value)

    def save(self):
        (self.evidence/'case-trace.json').write_text(json.dumps(public(self.record),ensure_ascii=False,indent=2,default=str)+'\n',encoding='utf-8')

    def upload(self,path,name=None):
        path=Path(path);name=name or path.name;raw=path.read_bytes()
        boundary='atlas-depth-'+uuid.uuid4().hex
        body=[]
        for key,value in (('request_id',uuid.uuid4().hex),('local_processing','true')):
            body.append((f'--{boundary}\r\nContent-Disposition: form-data; name="{key}"\r\n\r\n{value}\r\n').encode())
        body.extend([(f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="{name}"\r\nContent-Type: application/octet-stream\r\n\r\n').encode('utf-8'),raw,f'\r\n--{boundary}--\r\n'.encode()])
        code,source=self.host.call('POST',f'/api/projects/{self.project_id}/attachments',data=b''.join(body),headers={'Content-Type':'multipart/form-data; boundary='+boundary})
        if code!=200:raise AssertionError('Upload rejected: HTTP '+str(code))
        self.sources.append(source['id'])
        deadline=time.monotonic()+180
        while time.monotonic()<deadline:
            current=next((v for v in self.api('GET','/api/sources') if v['id']==source['id']),None)
            if current and current.get('status') in ('ready','failed'):break
            time.sleep(.5)
        else:raise TimeoutError('Input profiling did not finish')
        self.source_records.append(current)
        self.record['sources'].append({'id':current['id'],'name':name,'bytes':len(raw),
            'sha256':hashlib.sha256(raw).hexdigest(),'probe_status':current.get('status'),
            'profile':current.get('profile'),'probe_error':current.get('error')})
        self.save()
        # A workbook needing header selection remains inspectable by the agent.
        # The harness must not silently choose its business table boundaries.
        return current

    def catalog(self):
        rows=[];offset=0
        while True:
            value=self.api('GET',f'/api/projects/{self.project_id}/artifacts?'+urlencode({'offset':offset,'limit':100}))
            rows.extend(value['items'])
            if value.get('next_offset') is None:return rows
            offset=value['next_offset']

    def chat(self,prompt,timeout=None):
        if self.baseline_ids is None:self.baseline_ids={v['id'] for v in self.catalog()}
        value={'text':prompt,'interaction_mode':'work','source_ids':self.sources,
               'model_id':self.model_id,'request_id':uuid.uuid4().hex}
        if self.turns and self.turns[-1].get('session_id'):value['session_id']=self.turns[-1]['session_id']
        turn=self.api('POST',f'/api/projects/{self.project_id}/messages',value)
        started=time.monotonic();deadline=started+(timeout or self.timeout)
        while time.monotonic()<deadline:
            turn=self.api('GET',f'/api/projects/{self.project_id}/turns/{turn["id"]}')
            if turn.get('state') in ('done','failed'):break
            time.sleep(1)
        terminal=turn.get('state') in ('done','failed')
        self.turns.append(turn)
        contract=turn.get('task_contract',{});budget=contract.get('budgets',{})
        if contract.get('budget_policy')!='unlimited' or any(budget.get(key) is not None
                for key in ('model_calls','input_characters','output_tokens','duration_ms')):
            raise AssertionError('Expected the owner interactive default policy; no test-only budget override is permitted')
        self.record['default_budget_policy']='unlimited'
        usage=turn.get('model_usage',{})
        self.record['real_model_evaluated']|=usage.get('calls',0)>0
        self.record['turns'].append({'prompt':prompt,'terminal':terminal,'elapsed_seconds':round(time.monotonic()-started,2),
            'turn':public(turn)})
        self.save()
        if not terminal:raise TimeoutError('Real model task did not reach a terminal state')
        return turn

    def collect(self):
        folder=self.evidence/'outputs';folder.mkdir(exist_ok=True)
        artifacts=[];seen=set()
        def download(record,format,url):
            if (url,format) in seen:return
            seen.add((url,format))
            if not url.startswith('/api/'):raise AssertionError('Artifact download must stay on the test host')
            request=Request(self.host.url+url,headers={'Authorization':'Bearer '+self.host.token})
            try:
                with build_opener(ProxyHandler({})).open(request,timeout=30) as response:
                    code=response.status;raw=response.read(80*1024*1024+1)
            except HTTPError as error:code,raw=error.code,error.read()
            if code!=200 or not isinstance(raw,bytes):
                artifacts.append({**record,'format':format,'download_status':code,'path':None});return
            if len(raw)>80*1024*1024:raise AssertionError('Generated artifact exceeded acceptance download limit')
            digest=hashlib.sha256(raw).hexdigest();path=folder/(digest[:20]+'.'+format)
            path.write_bytes(raw)
            artifacts.append({**record,'format':format,'path':path,'bytes':len(raw),'download_sha256':digest,'download_url':url})
        for item in self.catalog():
            if item.get('kind')=='source' or item['id'] in (self.baseline_ids or set()):continue
            for option in item.get('downloads',[]):
                format=option.get('format','')
                if format in ('csv','tsv','xlsx','json','jsonl','docx','pptx','md','txt','handoff','zip','pdf'):
                    download(item,format,option['url'])
        pipelines=[p for p in self.api('GET','/api/pipelines') if p.get('project_id')==self.project_id]
        for pipeline in pipelines:
            run=pipeline.get('candidate_run') or pipeline.get('published_run')
            if not run:continue
            details=self.api('GET','/api/runs/'+run)
            download({'id':pipeline['id'],'kind':'application_csv','name':pipeline.get('name'),
                'status':'candidate' if pipeline.get('candidate_run') else 'adopted','version':pipeline.get('commit'),
                'source_refs':details.get('manifest',{}).get('inputs',[]),'manifest':details.get('manifest')},
                'csv',f'/api/pipelines/{pipeline["id"]}/export?'+urlencode({'run_id':run}))
        self.record['pipelines']=pipelines
        self.record['collections'].append(public(artifacts));self.save()
        return artifacts

    def tables(self,artifacts=None):
        result=[]
        for artifact in artifacts if artifacts is not None else self.collect():
            if not artifact.get('path'):continue
            try:
                for table in table_rows(artifact['path'],artifact['format']):result.append({**table,'artifact':artifact})
            except (ValueError,OSError,zipfile.BadZipFile):continue
        return result
