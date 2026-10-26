"""Synthetic XLSX acceptance fixtures for source layout and value integrity."""
from datetime import datetime
from decimal import Decimal
from pathlib import Path
import xml.etree.ElementTree as ET
import zipfile

import pytest
from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.worksheet.table import Table

from agent_service.workbook_reader import WorkbookReadError, inspect_workbook, read_table

NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'


def save(tmp_path, rows, name='source.xlsx'):
    wb = Workbook()
    ws = wb.active
    ws.title = '明细'
    for row in rows:
        ws.append(row)
    path = tmp_path / name
    wb.save(path)
    return path


def patch_cells(path, updates, part='xl/worksheets/sheet1.xml'):
    with zipfile.ZipFile(path) as archive:
        payload = {name: archive.read(name) for name in archive.namelist()}
    document = ET.fromstring(payload[part])
    for coordinate, raw in updates.items():
        cell = document.find('.//' + NS + 'c[@r="' + coordinate + '"]')
        assert cell is not None
        cell.set('t', raw.get('type', 'n'))
        value = cell.find(NS + 'v')
        if value is None:
            value = ET.SubElement(cell, NS + 'v')
        value.text = raw.get('value')
    payload[part] = ET.tostring(document)
    with zipfile.ZipFile(path, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        for name, data in payload.items():
            archive.writestr(name, data)


def test_multilevel_headers_use_merge_anchors_without_shifting_values(tmp_path):
    wb = Workbook(); ws = wb.active; ws.title = '底稿'
    ws['A1'] = '地区'; ws.merge_cells('A1:A2')
    ws['B1'] = '2025年'; ws.merge_cells('B1:C1')
    ws['B2'] = '收入（万元）'; ws['C2'] = '占比'
    ws.append(['合肥', 12.123456789012, 0.125]); ws['C3'].number_format = '0.0%'
    path = tmp_path / 'merged.xlsx'; wb.save(path)
    result = read_table(path, sheet='底稿', header_rows=[1, 2])
    assert [f['label'] for f in result['fields']] == ['地区', '2025年 / 收入（万元）', '2025年 / 占比']
    assert result['rows'] == [{'c_1': '合肥', 'c_2': Decimal('12.123456789012'), 'c_3': Decimal('0.125')}]
    assert result['fields'][1]['header_coordinates'][0]['source_coordinate'] == 'B1'
    assert result['fields'][2]['header_coordinates'][0]['source_coordinate'] == 'B1'
    assert result['fields'][1]['unit'] == '万元'
    assert result['fields'][2]['unit'] == '%'
    assert result['fields'][2]['stored_semantics'] == 'ratio'
    assert result['source_rows'] == [3]


def test_merged_header_without_explicit_boundary_is_rejected(tmp_path):
    wb = Workbook(); ws = wb.active
    ws['A1'] = '报告标题'; ws.merge_cells('A1:B1')
    ws.append(['地区', '收入']); ws.append(['合肥', 10])
    path = tmp_path / 'title.xlsx'; wb.save(path)
    with pytest.raises(WorkbookReadError) as exc:
        read_table(path)
    assert exc.value.code == 'HEADER_SELECTION_REQUIRED'
    assert read_table(path, header_rows=[2])['rows'][0] == {'c_1': '合肥', 'c_2': Decimal('10')}


def test_multi_sheet_selection_must_be_explicit(tmp_path):
    wb = Workbook(); wb.active.title = '订单'; wb.active.append(['金额']); wb.active.append([10])
    ws = wb.create_sheet('退款'); ws.append(['金额']); ws.append([2])
    path = tmp_path / 'sheets.xlsx'; wb.save(path)
    info = inspect_workbook(path)
    assert [s['name'] for s in info['sheets']] == ['订单', '退款']
    with pytest.raises(WorkbookReadError) as exc:
        read_table(path)
    assert exc.value.code == 'SHEET_SELECTION_REQUIRED'
    assert read_table(path, sheet='退款')['rows'] == [{'c_1': Decimal('2')}]


def test_named_table_resolves_sheet_range_and_excludes_totals(tmp_path):
    wb = Workbook(); wb.active.title = '说明'; wb.active['A1'] = '说明'
    ws = wb.create_sheet('数据'); ws.append(['其他']); ws.append([])
    ws.append(['客户', '金额']); ws.append(['甲', 10]); ws.append(['乙', 20]); ws.append(['合计', 30])
    table = Table(displayName='Revenue', ref='A3:B6', totalsRowCount=1)
    ws.add_table(table)
    path = tmp_path / 'named.xlsx'; wb.save(path)
    result = read_table(path, table='Revenue')
    assert result['selection']['sheet'] == '数据'
    assert result['source_rows'] == [4, 5]
    assert result['rows'][1]['c_2'] == Decimal('20')


def test_vertical_text_merge_requires_explicit_label_policy(tmp_path):
    wb = Workbook(); ws = wb.active
    ws.append(['区域', '金额']); ws.append(['华东', 10]); ws.append([None, 20]); ws.merge_cells('A2:A3')
    path = tmp_path / 'labels.xlsx'; wb.save(path)
    with pytest.raises(WorkbookReadError) as exc:
        read_table(path)
    assert exc.value.code == 'WORKBOOK_MERGE_AMBIGUOUS'
    result = read_table(path, merged_label_columns=['c_1'])
    assert result['rows'] == [{'c_1': '华东', 'c_2': Decimal('10')}, {'c_1': '华东', 'c_2': Decimal('20')}]
    assert result['special_cells']['3']['c_1']['source_coordinate'] == 'A2'
    assert result['special_cells']['3']['c_1']['propagation'] == 'explicit_vertical_text_label'


def test_merged_numeric_detail_is_never_forward_filled(tmp_path):
    wb = Workbook(); ws = wb.active
    ws.append(['客户', '金额']); ws.append(['甲', 10]); ws.append(['乙', None]); ws.merge_cells('B2:B3')
    path = tmp_path / 'measure.xlsx'; wb.save(path)
    with pytest.raises(WorkbookReadError) as exc:
        read_table(path, merged_label_columns=['B'])
    assert exc.value.code == 'MERGED_MEASURE_UNSAFE'


def test_horizontal_detail_merge_is_not_turned_into_two_fields(tmp_path):
    wb = Workbook(); ws = wb.active
    ws.append(['地区', '产品']); ws.append(['华东', None]); ws.merge_cells('A2:B2')
    path = tmp_path / 'horizontal.xlsx'; wb.save(path)
    with pytest.raises(WorkbookReadError) as exc:
        read_table(path, merged_label_columns=['A', 'B'])
    assert exc.value.code == 'WORKBOOK_MERGE_AMBIGUOUS'


def test_merge_crossing_header_data_is_rejected(tmp_path):
    wb = Workbook(); ws = wb.active
    ws.append(['地区', '金额']); ws.append([None, 10]); ws.merge_cells('A1:A2')
    path = tmp_path / 'crossing.xlsx'; wb.save(path)
    with pytest.raises(WorkbookReadError) as exc:
        read_table(path, header_rows=[1], merged_label_columns=['A'])
    assert exc.value.code == 'WORKBOOK_MERGE_AMBIGUOUS'


def test_formula_without_cached_value_blocks_execution_but_not_probe(tmp_path):
    path = save(tmp_path, [['金额', '两倍'], [10, '=A2*2']])
    with pytest.raises(WorkbookReadError) as exc:
        read_table(path)
    assert exc.value.code == 'FORMULA_CACHE_MISSING'
    assert exc.value.details == {'coordinate': 'B2'}
    preview = read_table(path, require_formula_cache=False)
    assert preview['rows'][0]['c_2'] is None
    assert not preview['executable_values_complete']
    assert preview['special_cells']['2']['c_2']['formula']['references'] == ['A2']
    assert preview['special_cells']['2']['c_2']['cache_status'] == 'missing'


def test_missing_formula_only_row_remains_visible_in_probe(tmp_path):
    path = save(tmp_path, [['结果'], ['=1+1']])
    preview = read_table(path, require_formula_cache=False)
    assert preview['source_rows'] == [2] and preview['rows'] == [{'c_1': None}]
    assert not preview['executable_values_complete']


def test_formula_outside_explicit_selected_table_does_not_block_reading(tmp_path):
    path = save(tmp_path, [['=TODAY()'], [], ['客户', '金额'], ['甲', 10]])
    result = read_table(path, header_rows=[3])
    assert result['rows'] == [{'c_1': '甲', 'c_2': Decimal('10')}]
    assert result['executable_values_complete']


def test_formula_header_cache_missing_is_visible_in_probe(tmp_path):
    path = save(tmp_path, [['="金额"'], [10]])
    preview = read_table(path, require_formula_cache=False)
    assert not preview['executable_values_complete']
    assert preview['header_cells']['A1']['cache_status'] == 'missing'
    assert preview['fields'][0]['semantic_status'] == 'requires_review'


def test_formula_cache_and_formula_knowledge_are_preserved_separately(tmp_path):
    path = save(tmp_path, [['金额', '两倍'], [Decimal('0.1'), '=A2*2']])
    patch_cells(path, {'B2': {'value': '0.2'}})
    result = read_table(path)
    assert result['rows'][0]['c_2'] == Decimal('0.2')
    cell = result['special_cells']['2']['c_2']
    assert cell['formula']['formula'] == '=A2*2'
    assert cell['cache_status'] == 'available' and cell['cache_freshness'] == 'unknown'
    assert result['formula_cache_freshness'] == 'unknown'
    assert 'no_formula_recalculation' in result['scope']


def test_cached_formula_error_is_not_a_successful_number(tmp_path):
    path = save(tmp_path, [['结果'], ['=1/0']])
    patch_cells(path, {'A2': {'value': '#DIV/0!', 'type': 'e'}})
    with pytest.raises(WorkbookReadError) as exc:
        read_table(path)
    assert exc.value.code == 'FORMULA_CACHED_ERROR'
    preview = read_table(path, require_formula_cache=False)
    assert preview['rows'][0]['c_1'] is None and not preview['executable_values_complete']


def test_raw_numeric_lexical_precision_survives_python_float_reader(tmp_path):
    path = save(tmp_path, [['编号'], [123456789012345678]])
    patch_cells(path, {'A2': {'value': '123456789012345678'}})
    result = read_table(path)
    assert result['rows'][0]['c_1'] == Decimal('123456789012345678')
    assert any(w['code'] == 'EXCEL_PRECISION_LIMIT' for w in result['warnings'])
    assert 'no recovery inferred' in result['special_cells']['2']['c_1']['numeric_precision_warning']


def test_dates_text_identifiers_percent_and_units_keep_source_semantics(tmp_path):
    wb = Workbook(); ws = wb.active
    ws.append(['日期', '客户编号', '占比', '金额(万元)'])
    ws.append([datetime(2026, 10, 21, 13, 30), '000012345678901234567', 0.15, 2])
    ws['A2'].number_format = 'yyyy-mm-dd hh:mm'; ws['C2'].number_format = '0.0%'
    path = tmp_path / 'semantics.xlsx'; wb.save(path)
    result = read_table(path)
    assert result['rows'][0] == {'c_1': '2026-10-21T13:30:00', 'c_2': '000012345678901234567', 'c_3': Decimal('0.15'), 'c_4': Decimal('2')}
    assert result['fields'][0]['observed_kinds'] == ['date_or_time']
    assert result['fields'][2]['display_multiplier'] == '100'
    assert result['fields'][3]['unit'] == '万元'


def test_indented_text_produces_layout_hierarchy_not_business_mapping(tmp_path):
    wb = Workbook(); ws = wb.active
    for row in [['组织', '人数'], ['华东', 10], ['合肥办事处', 4], ['南京办事处', 6], ['销售组', 3]]:
        ws.append(row)
    ws['A3'].alignment = Alignment(indent=1); ws['A4'].alignment = Alignment(indent=1)
    ws['A5'].alignment = Alignment(indent=2)
    path = tmp_path / 'indent.xlsx'; wb.save(path)
    result = read_table(path)
    assert result['special_cells']['5']['c_1']['hierarchy_path'] == ['华东', '南京办事处', '销售组']
    assert result['special_cells']['5']['c_1']['hierarchy_semantics'] == 'layout_candidate_not_business_mapping'


def test_hidden_rows_are_preserved_or_excluded_by_explicit_policy(tmp_path):
    wb = Workbook(); ws = wb.active
    for row in [['地区', '金额'], ['甲', 10], ['乙', 20], ['丙', 30]]: ws.append(row)
    ws.row_dimensions[3].hidden = True
    path = tmp_path / 'hidden.xlsx'; wb.save(path)
    include = read_table(path)
    assert include['source_rows'] == [2, 3, 4]
    assert 'hidden_in_source' in include['row_roles']['3']
    assert read_table(path, hidden_rows='exclude')['source_rows'] == [2, 4]
    with pytest.raises(WorkbookReadError) as exc:
        read_table(path, hidden_rows='error')
    assert exc.value.code == 'HIDDEN_ROWS_PRESENT'


def test_subtotal_candidate_is_not_silently_summed_as_detail(tmp_path):
    path = save(tmp_path, [['区域', '收入'], ['甲', 10], ['乙', 20], ['合计', 30]])
    result = read_table(path)
    assert result['row_roles']['4'] == ['subtotal_label_candidate']
    assert len(result['rows']) == 3
    with pytest.raises(WorkbookReadError) as exc:
        read_table(path, subtotal_policy='error')
    assert exc.value.code == 'SUBTOTAL_ROWS_PRESENT'


def test_duplicate_and_missing_headers_keep_both_column_coordinates(tmp_path):
    path = save(tmp_path, [['金额', '金额', None], [10, 20, 30]])
    result = read_table(path)
    assert result['rows'][0] == {'c_1': Decimal('10'), 'c_2': Decimal('20'), 'c_3': Decimal('30')}
    assert [f['label'] for f in result['fields']] == ['金额', '金额 [B]', '未命名列C']
    assert result['fields'][1]['semantic_status'] == result['fields'][2]['semantic_status'] == 'requires_review'


def test_range_and_noncontiguous_source_rows_remain_reconstructable(tmp_path):
    wb = Workbook(); ws = wb.active
    ws['C3'] = '客户'; ws['D3'] = '金额'; ws['C4'] = '甲'; ws['D4'] = 10
    ws['C6'] = '乙'; ws['D6'] = 20
    path = tmp_path / 'range.xlsx'; wb.save(path)
    result = read_table(path, cell_range='C3:D6')
    assert result['source_rows'] == [4, 6]
    assert result['fields'][0]['excel_column'] == 'C'
    assert result['rows'][1]['c_2'] == Decimal('20')


def test_row_limit_and_metadata_limit_cannot_fake_full_read(tmp_path):
    path = save(tmp_path, [['金额', '两倍'], [1, '=A2*2'], [2, '=A3*2']])
    patch_cells(path, {'B2': {'value': '2'}, 'B3': {'value': '4'}})
    limited = read_table(path, limit=1)
    assert limited['truncated'] and not limited['executable_values_complete']
    with pytest.raises(WorkbookReadError) as exc:
        read_table(path, provenance_limit=0)
    assert exc.value.code == 'WORKBOOK_METADATA_LIMIT'
    preview = read_table(path, require_formula_cache=False, provenance_limit=0)
    assert not preview['metadata_complete'] and not preview['executable_values_complete']


def test_date_columns_use_coordinate_metadata_without_exploding_exception_budget(tmp_path):
    path = save(tmp_path, [['日期', '金额']] + [[datetime(2026, 10, 21), i] for i in range(1000)])
    result = read_table(path, provenance_limit=0)
    assert len(result['rows']) == 1000 and result['metadata_complete']
    assert result['source_rows'][-1] == 1001 and result['fields'][0]['observed_kinds'] == ['date_or_time']


def test_mixed_percentage_and_plain_number_cannot_share_fraction_semantics(tmp_path):
    wb = Workbook(); ws = wb.active; ws.append(['占比']); ws.append([0.15]); ws.append([15])
    ws['A2'].number_format = '0%'
    path = tmp_path / 'mixed-percent.xlsx'; wb.save(path)
    with pytest.raises(WorkbookReadError) as exc:
        read_table(path)
    assert exc.value.code == 'MIXED_PERCENTAGE_REPRESENTATION'
    preview = read_table(path, require_formula_cache=False)
    assert preview['rows'] == [{'c_1': Decimal('0.15')}, {'c_1': Decimal('15')}]
    assert not preview['executable_values_complete']
    assert preview['fields'][0]['stored_semantics'] == 'ambiguous_percentage_representation'
    assert preview['fields'][0]['semantic_status'] == 'requires_review'


def test_region_discovery_finds_top_offset_title_and_actual_header(tmp_path):
    wb = Workbook(); ws = wb.active; ws.title='报告'
    ws['B5']='2026年销售报告';ws.merge_cells('B5:D5')
    ws.append([])
    for col,value in zip(('B','C','D'),('地区','金额','占比')):ws[col+'6']=value
    ws['B7']='合肥';ws['C7']=10;ws['D7']=0.2
    ws['B8']='南京';ws['C8']=20;ws['D8']=0.4
    path=tmp_path/'offset.xlsx';wb.save(path)
    info=inspect_workbook(path)
    candidate=next(c for c in info['candidate_regions'] if c['range']=='B6:D8')
    assert candidate['header_rows']==[6] and candidate['data_start_row']==7
    assert candidate['proposed_options']['sheet']=='报告'
    assert candidate['context_rows']==[5] and candidate['needs_confirmation']
    assert [c['text'] for c in candidate['header_preview']]==['地区','金额','占比']
    extracted=read_table(path,**candidate['proposed_options'])
    assert extracted['source_rows']==[7,8] and extracted['rows'][0]['c_1']=='合肥'


def test_region_discovery_keeps_vertical_tables_separate(tmp_path):
    path=save(tmp_path,[['客户','金额'],['甲',10],[],[],['部门','人数'],['财务',3],['研发',5]])
    info=inspect_workbook(path)
    assert {c['range'] for c in info['candidate_regions']}=={'A1:B2','A5:B7'}
    assert all(c['needs_confirmation'] for c in info['candidate_regions'])
    assert all(any('未自动拼接' in r for r in c['reasons']) for c in info['candidate_regions'])


def test_region_discovery_proposes_side_by_side_regions_without_auto_selecting(tmp_path):
    path=save(tmp_path,[['客户','收入',None,None,'部门','人数'],['甲',10,None,None,'研发',3],['乙',20,None,None,'财务',2]])
    info=inspect_workbook(path)
    assert {c['range'] for c in info['candidate_regions']}=={'A1:B3','E1:F3'}
    assert all(c['confidence']=='layout_ambiguous' and c['needs_confirmation'] for c in info['candidate_regions'])


def test_region_discovery_proposes_multilevel_merged_headers(tmp_path):
    wb=Workbook();ws=wb.active;ws['C10']='区域';ws.merge_cells('C10:C11')
    ws['D10']='2026年';ws.merge_cells('D10:E10')
    ws['D11']='收入';ws['E11']='成本';ws['C12']='华东';ws['D12']=10;ws['E12']=3
    path=tmp_path/'deepheaders.xlsx';wb.save(path)
    info=inspect_workbook(path)
    candidate=next(c for c in info['candidate_regions'] if c['range']=='C10:E12')
    assert candidate['header_rows']==[10,11] and candidate['data_start_row']==12
    result=read_table(path,**candidate['proposed_options'])
    assert result['fields'][1]['header_path']==['2026年','收入']


def test_region_discovery_prefers_native_table_evidence(tmp_path):
    wb=Workbook();ws=wb.active;ws.append(['客户','金额']);ws.append(['甲',10]);ws.add_table(Table(displayName='Orders',ref='A1:B2'))
    path=tmp_path/'native-table.xlsx';wb.save(path)
    regions=inspect_workbook(path)['candidate_regions']
    assert len(regions)==1 and regions[0]['confidence']=='source_declared_table'
    assert regions[0]['proposed_options']['table']=='Orders' and regions[0]['needs_confirmation']


@pytest.mark.parametrize('kwargs', [
    {'limit': True}, {'limit': 0}, {'header_rows': [1, 3]}, {'header_rows': [True]},
    {'data_start_row': 1}, {'cell_range': 'A:B'}, {'cell_range': 'A1:ZZ1'},
    {'merged_label_columns': ['c_999']}, {'hidden_rows': 'guess'},
])
def test_invalid_selection_never_guesses_structure(tmp_path, kwargs):
    path = save(tmp_path, [['字段'], ['值']])
    with pytest.raises(WorkbookReadError):
        read_table(path, **kwargs)
