import hashlib
import json

import pytest

from agent_service import business_reports as reports, project_data, workspace
from agent_service.config import DATA
from agent_service.execution import process_once
from agent_service.jobs import get_job
from agent_service.sources import revision
from agent_service.store import now, put, seal, uid


def schema(period='2026-09'):
    return {'columns':[
        {'name':'office','data_type':'text','meaning':'办事处'},
        {'name':'actual','data_type':'decimal','meaning':'本期净收入','unit':'CNY','period':period,'aggregation':'additive'},
        {'name':'budget','data_type':'decimal','meaning':'本期净收入预算','unit':'CNY','period':period,'aggregation':'additive'}],
        'grain':['office'],'meaning':'每个办事处一行的本期实际和预算'}


def table():
    return {'schema':schema(),'rows':[{'office':'东区','actual':'9007199254740.000001','budget':'9007199254740.000000'},
                                      {'office':'西区','actual':'100.000001','budget':'100.000000'}],
            'total':2,'complete_result_available':True}


def recipe(current=None):
    current=current or schema()
    return {'version':1,'title':'经营月报','period':'2026-09','expected_schema_sha256':reports.schema_fingerprint(current),
            'metrics':[{'id':'actual','label':'净收入','kind':'sum','column':'actual'},
                       {'id':'variance','label':'净收入预算差异','kind':'variance','column':'actual',
                        'baseline_column':'budget','rounding_scale':12}],
            'assumptions':['预算使用已核准版本；收入增长的原因尚待业务核实。']}


def test_exact_actual_budget_assertions_and_no_causal_explanation():
    value=reports.evaluate(recipe(),table())
    assert value['assertions'][0]['value']=='9007199254840.000002'
    assert value['assertions'][1]['value']=='0.000002'
    assert value['assertions'][1]['relative_change']=='0.000000000000'
    assert value['assertions'][1]['relative_change_formula']=='(actual-baseline)/baseline'
    assert value['selected_rows']==2 and value['input_rows']==2
    assert value['no_causal_inference'] is True


def test_preview_and_duplicate_business_grain_cannot_become_a_report():
    partial=table();partial['rows']=partial['rows'][:1]
    with pytest.raises(ValueError,match='完整'):reports.evaluate(recipe(),partial)
    duplicate=table();duplicate['rows'][1]['office']='东区'
    with pytest.raises(ValueError,match='重复'):reports.evaluate(recipe(),duplicate)


def test_semantic_drift_and_period_mismatch_stop_reuse():
    current=table();current['schema']['columns'][1]['unit']='USD'
    with pytest.raises(ValueError,match='口径已变化'):reports.evaluate(recipe(),current)
    current=table();current['schema']['columns'][2]['period']='2026-08'
    with pytest.raises(ValueError,match='期间'):reports.evaluate(recipe(),current)
    assert reports.schema_fingerprint(schema('2026-08'))==reports.schema_fingerprint(schema())


def test_null_numeric_and_no_selected_period_are_not_silently_zero():
    current=table();current['rows'][0]['actual']=None
    with pytest.raises(ValueError,match='缺失'):reports.evaluate(recipe(),current)
    r=recipe();r['filters']={'office':'无此办事处'}
    with pytest.raises(ValueError,match='没有数据'):reports.evaluate(r,table())


def test_zero_baseline_and_explicit_non_terminating_ratio_policy():
    current=table();current['rows']=[{'office':'东区','actual':'1','budget':'0'}];current['total']=1
    r=recipe();r['metrics']=r['metrics'][1:]
    with pytest.raises(ValueError,match='基数为0'):reports.evaluate(r,current)
    r['metrics'][0]['zero_policy']='null'
    assert reports.evaluate(r,current)['assertions'][0]['relative_change']=='不适用（比较基数为0）'
    current['rows'][0]['budget']='3';r['metrics'][0].update(kind='ratio',rounding_scale=None)
    with pytest.raises(ValueError,match='保留位数'):reports.evaluate(r,current)
    r['metrics'][0]['rounding_scale']=4
    assert reports.evaluate(r,current)['assertions'][0]['value']=='0.3333'


def test_percentage_points_not_fraction_or_summed_percentage():
    current=table();current['rows']=[{'office':'东区','actual':'0.15','budget':'0.10'}];current['total']=1
    for column in current['schema']['columns'][1:]:
        column.update(unit='1',value_format='percent_fraction',aggregation='non_additive')
    r=recipe(current['schema']);r['metrics']=r['metrics'][1:];r['metrics'][0]['reduction']='value'
    assertion=reports.evaluate(r,current)['assertions'][0]
    assert assertion['value']=='5.00' and assertion['unit']=='百分点' and assertion['relative_change']=='0.500000000000'
    r['metrics'][0]['reduction']='sum'
    with pytest.raises(ValueError,match='百分比'):reports.evaluate(r,current)
    current['schema']['columns'][2]['value_format']='percent_points';r=recipe(current['schema']);r['metrics']=r['metrics'][1:]
    r['metrics'][0]['reduction']='value'
    with pytest.raises(ValueError,match='百分比表示'):reports.evaluate(r,current)


def test_period_column_requires_explicit_month_values_and_filters_all_complete_rows():
    current=table()
    current['schema']['columns'].append({'name':'month','data_type':'text','meaning':'会计期间'})
    for column in current['schema']['columns'][1:3]:column['period']=''
    current['schema']['grain']=['office','month']
    current['rows'][0]['month']='2026-09';current['rows'][1]['month']='2026-08'
    r=recipe(current['schema']);r['period_column']='month'
    result=reports.evaluate(r,current)
    assert result['input_rows']==2 and result['selected_rows']==1
    assert result['assertions'][0]['value']=='9007199254740.000001'


def test_one_row_value_never_takes_first_row_implicitly():
    r=recipe();r['metrics']=[{'id':'one','label':'单值','kind':'value','column':'actual'}]
    with pytest.raises(ValueError,match='恰好一行'):reports.evaluate(r,table())


def test_row_count_cannot_label_unknown_or_other_month_as_target_period():
    r=recipe();r['metrics']=[{'id':'count','label':'办事处数量','kind':'row_count'}]
    current=table()
    assert reports.evaluate(r,current)['assertions'][0]['value']=='2'
    for column in current['schema']['columns'][1:]:column['period']=''
    with pytest.raises(ValueError,match='期间范围'):reports.evaluate(r,current)
    current['schema']['columns'][1]['period']='2026-09';current['schema']['columns'][2]['period']='2026-08'
    with pytest.raises(ValueError,match='期间范围'):reports.evaluate(r,current)


def real_result(sql=False):
    pid=put('project',{'name':'业务简报验收','created_at':now()},'p_'+uid())['id']
    raw='office,actual,budget\nEast,100.000001,100\nWest,200.000001,200\n'.encode()
    name=uid()+'.csv';path=DATA/'uploads'/name;path.parent.mkdir(parents=True,exist_ok=True);path.write_bytes(raw)
    source=revision(put('source',{'project_id':pid,'name':'month.csv','kind':'file',
        'secret_config':seal({'stored_name':name}),'content_hash':hashlib.sha256(raw).hexdigest(),
        'bytes':len(raw),'status':'ready','created_at':now()},'s_'+uid()))
    doc={'version':1,'plan':{'name':'月末实际预算','inputs':{'detail':schema()},'stages':[
        {'id':'keep','kind':'select','description':'保留全部已声明月末结果','input':'detail','columns':['office','actual','budget']}],
        'output':'keep'},'bindings':{'detail':{'source_id':source['id'],'column_map':{'c_1':'office','c_2':'actual','c_3':'budget'}}}}
    if sql:
        doc['plan']['stages']=[{'id':'keep','kind':'sql','description':'自定义已声明SQL输出口径',
            'inputs':['detail'],'sql':'SELECT office, actual, budget FROM detail','output_schema':schema()}]
    workspace.write_file(pid,'workflows/month.json',json.dumps(doc))
    workflow=project_data.register(pid,'workflows/month.json')
    run=project_data.enqueue_workflow(pid,workflow['workflow_id']);assert process_once()
    assert get_job(run['id'])['status']=='validated'
    return pid,run['id']


def test_real_workflow_report_is_reusable_hashed_and_project_scoped():
    pid,rid=real_result();info=reports.inspect_result(pid,rid)
    r=recipe();r['expected_schema_sha256']=info['schema_sha256']
    r['title']='经营月报 <script>alert(1)</script>'
    value=reports.render(pid,rid,r)
    assert value['assertions'][0]['value']=='300.000002'
    assert value['recipe_commit'] and value['complete']
    output=reports.report_file(pid,value['id']);text=output.read_text()
    assert '300.000002' in text and '\\<script\\>' in text
    assert '业务原因和后续判断需要另附依据' in text
    assert reports.render_saved(pid,rid,'reports/month-end.json')['id']==value['id']
    manifest=json.loads(reports.report_file(pid,value['id'],'manifest').read_text())
    assert manifest['handoff_sha256']==info['handoff_sha256']
    assert manifest['evaluation']['input_rows']==2
    assert '300.000002' not in workspace.read_file(pid,'reports/month-end.json')
    other=put('project',{'name':'other','created_at':now()},'p_'+uid())['id']
    with pytest.raises(ValueError,match='不属于'):reports.report_file(other,value['id'])
    output.write_text('changed')
    with pytest.raises(ValueError,match='完整性'):reports.report_file(pid,value['id'])
    with pytest.raises(ValueError,match='完整性'):reports.render_saved(pid,rid,'reports/month-end.json')


def test_recipe_paths_and_raw_unsupported_fields_fail_before_write():
    pid,rid=real_result();r=recipe();r['expected_schema_sha256']=reports.inspect_result(pid,rid)['schema_sha256']
    with pytest.raises(ValueError,match='只能操作'):reports.render(pid,rid,r,recipe_path='../escape.json')
    with pytest.raises(ValueError):reports.render(pid,rid,{**r,'invented_summary':'revenue doubled due to management'})


def test_sql_result_requires_genuine_server_adoption_before_brief():
    pid,rid=real_result(sql=True);job=get_job(rid)
    assert job['data_result']['requires_semantic_review'] is True
    with pytest.raises(ValueError,match='待核对'):reports.inspect_result(pid,rid)
    with pytest.raises(ValueError,match='待核对'):reports.render(pid,rid,recipe())
    project_data.publish_result(pid,rid,job['commit'])
    assert get_job(rid)['status']=='succeeded'
    info=reports.inspect_result(pid,rid)
    assert info['server_adopted_result'] and info['requires_semantic_review']
    r=recipe();r['expected_schema_sha256']=info['schema_sha256']
    result=reports.render(pid,rid,r)
    manifest=json.loads(reports.report_file(pid,result['id'],'manifest').read_text())
    assert manifest['requires_semantic_review'] is True and manifest['server_adopted_result'] is True


def test_office_exports_reuse_exact_frozen_assertions_and_verify_saved_hashes():
    from docx import Document
    from pptx import Presentation
    pid,rid=real_result();r=recipe();r['expected_schema_sha256']=reports.inspect_result(pid,rid)['schema_sha256']
    # Seven main metrics plus a comparison force multi-page native slide tables.
    r['metrics']=[{'id':'m'+str(i),'label':'净收入'+str(i),'kind':'sum','column':'actual'} for i in range(7)]+r['metrics'][1:]
    report=reports.render(pid,rid,r);manifest=json.loads(reports.report_file(pid,report['id'],'manifest').read_text())
    expected=reports._assertion_rows(manifest)
    docx=reports.report_file(pid,report['id'],'docx');docx_bytes=docx.read_bytes()
    assert [[c.text for c in row.cells] for row in list(Document(docx).tables[0].rows)[1:]]==expected
    deck=reports.report_file(pid,report['id'],'pptx');deck_bytes=deck.read_bytes()
    tables=[shape.table for slide in Presentation(deck).slides for shape in slide.shapes if shape.has_table]
    assert len(tables)>=2 and all(len(table.rows)<=7 for table in tables)
    assert [[c.text for c in row.cells] for table in tables for row in list(table.rows)[1:]]==expected
    # Reopening/re-rendering the saved recipe must retain cache hashes.
    assert reports.render_saved(pid,rid,'reports/month-end.json')['id']==report['id']
    assert reports.report_file(pid,report['id'],'docx').read_bytes()==docx_bytes
    assert reports.report_file(pid,report['id'],'pptx').read_bytes()==deck_bytes
    deck.write_bytes(deck_bytes+b'changed')
    with pytest.raises(ValueError,match='完整性'):reports.report_file(pid,report['id'],'pptx')
