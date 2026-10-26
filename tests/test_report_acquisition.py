"""Direct report acquisition uses mocked transport; no live report/account tested."""
import io
import json
from pathlib import Path
import httpx
import pytest
from pydantic import ValidationError
from openpyxl import Workbook
from agent_service import report_acquisition as reports


def request(**values):
    return reports.ReportRequest(scope_id='project_connector_owner',request_id='export_1',url='https://reports.example/export',format='csv',**values)


def client(monkeypatch,handler):
    monkeypatch.setattr(reports,'safe_api_url',lambda u:'203.0.113.20')
    def factory(**kw):return httpx.Client(transport=httpx.MockTransport(handler),**kw)
    return factory


def write(tmp_path,data,name='report.csv'):
    path=tmp_path/name;path.write_bytes(data);return path


def test_requested_filters_are_not_verified_without_export_checks(tmp_path,monkeypatch):
    seen=[]
    def respond(req):
        seen.append(req)
        return httpx.Response(200,content=b'office,amount\nShanghai,10.123456789012\n',headers={'Content-Type':'text/csv'})
    factory=client(monkeypatch,respond)
    r=request(filters={'office':'Shanghai'},context={'unit':'CNY','period':'2026-10'})
    receipt=reports.acquire_report(r,tmp_path,headers={'Authorization':'Bearer disposable-secret'},client_factory=factory)
    assert not receipt['filters_verified'] and not receipt['context_verified']
    assert receipt['filters_requested']=={'office':'Shanghai'}
    assert receipt['rows']==1 and receipt['completion_verified']
    assert seen[0].url.host=='203.0.113.20' and seen[0].headers['Host']=='reports.example'
    assert seen[0].extensions['sni_hostname']=='reports.example'
    assert seen[0].url.params['office']=='Shanghai'
    assert 'disposable-secret' not in json.dumps(receipt)
    assert 'authorization' not in json.dumps(receipt).lower()


def test_filters_verified_only_with_full_row_coverage(tmp_path):
    file=write(tmp_path,b'office,amount\nShanghai,10\nShanghai,20\n')
    r=request(filters={'office':'Shanghai'},filter_columns={'office':'office'})
    result=reports.validate_download(file,r)
    assert result['filters_verified'] and result['filter_verification'][0]['rows_checked']==2
    file.write_bytes(b'office,amount\nShanghai,10\nBeijing,20\n')
    with pytest.raises(reports.ReportError,match='筛选条件'):reports.validate_download(file,r)


def test_empty_rows_do_not_prove_filter_application(tmp_path):
    file=write(tmp_path,b'office,amount\n')
    r=request(filters={'office':'Shanghai'},filter_columns={'office':'office'},allow_empty=True,expected_rows=0)
    result=reports.validate_download(file,r)
    assert result['rows']==0 and not result['filters_verified']


def test_json_echo_checks_context_and_filter_scope(tmp_path):
    body={'meta':{'office':['Shanghai','Beijing'],'unit':'CNY','period':'2026-10'},'rows':[{'office':'Shanghai','amount':12}]}
    file=write(tmp_path,json.dumps(body).encode(),'report.json')
    r=reports.ReportRequest(scope_id='p',request_id='r',url='https://reports.example/export',format='json',
        filters={'office':['Shanghai','Beijing']},context={'unit':'CNY','period':'2026-10'},json_rows_path='rows',
        filter_echo_paths={'office':'meta.office'},context_echo_paths={'unit':'meta.unit','period':'meta.period'},expected_columns=['office','amount'])
    result=reports.validate_download(file,r)
    assert result['filters_verified'] and result['context_verified']
    body['meta']['unit']='USD';file.write_text(json.dumps(body))
    with pytest.raises(reports.ReportError,match='上下文'):reports.validate_download(file,r)


def test_idempotence_checks_actual_file_and_request_identity(tmp_path,monkeypatch):
    calls=[]
    def respond(req):calls.append(req);return httpx.Response(200,content=b'a\n1\n')
    factory=client(monkeypatch,respond);r=request()
    first=reports.acquire_report(r,tmp_path,client_factory=factory)
    assert reports.acquire_report(r,tmp_path,client_factory=factory)['reused']
    assert len(calls)==1
    with pytest.raises(reports.ReportError,match='规则已变化'):
        reports.acquire_report(r.model_copy(update={'filters':{'office':'Shanghai'}}),tmp_path,client_factory=factory)
    (tmp_path/first['file']).write_bytes(b'a\n2\n')
    with pytest.raises(reports.ReportError,match='内容已变化'):reports.acquire_report(r,tmp_path,client_factory=factory)


def test_principal_scope_separates_cache_identity(tmp_path,monkeypatch):
    factory=client(monkeypatch,lambda req:httpx.Response(200,content=b'a\n1\n'))
    first=reports.acquire_report(request(),tmp_path,client_factory=factory)
    second=reports.acquire_report(request().model_copy(update={'scope_id':'other_owner'}),tmp_path,client_factory=factory)
    assert first['file']!=second['file'] and not second['reused']


@pytest.mark.parametrize('body,ctype',[(b'<html>Sign in</html>','application/octet-stream'),(b'a\n1\n','text/html'),(b'{"error":"sign in"}','application/json')])
def test_login_or_error_page_never_promotes_report(tmp_path,monkeypatch,body,ctype):
    factory=client(monkeypatch,lambda req:httpx.Response(200,content=body,headers={'Content-Type':ctype}))
    with pytest.raises(reports.ReportError):reports.acquire_report(request(),tmp_path,client_factory=factory)
    assert not list(tmp_path.glob('*.csv')) and not list(tmp_path.glob('*.receipt.json'))
    assert not [p for p in tmp_path.iterdir() if p.name.startswith('.')]


def test_http_retry_is_bounded_and_only_get(tmp_path,monkeypatch):
    calls=[]
    def respond(req):
        calls.append(req)
        if len(calls)<3:return httpx.Response(503)
        return httpx.Response(200,content=b'a\n1\n')
    factory=client(monkeypatch,respond)
    monkeypatch.setattr(reports.time,'sleep',lambda n:None)
    receipt=reports.acquire_report(request(max_attempts=3),tmp_path,client_factory=factory)
    assert receipt['attempts']==3 and all(c.method=='GET' for c in calls)


def test_redirect_not_followed_and_auth_failure_not_retried(tmp_path,monkeypatch):
    calls=[]
    def respond(req):calls.append(req);return httpx.Response(302,headers={'Location':'http://169.254.169.254/latest'})
    factory=client(monkeypatch,respond)
    with pytest.raises(reports.ReportError):reports.acquire_report(request(),tmp_path,client_factory=factory)
    assert len(calls)==1


def test_interrupted_length_never_promotes_partial_file(tmp_path,monkeypatch):
    factory=client(monkeypatch,lambda req:httpx.Response(200,content=b'a\n1\n',headers={'Content-Length':'100'}))
    with pytest.raises(reports.ReportError,match='下载未完成'):
        reports.acquire_report(request(max_attempts=1),tmp_path,client_factory=factory)
    assert not list(tmp_path.glob('*.csv')) and not list(tmp_path.glob('*.receipt.json'))


def test_stream_byte_cap_enforced_before_validation(tmp_path,monkeypatch):
    factory=client(monkeypatch,lambda req:httpx.Response(200,content=b'amount\n12345\n'))
    with pytest.raises(reports.ReportError,match='体积预算'):
        reports.acquire_report(request(max_bytes=5),tmp_path,client_factory=factory)
    assert not list(tmp_path.glob('*.csv'))


def test_xlsx_signature_sheet_columns_row_count(tmp_path):
    wb=Workbook();ws=wb.active;ws.title='明细';ws.append(['办事处','金额']);ws.append(['上海',10]);ws.append(['北京',20]);wb.create_sheet('口径说明')
    file=tmp_path/'report.xlsx';wb.save(file)
    r=reports.ReportRequest(scope_id='p',request_id='r',url='https://reports.example/export',format='xlsx',
        xlsx_sheet='明细',expected_columns=['办事处','金额'],expected_rows=2)
    result=reports.validate_download(file,r)
    assert result['rows']==2 and result['selection']['sheet']=='明细'
    assert result['selection']['available_sheets']==['明细','口径说明']
    with pytest.raises(reports.ReportError,match='页签'):
        reports.validate_download(file,r.model_copy(update={'xlsx_sheet':'不存在'}))
    with pytest.raises(reports.ReportError,match='行数'):
        reports.validate_download(file,r.model_copy(update={'expected_rows':3}))


def test_no_silent_truncation_or_duplicate_headers(tmp_path):
    file=write(tmp_path,b'a,b\n1,2\n3,4\n')
    with pytest.raises(reports.ReportError,match='行数预算'):reports.validate_download(file,request(max_rows=1))
    file.write_bytes(b'a,a\n1,2\n')
    with pytest.raises(reports.ReportError,match='表头'):reports.validate_download(file,request())


def test_signed_transport_secret_not_persisted_or_hashed(tmp_path,monkeypatch):
    factory=client(monkeypatch,lambda req:httpx.Response(200,content=b'a\n1\n'))
    a=request().model_copy(update={'url':'https://reports.example/export?report=R1&token=first-secret'})
    b=a.model_copy(update={'url':'https://reports.example/export?report=R1&token=second-secret'})
    assert reports.request_digest(a)==reports.request_digest(b)
    result=reports.acquire_report(a,tmp_path,client_factory=factory)
    assert 'first-secret' not in json.dumps(result) and 'token' not in json.dumps(result)
    with pytest.raises(ValidationError):request(filters={'token':'wrong-place'})


def test_ssrf_guard_runs_before_http_even_with_pinned_client(tmp_path,monkeypatch):
    called=[]
    def block(url):raise reports.ReportError('禁止访问云元数据地址','NETWORK_BLOCKED')
    monkeypatch.setattr(reports,'safe_api_url',block)
    def factory(**kw):called.append(kw);raise AssertionError('transport must not open')
    with pytest.raises(reports.ReportError,match='云元数据'):
        reports.acquire_report(request(),tmp_path,client_factory=factory)
    assert not called and not list(tmp_path.glob('*.csv'))


def test_xlsx_download_validation_uses_content_not_temporary_filename(tmp_path):
    wb=Workbook();wb.active.append(['id','amount']);wb.active.append(['A',12])
    content=io.BytesIO();wb.save(content)
    file=write(tmp_path,content.getvalue(),'report.xlsx.part-uuid')
    r=reports.ReportRequest(scope_id='p',request_id='r',url='https://reports.example/export',format='xlsx',expected_columns=['id','amount'])
    assert reports.validate_download(file,r)['rows']==1


def test_json_decimal_echo_is_exact_and_nonfinite_values_rejected(tmp_path):
    file=write(tmp_path,b'{"meta":{"threshold":0.1},"rows":[{"amount":1}]}','r.json')
    r=reports.ReportRequest(scope_id='p',request_id='r',url='https://reports.example/export',format='json',
        filters={'threshold':0.1},filter_echo_paths={'threshold':'meta.threshold'},json_rows_path='rows')
    assert reports.validate_download(file,r)['filters_verified']
    file.write_bytes(b'{"meta":{"threshold":NaN},"rows":[{"amount":1}]}')
    with pytest.raises(reports.ReportError,match='有效 JSON'):reports.validate_download(file,r)


@pytest.mark.parametrize('body',[b'<form action="login">\n<input name="username">\n</form>',b'{\n  "error": "access denied"\n}'])
def test_html_fragment_or_pretty_json_is_not_a_one_column_csv(tmp_path,body):
    file=write(tmp_path,body)
    with pytest.raises(reports.ReportError):reports.validate_download(file,request())
