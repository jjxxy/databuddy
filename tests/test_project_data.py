"""Actual project CSV -> queued workflow -> complete handoff lifecycle acceptance."""
import copy
import hashlib
import json
import time
from decimal import Decimal, localcontext
from pathlib import Path

import pytest
from sqlalchemy import update
from agent_service import project_data as data, workspace
from agent_service.config import DATA
from agent_service.data_workflows import digest
from agent_service.execution import process_once
from agent_service.jobs import get_job, list_jobs, jobs, recover_expired
from agent_service.sources import revision
from agent_service.store import put, get, patch, uid, seal, now, engine


COUNT=215
TOTAL='23220.000000000215'


def project():
    return put('project',{'name':'完整加工交接验收','created_at':now()},'p_'+uid())['id']


def csv_bytes(count=COUNT,amount_offset=0,duplicate=False):
    rows=['id,office,amount,note']
    for i in range(1,count+1):
        ident='ID1' if duplicate and i==count else 'ID'+str(i)
        rows.append(f'{ident},{"East" if i%2==0 else "West"},{i+amount_offset}.000000000001,note-{i}')
    return ('\n'.join(rows)+'\n').encode()


def source(pid,content=None):
    content=content or csv_bytes()
    name=uid()+'.csv';path=DATA/'uploads'/name;path.parent.mkdir(parents=True,exist_ok=True);path.write_bytes(content)
    value=put('source',{'project_id':pid,'name':'orders.csv','kind':'file',
        'secret_config':seal({'stored_name':name}),'content_hash':hashlib.sha256(content).hexdigest(),
        'bytes':len(content),'status':'ready','created_at':now()},'s_'+uid())
    return revision(value)


def replace_source(sid,content):
    name=uid()+'.csv';path=DATA/'uploads'/name;path.write_bytes(content)
    updated=patch('source',sid,name='orders.csv',secret_config=seal({'stored_name':name}),
        content_hash=hashlib.sha256(content).hexdigest(),bytes=len(content),status='ready')
    return revision(updated)


def column(name,kind='text',unit='',aggregation='unknown'):
    return {'name':name,'data_type':kind,'meaning':name,'unit':unit,'aggregation':aggregation}


def row_schema():
    return {'columns':[column('id'),column('office'),column('amount','decimal','CNY','additive'),column('note')],
            'grain':['id'],'meaning':'One order per unique order id'}


def document(sid,exact_total=False):
    contract={'version':1,'min_rows':COUNT,'max_rows':COUNT,'required':['id','office','amount','note'],'unique':[['id']]}
    if exact_total:contract['totals']={'amount':{'expected':TOTAL,'tolerance':'0'}}
    return {'version':1,'plan':{'version':1,'name':'Complete order detail','inputs':{'orders':row_schema()},
        'stages':[{'id':'complete_rows','kind':'select','description':'Preserve every declared detail row',
                   'input':'orders','columns':['id','office','amount','note']}], 'output':'complete_rows'},
        'bindings':{'orders':{'source_id':sid,'column_map':{'c_1':'id','c_2':'office','c_3':'amount','c_4':'note'}}},
        'contract':contract}


def registered(pid,doc,path='workflows/orders.json'):
    workspace.write_file(pid,path,json.dumps(doc,ensure_ascii=False))
    value=data.register(pid,path)
    return value['workflow_id']


def run(pid,wid,**kw):
    value=data.enqueue_workflow(pid,wid,**kw)
    assert process_once(),value
    current=get_job(value['id'])
    assert current['status'] in ('validated','succeeded'),current
    assert current.get('validated') and current['records']>0,current
    return current


def all_rows(pid,rid):
    rows=[];offset=0
    while True:
        page=data.read_result(pid,rid,offset=offset,limit=100)
        rows.extend(page['rows'])
        if page['next_offset'] is None:return rows,page
        offset=page['next_offset']


def amount_sum(rows):
    with localcontext() as ctx:
        ctx.prec=80
        return sum((Decimal(r['amount']) for r in rows),Decimal(0))


def test_real_csv_queue_full_handoff_and_downstream_agent_binding():
    pid=project();s=source(pid);wid=registered(pid,document(s['id'],exact_total=True))
    result=run(pid,wid)
    assert get('data_workflow',wid)['published_run'] is None
    first=data.read_result(pid,result['id'],limit=100)
    assert len(first['rows'])==100 and first['total']==COUNT
    assert first['next_offset']==100 and first['page_is_complete_result'] is False
    rows,page=all_rows(pid,result['id'])
    assert len(rows)==COUNT and page['next_offset'] is None
    assert rows[-1]['id']=='ID215' and amount_sum(rows)==Decimal(TOTAL)
    published=data.publish_result(pid,result['id'],result['commit'])
    assert published['published_run']==result['id']
    # A second workflow consumes the immutable full result, not a preview page.
    downstream={'version':1,'plan':{'name':'Office revenue from complete handoff','inputs':{'detail':row_schema()},
        'stages':[{'kind':'aggregate','id':'office_revenue','description':'Aggregate all rows at office grain','input':'detail',
            'by':['office'],'meaning':'Office order counts and exact revenue','aggregations':[
                {'function':'sum','column':'amount','output':column('revenue','decimal','CNY','additive')},
                {'function':'count','output':column('order_count','integer','1','additive')},
                {'function':'concat','column':'note','output':column('notes'),'order_by':['id'],'delimiter':'|'}]}],
        'output':'office_revenue'},'bindings':{'detail':{'result_id':result['id']}}}
    downstream_id=registered(pid,downstream,path='workflows/office-revenue.json')
    next_result=run(pid,downstream_id);totals,_=all_rows(pid,next_result['id'])
    assert next_result['records']==2 and sum(r['order_count'] for r in totals)==COUNT
    assert sum((Decimal(r['revenue']) for r in totals),Decimal(0))==Decimal(TOTAL)
    assert sum(len(r['notes'].split('|')) for r in totals)==COUNT
    assert any('note-215' in r['notes'] for r in totals)
    folder,meta=data.result_folder(pid,next_result['id'])
    assert meta['complete'] and meta['records']==2
    receipts=json.loads((folder/'stage-receipts.json').read_text())
    assert receipts[-1]['inputs'][0]['rows']==COUNT


def test_frozen_draft_changes_neither_queued_input_snapshot_nor_formal_plan():
    pid=project();s=source(pid);doc=document(s['id']);wid=registered(pid,doc)
    queued=data.enqueue_workflow(pid,wid)
    old_source_revision=get('source',s['id'])['revision']
    replace_source(s['id'],csv_bytes(count=COUNT+1))
    # A source replacement allocates a new immutable file; queued run keeps old bytes.
    assert process_once();result=get_job(queued['id'])
    assert result['status']=='validated',result
    assert result['records']==COUNT and result['source_revisions'][s['id']]==old_source_revision
    data.publish_result(pid,result['id'],result['commit'])
    # Replace the draft with a filter. Formal rerun must still execute the old plan.
    draft=copy.deepcopy(doc);draft['plan']['name']='New unadopted office filter'
    draft['plan']['stages']=[{'kind':'filter','id':'only_east','description':'Candidate-only office filter','input':'orders',
                              'conditions':[{'column':'office','operator':'eq','value':'East'}]}]
    draft['plan']['output']='only_east';draft['contract']={'version':1,'min_rows':1}
    workspace.write_file(pid,'workflows/orders.json',json.dumps(draft))
    data.register(pid,'workflows/orders.json')
    w=get('data_workflow',wid)
    assert w['active_sha256']==result['commit'] and w['draft_sha256']!=result['commit']
    # Restore a valid 215-row source; no new candidate draft is adopted implicitly.
    replace_source(s['id'],csv_bytes(amount_offset=1))
    again=run(pid,wid,use_published=True)
    assert again['commit']==result['commit'] and again['records']==COUNT
    assert get('data_workflow',wid)['published_run']==again['id']
    rows,_=all_rows(pid,again['id'])
    assert amount_sum(rows)==Decimal(TOTAL)+COUNT


def test_candidate_replaced_by_new_draft_cannot_be_adopted():
    pid=project();s=source(pid);doc=document(s['id']);wid=registered(pid,doc)
    result=run(pid,wid)
    changed=copy.deepcopy(doc);changed['plan']['name']='Changed draft'
    workspace.write_file(pid,'workflows/orders.json',json.dumps(changed))
    with pytest.raises(ValueError,match='替代'):
        data.publish_result(pid,result['id'],result['commit'])
    assert get('data_workflow',wid)['published_run'] is None
    with pytest.raises(ValueError,match='版本'):
        data.publish_result(pid,result['id'],'0'*64)


def test_watch_sources_refreshes_complete_result_and_failure_retains_formal_data():
    pid=project();s=source(pid);wid=registered(pid,document(s['id']))
    initial=run(pid,wid);data.publish_result(pid,initial['id'],initial['commit'])
    data.schedule(pid,wid,watch_sources=True)
    replace_source(s['id'],csv_bytes(amount_offset=2))
    data.queue_due();pending=list_jobs(target=wid)[0]
    assert pending['status']=='queued'
    data.queue_due();assert list_jobs(target=wid)[0]['id']==pending['id']
    assert process_once();refreshed=get_job(pending['id'])
    assert refreshed['status']=='succeeded',refreshed
    w=get('data_workflow',wid)
    assert w['published_run']==refreshed['id'] and w['source_revisions'][s['id']]==get('source',s['id'])['revision']
    rows,_=all_rows(pid,refreshed['id'])
    assert len(rows)==COUNT and amount_sum(rows)==Decimal(TOTAL)+COUNT*2
    replace_source(s['id'],csv_bytes(duplicate=True))
    data.queue_due();failed_id=list_jobs(target=wid)[0]['id'];assert process_once()
    failed=get_job(failed_id);assert failed['status']=='failed',failed
    assert failed['error_code']=='GRAIN_DUPLICATE'
    assert failed['workflow_failure']['stage_id']=='input:orders'
    assert failed['workflow_failure']['code']=='GRAIN_DUPLICATE'
    assert get('data_workflow',wid)['published_run']==refreshed['id']
    previous,_=all_rows(pid,refreshed['id']);assert amount_sum(previous)==Decimal(TOTAL)+COUNT*2
    job_count=len(list_jobs(target=wid));data.queue_due();data.queue_due()
    assert len(list_jobs(target=wid))==job_count


def test_contract_failure_never_exposes_a_verified_result_or_replaces_formal():
    pid=project();s=source(pid);doc=document(s['id'],exact_total=True);wid=registered(pid,doc)
    initial=run(pid,wid);data.publish_result(pid,initial['id'],initial['commit'])
    replace_source(s['id'],csv_bytes(amount_offset=1))
    queued=data.enqueue_workflow(pid,wid,use_published=True);assert process_once()
    failed=get_job(queued['id'])
    assert failed['status']=='failed' and failed['error_code']=='DATA_CONTRACT_FAILED',failed
    assert get('data_workflow',wid)['published_run']==initial['id']
    with pytest.raises(ValueError,match='未验证'):
        data.result_folder(pid,failed['id'])


def test_failed_step_preserves_specific_error_and_completed_receipt_without_publishing():
    pid=project();secret='private-order-note-should-not-appear-in-diagnostics'
    s=source(pid,csv_bytes().replace(b'note-215',secret.encode()))
    doc=document(s['id']);wid=registered(pid,doc)
    initial=run(pid,wid);data.publish_result(pid,initial['id'],initial['commit'])
    doc['plan']['stages'].append({'kind':'select','id':'broken_report','description':'Missing business field needs repair',
                                  'input':'complete_rows','columns':['id','office','missing_amount','note']})
    doc['plan']['output']='broken_report'
    workspace.write_file(pid,'workflows/orders.json',json.dumps(doc));data.register(pid,'workflows/orders.json')
    queued=data.enqueue_workflow(pid,wid);assert process_once()
    failed=get_job(queued['id'])
    assert failed['status']=='failed' and failed['error_code']=='MISSING_COLUMN',failed
    assert failed['workflow_failure']=={'stage_id':'broken_report','code':'MISSING_COLUMN',
                                       'message':'Operation references missing or duplicate columns'}
    assert failed['error']==failed['workflow_failure']['message']
    assert secret not in json.dumps(failed['workflow_failure'])+failed['error']
    receipts=json.loads((DATA/failed['artifact_dir']/'stage-receipts.json').read_text())
    assert [(r['stage_id'],r['output_rows']) for r in receipts]==[('complete_rows',COUNT)]
    assert failed['workflow_steps']==1 and not failed.get('validated')
    assert get('data_workflow',wid)['published_run']==initial['id']
    assert len(all_rows(pid,initial['id'])[0])==COUNT
    with pytest.raises(ValueError,match='未验证'):data.result_folder(pid,failed['id'])


def test_source_and_result_bindings_are_project_scoped():
    owner=project();other=project();s=source(owner)
    with pytest.raises(ValueError,match='当前项目'):
        registered(other,document(s['id']))
    wid=registered(owner,document(s['id']));result=run(owner,wid)
    with pytest.raises(ValueError,match='当前项目'):
        data.read_result(other,result['id'])
    foreign_doc=document(s['id']);foreign_doc['bindings']={'orders':{'result_id':result['id']}}
    with pytest.raises(ValueError,match='当前项目'):
        registered(other,foreign_doc)


def test_queue_active_target_and_scheduled_key_deduplication_are_idempotent():
    pid=project();s=source(pid);wid=registered(pid,document(s['id']))
    first=data.enqueue_workflow(pid,wid,scheduled_key='integration-'+uid())
    second=data.enqueue_workflow(pid,wid)
    assert second['id']==first['id']
    assert process_once();complete=get_job(first['id']);assert complete['status']=='validated',complete
    same=data.enqueue_workflow(pid,wid,scheduled_key=first['scheduled_key'])
    assert same['id']==first['id'] and len(list_jobs(target=wid))==1
    assert data.register(pid,'workflows/orders.json')['workflow_id']==wid


def test_handoff_artifact_tampering_blocks_publish_and_read():
    pid=project();s=source(pid);wid=registered(pid,document(s['id']));result=run(pid,wid)
    folder,meta=data.result_folder(pid,result['id'])
    rows_file=folder/'handoff'/'rows.jsonl';rows_file.write_text(rows_file.read_text()+'{}\n')
    with pytest.raises(ValueError,match='checksum|校验'):
        data.publish_result(pid,result['id'],result['commit'])
    with pytest.raises(ValueError,match='checksum|校验'):
        data.read_result(pid,result['id'])
    assert get('data_workflow',wid)['published_run'] is None


def test_worker_lease_loss_before_validation_cannot_publish(monkeypatch):
    pid=project();s=source(pid);wid=registered(pid,document(s['id']));initial=run(pid,wid)
    data.publish_result(pid,initial['id'],initial['commit'])
    queued=data.enqueue_workflow(pid,wid,use_published=True)
    original=data.execute
    def expire_lease(*args,**kwargs):
        result=original(*args,**kwargs)
        with engine.begin() as c:
            c.execute(update(jobs).where(jobs.c.id==queued['id']).values(lease_until=time.time()-1))
        return result
    monkeypatch.setattr(data,'execute',expire_lease)
    assert process_once()
    assert get('data_workflow',wid)['published_run']==initial['id']
    recover_expired()
    assert get_job(queued['id'])['status']=='failed'


def test_auto_publication_rechecks_lease_after_artifact_and_plan_checks(monkeypatch):
    pid=project();s=source(pid);doc=document(s['id']);wid=registered(pid,doc)
    initial=run(pid,wid);data.publish_result(pid,initial['id'],initial['commit'])
    changed=copy.deepcopy(doc);changed['plan']['name']='Candidate with a distinct frozen code version'
    workspace.write_file(pid,'workflows/orders.json',json.dumps(changed));data.register(pid,'workflows/orders.json')
    candidate=run(pid,wid)
    clock=[time.time()]
    with engine.begin() as c:
        c.execute(update(jobs).where(jobs.c.id==candidate['id']).values(status='running',lease_token='final-fence-test',lease_until=clock[0]+10))
    original_load=data.load_document
    def slow_plan_check(*args,**kwargs):
        result=original_load(*args,**kwargs)
        clock[0]+=30
        return result
    monkeypatch.setattr(data,'load_document',slow_plan_check)
    monkeypatch.setattr(data.time,'time',lambda:clock[0])
    try:
        with pytest.raises(ValueError,match='租约失效'):
            data.publish_result(pid,candidate['id'],candidate['commit'],lease_token='final-fence-test')
        assert get('data_workflow',wid)['published_run']==initial['id']
        assert get_job(candidate['id'])['status']=='running'
    finally:
        with engine.begin() as c:
            c.execute(update(jobs).where(jobs.c.id==candidate['id']).values(status='validated',lease_token=candidate['lease_token'],lease_until=candidate['lease_until']))


def test_column_rename_cannot_collide_with_an_unchanged_input_column():
    pid=project();s=source(pid)
    with pytest.raises(ValueError,match='映射|重复|collision'):
        data.read_input(pid,{'source_id':s['id'],'column_map':{'c_1':'c_2'}})


def test_upstream_sql_semantic_review_flag_survives_full_result_handoff():
    pid=project();s=source(pid);doc=document(s['id'])
    doc['plan']['stages']=[{'kind':'sql','id':'sql_rows','description':'Explicit SQL outputs need semantic review',
        'inputs':['orders'],'sql':'SELECT id, office, amount, note FROM orders ORDER BY id','output_schema':row_schema()}]
    doc['plan']['output']='sql_rows'
    wid=registered(pid,doc);upstream=run(pid,wid)
    _,meta=data.result_folder(pid,upstream['id']);assert meta['requires_semantic_review'] is True
    downstream=document(s['id']);downstream['bindings']={'orders':{'result_id':upstream['id']}}
    downstream_id=registered(pid,downstream,path='workflows/downstream-review.json')
    result=run(pid,downstream_id);folder,meta=data.result_folder(pid,result['id'])
    assert meta['records']==COUNT and meta['requires_semantic_review'] is True
    page=data.read_result(pid,result['id'],limit=100)
    assert page['context']['requires_semantic_review'] is True
    assert page['context']['output_semantics']['grain']==['id']
    assert next(c for c in page['context']['output_semantics']['columns'] if c['name']=='amount')['unit']=='CNY'


@pytest.mark.parametrize('changed', ['unit','period','aggregation','meaning','value_format','grain'])
def test_downstream_binding_cannot_silently_override_upstream_semantics(changed):
    pid=project();s=source(pid);wid=registered(pid,document(s['id']));upstream=run(pid,wid)
    downstream=document(s['id']);downstream['bindings']={'orders':{'result_id':upstream['id']}}
    schema=downstream['plan']['inputs']['orders']
    amount=next(c for c in schema['columns'] if c['name']=='amount')
    if changed=='unit':amount['unit']='USD'
    elif changed=='period':amount['period']='2026-01'
    elif changed=='aggregation':amount['aggregation']='non_additive'
    elif changed=='meaning':amount['meaning']='Profit after costs'
    elif changed=='value_format':amount['value_format']='percent_fraction'
    else:schema['grain']=['note']
    down_id=registered(pid,downstream,path='workflows/downstream-conflict.json')
    queued=data.enqueue_workflow(pid,down_id);assert process_once()
    failed=get_job(queued['id'])
    assert failed['status']=='failed',failed
    assert not failed.get('validated') and get('data_workflow',down_id)['published_run'] is None
    with pytest.raises(ValueError):data.result_folder(pid,failed['id'])


def test_complete_result_excel_export_can_be_downloaded_more_than_once():
    from openpyxl import load_workbook
    pid=project();s=source(pid);wid=registered(pid,document(s['id']));result=run(pid,wid)
    first=data.export_result(pid,result['id'],precision_policy='text_exact')
    content=first.read_bytes();second=data.export_result(pid,result['id'],precision_policy='text_exact')
    assert first.read_bytes()==content and second.is_file()
    book=load_workbook(second,read_only=True,data_only=False)
    try:
        rows=list(book['结果'].iter_rows(values_only=True))
        assert len(rows)==COUNT+1
        assert rows[0]==('id','office','amount','note')
        assert rows[-1][0]=='ID215'
        assert sum((Decimal(str(r[2])) for r in rows[1:]),Decimal(0))==Decimal(TOTAL)
    finally:book.close()


@pytest.mark.parametrize('review_state', ['unreferenced','proposed','approved','changed_after_approval'])
def test_reclassification_review_is_bound_to_exact_mapping_fields_and_payload(review_state):
    from agent_service import semantic_service
    pid=project();s=source(pid);doc=document(s['id'])
    mapping={'id':'office-region-map','version':1,'source_field_id':'office','target_field_id':'region',
             'entries':[{'source':'East','target':'East-region'},{'source':'West','target':'West-region'}],
             'unmapped_policy':'error'}
    proposal=semantic_service.propose_mapping(pid,mapping)
    stage={'kind':'reclassify','id':'regions','description':'Office to region is a business rule','input':'orders',
           'column':'office','output':column('region'),'mapping':{'East':'East-region','West':'West-region'}}
    doc['plan']['stages']=[stage];doc['plan']['output']='regions'
    if review_state!='unreferenced':doc['reviewed_mapping_ids']=[proposal['id']]
    if review_state in ('approved','changed_after_approval'):
        semantic_service.review_mapping_record(pid,proposal['id'],proposal['content_sha256'],actor='authenticated-test-owner')
    if review_state=='changed_after_approval':stage['mapping']['East']='Different-region'
    wid=registered(pid,doc);queued=data.enqueue_workflow(pid,wid);assert process_once()
    result=get_job(queued['id'])
    if review_state in ('proposed','changed_after_approval'):
        assert result['status']=='failed' and not result.get('validated'),result
        assert get('data_workflow',wid)['published_run'] is None
    else:
        assert result['status']=='validated',result
        _,meta=data.result_folder(pid,result['id'])
        assert meta['requires_semantic_review'] is (review_state=='unreferenced')
        rows,_=all_rows(pid,result['id'])
        assert len(rows)==COUNT and {r['region'] for r in rows}=={'East-region','West-region'}


def test_excel_percentage_requires_fraction_declaration_before_workflow_execution():
    from openpyxl import Workbook
    pid=project();name=uid()+'.xlsx';path=DATA/'uploads'/name;path.parent.mkdir(parents=True,exist_ok=True)
    book=Workbook();sheet=book.active;sheet.title='Percent'
    sheet.append(['id','share']);sheet.append(['A',0.15]);sheet['B2'].number_format='0.0%'
    book.save(path);book.close();content=path.read_bytes()
    s=put('source',{'project_id':pid,'name':'shares.xlsx','kind':'file','secret_config':seal({'stored_name':name}),
                   'content_hash':hashlib.sha256(content).hexdigest(),'bytes':len(content),'status':'ready','created_at':now()},'s_'+uid())
    s=revision(s)
    percent=column('share','decimal','%','non_additive')
    doc={'version':1,'plan':{'name':'Percentage source context','inputs':{'table':{'columns':[column('id'),percent],
                  'grain':['id'],'meaning':'One share observation'}},
            'stages':[{'kind':'select','id':'copied','description':'Keep exact stored percentage value','input':'table','columns':['id','share']}],
            'output':'copied'},'bindings':{'table':{'source_id':s['id'],'column_map':{'c_1':'id','c_2':'share'},
                'options':{'sheet':'Percent','header_rows':[1],'data_start_row':2,'data_end_row':2}}}}
    wid=registered(pid,doc);queued=data.enqueue_workflow(pid,wid);assert process_once()
    rejected=get_job(queued['id']);assert rejected['status']=='failed',rejected
    percent.update(value_format='percent_fraction',number_format='0.0%',label='Revenue share')
    workspace.write_file(pid,'workflows/orders.json',json.dumps(doc));data.register(pid,'workflows/orders.json')
    result=run(pid,wid);page=data.read_result(pid,result['id'])
    assert page['rows']==[{'id':'A','share':'0.15'}]
    field=next(c for c in page['schema'] if c['name']=='share')
    assert field['value_format']=='percent_fraction' and field['unit']=='%' and field['number_format']=='0.0%'


def test_cleanup_keeps_formal_candidate_draft_shared_report_and_transitive_inputs(monkeypatch,request):
    from agent_service import maintenance
    from agent_service.jobs import save,request_cancel
    from agent_service.store import delete
    pid=project();owned_runs=set();folders={};cleanup_items=[];cleanup_active=[]
    def cleanup():
        for ident in cleanup_active:request_cancel(ident)
        for kind,ident in cleanup_items:delete(kind,ident)
    request.addfinalizer(cleanup)
    def saved(kind,value,ident=None):
        result=put(kind,value,ident);cleanup_items.append((kind,result['id']));return result
    def historical(*upstream,status='succeeded',target=None,recent=False):
        doc=document('source-not-required-for-retention')
        if upstream:
            doc['bindings']={str(i):{'result_id':ident} for i,ident in enumerate(upstream)}
        rid=uid();folder=DATA/'runs'/rid/'attempt-1';folder.mkdir(parents=True)
        (folder/'provenance.txt').write_text('Immutable input/result evidence')
        save({'id':rid,'kind':'data_workflow','target':target or 'historical_'+uid(),'status':status,
              'project_id':pid,'created_at':now() if recent else '2000-01-01T00:00:00+00:00',
              'workflow_document':doc,'artifact_dir':str(folder.relative_to(DATA))})
        owned_runs.add(rid);folders[rid]=folder
        return rid

    first=historical();second=historical(first)
    wid='wf_'+uid();formal=historical(second,target=wid)
    candidate_input=historical();candidate=historical(candidate_input,status='validated',target=wid)
    draft_input=historical(second)
    draft=document('source-not-required-for-retention');draft['bindings']={'orders':{'result_id':draft_input}}
    workspace.write_file(pid,'workflows/cleanup.json',json.dumps(draft))
    saved('data_workflow',{'project_id':pid,'path':'workflows/cleanup.json','published_run':formal,
                        'candidate_run':candidate,'active_document':get_job(formal)['workflow_document']},wid)
    shared_input=historical();shared=historical(shared_input)
    saved('data_share',{'project_id':pid,'result_id':shared,'state':'active','expires_at':'2099-01-01T00:00:00+00:00'})
    report_input=historical();reported=historical(report_input)
    saved('business_report',{'project_id':pid,'run_id':reported})
    active_input=historical();active=historical(active_input,status='queued')
    cleanup_active.append(active)
    recent_input=historical();recent=historical(recent_input,status='failed',recent=True)
    expired=historical();revoked=historical();unreferenced=historical()
    saved('data_share',{'project_id':pid,'result_id':expired,'state':'active','expires_at':'2000-01-01T00:00:00+00:00'})
    saved('data_share',{'project_id':pid,'result_id':revoked,'state':'revoked','expires_at':'2099-01-01T00:00:00+00:00'})

    # Limit deletion to this fixture; other acceptance cases share the test DB.
    original=maintenance.list_jobs
    def scoped(limit=100,offset=0,target=None,status=None):
        values=[r for r in original(100000,target=target,status=status) if r['id'] in owned_runs]
        return values[offset:offset+limit]
    monkeypatch.setattr(maintenance,'list_jobs',scoped)
    expected={expired,revoked,unreferenced}
    preview=maintenance.retention(days=1,keep_success=1,apply=False)
    assert {r['run_id'] for r in preview['candidates']}==expected
    assert all(folder.exists() for folder in folders.values())
    removed=maintenance.retention(days=1,keep_success=1,apply=True)
    assert {r['run_id'] for r in removed['candidates']}==expected
    for rid,folder in folders.items():
        assert folder.exists() is (rid not in expected),rid
        if rid in expected:assert get_job(rid)['artifact_removed'] is True
    assert {first,second,formal,candidate_input,candidate,draft_input,shared_input,shared,
            report_input,reported,active_input,active,recent_input,recent}==owned_runs-expected


def test_cleanup_refuses_to_delete_when_saved_workflow_dependencies_are_unreadable(monkeypatch):
    from agent_service import maintenance
    from agent_service.jobs import save
    from agent_service.store import delete
    pid=project();rid=uid();folder=DATA/'runs'/rid/'attempt-1';folder.mkdir(parents=True)
    (folder/'evidence.txt').write_text('Must remain until draft dependencies can be checked')
    value=save({'id':rid,'kind':'data_workflow','target':'old_'+uid(),'status':'failed','project_id':pid,
                'created_at':'2000-01-01T00:00:00+00:00','workflow_document':document('unused-source')})
    path='workflows/temporarily-edited.json';workspace.write_file(pid,path,'{"bindings":')
    workflow=put('data_workflow',{'project_id':pid,'path':path},'wf_'+uid())
    monkeypatch.setattr(maintenance,'list_jobs',lambda limit=100,offset=0,target=None,status=None:
                        [] if target or status else [value])
    try:
        with pytest.raises(ValueError,match='未执行历史清理'):
            maintenance.retention(days=0,keep_success=1,apply=True)
        assert folder.is_dir() and not get_job(rid).get('artifact_removed')
    finally:
        delete('data_workflow',workflow['id'])
