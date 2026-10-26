import hashlib
import json
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path
import zipfile
import pytest
from openpyxl import load_workbook
from agent_service.result_delivery import (
    DeliveryError, write_handoff, verify_handoff, read_handoff, export_xlsx,
)

SCHEMA=[{'name':'office','label':'办事处','type':'string','grain':['office','month']},
        {'name':'amount','label':'净收入','type':'decimal','unit':'CNY','period':'2026-10','aggregation':'additive'}]


def test_complete_handoff_never_uses_preview_and_preserves_meaning(tmp_path):
    rows=[{'office':'O'+str(i),'amount':Decimal(str(i)+'.000000000001')} for i in range(603)]
    path=tmp_path/'handoff'
    context={'grain':['office','month'],'period':'2026-10','metric':'amount minus refund',
             'source_filters':{'status':'paid'},'source_sha256':'a'*64}
    provenance=[{'row_index':i+1,'sheet':'明细','cell':'C'+str(i+2)} for i in range(603)]
    receipt=write_handoff(path,iter(rows),SCHEMA,context,provenance,[{'stage':'aggregate','input_rows':999,'output_rows':603}])
    assert receipt['row_count']==603 and receipt['provenance_count']==603 and receipt['complete']
    assert verify_handoff(path,receipt['manifest_sha256'])['verified']
    page=read_handoff(path,offset=600,limit=100,expected_manifest_sha256=receipt['manifest_sha256'])
    assert page['rows']==[{'office':'O600','amount':'600.000000000001'},
                          {'office':'O601','amount':'601.000000000001'},
                          {'office':'O602','amount':'602.000000000001'}]
    assert page['schema']==SCHEMA and page['context']==context
    assert page['next_offset'] is None and not page['page_is_complete_result']
    all_rows=read_handoff(path,limit=1000,native_types=True)
    assert all_rows['rows']==rows and all_rows['page_is_complete_result']


def test_exact_nested_types_markers_and_large_integer_roundtrip(tmp_path):
    schema=[{'name':'value'}]
    value={'decimal':Decimal('9999999999999999999999999.123456789012'),
           'integer':98765432109876543210987654321,'date':date(2026,10,21),
           'datetime':datetime(2026,10,21,9,30,tzinfo=timezone.utc),
           'user_marker':{'$atlas':'decimal','value':'user text, not an actual decimal'}}
    path=tmp_path/'handoff';write_handoff(path,[{'value':value}],schema)
    assert read_handoff(path,native_types=True)['rows'][0]['value']==value
    safe=read_handoff(path)['rows'][0]['value']
    assert safe['integer']=='98765432109876543210987654321'
    assert safe['decimal']=='9999999999999999999999999.123456789012'
    assert safe['user_marker']==value['user_marker']


def test_tampered_rows_or_manifest_are_rejected(tmp_path):
    path=tmp_path/'handoff';r=write_handoff(path,[{'office':'A','amount':Decimal('12.5')}],SCHEMA)
    original=(path/'rows.jsonl').read_bytes()
    (path/'rows.jsonl').write_bytes(original.replace(b'12.5',b'13.5'))
    with pytest.raises(DeliveryError,match='checksum'):verify_handoff(path)
    (path/'rows.jsonl').write_bytes(original)
    manifest=json.loads((path/'manifest.json').read_text());manifest['row_count']=2
    (path/'manifest.json').write_text(json.dumps(manifest))
    with pytest.raises(DeliveryError,match='manifest checksum'):verify_handoff(path,r['manifest_sha256'])
    with pytest.raises(DeliveryError,match='count mismatch'):verify_handoff(path)


def test_unknown_schema_missing_keys_and_nonfinite_values_fail_atomically(tmp_path):
    for value in [{'office':'A'}, {'office':'A','amount':Decimal('NaN')}, {'office':'A','amount':float('inf')}]:
        path=tmp_path/str(len(list(tmp_path.iterdir())))
        with pytest.raises(DeliveryError):write_handoff(path,[value],SCHEMA)
        assert not path.exists()
    path=tmp_path/'empty';r=write_handoff(path,[],SCHEMA)
    assert r['row_count']==0 and read_handoff(path)['rows']==[]
    with pytest.raises(DeliveryError,match='immutable'):write_handoff(path,[],SCHEMA)


def test_handoff_symlink_and_path_mutation_are_rejected(tmp_path):
    path=tmp_path/'handoff';write_handoff(path,[],SCHEMA)
    real=path/'context.json';copy=tmp_path/'outside.json';copy.write_bytes(real.read_bytes());real.unlink()
    try:real.symlink_to(copy)
    except OSError:
        # Non-admin Windows accounts may not create symlinks. A hard link still
        # tests rejection of an artifact aliased outside its immutable bundle.
        import os
        if os.name!='nt':raise
        os.link(copy,real)
    with pytest.raises(DeliveryError):verify_handoff(path)
    for offset,limit in [(True,1),(0,True),(-1,5),(0,1001)]:
        with pytest.raises(DeliveryError):read_handoff(path,offset,limit)


def test_excel_all_rows_typed_values_formulas_are_text_and_controls_pass(tmp_path):
    schema=[{'name':'id','label':'标识','type':'string'}, {'name':'amount','label':'金额','type':'decimal','unit':'CNY','grain':['id']}]
    texts=['=HYPERLINK("https://evil.test","open")','+SUM(A1:A2)','-10','@cmd','plain']
    rows=[{'id':texts[i%len(texts)],'amount':Decimal('-12.25')} for i in range(603)]
    file=tmp_path/'result.xlsx'
    receipt=export_xlsx(iter(rows),schema,file,layout={'start_cell':'B3','columns':['amount','id'],'sheet_name':'报表'},
                         context={'period':'2026-10'},provenance=[{'source':'input.xlsx','sheet':'明细','cell':'C2'}],
                         controls=[{'kind':'row_count','expected':603},{'kind':'sum','column':'amount','expected':'-7386.75'}])
    assert receipt['row_count']==603 and receipt['complete'] and not receipt['formula_recalculation']
    book=load_workbook(file,data_only=False)
    sheet=book['报表'];assert sheet.max_row==606
    assert sheet['B4'].value==-12.25 and sheet['B4'].data_type=='n'
    for i,text in enumerate(texts,4):
        assert sheet['C'+str(i)].value==text and sheet['C'+str(i)].data_type=='s'
    assert sheet['C606'].value==texts[602%5]
    assert book['语义说明']['D3'].value=='CNY'
    assert any('input.xlsx' in str(c.value) for row in book['来源与校验'] for c in row)
    book.close()


def test_excel_decimal15digit_limit_is_explicit(tmp_path):
    high=Decimal('1234567890123456.123456789012')
    rows=[{'office':'A','amount':high}]
    fail=tmp_path/'fail.xlsx'
    with pytest.raises(DeliveryError,match='precision'):export_xlsx(rows,SCHEMA,fail)
    assert not fail.exists()
    okay=tmp_path/'exact.xlsx'
    r=export_xlsx(rows,SCHEMA,okay,precision_policy='text_exact')
    book=load_workbook(okay,data_only=False);cell=book['结果']['B2']
    assert cell.value==str(high) and cell.data_type=='s'
    assert r['precision_notes'][0]['reason']=='exact_numeric_text' and r['precision_notes'][0]['exact_value']==str(high)
    book.close()


def test_excel_complete_layout_and_large_cell_fail_without_truncation(tmp_path):
    with pytest.raises(DeliveryError,match='every declared column'):export_xlsx([],SCHEMA,tmp_path/'subset.xlsx',layout={'columns':['office']})
    with pytest.raises(DeliveryError,match='cell text limit'):export_xlsx([{'office':'a'*32768,'amount':Decimal(1)}],SCHEMA,tmp_path/'long.xlsx')
    assert not (tmp_path/'long.xlsx').exists()
    with pytest.raises(DeliveryError,match='worksheet limits'):export_xlsx([],SCHEMA,tmp_path/'columns.xlsx',layout={'start_cell':'XFD1'})
    with pytest.raises(DeliveryError,match='control failed'):export_xlsx([{'office':'A','amount':Decimal(1)}],SCHEMA,tmp_path/'wrong.xlsx',controls=[{'kind':'sum','column':'amount','expected':'2'}])
    assert not (tmp_path/'wrong.xlsx').exists()


def test_excel_numerical_range_and_negative_number_string_type(tmp_path):
    schema=[{'name':'x','type':'decimal'}]
    with pytest.raises(DeliveryError):export_xlsx([{'x':Decimal('1e-400')}],schema,tmp_path/'underflow.xlsx')
    r=export_xlsx([{'x':'-123.45'}],schema,tmp_path/'numeric.xlsx')
    book=load_workbook(tmp_path/'numeric.xlsx');assert book['结果']['A2'].value==-123.45;assert book['结果']['A2'].data_type=='n';book.close()
    with pytest.raises(DeliveryError):export_xlsx([],schema,tmp_path/'numeric.xlsx')


def test_excel_saved_numeric_xml_roundtrip_matches_values(tmp_path):
    schema=[{'name':'n','type':'decimal'}]
    numbers=['1.00000000000001','0.000123456789012345','12345678901234500000','1e307','-0.000000000000001']
    receipt=export_xlsx([{'n':Decimal(n)} for n in numbers],schema,tmp_path/'numbers.xlsx')
    assert receipt['numeric_xml_verified']
    with zipfile.ZipFile(tmp_path/'numbers.xlsx') as archive:
        from xml.etree import ElementTree as ET
        root=ET.fromstring(archive.read('xl/worksheets/sheet1.xml'))
        values=[Decimal(c.find('{*}v').text) for c in root.findall('.//{*}c') if c.get('t')=='n']
        assert values==[Decimal(n) for n in numbers]


def test_excel_serialization_loss_is_detected_before_publication(tmp_path,monkeypatch):
    from agent_service import result_delivery as delivery
    original=delivery._verify_excel_numeric_xml
    def tamper(path,sha,count):
        with zipfile.ZipFile(path) as source:
            items={n:source.read(n) for n in source.namelist()}
        items['xl/worksheets/sheet1.xml']=items['xl/worksheets/sheet1.xml'].replace(b'<v>12.25</v>',b'<v>12.26</v>')
        with zipfile.ZipFile(path,'w') as target:
            for name,body in items.items():target.writestr(name,body)
        return original(path,sha,count)
    monkeypatch.setattr(delivery,'_verify_excel_numeric_xml',tamper)
    path=tmp_path/'tamper.xlsx'
    with pytest.raises(DeliveryError,match='serialization changed'):export_xlsx([{'office':'A','amount':Decimal('12.25')}],SCHEMA,path)
    assert not path.exists()


def test_missing_numeric_value_is_not_reconciled_as_zero(tmp_path):
    with pytest.raises(DeliveryError,match='missing numeric'):
        export_xlsx([{'office':'A','amount':None}],SCHEMA,tmp_path/'null.xlsx',
                    controls=[{'kind':'sum','column':'amount','expected':'0'}])
    assert not (tmp_path/'null.xlsx').exists()


def test_interrupted_handoff_has_no_success_manifest_or_partial_destination(tmp_path):
    def fail_after_one_row():
        yield {'office':'A','amount':Decimal(1)}
        raise RuntimeError('upstream stage failed')
    path=tmp_path/'result'
    with pytest.raises(RuntimeError,match='upstream'):write_handoff(path,fail_after_one_row(),SCHEMA)
    assert not path.exists() and not list(tmp_path.iterdir())


def test_complete_consumer_read_crosses_pages_with_one_integrity_pass(tmp_path,monkeypatch):
    from agent_service import result_delivery as delivery
    path=tmp_path/'complete';rows=[{'office':'O'+str(i),'amount':Decimal(i)+Decimal('0.000000000001')} for i in range(1203)]
    receipt=write_handoff(path,rows,SCHEMA,context={'grain':['office']},provenance=[{'source':'s'}],stages=[{'step':'complete'}])
    original=delivery.verify_handoff;calls=[]
    def verify_once(*a,**kw):
        calls.append(True);return original(*a,**kw)
    monkeypatch.setattr(delivery,'verify_handoff',verify_once)
    value=delivery.read_all_handoff(path,receipt['manifest_sha256'],native_types=True)
    assert len(calls)==1 and value['rows']==rows and value['total']==1203
    assert value['provenance']==[{'source':'s'}] and value['stages']==[{'step':'complete'}]
    assert value['page_is_complete_result']
    with pytest.raises(DeliveryError,match='row budget'):
        delivery.read_all_handoff(path,receipt['manifest_sha256'],max_rows=1202)


def test_text_exact_handles_actual_serializer_precision_without_rounding(tmp_path):
    schema=[{'name':'n','data_type':'decimal'}]
    numbers=[Decimal(i)+Decimal('0.000000000001') for i in range(1,216)]
    file=tmp_path/'actual-serializer.xlsx'
    result=export_xlsx([{'n':n} for n in numbers],schema,file,precision_policy='text_exact')
    assert result['numeric_xml_verified']
    book=load_workbook(file,data_only=False)
    assert len(list(book['结果'].iter_rows(min_row=2)))==215
    for i,number in enumerate(numbers,2):
        value=book['结果'].cell(i,1).value
        assert Decimal(str(value))==number
    assert book['结果']['A65'].value=='64.000000000001'
    assert book['结果']['A65'].data_type=='s'
    book.close()
    with pytest.raises(DeliveryError,match='precision'):
        export_xlsx([{'n':Decimal('64.000000000001')}],schema,tmp_path/'strict-serializer.xlsx')


def test_blank_optional_label_and_number_format_keep_useful_defaults(tmp_path):
    file=tmp_path/'blank-optional.xlsx'
    export_xlsx([{'office':'A','amount':Decimal('12.25')}],
                [{'name':'office','label':''},{'name':'amount','label':'','data_type':'decimal','number_format':''}],file)
    book=load_workbook(file)
    assert [c.value for c in book['结果'][1]]==['office','amount']
    assert book['结果']['B2'].number_format=='#,##0.############'
    assert book['语义说明']['B3'].value=='amount'
    book.close()
