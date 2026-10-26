"""Real saved assets, exact full-result paging and safe authenticated previews."""
import hashlib
import io
import json
from decimal import Decimal

import pytest
from fastapi import FastAPI,Header,HTTPException
from fastapi.testclient import TestClient
from openpyxl import Workbook,load_workbook
from sqlalchemy import select

from agent_service import workspace_artifacts as assets,workspace,office,agent_handoff
from agent_service.config import DATA
from agent_service.jobs import jobs,engine
from agent_service.store import all_items,delete,get,patch,put,uid,now
from test_agent_handoff import project,turn,source,result


@pytest.fixture(autouse=True)
def clean():
    kinds=('source','source_version','project','turn','conversation','data_workflow','pipeline','office_artifact','business_report','workspace_artifact','workspace_artifact_export','agent_handoff_import','agent_handoff_artifact_import','agent_handoff_result_import','agent_imported_result','agent_task_artifact','metric')
    before={kind:{v['id'] for v in all_items(kind)} for kind in kinds}
    with engine.connect() as connection:previous_jobs=set(connection.execute(select(jobs.c.id)).scalars())
    yield
    with engine.begin() as connection:
        query=jobs.delete()
        if previous_jobs:query=query.where(jobs.c.id.not_in(previous_jobs))
        connection.execute(query)
    for kind in kinds:
        for item in all_items(kind):
            if item['id'] not in before[kind]:delete(kind,item['id'])


@pytest.fixture
def api():
    app=FastAPI()
    def auth(authorization:str|None=Header(None)):
        if authorization!='Bearer test':raise HTTPException(401,'Unauthenticated')
    assets.install(app,auth)
    return TestClient(app),{'Authorization':'Bearer test'}


def entry(pid,kind):
    found=assets.catalog(pid,kind)['items'];assert found
    return found[0]


def test_catalog_pages_actual_assets_and_kind_filter_without_preview_count(api):
    pid=project()
    for index in range(5):workspace.write_file(pid,f'reports/{index}.md','# Revenue\n'+str(index))
    source(pid)
    first=assets.catalog(pid,limit=2)
    assert first['total']==6 and len(first['items'])==2 and first['next_offset']==2
    second=assets.catalog(pid,offset=2,limit=2);third=assets.catalog(pid,offset=4,limit=2)
    assert len({v['id'] for page in (first,second,third) for v in page['items']})==6
    assert third['next_offset'] is None and not first['page_is_complete_catalog']
    assert assets.catalog(pid,'code')['total']==5
    assert all(v['kind']=='code' and not v['integrity_verified'] for v in assets.catalog(pid,'code')['items'])
    client,headers=api
    assert client.get(f'/api/projects/{pid}/artifacts').status_code==401
    response=client.get(f'/api/projects/{pid}/artifacts?kind=code&offset=2&limit=2',headers=headers)
    assert response.status_code==200 and response.json()['total']==5
    assert client.get(f'/api/projects/{pid}/artifacts?kind=unknown',headers=headers).status_code==409


def test_safe_markdown_text_paging_versions_and_cross_project_scope(api):
    pid=project();text='# Revenue\n<script>window.BAD=true</script>\n'+'1234567890'*100
    workspace.write_file(pid,'report.md',text);old=entry(pid,'code')
    first=assets.preview(pid,old['id'],limit=50,text_limit=110)
    assert first['render']=='markdown' and first['text']==text[:110] and first['next_offset']==110
    assert first['total']==len(text) and first['unit']=='characters' and first['active_content_executed'] is False
    assert first['artifact']['integrity_verified'] and not first['artifact']['financial_verified']
    workspace.write_file(pid,'report.md','Changed report');new=entry(pid,'code')
    assert old['id']!=new['id'] and assets.preview(pid,old['id'],text_limit=2000)['text']==text
    client,headers=api
    assert client.get(f'/api/projects/{project()}/artifacts/{old["id"]}/preview',headers=headers).status_code==404
    response=client.get(old['download_url'],headers=headers)
    assert response.status_code==200 and response.content==text.encode()
    assert response.headers['x-artifact-sha256']==hashlib.sha256(response.content).hexdigest()
    assert response.headers['x-content-type-options']=='nosniff'
    assert 'private' not in json.dumps(client.get(f'/api/projects/{pid}/artifacts',headers=headers).json())


def test_exact_full_result_pages_chart_and_real_xlsx_download(api):
    pid=project();rid,rows,schema=result(pid,count=140)
    item=entry(pid,'result');first=assets.preview(pid,item['id'],limit=50)
    assert first['total']==140 and len(first['rows'])==50 and first['next_offset']==50
    assert first['columns'][1]['unit']=='CNY' and first['context']['output_semantics']==schema
    assert first['rows'][0]['amount']=='7.500000000001' and first['rows'][0]['day']=='2026-10-01'
    assert not first['page_is_complete_result'] and first['complete_result_available']
    values=[];offset=0
    while True:
        page=assets.preview(pid,item['id'],offset,50);values.extend(page['rows'])
        if page['next_offset'] is None:break
        offset=page['next_offset']
    assert len(values)==140 and sum(Decimal(v['amount']) for v in values)==Decimal('1050.000000000140')
    chart=assets.preview(pid,item['id'],chart_x='paid',chart_y='amount')['chart']
    assert chart['source_rows']==140 and chart['points']==[{'label':'True','value':'1050.000000000140'}]
    client,headers=api
    response=client.get(item['download_url']+'?format=xlsx',headers=headers)
    assert response.status_code==200,response.text
    assert response.headers['x-artifact-sha256']==hashlib.sha256(response.content).hexdigest()
    book=load_workbook(io.BytesIO(response.content),data_only=False)
    assert book.worksheets[0].max_row==141 and str(book.worksheets[0]['B141'].value)=='7.500000000001'
    book.close()


def test_external_agent_result_needs_semantic_review_even_with_valid_full_hash():
    original=project();rid,_,_=result(original);target=project()
    accepted=agent_handoff.import_bundle(target,agent_handoff.build_input(turn(original),[],[rid]),scope_id='peer',request_id='artifact-review')
    item=entry(target,'result')
    assert item['status']=='needs_review' and item['semantic_status']=='requires_review' and not item['financial_verified']
    page=assets.preview(target,item['id'])
    assert page['artifact']['integrity_verified'] and page['artifact']['semantic_status']=='requires_review'
    assert page['context']['output_semantics']['columns'][1]['unit']=='CNY'


def test_percentages_and_unknown_additivity_are_not_summed_in_charts():
    rows=[{'office':'E','percent':'0.15'}]
    columns=[{'name':'office'},{'name':'percent','aggregation':'non_additive','value_format':'percent_fraction'}]
    with pytest.raises(ValueError,match='不能猜测'):assets._chart(rows,columns,'office','percent')
    columns[1].update(aggregation='unknown',value_format='native')
    with pytest.raises(ValueError):assets._chart(rows,columns,'office','percent')
    columns[1]['aggregation']='additive';rows[0]['percent']=None
    with pytest.raises(ValueError,match='未补零'):assets._chart(rows,columns,'office','percent')
    chart=assets._chart([{'office':str(i),'percent':'1.000000000001'} for i in range(130)],columns,'office','percent')
    assert chart['source_rows']==130 and chart['total_categories']==130 and chart['shown_categories']==100 and not chart['complete']


def test_source_revision_selection_archive_project_trash_and_byte_tamper(api):
    pid=project();s=source(pid);item=entry(pid,'source');client,headers=api
    expected=assets.file(pid,item['id'])[0].read_bytes()
    response=client.get(item['download_url'],headers=headers)
    assert response.content==expected
    from agent_service.project_data import source_path
    original_path=source_path(s);original_path.write_bytes(b'bad')
    assert client.get(item['download_url'],headers=headers).status_code==409
    original_path.write_bytes(expected);patch('source',s['id'],archived=True)
    assert client.get(item['preview_url'],headers=headers).status_code==410
    patch('source',s['id'],archived=False);patch('project',pid,deleted_at=now())
    assert client.get(f'/api/projects/{pid}/artifacts',headers=headers).status_code==410
    assert client.get(item['download_url'],headers=headers).status_code==410


def test_original_xlsx_multiple_sheets_merged_grid_formats_and_formula_text():
    pid=project();book=Workbook();ws=book.active;ws.title='空行和合并';ws['B2']='2026收入';ws.merge_cells('B2:D2');ws['B3']='东区';ws['C3']='100.000000000001';ws['D3']='=SUM(C3:C3)';ws['C3'].number_format='0.00%'
    second=book.create_sheet('明细');second.append(['id','收入']);second.append(['000001','9007199254740993.0001'])
    raw=io.BytesIO();book.save(raw);source(pid,raw.getvalue(),'原始.xlsx');item=entry(pid,'source')
    page=assets.preview(pid,item['id'],sheet='空行和合并')
    assert page['sheet_names']==['空行和合并','明细'] and page['current_sheet']=='空行和合并'
    assert page['rows'][1]['c_2']=='2026收入' and page['rows'][1]['c_3'] is None
    assert page['rows'][2]['c_4']=='=SUM(C3:C3)' and page['cell_formats']['C3']=='0.00%'
    assert page['formula_recalculation'] is False and not page['business_semantics_approved']
    assert assets.preview(pid,item['id'],sheet='明细')['rows'][1]['c_2']=='9007199254740993.0001'
    with pytest.raises(ValueError,match='页签不存在'):assets.preview(pid,item['id'],sheet='不存在')


@pytest.mark.parametrize('kind,extension',[('document','docx'),('presentation','pptx')])
def test_real_office_extracts_content_and_marks_layout_limit(kind,extension):
    pid=project();t=turn(pid)
    spec={'kind':kind,'title':'收入说明','sections':[{'heading':'核对结果','body':'精确收入 300.000000000002 CNY'}]}
    workspace.write_file(pid,'spec.json',json.dumps(spec,ensure_ascii=False));rendered=office.render(t,'spec.json','收入报告')
    item=entry(pid,'office');preview=assets.preview(pid,item['id'])
    assert item['format']==extension and preview['render']=='document' and preview['layout_fidelity'] is False
    assert '300.000000000002' in json.dumps(preview,ensure_ascii=False)
    assert preview['unit']=='blocks' and preview['total']>=2
    path,metadata=assets.file(pid,item['id'])
    assert hashlib.sha256(path.read_bytes()).hexdigest()==metadata['download_sha256']==rendered['sha256']


def test_docx_table_grid_keeps_empty_and_merged_anchor_positions(tmp_path):
    from docx import Document
    document=Document();document.add_heading('证据',0);table=document.add_table(rows=2,cols=3)
    table.cell(0,0).text='办事处';table.cell(0,1).text='金额';table.cell(0,2).text='备注'
    table.cell(1,0).text='东区';table.cell(1,1).text='100.000000000001';table.cell(1,2).text='待确认'
    raw=io.BytesIO();document.save(raw);pid=project();source(pid,raw.getvalue(),'表格.docx')
    preview=assets.preview(pid,entry(pid,'source')['id']);tables=[b for b in preview['blocks'] if b['type']=='table']
    assert len(tables)==1 and tables[0]['rows'][1]['c_2']=='100.000000000001'
    assert preview['layout_fidelity'] is False


def test_origin_chat_deleted_does_not_delete_independent_office_asset():
    pid=project();sid='c_'+uid();put('conversation',{'project_id':pid,'name':'报告工作','created_at':now()},sid);t=turn(pid,session_id=sid)
    workspace.write_file(pid,'spec.json',json.dumps({'kind':'document','title':'报告','sections':[{'body':'123'}]}));own=office.render(t,'spec.json','报告')
    patch('office_artifact',own['artifact_id'],origin_turn_id=t['id'])
    assert entry(pid,'office')['origin_turn_id']==t['id']
    patch('conversation',sid,deleted_at=now())
    item=entry(pid,'office');assert item['origin_turn_id'] is None
    assert assets.preview(pid,item['id'])['render']=='document'


def test_json_large_integer_decimal_and_unrecognized_binary_have_honest_views():
    pid=project();source(pid,b'[{"id":9007199254740993,"amount":0.10000000000000000001}]','values.json')
    page=assets.preview(pid,entry(pid,'source')['id'])
    assert page['rows']==[{'id':'9007199254740993','amount':'0.10000000000000000001'}]
    source(pid,b'binary data','legacy.bin');item=next(v for v in assets.catalog(pid,'source')['items'] if v['format']=='bin')
    assert assets.preview(pid,item['id'])['render']=='download_only'


def test_pdf_download_is_authenticated_and_html_is_never_executed(api):
    pid=project();source(pid,b'%PDF-1.4\n%%EOF','report.pdf');pdf=entry(pid,'source');client,headers=api
    preview=assets.preview(pid,pdf['id']);assert preview['render']=='pdf' and preview['url'].startswith('/api/projects/'+pid+'/')
    assert client.get(preview['url']).status_code==401
    assert client.get(preview['url'],headers=headers).content.startswith(b'%PDF-')
    raw=b'<script>fetch("https://bad.example/")</script>';source(pid,raw,'unsafe.html');html=next(v for v in assets.catalog(pid,'source')['items'] if v['format']=='html')
    preview=assets.preview(pid,html['id']);assert preview['render']=='text' and not preview['active_content_executed']
    response=client.get(html['download_url'],headers=headers);assert response.status_code==200
    assert response.headers['content-type'].startswith('text/plain') and 'attachment' in response.headers['content-disposition']


def test_business_report_groups_office_variants_and_preserves_exact_assertions(api):
    from agent_service import business_reports as reports
    from test_business_reports import real_result,recipe
    pid,rid=real_result();r=recipe();r['expected_schema_sha256']=reports.inspect_result(pid,rid)['schema_sha256']
    report=reports.render(pid,rid,r);item=entry(pid,'report')
    assert assets.catalog(pid,'report')['total']==1 and {v['format'] for v in item['downloads']}=={'docx','pptx','md','json'}
    text=assets.preview(pid,item['id']);assert text['render']=='markdown' and '300.000002' in text['text']
    client,headers=api
    for variant in ('docx','pptx','md','json'):
        response=client.get(item['download_url']+'?format='+variant,headers=headers)
        assert response.status_code==200,(variant,response.text)
        assert hashlib.sha256(response.content).hexdigest()==response.headers['x-artifact-sha256']
    assert 'not_confirmed'==item['metric_status'] and not item['financial_verified']
    path=reports.report_file(pid,report['id']);path.write_text('tampered')
    assert client.get(item['preview_url'],headers=headers).status_code==409


def test_published_pipeline_true_parquet_pages_full_chart_and_exact_xlsx(api,monkeypatch):
    from agent_service.codegen import generate
    from agent_service.execution import enqueue,process_once
    pid=project();s=source(pid,b'id,office,amount\nA,E,"1,000.000000000001"\nB,E,200.000000000001\nC,W,0.000000000001\n')
    enqueue('probe',s['id']);assert process_once();s=get('source',s['id']);assert s['status']=='ready'
    spec=dict(s['proposed_spec'],primary_key='c_1',group_by='c_2',amount_field='c_3',currency='CNY',metric_definition='每笔收入明细金额求和')
    p=generate(s,'月末应用',spec,project_id=pid,release_mode='auto',application={'title':'月末应用','layout':'summary','chart':'bar'})
    enqueue('pipeline',p['id'],p['commit']);assert process_once();p=get('pipeline',p['id']);assert p['published_run']
    item=entry(pid,'application');page=assets.preview(pid,item['id'],limit=2)
    assert page['total']==3 and len(page['rows'])==2 and page['next_offset']==2
    chart=assets.preview(pid,entry(pid,'chart')['id'])['chart']
    assert chart['source_rows']==3 and chart['points']==[{'label':'E','value':'1200.000000000002'},{'label':'W','value':'1E-12'}]
    assert item['metric_status']=='draft' and item['semantic_status']=='requires_review'
    original_put=assets.put
    def crash_after_publish(kind,value,*args,**kwargs):
        if kind=='workspace_artifact_export' and value.get('state')=='committed':raise RuntimeError('simulated receipt commit crash')
        return original_put(kind,value,*args,**kwargs)
    monkeypatch.setattr(assets,'put',crash_after_publish)
    with pytest.raises(RuntimeError,match='simulated'):assets.file(pid,item['id'],'xlsx')
    pending=get('workspace_artifact_export',item['id']+'_xlsx');assert pending['state']=='pending' and pending['pending_sha256']
    monkeypatch.setattr(assets,'put',original_put)
    path,metadata=assets.file(pid,item['id'],'xlsx');book=load_workbook(path,data_only=False)
    sheet=book.worksheets[0];header=[c.value for c in sheet[1]];position=header.index('amount') if 'amount' in header else header.index('c_3') if 'c_3' in header else header.index('金额')
    values=[str(row[position]) for row in sheet.iter_rows(min_row=2,values_only=True)]
    assert sum(Decimal(v.replace(',','')) for v in values)==Decimal('1200.000000000003') and sheet.max_row==4
    book.close();assert hashlib.sha256(path.read_bytes()).hexdigest()==metadata['download_sha256']
    first=path.read_bytes();assert assets.file(pid,item['id'],'xlsx')[0].read_bytes()==first
    patch('pipeline',p['id'],archived=True)
    with pytest.raises(assets.ArtifactGone):assets.preview(pid,item['id'])


def test_xlsx_xml_numeric_text_is_preserved_before_binary_float_conversion():
    import zipfile
    book=Workbook();book.active['A1']='amount';book.active['A2']=1.2;buffer=io.BytesIO();book.save(buffer)
    altered=io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(buffer.getvalue())) as original,zipfile.ZipFile(altered,'w') as output:
        for info in original.infolist():
            raw=original.read(info.filename)
            if info.filename=='xl/worksheets/sheet1.xml':raw=raw.replace(b'<v>1.2</v>',b'<v>9007199254740993.000000000001</v>')
            output.writestr(info,raw)
    pid=project();source(pid,altered.getvalue(),'exact.xlsx');page=assets.preview(pid,entry(pid,'source')['id'])
    assert page['rows'][1]['c_1']=='9007199254740993.000000000001'
    assert page['raw_numeric_values']['A2']=='9007199254740993.000000000001'


def test_remote_office_and_markdown_outputs_enter_catalog_with_exact_safe_previews(api):
    from agent_service.actor_context import actor_scope
    parent=turn();pid=parent['project_id'];remote=project();remote_turn=turn(remote)
    workspace.write_file(remote,'spec.json',json.dumps({'kind':'document','title':'远端经营报告','sections':[{'body':'精确收入300.000000000002 CNY'}]},ensure_ascii=False))
    rendered=office.render(remote_turn,'spec.json','远端经营报告');path,_=office.artifact(remote,rendered['artifact_id'])
    bundle=agent_handoff._bundle([agent_handoff._item('external-office','artifact','经营报告.docx',path.read_bytes()),
        agent_handoff._item('external-markdown','artifact','经营小结.md',b'# Revenue\n300.000000000002\n<script>alert(1)</script>')])
    accepted=agent_handoff.import_artifacts(parent,bundle,scope_id='mesh:catalog-test',request_id='output-1')
    entries=assets.catalog(pid,'agent')['items'];assert len(entries)==2
    assert {v['resource_id'] for v in all_items('workspace_artifact') if v.get('project_id')==pid}==set(accepted['artifact_ids'])
    client,headers=api
    for item in entries:
        assert item['status']=='needs_review' and item['semantic_status']=='requires_review' and not item['financial_verified']
        assert item['origin_turn_id']==parent['id'] and item['source_refs'][0]['task_id']==parent['id']
        preview=client.get(item['preview_url'],headers=headers);assert preview.status_code==200
        value=preview.json();assert value['artifact']['integrity_verified'] and not value['artifact']['financial_verified']
        assert '300.000000000002' in json.dumps(value)
        if item['format']=='docx':assert value['render']=='document' and not value['layout_fidelity']
        else:assert value['render']=='markdown' and not value['active_content_executed']
        download=client.get(item['download_url'],headers=headers);assert download.status_code==200
        assert hashlib.sha256(download.content).hexdigest()==item['sha256']==download.headers['x-artifact-sha256']
        assert client.get(item['download_url']).status_code==401
        assert client.get(item['preview_url'].replace(pid,remote),headers=headers).status_code==404


def test_default_catalog_prioritizes_business_outputs_over_new_internal_specs_and_sources():
    pid=project();parent=turn(pid);rid,_,_=result(pid)
    bundle=agent_handoff._bundle([agent_handoff._item('summary','artifact','经营报告.md',b'# Revenue\n300.000000000002')])
    agent_handoff.import_artifacts(parent,bundle,scope_id='mesh:priority',request_id='report')
    for index in range(4):workspace.write_file(pid,f'reports/new-spec-{index}.json',json.dumps({'internal':'new specification'}))
    source(pid)
    first=assets.catalog(pid,limit=2)
    assert first['total']==7 and [v['kind'] for v in first['items']]==['agent','result']
    assert assets.catalog(pid,'code')['total']==4 and assets.catalog(pid,'source')['total']==1
    assert assets.catalog(pid,offset=6,limit=2)['next_offset'] is None


def replace_source_snapshot(value,raw,name=None):
    from pathlib import Path
    from agent_service.config import DATA
    from agent_service.sources import revision
    from agent_service.store import seal
    name=name or value['name'];stored=uid()+Path(name).suffix
    (DATA/'uploads'/stored).write_bytes(raw)
    return revision(patch('source',value['id'],name=name,bytes=len(raw),
        content_hash=hashlib.sha256(raw).hexdigest(),secret_config=seal({'stored_name':stored})))


def test_direct_uploaded_source_descriptor_opens_real_workbook_sheets_and_text(api):
    pid=project();client,headers=api
    book=Workbook();book.active.title='收入';book.active.append(['金额','口径'])
    book.active.append(['9007199254740993.000000000001','<script>do_not_run()</script>'])
    book.create_sheet('预算').append(['预算','含税']);buffer=io.BytesIO();book.save(buffer);book.close()
    value=source(pid,buffer.getvalue(),'业务源表.xlsx')
    route=f'/api/projects/{pid}/sources/{value["id"]}/artifact?revision=1'
    assert client.get(route).status_code==401
    response=client.get(route,headers=headers);assert response.status_code==200,response.text
    item=response.json();assert item['kind']=='source' and item['integrity_verified']
    assert item['source_refs']==[{'source_id':value['id'],'revision':1}]
    assert 'secret_config' not in response.text and 'private' not in item
    page=client.get(item['preview_url'],headers=headers).json()
    assert page['sheet_names']==['收入','预算'] and page['rows'][1]['c_1']=='9007199254740993.000000000001'
    budget=client.get(item['preview_url']+'?sheet=预算',headers=headers).json()
    assert budget['current_sheet']=='预算' and budget['rows'][0]['c_1']=='预算'
    download=client.get(item['download_url'],headers=headers)
    assert download.status_code==200 and download.content==buffer.getvalue()
    assert client.get(item['preview_url']).status_code==401
    text='口径：负数保留；金额单位不得猜测。\n<script>do_not_run()</script>'
    note=source(pid,text.encode(),'口径.txt')
    descriptor=client.get(f'/api/projects/{pid}/sources/{note["id"]}/artifact',headers=headers).json()
    page=client.get(descriptor['preview_url'],headers=headers).json()
    assert page['render']=='text' and page['text']==text and page['active_content_executed'] is False


def test_uploaded_source_descriptor_freezes_explicit_revision_and_never_uses_latest(api):
    pid=project();client,headers=api;old=source(pid,'原口径：含税。'.encode(),'原口径.txt')
    route=f'/api/projects/{pid}/sources/{old["id"]}/artifact'
    first=client.get(route+'?revision=1',headers=headers).json()
    current=replace_source_snapshot(old,'新口径：不含税。'.encode(),'新口径.txt')
    latest=client.get(route,headers=headers).json()
    frozen=client.get(route+'?revision=1',headers=headers).json()
    assert current['revision']==2 and latest['name']=='新口径.txt'
    assert first['id']==frozen['id']!=latest['id'] and frozen['name']=='原口径.txt'
    assert frozen['source_refs'][0]['revision']==1 and latest['source_refs'][0]['revision']==2
    assets.catalog(pid,'source')
    assert client.get(frozen['preview_url'],headers=headers).json()['text']=='原口径：含税。'
    assert client.get(frozen['download_url'],headers=headers).content=='原口径：含税。'.encode()
    assert client.get(latest['preview_url'],headers=headers).json()['text']=='新口径：不含税。'
    assert client.get(route+'?revision=99',headers=headers).status_code==404


def test_identical_source_bytes_in_new_revision_do_not_overwrite_prior_descriptor(api):
    pid=project();client,headers=api;old=source(pid,b'unchanged business text','old.txt')
    route=f'/api/projects/{pid}/sources/{old["id"]}/artifact'
    first=client.get(route+'?revision=1',headers=headers).json()
    replace_source_snapshot(old,b'unchanged business text','renamed.txt')
    latest=client.get(route+'?revision=2',headers=headers).json();assets.catalog(pid,'source')
    assert first['id']!=latest['id'] and first['sha256']==latest['sha256']
    old_again=client.get(first['preview_url'],headers=headers).json()['artifact']
    assert old_again['name']=='old.txt' and old_again['source_refs'][0]['revision']==1


def test_source_descriptor_scope_archives_and_invalid_revision_reject_without_fallback(api):
    pid=project();client,headers=api;value=source(pid,b'policy','policy.txt')
    route=f'/api/projects/{pid}/sources/{value["id"]}/artifact'
    assert client.get(route.replace(pid,project()),headers=headers).status_code==404
    assert client.get(route+'?revision=0',headers=headers).status_code==409
    assert client.get(route+'?revision=-1',headers=headers).status_code==409
    assert client.get(route+'?revision=invalid',headers=headers).status_code==422
    patch('source',value['id'],archived=True)
    assert client.get(route,headers=headers).status_code==410
    patch('source',value['id'],archived=False);patch('project',pid,deleted_at=now())
    assert client.get(route,headers=headers).status_code==410


@pytest.mark.parametrize('damage',['missing','tamper'])
def test_source_descriptor_checks_missing_and_tampered_bytes_before_returning(api,damage):
    from agent_service.project_data import source_path
    pid=project();client,headers=api;value=source(pid,b'original business text','policy.txt')
    path=source_path(value)
    if damage=='missing':path.unlink()
    else:path.write_bytes(b'unauthorized replacement')
    response=client.get(f'/api/projects/{pid}/sources/{value["id"]}/artifact?revision=1',headers=headers)
    assert response.status_code==409 and 'preview_url' not in response.text


def test_historical_source_version_project_or_archive_tamper_is_rejected(api):
    pid=project();client,headers=api;value=source(pid,b'original','policy.txt')
    replace_source_snapshot(value,b'current')
    route=f'/api/projects/{pid}/sources/{value["id"]}/artifact?revision=1'
    patch('source_version',value['revision_id'],project_id=project())
    assert client.get(route,headers=headers).status_code==404
    patch('source_version',value['revision_id'],project_id=pid,archived=True)
    assert client.get(route,headers=headers).status_code==410


def test_write_only_workbook_without_dimension_has_real_rows_and_empty_sheet_stays_empty():
    import zipfile
    pid=project();book=Workbook(write_only=True);sheet=book.create_sheet('流水')
    sheet.append(['金额','公式']);sheet.append(['9007199254740993.000000000001','=1+2'])
    book.create_sheet('空页签');buffer=io.BytesIO();book.save(buffer);book.close()
    with zipfile.ZipFile(io.BytesIO(buffer.getvalue())) as archive:
        assert b'<dimension' not in archive.read('xl/worksheets/sheet1.xml')
    value=source(pid,buffer.getvalue(),'streamed.xlsx');item=assets.source_artifact(pid,value['id'])
    page=assets.preview(pid,item['id'],limit=1)
    assert page['total']==2 and len(page['columns'])==2 and page['next_offset']==1
    second=assets.preview(pid,item['id'],offset=1,limit=1)
    assert second['rows']==[{'c_1':'9007199254740993.000000000001','c_2':'=1+2'}]
    assert second['total']==2 and second['next_offset'] is None and not second['formula_recalculation']
    empty=assets.preview(pid,item['id'],sheet='空页签')
    assert empty['total']==0 and empty['rows']==[] and empty['columns']==[]


def test_underreported_workbook_dimension_does_not_hide_coordinates_formulas_or_raw_numbers():
    import re,zipfile
    pid=project();book=Workbook();sheet=book.active
    sheet['A1']='原始合并锚点';sheet.merge_cells('A1:E1')
    sheet['B2']=1.2;sheet['D4']='=SUM(B2:B3)';sheet['C5']='原始坐标C5'
    original=io.BytesIO();book.save(original);book.close();altered=io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(original.getvalue())) as archive,zipfile.ZipFile(altered,'w') as output:
        for entry in archive.infolist():
            raw=archive.read(entry.filename)
            if entry.filename=='xl/worksheets/sheet1.xml':
                raw=re.sub(rb'<dimension ref="[^"]+"',b'<dimension ref="A1:A1"',raw)
                raw=raw.replace(b'<v>1.2</v>',b'<v>9007199254740993.000000000001</v>')
            output.writestr(entry,raw)
    value=source(pid,altered.getvalue(),'stale-dimension.xlsx');item=assets.source_artifact(pid,value['id'])
    page=assets.preview(pid,item['id'])
    assert page['total']==5 and len(page['columns'])==4 and len(page['rows'])==5
    assert page['rows'][0]['c_1']=='原始合并锚点' and page['rows'][0]['c_2'] is None
    assert page['rows'][1]['c_2']=='9007199254740993.000000000001'
    assert page['raw_numeric_values']['B2']=='9007199254740993.000000000001'
    assert page['rows'][3]['c_4']=='=SUM(B2:B3)' and page['rows'][4]['c_3']=='原始坐标C5'


def test_real_finance_phase1_workbook_all_four_sheet_previews_have_actual_content():
    from pathlib import Path
    fixture=Path(__file__).resolve().parents[1]/'acceptance/finance_monthly/inputs/phase1/2026-09_业务源表.xlsx'
    pid=project();value=source(pid,fixture.read_bytes(),fixture.name);item=assets.source_artifact(pid,value['id'])
    expected={'收入底稿':(21,9),'成本底稿':(15,9),'预算':(7,5),'省份映射':(10,4)}
    for sheet,(rows,columns) in expected.items():
        page=assets.preview(pid,item['id'],sheet=sheet)
        assert page['total']==rows and len(page['rows'])==rows and len(page['columns'])==columns
        assert any(cell is not None for row in page['rows'] for cell in row.values()),sheet
        assert page['current_sheet']==sheet and page['sheet_names']==list(expected)


@pytest.mark.parametrize('row,column',[(1,201),(6,1)])
def test_underreported_dimension_cannot_bypass_actual_grid_budget(monkeypatch,row,column):
    import re,zipfile
    monkeypatch.setattr(assets,'MAX_ROWS',5)
    pid=project();book=Workbook();book.active.cell(row,column,'outside the preview budget')
    original=io.BytesIO();book.save(original);book.close();altered=io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(original.getvalue())) as archive,zipfile.ZipFile(altered,'w') as output:
        for entry in archive.infolist():
            raw=archive.read(entry.filename)
            if entry.filename=='xl/worksheets/sheet1.xml':
                raw=re.sub(rb'<dimension ref="[^"]+"',b'<dimension ref="A1:A1"',raw)
            output.writestr(entry,raw)
    value=source(pid,altered.getvalue(),'underreported-budget.xlsx');item=assets.source_artifact(pid,value['id'])
    with pytest.raises(assets.ArtifactError,match='超过完整预览预算'):
        assets.preview(pid,item['id'])
