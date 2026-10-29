"""Controller/runner protocol and fault tests. FakeDocker does not prove isolation."""
import base64
import io
import json
import tarfile

import httpx
import pytest
from fastapi import FastAPI, Header, HTTPException
from fastapi.testclient import TestClient

from work_runner.protocol import Command, InputFile, digest, pack, unpack, file_path
from work_runner.service import Engine, create_app
from agent_service import agent, agent_context, provider, tasks, work_runtime as work, workspace_artifacts
from agent_service.store import get, patch, put, uid, now, all_items, delete
from agent_service.actor_context import actor_scope


@pytest.fixture(autouse=True)
def clean_test_records():
    kinds=('work_execution','work_artifact','workspace_artifact','installed_skill','checkpoint','turn','conversation','source','project')
    before={kind:{value['id'] for value in all_items(kind)} for kind in kinds}
    yield
    for kind in kinds:
        for value in all_items(kind):
            if value['id'] not in before[kind]:delete(kind,value['id'])


class FakeDocker:
    """Only a daemon-protocol fixture, never an application execution backend."""
    def __init__(self):
        self.containers = {}
        self.execs = {}
        self.launches = 0
        self.output = {'result.csv': b'customer,amount\nA,118.1500\n'}
        self.fail_snapshot = False
        self.exit_code = 0
        self.keep_running = False

    def ready(self):
        return {'runtime': 'runsc', 'image_id': 'fixture-only', 'network': 'none'}

    def create(self, job):
        cid = job['execution_id']
        self.containers[cid] = {}
        return cid

    def start(self, cid):
        pass

    def stage(self, cid, archive):
        self.containers[cid] = unpack(archive)

    def prepare_exec(self, cid, command, timeout):
        eid = digest(cid.encode())
        self.execs[eid] = {'Running': False, 'ExitCode': self.exit_code, 'container': cid}
        return eid

    def start_exec(self, eid):
        self.launches += 1
        self.execs[eid]['Running'] = self.keep_running
        self.containers[self.execs[eid]['container']].update(self.output)

    def inspect_exec(self, eid):
        return self.execs[eid]

    def logs(self, cid):
        return {'stdout': 'completed fixture\n', 'stderr': ''}

    def quiesce(self, cid):
        pass

    def files(self, cid):
        if self.fail_snapshot:
            raise ValueError('malicious archive')
        return self.containers[cid]

    def remove(self, cid):
        self.containers.pop(cid, None)


def command(*, eid=None, wid=None, text='python analysis.py', files=None):
    return Command(execution_id=eid or digest(uid().encode()), workspace_id=wid or digest(b'workspace'),
                   command=text, files=files or [])


def finish(engine, value):
    engine.step(value.execution_id)
    engine.step(value.execution_id)
    return engine.get(value.execution_id)


@pytest.mark.parametrize('path', ['/etc/passwd', '../secret', 'a/../b', 'a//b', './x', 'a\\b', '.atlas-key', 'a\x00b'])
def test_paths_refuse_host_and_internal_names(path):
    with pytest.raises(ValueError):
        file_path(path)


def test_archive_never_extracts_links_or_duplicate_paths():
    for kind in (tarfile.SYMTYPE, tarfile.LNKTYPE, tarfile.FIFOTYPE):
        raw = io.BytesIO()
        with tarfile.open(fileobj=raw, mode='w') as archive:
            info = tarfile.TarInfo('escape'); info.type = kind; info.linkname = '/etc/passwd'; archive.addfile(info)
        with pytest.raises(ValueError):
            unpack(raw.getvalue())
    assert unpack(pack({'outputs/中文.csv': b'exact,118.1500\n'})) == {'outputs/中文.csv': b'exact,118.1500\n'}


def test_invalid_file_digest_and_request_id_reuse(tmp_path):
    with pytest.raises(ValueError):
        InputFile(path='input.csv', content_base64=base64.b64encode(b'wrong').decode(), sha256=digest(b'expected'))
    engine = Engine(tmp_path, FakeDocker()); value = command()
    assert engine.submit(value) == engine.submit(value)
    with pytest.raises(ValueError, match='different request'):
        engine.submit(value.model_copy(update={'command':'different'}))
    assert engine.backend.launches == 0


def test_runner_http_auth_precedes_body_and_scopes_files(tmp_path):
    engine = Engine(tmp_path, FakeDocker()); client = TestClient(create_app(engine, 'r' * 40))
    assert client.post('/v1/executions', content='{malformed').status_code == 401
    headers = {'Authorization':'Bearer ' + 'r' * 40}
    assert client.post('/v1/executions', headers=headers, json={'secret':'do-not-reflect-me'}).json() == {'detail':'Invalid execution request; input was not executed'}
    value = command(); assert client.post('/v1/executions', headers=headers, json=value.model_dump()).status_code == 200
    finish(engine, value)
    endpoint = '/v1/executions/' + value.execution_id + '/file'
    assert client.get(endpoint, headers=headers, params={'workspace_id':digest(b'other'), 'path':'result.csv'}).status_code == 404
    response = client.get(endpoint, headers=headers, params={'workspace_id':value.workspace_id, 'path':'result.csv'})
    assert base64.b64decode(response.json()['content_base64']) == engine.backend.output['result.csv']


def test_restart_reattaches_running_execution_without_replay(tmp_path):
    backend = FakeDocker(); backend.keep_running = True
    engine = Engine(tmp_path, backend); value = command(); engine.submit(value); engine.step(value.execution_id)
    assert engine.get(value.execution_id)['state'] == 'running'
    recovered = Engine(tmp_path, backend); backend.execs[next(iter(backend.execs))]['Running'] = False
    recovered.step(value.execution_id)
    assert recovered.get(value.execution_id)['state'] == 'succeeded' and backend.launches == 1
    recovered.submit(value); recovered.step(value.execution_id)
    assert backend.launches == 1


def test_failed_checkpoint_keeps_previous_workspace_and_unknown_start_never_replays(tmp_path):
    backend = FakeDocker(); engine = Engine(tmp_path, backend); first = command(); engine.submit(first); finish(engine, first)
    previous = engine.latest(first.workspace_id)['execution_id']
    backend.fail_snapshot = True
    second = command(wid=first.workspace_id); engine.submit(second); assert finish(engine, second)['state'] == 'failed'
    assert engine.latest(first.workspace_id)['execution_id'] == previous
    third = command(wid=first.workspace_id); engine.submit(third)
    current = engine.get(third.execution_id); current.update(state='starting', container_id='orphan'); engine.save(current)
    Engine(tmp_path, backend).step(third.execution_id)
    assert engine.get(third.execution_id)['state'] == 'unknown' and backend.launches == 2


def test_cancel_and_workspace_serialization(tmp_path):
    backend = FakeDocker(); backend.keep_running = True
    engine = Engine(tmp_path, backend); value = command(); engine.submit(value)
    next_command=command(wid=value.workspace_id);engine.submit(next_command)
    engine.step(value.execution_id);engine.step(next_command.execution_id)
    assert engine.get(next_command.execution_id)['state']=='queued' and backend.launches==1
    engine.step(value.execution_id); engine.cancel(value.execution_id, value.workspace_id); engine.step(value.execution_id)
    assert engine.get(value.execution_id)['state'] == 'cancelled' and not backend.containers
    engine.step(next_command.execution_id)
    assert engine.get(next_command.execution_id)['state']=='running' and backend.launches==2


@pytest.fixture
def controller(tmp_path, monkeypatch):
    monkeypatch.setenv('ATLAS_WORK_RUNNER_URL', 'https://runner.example.test')
    monkeypatch.setenv('ATLAS_WORK_RUNNER_TOKEN', 'r' * 40)
    monkeypatch.setattr(provider, 'settings', lambda:{'share_data':True})
    backend = FakeDocker(); engine = Engine(tmp_path, backend)
    remote = TestClient(create_app(engine, 'r' * 40))
    def request(method, path, payload=None, params=None):
        response = remote.request(method, '/v1/' + path, headers={'Authorization':'Bearer ' + 'r' * 40}, json=payload, params=params)
        if response.status_code == 404:raise KeyError(path)
        response.raise_for_status(); return response.json()
    monkeypatch.setattr(work, '_request', request)
    pid = 'p_' + uid(); put('project', {'name':'Work mode acceptance'}, pid)
    sid = 's_' + uid(); put('conversation', {'project_id':pid, 'actor_id':'owner','name':'Work'}, sid)
    turn = put('turn', {'project_id':pid,'session_id':sid,'actor_id':'owner','text':'运行脚本并生成结果文件',
        'interaction_mode':'work','state':'agent','created_at':now()}, 't_' + uid())
    return engine, backend, turn


def run_tool(turn, name, arguments):
    agent._execute_node({'turn_id':turn['id'], 'step':{'kind':'tool','tool':name,'arguments':arguments}})
    return get('turn', turn['id'])


def test_agent_background_resume_deliver_and_catalog_without_model(controller):
    engine, backend, turn = controller
    turn = run_tool(turn, 'run_work_command', {'command':'python analysis.py'})
    assert turn['state'] == 'waiting_work_command' and turn['agent_trace'][-1]['status'] == 'started'
    eid = turn['work_execution_id']; engine.step(eid); engine.step(eid)
    assert work.resume(turn)
    turn = get('turn', turn['id'])
    assert turn['state'] == 'agent' and turn['agent_trace'][-1]['observation']['exit_code'] == 0
    assert tasks.verify(turn)['missing'] == ['file']
    turn = run_tool(turn, 'publish_work_file', {'execution_id':eid, 'path':'result.csv'})
    receipt = turn['agent_trace'][-1]['observation']; assert receipt['status'] == 'needs_review'
    assert tasks.verify(turn)['verified']
    path, metadata = work.artifact(turn['project_id'], receipt['artifact_id'])
    assert path.read_bytes() == backend.output['result.csv'] and not metadata['financial_approval']
    assets = workspace_artifacts.catalog(turn['project_id'], 'work')['items']
    assert len(assets) == 1 and not assets[0]['financial_verified']
    resolved=work.execute(turn,'work_environment_status',{'artifact_id':assets[0]['id']})
    assert resolved['artifact']['execution_id']==eid and resolved['artifact']['path']=='result.csv'
    assert resolved['inputs_match_current'] is True and resolved['artifact']['financial_approval'] is False
    preview = workspace_artifacts.preview(turn['project_id'], assets[0]['id'])
    assert preview['rows'][0]['c_2'] == '118.1500'
    app = FastAPI()
    def auth(authorization:str=Header('')):
        if authorization!='Bearer test':raise HTTPException(401)
    work.install(app, auth)
    client = TestClient(app); url = receipt['download_url']
    assert client.get(url).status_code == 401
    assert client.get(url, headers={'Authorization':'Bearer test'}).content == path.read_bytes()
    with actor_scope('u_' + 'a' * 64):
        with pytest.raises(KeyError):work.artifact(turn['project_id'], receipt['artifact_id'])


def test_failed_program_requires_successful_repair(controller):
    engine, backend, turn = controller; backend.exit_code = 1
    turn = patch('turn', turn['id'], text='运行脚本')
    turn = run_tool(turn, 'run_work_command', {'command':'python broken.py'})
    engine.step(turn['work_execution_id']); engine.step(turn['work_execution_id']); work.resume(turn)
    turn = get('turn', turn['id'])
    assert not tasks.verify(turn)['verified']
    backend.exit_code = 0
    turn = run_tool(turn, 'run_work_command', {'command':'python repaired.py'})
    engine.step(turn['work_execution_id']); engine.step(turn['work_execution_id']); work.resume(turn)
    turn = get('turn', turn['id'])
    assert tasks.verify(turn)['verified'] and not agent._unresolved_failures(turn['agent_trace'])


def test_lost_submit_response_recovers_same_command_and_stop_reconciles(controller, monkeypatch):
    engine, backend, turn = controller
    request = work._request
    def lost_response(method, path, **kwargs):
        response = request(method, path, **kwargs)
        if method == 'POST' and path == 'executions':raise httpx.ReadTimeout('lost acknowledgement')
        return response
    monkeypatch.setattr(work, '_request', lost_response)
    turn = run_tool(turn, 'run_work_command', {'command':'python analysis.py'})
    assert turn['state'] == 'waiting_work_command'
    monkeypatch.setattr(work, '_request', request)
    eid = turn['work_execution_id']; engine.step(eid); engine.step(eid); work.resume(turn)
    assert backend.launches == 1
    backend.keep_running = True
    turn = run_tool(get('turn', turn['id']), 'run_work_command', {'command':'python long.py'})
    eid = turn['work_execution_id']; engine.step(eid)
    patch('turn', turn['id'], outcome='cancelled', state='done'); work.reconcile(); engine.step(eid)
    assert engine.get(eid)['state'] == 'cancelled'


def test_disabled_runner_chat_and_human_veto_stay_enforced(controller, monkeypatch):
    _, _, turn = controller
    monkeypatch.delenv('ATLAS_WORK_RUNNER_TOKEN')
    assert 'run_work_command' not in agent_context.allowed(turn)
    monkeypatch.setenv('ATLAS_WORK_RUNNER_TOKEN', 'r' * 40)
    chat = dict(turn, interaction_mode='chat')
    assert 'run_work_command' not in agent_context.allowed(chat)
    with pytest.raises(ValueError):agent.authorize(chat, 'run_work_command', {'command':'true'})
    with pytest.raises(ValueError):agent.authorize(dict(turn,text='不要执行任何操作，只解释代码'), 'run_work_command', {'command':'true'})


def test_office_extension_cannot_forge_an_office_delivery(controller):
    engine, backend, turn = controller; backend.output={'fake.xlsx':b'not a workbook'}
    turn = run_tool(turn, 'run_work_command', {'command':'generate workbook'})
    eid=turn['work_execution_id']; engine.step(eid); engine.step(eid); work.resume(turn)
    turn=run_tool(get('turn',turn['id']), 'publish_work_file', {'execution_id':eid,'path':'fake.xlsx'})
    assert turn['agent_trace'][-1]['status']=='failed'


@pytest.mark.parametrize('extension',['xlsx','docx','pptx','png'])
def test_real_binary_deliverable_can_be_downloaded_and_previewed(controller,extension):
    engine,backend,turn=controller;content=io.BytesIO()
    if extension=='xlsx':
        from openpyxl import Workbook
        book=Workbook();book.active.append(['amount']);book.active.append(['9007199254740993.00000001']);book.save(content)
    elif extension=='docx':
        from docx import Document
        document=Document();document.add_paragraph('Actual generated report');document.save(content)
    elif extension=='pptx':
        from pptx import Presentation
        deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[1]);slide.shapes.title.text='Actual generated slide';deck.save(content)
    else:
        from PIL import Image
        Image.new('RGB',(32,32),'green').save(content,format='PNG')
    backend.output={'outputs/deliverable.'+extension:content.getvalue()}
    turn=run_tool(turn,'run_work_command',{'command':'python generate.py'})
    eid=turn['work_execution_id'];engine.step(eid);engine.step(eid);work.resume(turn)
    turn=run_tool(get('turn',turn['id']),'publish_work_file',{'execution_id':eid,'path':'outputs/deliverable.'+extension})
    assert turn['agent_trace'][-1]['status']=='done',turn['agent_trace'][-1].get('observation')
    receipt=turn['agent_trace'][-1]['observation'];file,metadata=work.artifact(turn['project_id'],receipt['artifact_id'])
    assert file.read_bytes()==content.getvalue() and metadata['financial_approval'] is False
    assert file.suffix=='.'+extension
    entry=workspace_artifacts.catalog(turn['project_id'],'work')['items'][0]
    preview=workspace_artifacts.preview(turn['project_id'],entry['id'])
    assert preview['render'] in ('table','document','image')
    if extension=='xlsx':assert '9007199254740993.00000001' in json.dumps(preview)


def test_real_guest_command_log_cap_and_timeout_without_isolation_claim(tmp_path):
    import shlex
    import sys
    from work_runner.guest import execute
    code='import sys;sys.stdout.write("a"*100000);sys.stderr.write("b"*100000)'
    assert execute(shlex.quote(sys.executable)+' -c '+shlex.quote(code),5,workspace=tmp_path,logs=tmp_path)==0
    assert (tmp_path/'atlas-stdout').stat().st_size==64000
    assert (tmp_path/'atlas-stderr').stat().st_size==64000
    assert execute(shlex.quote(sys.executable)+' -c '+shlex.quote('import time;time.sleep(30)'),1,workspace=tmp_path,logs=tmp_path)==124


def test_binary_attachment_is_explicit_versioned_and_byte_exact(controller):
    engine, _, turn=controller
    from agent_service.config import DATA
    from agent_service.store import seal
    from agent_service.conversation import snapshot_attachment
    content=b'original binary\x00\xff'
    sid='s_'+uid(); name=uid()+'.xlsx'; (DATA/'uploads').mkdir(exist_ok=True); path=DATA/'uploads'/name; path.write_bytes(content)
    source=put('source',{'project_id':turn['project_id'],'kind':'file','name':'底稿.xlsx','status':'ready',
        'revision':1,'content_hash':digest(content),'bytes':len(content),'secret_config':seal({'stored_name':name})},sid)
    turn=patch('turn',turn['id'],source_ids=[sid],attached_sources=[snapshot_attachment(source)])
    turn=run_tool(turn,'run_work_command',{'command':'python analysis.py','source_ids':[sid]})
    remote=engine.get(turn['work_execution_id']); assert len(remote['files'])==1
    assert remote['files'][0]['path']=='inputs/'+sid+'/底稿.xlsx'
    assert base64.b64decode(remote['files'][0]['content_base64'])==content
    path.write_bytes(b'tampered')
    # A different command cannot silently substitute modified source bytes.
    engine.cancel(remote['execution_id'],remote['workspace_id']);engine.step(remote['execution_id']);work.resume(turn)
    turn=run_tool(get('turn',turn['id']),'run_work_command',{'command':'python corrected.py','source_ids':[sid]})
    assert turn['agent_trace'][-1]['status']=='failed' and engine.backend.launches==0


def test_runner_storage_limit_fails_closed(tmp_path,monkeypatch):
    monkeypatch.setenv('ATLAS_RUNNER_STORAGE_MB','1')
    engine=Engine(tmp_path,FakeDocker())
    with pytest.raises(ValueError,match='storage quota'):
        engine.submit(command())


def test_guest_links_cannot_be_published_from_a_partial_snapshot(controller):
    engine,backend,turn=controller;backend.fail_snapshot=True
    turn=run_tool(turn,'run_work_command',{'command':'python analysis.py'})
    eid=turn['work_execution_id'];engine.step(eid);engine.step(eid);work.resume(turn)
    turn=run_tool(get('turn',turn['id']),'publish_work_file',{'execution_id':eid,'path':'result.csv'})
    assert turn['agent_trace'][-1]['status']=='failed'


def test_cloud_work_guidance_is_provisioned_only_with_execution_configuration(controller,monkeypatch):
    from agent_service import extensions
    configured=extensions.provision_presets()
    assert any(value['id']=='skp_cloud_work' for value in configured['skills'])
    monkeypatch.delenv('ATLAS_WORK_RUNNER_TOKEN')
    unconfigured=extensions.provision_presets()
    assert not any(value['id']=='skp_cloud_work' for value in unconfigured['skills'])


def test_changed_project_code_cannot_reuse_old_success(controller):
    from agent_service import workspace
    engine,_,turn=controller
    turn=patch('turn',turn['id'],text='执行脚本')
    workspace.write_file(turn['project_id'],'analysis.py','print("original")')
    turn=run_tool(turn,'run_work_command',{'command':'python analysis.py','project_files':['analysis.py']})
    eid=turn['work_execution_id'];engine.step(eid);engine.step(eid);work.resume(turn)
    turn=get('turn',turn['id']);assert tasks.verify(turn)['verified']
    workspace.write_file(turn['project_id'],'analysis.py','print("changed")')
    assert 'execution' in tasks.verify(turn)['missing']


def test_general_web_code_is_versioned_and_previewed_as_text(controller):
    from agent_service import workspace
    _,_,turn=controller
    for name in ('page.html','app.js','style.css','tools/verify.sh'):
        workspace.write_file(turn['project_id'],name,'untrusted <script>window.evil=true</script>')
    entries=workspace_artifacts.catalog(turn['project_id'],'code')['items']
    html=next(value for value in entries if value['name']=='page.html')
    preview=workspace_artifacts.preview(turn['project_id'],html['id'])
    assert preview['render']!='sandboxed_html' and 'untrusted' in preview['text']
    workspace.checkpoint(turn['project_id'],turn['id'],'test code checkpoint')
    assert workspace.history(turn['project_id'])[0]['commit']


def test_docker_launch_contract_and_missing_gvisor_fail_closed():
    from work_runner.docker_backend import DockerBackend
    requests=[]
    runtime={'runsc':{}}
    def respond(request):
        requests.append(request)
        path=request.url.path
        if path.endswith('/info'):return httpx.Response(200,json={'Runtimes':runtime,'MemoryLimit':True,'SwapLimit':True,'PidsLimit':True,'CpuCfsQuota':True})
        if '/images/' in path:return httpx.Response(200,json={'Id':'sha256:fixture'})
        if path.endswith('/containers/create'):return httpx.Response(201,json={'Id':'container'})
        raise AssertionError(path)
    backend=DockerBackend('registry.test/workspace@sha256:'+'a'*64)
    backend.client.close()
    backend.client=httpx.Client(transport=httpx.MockTransport(respond),base_url='http://docker/v1.45')
    try:
        value=command();assert backend.create(value.model_dump())=='container'
        payload=json.loads(requests[-1].content);host=payload['HostConfig']
        assert host['Runtime']=='runsc' and host['NetworkMode']=='none' and host['ReadonlyRootfs']
        assert host['CapDrop']==['ALL'] and host['PidsLimit']==128 and host['Memory']==host['MemorySwap']
        assert 'Binds' not in host and 'Volumes' not in payload
        assert not any(name in '\n'.join(payload['Env']) for name in ('APP_TOKEN','LLM_API_KEY','ATLAS_WORK_RUNNER_TOKEN'))
        runtime.clear()
        with pytest.raises(ValueError,match='runsc'):backend.create(value.model_dump())
    finally:backend.client.close()


def test_guest_transfer_roundtrip_and_link_rejection(tmp_path):
    from work_runner.transfer import restore,export
    inputs={'inputs/底稿.csv':b'amount\n9007199254740993.00000001\n','analyze.py':b'print("ok")'}
    restore(io.BytesIO(pack(inputs)),root=tmp_path)
    result=io.BytesIO();export(result,root=tmp_path)
    assert unpack(result.getvalue())==inputs
    (tmp_path/'escape').symlink_to('/etc/passwd')
    with pytest.raises(ValueError,match='regular file'):export(io.BytesIO(),root=tmp_path)


def test_docker_duplex_stream_stages_and_reads_tmpfs_without_archive_api(monkeypatch):
    import queue
    import socket
    import struct
    import tempfile
    import threading
    from pathlib import Path
    from types import SimpleNamespace
    from work_runner import docker_backend
    from work_runner.docker_backend import DockerBackend
    with tempfile.TemporaryDirectory(prefix='atlas-stream-') as temporary:
        socket_path=str(Path(temporary)/'docker.sock')
        # This executor forbids AF_UNIX. Adapt only the test socket factory to
        # a real loopback duplex connection; production remains Unix-only.
        server=socket.socket(socket.AF_INET,socket.SOCK_STREAM);server.bind(('127.0.0.1',0));server.listen()
        address=server.getsockname()
        class Connection:
            def __init__(self):self.channel=socket.socket(socket.AF_INET,socket.SOCK_STREAM)
            def __getattr__(self,key):return getattr(self.channel,key)
            def connect(self,path):
                assert path==socket_path
                return self.channel.connect(address)
            def __enter__(self):return self
            def __exit__(self,*args):self.channel.close()
        monkeypatch.setattr(docker_backend,'socket',SimpleNamespace(AF_UNIX=socket.AF_UNIX,SOCK_STREAM=socket.SOCK_STREAM,
            SHUT_WR=socket.SHUT_WR,socket=lambda *args:Connection()))
        modes=queue.Queue();received=[];errors=[]
        inputs=pack({'a.txt':b'exact input'});output=pack({'result.csv':b'amount\n118.1500\n'})
        def serve():
            try:
                for _ in range(2):
                    connection,_=server.accept()
                    with connection:
                        raw=b''
                        while b'\r\n\r\n' not in raw:raw+=connection.recv(4096)
                        header,body=raw.split(b'\r\n\r\n',1)
                        assert b'POST /v1.45/exec/transfer/start HTTP/1.1' in header
                        size=int(next(line.split(b':',1)[1] for line in header.split(b'\r\n') if line.lower().startswith(b'content-length:')))
                        while len(body)<size:body+=connection.recv(4096)
                        connection.sendall(b'HTTP/1.1 101 UPGRADED\r\nConnection: Upgrade\r\nUpgrade: tcp\r\n\r\n')
                        mode=modes.get(timeout=3)
                        if mode=='import':
                            data=bytearray(body[size:])
                            while chunk:=connection.recv(65536):data.extend(chunk)
                            received.append(bytes(data))
                        else:
                            frame=b'\x01\0\0\0'+struct.pack('>I',len(output))+output
                            for offset in range(0,len(frame),7):connection.sendall(frame[offset:offset+7])
            except Exception as error:errors.append(error)
            finally:server.close()
        worker=threading.Thread(target=serve,daemon=True);worker.start()
        def respond(request):
            if request.method=='POST':
                payload=json.loads(request.content);modes.put(payload['Cmd'][-1])
                assert '/archive' not in request.url.path and payload['User']=='10001:10001'
                return httpx.Response(201,json={'Id':'transfer'})
            return httpx.Response(200,json={'Running':False,'ExitCode':0})
        backend=DockerBackend('registry.test/workspace@sha256:'+'a'*64,socket_path);backend.client.close()
        backend.client=httpx.Client(transport=httpx.MockTransport(respond),base_url='http://docker/v1.45')
        try:
            backend.stage('container',inputs)
            assert backend.files('container')=={'result.csv':b'amount\n118.1500\n'}
            worker.join(timeout=3)
            assert not worker.is_alive() and not errors and received==[inputs]
        finally:backend.client.close()
