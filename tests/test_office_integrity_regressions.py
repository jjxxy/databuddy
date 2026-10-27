"""Red regressions from synthetic Office precision/cache-integrity counterexamples.

Each test has its own metadata database and files; no model or user profile.
"""
import hashlib
from decimal import Decimal
from xml.etree import ElementTree as ET
import zipfile

import pytest
from sqlalchemy import create_engine

from agent_service import office,store,workspace


@pytest.fixture
def isolated_office(monkeypatch,tmp_path):
    monkeypatch.delenv('ATLAS_LOCAL_WORKSPACE',raising=False)
    data=tmp_path/'isolated-office';data.mkdir()
    engine=create_engine('sqlite:///'+str(data/'metadata.sqlite'))
    store.metadata.create_all(engine)
    monkeypatch.setattr(store,'engine',engine)
    monkeypatch.setattr(store,'DATA',data)
    monkeypatch.setattr(office,'DATA',data)
    monkeypatch.setattr(workspace,'DATA',data)
    pid='p_'+store.uid()
    store.put('project',{'name':'Synthetic Office integrity audit'},pid)
    turn=store.put('turn',{'project_id':pid,'text':'请生成Excel和经营汇报文件','interaction_mode':'work',
        'state':'agent','created_at':store.now()},'t_'+store.uid())
    yield turn,data
    engine.dispose()


def numeric_spec(token):
    # Keep the original JSON numeric token, before any binary float conversion.
    return ('{"kind":"spreadsheet","title":"Synthetic exact values","sheets":['
            '{"name":"Values","columns":["amount"],"rows":[['+token+']]}]}')


def stored_value(path):
    ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with zipfile.ZipFile(path) as archive:
        sheet=ET.fromstring(archive.read('xl/worksheets/sheet1.xml'))
    cell=sheet.find('.//m:c[@r="A2"]',ns)
    if cell.get('t')=='inlineStr':return ''.join(node.text or '' for node in cell.findall('.//m:t',ns))
    value=cell.find('m:v',ns)
    return value.text if value is not None else None


@pytest.mark.parametrize('number',['9007199254740993','9007199254740.000001','1e309'])
def test_xlsx_numeric_atoms_are_exact_or_rejected_explicitly(isolated_office,number):
    turn,_=isolated_office
    workspace.write_file(turn['project_id'],'exact.json',numeric_spec(number))
    try:result=office.render(turn,'exact.json','exact-values')
    except ValueError as error:
        # Explicit refusal is safe; silently changing a value or blanking it is not.
        assert any(word in str(error) for word in ('精度','有限','范围','非有限','exact','finite'))
        return
    file,_=office.artifact(turn['project_id'],result['artifact_id'])
    value=stored_value(file)
    assert value is not None,'A finite original numeric atom became an empty Excel cell'
    assert Decimal(value)==Decimal(number),'Export silently changed the original JSON numeric atom'


@pytest.mark.parametrize('number',['NaN','Infinity','-Infinity'])
def test_nonfinite_numeric_constants_are_rejected(isolated_office,number):
    turn,_=isolated_office
    workspace.write_file(turn['project_id'],'nonfinite.json',numeric_spec(number))
    with pytest.raises(ValueError):office.render(turn,'nonfinite.json','nonfinite-values')


def rendered_document(turn):
    spec='{"kind":"document","title":"Synthetic source title","sections":[{"body":"Amount 118.15 CNY"}]}'
    workspace.write_file(turn['project_id'],'report.json',spec)
    result=office.render(turn,'report.json','synthetic-report')
    file,record=office.artifact(turn['project_id'],result['artifact_id'])
    return spec,result,file,record


def test_repeated_render_cannot_resign_registered_corrupt_bytes(isolated_office):
    turn,_=isolated_office
    _,original,file,record=rendered_document(turn)
    file.write_bytes(b'ISOLATED SYNTHETIC CORRUPT OFFICE BYTES')
    with pytest.raises(ValueError):office.artifact(turn['project_id'],original['artifact_id'])
    with pytest.raises(ValueError):office.render(turn,'report.json','synthetic-report')
    assert store.get('office_artifact',original['artifact_id'])['sha256']==record['sha256']


def test_valid_but_unregistered_office_file_is_not_adopted_as_spec_result(isolated_office):
    from docx import Document
    turn,_=isolated_office
    _,original,file,record=rendered_document(turn)
    # Derive the current renderer's real ID, so this regression survives a
    # deliberate renderer-version bump instead of planting an obsolete path.
    store.delete('office_artifact',original['artifact_id'])
    document=Document();document.add_paragraph('UNRELATED SYNTHETIC CONTENT');document.save(file)
    assert zipfile.is_zipfile(file)
    with pytest.raises(ValueError):office.render(turn,'report.json','synthetic-report')
    with pytest.raises(KeyError):store.get('office_artifact',original['artifact_id'])


def test_existing_registration_must_still_match_spec_identity(isolated_office):
    turn,_=isolated_office
    _,original,file,record=rendered_document(turn)
    store.patch('office_artifact',original['artifact_id'],source_sha256='f'*64)
    with pytest.raises(ValueError):office.render(turn,'report.json','synthetic-report')
    assert store.get('office_artifact',original['artifact_id'])['source_sha256']=='f'*64


def test_cached_structural_validation_is_not_inferred_from_a_matching_hash(isolated_office):
    turn,_=isolated_office
    _,original,file,record=rendered_document(turn)
    file.write_bytes(b'SYNTHETIC LEGACY INVALID OOXML')
    store.patch('office_artifact',original['artifact_id'],sha256=hashlib.sha256(file.read_bytes()).hexdigest(),bytes=file.stat().st_size)
    with pytest.raises(ValueError):office.render(turn,'report.json','synthetic-report')


def test_safe_numbers_booleans_blanks_and_formula_text_remain_compatible(isolated_office):
    import json
    from openpyxl import load_workbook
    turn,_=isolated_office
    spec={'kind':'spreadsheet','title':'Compatibility','sheets':[{'name':'Values',
        'columns':['amount','boolean','blank','formula_text','literal_number'],
        'rows':[[118.15,True,None,'=SUM(1,2)','9007199254740993']]}]}
    workspace.write_file(turn['project_id'],'compatible.json',json.dumps(spec))
    result=office.render(turn,'compatible.json','compatible');path,_=office.artifact(turn['project_id'],result['artifact_id'])
    book=load_workbook(path,read_only=True,data_only=False)
    try:
        row=list(book.active.iter_rows(min_row=2,max_row=2))[0]
        assert Decimal(str(row[0].value))==Decimal('118.15') and row[0].data_type=='n'
        assert row[1].value is True and row[2].value is None
        assert row[3].value=='=SUM(1,2)' and row[3].data_type=='s'
        assert row[4].value=='9007199254740993' and row[4].data_type=='s'
    finally:book.close()
    assert result['precision_policy']=='text_exact' and result['precision_notes']==[]
    assert result['visual_validation'] is False and result['formula_recalculation'] is False


def test_exact_numeric_text_is_disclosed_and_cached_receipts_stay_unchanged(isolated_office):
    turn,_=isolated_office
    workspace.write_file(turn['project_id'],'exact.json',numeric_spec('9007199254740.000001'))
    first=office.render(turn,'exact.json','precise');_,before=office.artifact(turn['project_id'],first['artifact_id'])
    second=office.render(turn,'exact.json','precise');_,after=office.artifact(turn['project_id'],first['artifact_id'])
    assert first==second and before==after and len(store.all_items('office_artifact'))==1
    assert first['precision_policy']=='text_exact'
    assert first['precision_notes']==[{'sheet':'Values','cells':['A2'],'reason':'exact_numeric_text'}]


def test_maximum_high_precision_rows_fit_the_manifest_without_truncating_notes(isolated_office):
    turn,_=isolated_office
    text=('{"kind":"spreadsheet","title":"Complete precision notes","sheets":['
          '{"name":"Values","columns":["amount"],"rows":['+
          ','.join(['[9007199254740993]']*5000)+']}]}')
    assert len(text.encode())<256*1024
    workspace.write_file(turn['project_id'],'many-exact.json',text)
    result=office.render(turn,'many-exact.json','many-exact')
    assert result['precision_policy']=='text_exact'
    assert result['precision_notes']==[{'sheet':'Values','reason':'exact_numeric_text',
        'cells':['A'+str(row) for row in range(2,5002)]}]
    file,_=office.artifact(turn['project_id'],result['artifact_id'])
    assert stored_value(file)=='9007199254740993'
    manifest=workspace.read_file(turn['project_id'],'deliverables/many-exact.manifest.json')
    assert len(manifest.encode())<=256*1024


def test_wide_finite_scientific_values_keep_manifest_within_declared_input_budget(isolated_office):
    import json
    turn,_=isolated_office
    row='['+','.join(['1e309']*8)+']'
    text=('{"kind":"spreadsheet","title":"Finite wide exact values","sheets":['
        '{"name":"Values","columns":'+json.dumps(['amount'+str(i) for i in range(8)])+
        ',"rows":['+','.join([row]*5000)+']}]}')
    assert len(text.encode())<=256*1024
    workspace.write_file(turn['project_id'],'wide-exact.json',text)
    result=office.render(turn,'wide-exact.json','wide-exact')
    manifest=workspace.read_file(turn['project_id'],'deliverables/wide-exact.manifest.json')
    assert len(manifest.encode())<=256*1024
    assert result['precision_policy']=='text_exact'
    assert result['precision_notes_format']=='grouped_cells_or_column_rows_v1'
    assert result['precision_notes']==[{'sheet':'Values','reason':'exact_numeric_text',
        'column_rows':{column:[[2,5001]] for column in 'ABCDEFGH'}}]
    file,_=office.artifact(turn['project_id'],result['artifact_id'])
    assert Decimal(stored_value(file))==Decimal('1e309')


def test_large_precision_coordinate_ranges_preserve_holes_and_every_cell():
    original=[column+str(row) for row in range(2,4202) for index,column in enumerate('ABCDEFGHIJKL')
        if (row+index)%4!=0]
    notes=[{'sheet':'Values','reason':'exact_numeric_text','cells':list(original)}]
    office._compact_precision_notes(notes)
    assert 'cells' not in notes[0]
    restored=[]
    for column,rows in notes[0]['column_rows'].items():
        for item in rows:
            selected=range(item[0],item[1]+1) if isinstance(item,list) else [item]
            restored.extend(column+str(row) for row in selected)
    assert len(restored)==len(original) and set(restored)==set(original)


def test_manifest_budget_is_checked_before_independent_receipt_or_publication(isolated_office,monkeypatch):
    turn,data=isolated_office
    workspace.write_file(turn['project_id'],'budget.json','{"kind":"document","title":"Synthetic manifest budget"}')
    monkeypatch.setattr(office,'MANIFEST_MAX_BYTES',128)
    with pytest.raises(ValueError,match='未发布'):
        office.render(turn,'budget.json','budget-proof')
    assert store.all_items('office_artifact')==[]
    assert not [file for file in (data/'deliverables').rglob('*.docx') if not file.name.startswith('.')]


def test_git_report_checkpoint_preserves_chinese_under_windows_default_encoding(isolated_office):
    import json,os,subprocess,sys
    from agent_service.process_control import executable
    turn,data=isolated_office
    repository=data/'chinese-git';repository.mkdir()
    script='''
import json,sys
from pathlib import Path
from agent_service.codegen import git
repo=Path(sys.argv[1])
git(repo,'init','-b','main')
git(repo,'config','user.name','Synthetic Office audit')
git(repo,'config','user.email','fixture@example.invalid')
message='保存可复用业务简报方案'
output=git(repo,'commit','--allow-empty','-m',message)
print(json.dumps({'utf8_mode':sys.flags.utf8_mode,'chinese_message_preserved':message in output}))
'''
    env={key:os.environ[key] for key in ('SystemRoot','SYSTEMROOT','WINDIR','COMSPEC','TEMP','TMP','PATHEXT') if key in os.environ}
    env.update(DATA_DIR=str(data),DATABASE_URL='sqlite:///'+str(data/'metadata.sqlite'),
        APP_TOKEN='isolated-office-test-only-token',LLM_BASE_URL='',LLM_MODEL='',
        ATLAS_GIT_BINARY=executable('git'),PYTHONIOENCODING='utf-8',
        GIT_CONFIG_NOSYSTEM='1',GIT_CONFIG_GLOBAL=str(data/'empty-global.gitconfig'))
    result=subprocess.run([sys.executable,'-X','utf8=0','-c',script,str(repository)],env=env,
        capture_output=True,text=True,encoding='utf-8',timeout=30)
    assert result.returncode==0,result.stderr
    assert json.loads(result.stdout)=={'utf8_mode':0,'chinese_message_preserved':True}


def test_concurrent_identical_render_has_one_registered_result(isolated_office):
    from concurrent.futures import ThreadPoolExecutor
    turn,_=isolated_office
    workspace.write_file(turn['project_id'],'concurrent.json','{"kind":"document","title":"Concurrent synthetic report"}')
    with ThreadPoolExecutor(max_workers=2) as pool:
        results=list(pool.map(lambda _:office.render(turn,'concurrent.json','concurrent'),range(2)))
    assert results[0]==results[1] and len(store.all_items('office_artifact'))==1
    office.artifact(turn['project_id'],results[0]['artifact_id'])


@pytest.mark.parametrize('boundary',['pending_saved','file_published'])
def test_killed_generic_office_render_recovers_only_its_independent_pending_hash(isolated_office,boundary):
    import os,subprocess,sys
    turn,data=isolated_office
    workspace.write_file(turn['project_id'],'crash.json','{"kind":"document","title":"Isolated crash recovery"}')
    script='''
import os,sys
from agent_service import office
from agent_service.store import get
identity,boundary=sys.argv[1:]
original=office.put
def fault(kind,value,item_id=None):
    if kind=='office_artifact' and value.get('state')=='committed' and boundary=='file_published':os._exit(73)
    result=original(kind,value,item_id)
    if kind=='office_artifact' and value.get('state')=='pending' and boundary=='pending_saved':os._exit(73)
    return result
office.put=fault
office.render(get('turn',identity),'crash.json','crash-proof')
raise SystemExit('fault boundary was not reached')
'''
    env={key:os.environ[key] for key in ('SystemRoot','SYSTEMROOT','WINDIR','COMSPEC','TEMP','TMP','PATHEXT') if key in os.environ}
    env.update(DATA_DIR=str(data),DATABASE_URL='sqlite:///'+str(data/'metadata.sqlite'),
        APP_TOKEN='isolated-office-test-only-token',LLM_BASE_URL='',LLM_MODEL='',PYTHONUTF8='1')
    killed=subprocess.run([sys.executable,'-X','utf8','-c',script,turn['id'],boundary],env=env,
        capture_output=True,text=True,encoding='utf-8',timeout=30)
    assert killed.returncode==73,killed.stderr
    pending=store.all_items('office_artifact');assert len(pending)==1 and pending[0]['state']=='pending'
    trusted=pending[0]['sha256']
    recovered=office.render(turn,'crash.json','crash-proof')
    file,record=office.artifact(turn['project_id'],recovered['artifact_id'])
    assert record['state']=='committed' and record['sha256']==trusted==hashlib.sha256(file.read_bytes()).hexdigest()
