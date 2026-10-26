"""Existing source ingestion retains enhanced XLSX semantics through adapters."""
import hashlib
from pathlib import Path
import shutil

import pytest
from openpyxl import Workbook

from agent_service.config import DATA, MAX_ROWS
from agent_service.connectors import SourceError, read_source
from agent_service.store import seal, uid


def source(tmp_path, rows, *, options=None, sheet=None, second_sheet=False):
    book=Workbook();ws=book.active;ws.title='数据'
    for row in rows:ws.append(row)
    if second_sheet:
        other=book.create_sheet('其他');other.append(['字段']);other.append(['值'])
    file=tmp_path/(uid()+'.xlsx');book.save(file)
    uploads=DATA/'uploads';uploads.mkdir(exist_ok=True)
    shutil.copy2(file,uploads/file.name)
    cfg={'stored_name':file.name}
    if options is not None:cfg['xlsx_options']=options
    if sheet is not None:cfg['sheet']=sheet
    return {'id':'s_'+uid(),'kind':'file','name':'source.xlsx','secret_config':seal(cfg),
            'content_hash':hashlib.sha256(file.read_bytes()).hexdigest()},file


def test_plain_excel_keeps_normalized_names_and_rich_field_context(tmp_path):
    s,file=source(tmp_path,[['客户','金额(万元)'],['000123',2]])
    rows,fields,evidence,truncated=read_source(s)
    assert not truncated and rows[0]['c_1']=='000123' and rows[0]['c_2']=='2'
    assert len(rows[0]['record_hash'])==64 and rows[0]['_source_index']==1
    assert fields[0]['name']=='c_1' and fields[0]['excel_column']=='A'
    assert fields[1]['unit']=='万元' and fields[1]['header_path']==['金额(万元)']
    assert evidence['table_id'] and evidence['sheet']=='数据' and evidence['sheets']==['数据']
    assert evidence['source_rows']==[2] and evidence['complete']
    assert evidence['provenance']['file_sha256']==hashlib.sha256(file.read_bytes()).hexdigest()
    assert evidence['provenance']['columns']['c_2']['excel_column']=='B'
    assert '_normalized_fields' not in evidence


def test_legacy_first_sheet_is_explicitly_disclosed(tmp_path):
    s,file=source(tmp_path,[['金额'],[10]],second_sheet=True)
    rows,fields,evidence,truncated=read_source(s,limit=500)
    assert evidence['sheet']=='数据'
    assert any(w['code']=='LEGACY_FIRST_SHEET_SELECTED' for w in evidence['warnings'])
    assert evidence['sheets']==['数据','其他'] and len(rows)==1


def test_nested_read_options_sheet_overrides_legacy_default(tmp_path):
    s,_=source(tmp_path,[['金额'],[10]],options={'sheet':'其他'},sheet='数据',second_sheet=True)
    rows,fields,evidence,truncated=read_source(s)
    assert evidence['sheet']=='其他' and fields[0]['label']=='字段' and rows[0]['c_1']=='值'


def test_connector_multilevel_header_paths_and_units_survive(tmp_path):
    book=Workbook();ws=book.active;ws.title='底稿'
    ws['A1']='区域';ws.merge_cells('A1:A2');ws['B1']='2026年';ws.merge_cells('B1:C1')
    ws['B2']='金额(万元)';ws['C2']='同比';ws.append(['华东',2,0.125]);ws['C3'].number_format='0.0%'
    file=tmp_path/(uid()+'.xlsx');book.save(file)
    uploads=DATA/'uploads';uploads.mkdir(exist_ok=True);shutil.copy2(file,uploads/file.name)
    s={'kind':'file','id':'s_'+uid(),'name':'bottom.xlsx','secret_config':seal({'stored_name':file.name,'xlsx_options':{'sheet':'底稿','header_rows':[1,2]}}),
       'content_hash':hashlib.sha256(file.read_bytes()).hexdigest()}
    rows,fields,evidence,_=read_source(s)
    assert fields[1]['label']=='2026年 / 金额(万元)' and fields[1]['unit']=='万元'
    assert fields[2]['stored_semantics']=='ratio' and rows[0]['c_3']=='0.125'
    assert fields[2]['header_coordinates'][0]['source_coordinate']=='B1'
    assert evidence['provenance']['source_rows']==[3]


def test_formula_probe_is_incomplete_and_full_read_fails(tmp_path):
    s,_=source(tmp_path,[['金额','两倍'],[10,'=A2*2']])
    rows,fields,evidence,truncated=read_source(s,limit=500)
    assert rows[0]['c_2'] is None and not evidence['complete']
    assert not evidence['executable_values_complete']
    assert evidence['special_cells']['2']['c_2']['cache_status']=='missing'
    assert evidence['provenance']['special_cells']['2']['c_2']['formula']['formula']=='=A2*2'
    with pytest.raises(SourceError) as exc:
        read_source(s)
    assert exc.value.code=='FORMULA_CACHE_MISSING'


def test_probe_metadata_matches_returned_rows_and_truncation_is_honest(tmp_path):
    s,_=source(tmp_path,[['金额'],[1],[2],[3],[4]])
    rows,fields,evidence,truncated=read_source(s,limit=2)
    assert truncated and len(rows)==2 and evidence['source_rows']==[2,3]
    assert evidence['provenance']['source_rows']==[2,3] and not evidence['complete']


def test_option_injection_cannot_disable_full_read_formula_gate(tmp_path):
    s,_=source(tmp_path,[['结果'],['=1+1']],options={'require_formula_cache':False})
    with pytest.raises(SourceError):read_source(s)


def test_csv_stays_on_original_normalization_path(tmp_path):
    uploads=DATA/'uploads';uploads.mkdir(exist_ok=True)
    name=uid()+'.csv';(uploads/name).write_text('客户,金额\n甲,10\n')
    s={'id':'s_'+uid(),'kind':'file','name':'source.csv','secret_config':seal({'stored_name':name}),
       'content_hash':hashlib.sha256((uploads/name).read_bytes()).hexdigest()}
    rows,fields,evidence,truncated=read_source(s)
    assert rows[0]['c_1']=='甲' and rows[0]['c_2']=='10'
    assert fields==[{'name':'c_1','label':'客户'},{'name':'c_2','label':'金额'}]
    assert evidence['parser']=='csv' and not truncated
