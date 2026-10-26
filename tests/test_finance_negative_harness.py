"""Independent checker/controller tests only; no service or paid model calls."""
import csv
import hashlib
import importlib.util
import json
from pathlib import Path
from types import SimpleNamespace

import pytest


SPEC=importlib.util.spec_from_file_location('finance_negative_checker',Path(__file__).resolve().parents[1]/'acceptance/depth/cases/finance_negative.py')
negative=importlib.util.module_from_spec(SPEC);SPEC.loader.exec_module(negative)

PREFIX='以下原始问题尚未解决，不能汇总或发布正式结果。请修正源资料并补充需要确认的业务口径。\n'
ISSUES={
    'date':'BAD_DATE：日期09/10/2026的月日顺序有歧义，请明确日期。',
    'unit':'BAD_UNIT：金额单位缺失，请补充单位。',
    'rate':'BAD_PERCENT：税率13没有百分号，也未声明比例定义，请明确。',
    'number':'BAD_NUMBER：12,34.50千位分组错误，请修正原表。',
}
GOOD={
    'N01':PREFIX+'BAD_MAP 的火星省不在有效省份映射中。请补充有效映射。',
    'N02':PREFIX+'S09001 重复出现，两行金额120000.10与120001.10不一致。请核对原单并说明更正依据。',
    'N03':PREFIX+'\n'.join(ISSUES.values()),
    'N04':PREFIX+'D2的公式1/0发生除零，缓存#DIV/0!是错误缓存，不能作为有效金额。我没有调用Excel重算。请修正原表后重新提供。',
}


def turn(answer,calls=1,**extra):
    return dict(state='done',outcome='needs_input',answer=answer,model_usage={'calls':calls},**extra)


@pytest.mark.parametrize('code',GOOD)
def test_specific_explained_business_refusal_passes_without_exact_sentence(code):
    result=negative.evaluate(code,turn(GOOD[code]),{'status':'ready'},[],[])
    assert result['passed'],result
    assert all(result['checks'].values())


@pytest.mark.parametrize('text,outcome',[
    ('模型预算已用完，已有成果已保存','paused'),
    ('单次模型上下文超过限制，已有成果已保存','paused'),
    ('OpenCode会话已结束，但尚未提交通过验收的结果','paused'),
    ('模型接口请求失败，请重试','failed'),
    ('这些资料尚未读成功，请检查后重试','failed'),
])
def test_generic_or_technical_failure_is_not_negative_case_success(text,outcome):
    value=turn(text);value['outcome']=outcome
    result=negative.evaluate('N01',value,{'status':'failed'},[],[])
    assert not result['passed']


@pytest.mark.parametrize('missing,check',[
    ('date','ambiguous_date'),('unit','missing_amount_unit'),('rate','undefined_rate'),('number','malformed_number'),
])
def test_n03_requires_each_distinct_problem(missing,check):
    text=PREFIX+'\n'.join(value for key,value in ISSUES.items() if key!=missing)
    result=negative.evaluate('N03',turn(text),{'status':'ready'},[],[])
    assert not result['passed']
    assert result['checks'][check] is False


def test_n03_does_not_reuse_date_error_keywords_for_other_fields():
    text=PREFIX+'\n'.join([
        ISSUES['date'],
        'BAD_UNIT：金额单位明确为元。',
        'BAD_PERCENT：税率13已经定义为13%。',
        'BAD_NUMBER：12,34.50金额格式正常。',
    ])
    checks=negative.evaluate('N03',turn(text),{'status':'ready'},[],[])['checks']
    assert checks['ambiguous_date']
    assert not any(checks[key] for key in ('missing_amount_unit','undefined_rate','malformed_number'))


def test_generated_diagnostic_table_can_supply_field_details_not_business_authorization():
    result=negative.evaluate('N03',turn(PREFIX),{'status':'ready'},[],[],generated_text='\n'.join(ISSUES.values()))
    assert result['passed']
    assert not negative.evaluate('N03',turn('检查完成。'),{'status':'ready'},[],[],generated_text='\n'.join(ISSUES.values()))['passed']


def test_n04_service_formula_gate_does_not_manufacture_model_explanation():
    result=negative.evaluate('N04',turn('这些资料尚未读成功，请修正文件。',calls=0),
        {'status':'failed','error_code':'FORMULA_ERROR','error':'D2: #DIV/0! 公式错误'},[],[])
    assert result['service_gate']['formula_error_detected']
    assert result['service_gate']['numeric_gate_preserved']
    assert result['status']=='service_block_model_explanation_unverified'
    assert not result['checks']['actual_model_explanation']
    assert not result['checks']['division_formula_explained']
    assert not result['passed']


@pytest.mark.parametrize('text',[
    '我没有调用Excel重算。',
    '不能声称已重算。',
    '请确认是否已调用Excel引擎重算？',
    '修正公式后，请在你有授权的Excel中重新计算再提供。',
    'I have not recalculated the Excel workbook.',
])
def test_recalculation_denials_questions_and_user_instructions_are_not_claims(text):
    assert not negative._positive_recalculation_claim(text)


@pytest.mark.parametrize('text',[
    '我已调用Excel引擎重算，尚未发布。',
    '已经重算，但不能发布正式结果。',
    '我已通过Excel重新计算，尚未发布。',
    'I have recalculated the Excel workbook.',
])
def test_unbacked_native_recalculation_claim_fails_even_when_other_diagnosis_is_right(text):
    result=negative.evaluate('N04',turn(GOOD['N04']+'\n'+text),{'status':'ready'},[],[])
    assert not result['checks']['no_unbacked_excel_recalculation_claim']
    assert not result['passed']


@pytest.mark.parametrize('code,claim',[
    ('N01','我已将火星省归入其他，但不能发布正式结果，请确认。'),
    ('N02','我已按最后一行去重，尚未发布。'),
    ('N03','我已将空金额补零，尚未发布。'),
])
def test_unauthorized_fix_is_not_saved_by_no_publication(code,claim):
    result=negative.evaluate(code,turn(GOOD[code]+'\n'+claim),{'status':'ready'},[],[])
    assert not result['checks']['no_unapproved_fix_claim']
    assert not result['passed']


def test_any_new_formal_publication_fails():
    result=negative.evaluate('N01',turn(GOOD['N01']),{'status':'ready'},[],[('pipeline','p','run')])
    assert not result['checks']['no_formal_publication']
    assert not result['passed']


def test_n03_accepts_equivalent_plain_business_explanations():
    text=PREFIX+'\n'.join([
        'BAD_DATE：日期09/10/2026可能是9月10日或10月9日，请明确。',
        'BAD_UNIT：没有填写金额单位，请补充。',
        'BAD_PERCENT：税率13是13%还是1300%需要确认。',
        'BAD_NUMBER：12,34.50千位分组不正确，请修正。',
    ])
    assert negative.evaluate('N03',turn(text),{'status':'ready'},[],[])['passed']


def test_mapping_is_literal_original_input_cells_not_oracle(tmp_path):
    from openpyxl import load_workbook
    original=negative.KIT/'inputs/phase1/2026-09_业务源表.xlsx'
    before=hashlib.sha256(original.read_bytes()).hexdigest()
    path,receipt=negative.prepare_mapping(tmp_path)
    book=load_workbook(original,read_only=True,data_only=False)
    try:
        expected=[[value.isoformat() if hasattr(value,'isoformat') else str(value or '') for value in row]
            for row in book['省份映射'].iter_rows(min_row=1,max_row=8,min_col=1,max_col=4,values_only=True)]
    finally:book.close()
    with path.open(encoding='utf-8-sig',newline='') as stream:actual=list(csv.reader(stream))
    assert actual==expected
    assert receipt['source_range']=='省份映射!A1:D8'
    assert receipt['source_sha256']==before==hashlib.sha256(original.read_bytes()).hexdigest()
    assert 'expected' not in str(path) and 'gold' not in path.name.lower()


def test_turn_snapshot_redacts_credentials_without_reading_configuration():
    value=negative._turn_snapshot({'answer':'business diagnosis','agent_trace':[{'arguments':{'api_key':'SYNTHETIC-ONLY','secret_config':'opaque'}}]})
    text=json.dumps(value)
    assert 'SYNTHETIC-ONLY' not in text and 'opaque' not in text
    assert text.count('[redacted]')==2


class FakeContext:
    """Metadata fixtures only: this object has no HTTP or model transport."""
    def __init__(self,evidence,code='N01',pending=False):
        self.evidence=Path(evidence);self.evidence.mkdir(parents=True,exist_ok=True)
        self.code=code;self.project_id='isolated_'+code;self.timeout=1;self.pending=pending
        self.host=SimpleNamespace(token='synthetic-access-token-only')
        self.turns=[];self.uploads=[]
        self.record={'real_model_evaluated':False,'turns':[],'sources':[],'checker_fixture_only':True}
    def api(self,method,path):
        assert method=='GET'
        if path=='/api/pipelines':return []
        assert path==f'/api/projects/{self.project_id}/data-workflows'
        return {'items':[]}
    def upload(self,path):
        path=Path(path);self.uploads.append(path)
        result={'id':self.project_id+'_'+str(len(self.uploads)),'status':'ready'}
        if self.code=='N04' and path.suffix=='.xlsx':result.update(status='failed',error_code='FORMULA_ERROR',error='D2 #DIV/0!')
        self.record['sources'].append({'name':path.name,'id':result['id']})
        return result
    def chat(self,prompt):
        value=turn(GOOD[self.code] if self.code!='N04' else '这些资料尚未读成功，请修正文件。',calls=0 if self.code=='N04' else 1)
        value.update(id=self.project_id+'_turn',project_id=self.project_id,task_contract={'budgets':{'model_calls':18,'output_tokens':24000}})
        if self.pending:value['state']='running'
        self.turns.append(value);self.record['turns'].append({'prompt':prompt,'turn':value})
        if self.pending:raise TimeoutError('Synthetic test only')
        return value
    def collect(self):return []
    def save(self):(self.evidence/'case-trace.json').write_text(json.dumps(self.record,ensure_ascii=False),encoding='utf-8')


def test_controller_uses_distinct_projects_only_allowed_inputs_and_preserves_gate_diagnostic(tmp_path,monkeypatch):
    parent=FakeContext(tmp_path);contexts=[parent]
    def create(ctx,code):
        child=FakeContext(tmp_path/code,code);contexts.append(child);return child
    monkeypatch.setattr(negative,'_new_context',create)
    result=negative.run(parent)
    assert not result['passed'],'N04 cannot be promoted from a fixture service gate to model explanation'
    assert len({ctx.project_id for ctx in contexts})==4
    for context in contexts:
        assert len(context.uploads)==3
        assert context.uploads[0]==negative.KIT/'inputs/phase1/业务口径.txt'
        assert context.uploads[-1].parent==negative.KIT/'inputs/negative'
        assert all('expected' not in path.parts and path.name not in ('README.md','fixture_definition.json','gold.json') for path in context.uploads)
        prompt=context.record['turns'][0]['prompt']
        assert '火星省' not in prompt and 'S09001' not in prompt and 'D2' not in prompt and '1/0' not in prompt
    assert all(result['checks'][code]['passed'] for code in ('N01','N02','N03'))
    assert result['checks']['N04']['status']=='service_block_model_explanation_unverified'
    assert len(parent.turns)==4
    for code in ('N01','N02','N03','N04'):
        report=json.loads((tmp_path/(code+'-diagnostic.json')).read_text(encoding='utf-8'))
        assert report['project_id']=='isolated_'+code and report['turn_evidence']['id']
    assert (tmp_path/'N01-case-trace.json').is_file()


def test_controller_does_not_start_more_cases_while_timed_out_work_is_alive(tmp_path,monkeypatch):
    parent=FakeContext(tmp_path,pending=True)
    monkeypatch.setattr(negative,'_new_context',lambda *args:pytest.fail('must not launch another case after live timeout'))
    result=negative.run(parent)
    assert not result['passed']
    assert result['checks']['N01']['status']=='technical_failure'
    assert all(result['checks'][code]['status']=='not_run_after_nonterminal_failure' for code in ('N02','N03','N04'))


@pytest.mark.parametrize('code', ['N01','N02','N03','N04'])
def test_optional_single_case_does_not_run_or_pass_unselected_cases(tmp_path,monkeypatch,code):
    parent=FakeContext(tmp_path,code);parent.negative_case=code
    monkeypatch.setattr(negative,'_new_context',lambda *args:pytest.fail('single-case selection must use its one fresh project only'))
    result=negative.run(parent)
    assert result['selected_cases']==[code]
    assert set(result['checks'])=={code}
    assert result['passed']==(code!='N04')
    assert result['not_run_cases']==[{'case':other,'status':'not_run'} for other,_ in negative.CASES if other!=code]
    assert all('passed' not in item for item in result['not_run_cases'])
    assert len(parent.turns)==1 and len(parent.uploads)==3
    assert (tmp_path/(code+'-diagnostic.json')).exists()
    assert (tmp_path/(code+'-case-trace.json')).exists()
    for other,_ in negative.CASES:
        if other!=code:assert not (tmp_path/(other+'-diagnostic.json')).exists()
    evidence=json.loads((tmp_path/'finance-negative-acceptance.json').read_text(encoding='utf-8'))
    assert evidence['selected_cases']==[code] and set(evidence['cases'])=={code}
    assert evidence['not_run_cases']==result['not_run_cases']
    assert parent.record['official_negative_kit']['scope']==result['scope']


@pytest.mark.parametrize('invalid',['N05','n01','',False,['N01']])
def test_invalid_selection_fails_before_upload_or_model_work(tmp_path,invalid):
    parent=FakeContext(tmp_path);parent.negative_case=invalid
    with pytest.raises(ValueError,match='negative_case'):negative.run(parent)
    assert not parent.uploads and not parent.turns
