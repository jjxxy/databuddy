"""Project-scoped, versioned artifacts with safe, honest preview levels.

Generated HTML/JavaScript is never executed. Table pages preserve exact strings;
Office extraction is explicitly a content view, not a faithful page renderer.
"""
from __future__ import annotations

import csv
from datetime import date,datetime,timezone
from decimal import Decimal,localcontext
import hashlib
import io
import json
import mimetypes
import os
from pathlib import Path
import re
import stat

from fastapi import Depends,HTTPException
from fastapi.responses import FileResponse
from sqlalchemy import select

from .config import DATA,MAX_ROWS
from .store import get,put,all_items,now
from .locks import FileLock
from . import workspace
from .platform_files import open_regular_read,assert_plain_path

MAX_FILE_BYTES=256*1024*1024
KINDS={'code','office','report','result','application','chart','agent','source'}


class ArtifactError(ValueError):pass
class ArtifactGone(ArtifactError):pass


def _live(pid):
    value=get('project',pid)
    if value.get('deleted_at') or value.get('archived'):raise ArtifactGone('项目已移到回收站，恢复后才能查看成果')
    return value


def _gone(value):return bool(value.get('deleted_at') or value.get('archived'))


def _plain(value):
    if isinstance(value,(Decimal,date,datetime)):return str(value)
    if isinstance(value,int) and not isinstance(value,bool) and abs(value)>9007199254740991:return str(value)
    if isinstance(value,list):return [_plain(x) for x in value]
    if isinstance(value,dict):return {k:_plain(v) for k,v in value.items()}
    return value


def _hash(raw):return hashlib.sha256(raw).hexdigest()


def _root():
    root=DATA/'workspace-artifacts'
    if root.is_symlink():raise ArtifactError('成果存储目录不能是链接')
    root.mkdir(mode=0o700,parents=True,exist_ok=True)
    assert_plain_path(root,missing_ok=False)
    return root


def _info(path,expected=None):
    descriptor=open_regular_read(path)
    with os.fdopen(descriptor,'rb') as stream:
        before=os.fstat(stream.fileno())
        if not stat.S_ISREG(before.st_mode) or before.st_nlink!=1 or before.st_size>MAX_FILE_BYTES:raise ArtifactError('成果不是普通独立文件或超过预览/下载预算')
        digest=hashlib.sha256();size=0
        while chunk:=stream.read(1024*1024):
            digest.update(chunk);size+=len(chunk)
            if size>MAX_FILE_BYTES:raise ArtifactError('成果文件超过读取预算')
        after=os.fstat(stream.fileno())
        if (before.st_size,before.st_mtime_ns,before.st_ctime_ns)!=(after.st_size,after.st_mtime_ns,after.st_ctime_ns):raise ArtifactError('成果在读取时变化，请重新打开')
    checksum=digest.hexdigest()
    if expected and checksum!=expected:raise ArtifactError('成果文件与已验证摘要不一致，未返回内容')
    return checksum,size


def _snapshot(path,expected,name):
    """Materialize a verified immutable copy, avoiding later workspace edits."""
    checksum,size=_info(path,expected);suffix=Path(name).suffix.lower()
    if not re.fullmatch(r'\.[a-z0-9]{1,8}',suffix):suffix='.bin'
    target=_root()/(checksum+suffix)
    from .agent_handoff import _atomic_bytes,_read_bytes
    with FileLock(str(DATA/('workspace-artifact-'+checksum+'.lock')),timeout=20):
        if not target.exists():_atomic_bytes(target,_read_bytes(path,MAX_FILE_BYTES))
        _info(target,checksum)
    return target,checksum,size


def _origin(value,pid):
    identity=value.get('origin_turn_id') or value.get('turn_id')
    if not identity:return None
    try:turn=get('turn',identity)
    except KeyError:return None
    if turn.get('project_id')!=pid or _gone(turn):return None
    from .workspace_lifecycle import visible_turn
    if not visible_turn(turn):return None
    return identity


def _preview_level(format):
    if format in ('docx','pptx'):return 'content_extraction'
    if format in ('xlsx','csv','jsonl','parquet','data'):return 'table'
    if format=='pdf':return 'browser_pdf'
    if format in ('png','jpg','jpeg','webp','tiff'):return 'image'
    if format in ('md','txt','py','sql','json','yaml','yml','html','htm'):return 'safe_text'
    return 'download_only'


def _register(pid,kind,identity,name,checksum,*,format=None,version=None,bytes=None,source_refs=None,updated_at=None,status='saved',private=None,downloads=None,origin=None,sha256_scope='file',semantic_status='not_confirmed',metric_status='not_confirmed'):
    fmt=format or Path(name).suffix.lower().lstrip('.') or 'bin'
    aid='wa_'+_hash(json.dumps([pid,kind,identity,fmt,checksum],ensure_ascii=False,separators=(',',':')).encode())[:44]
    base=f'/api/projects/{pid}/artifacts/{aid}'
    record={'id':aid,'project_id':pid,'kind':kind,'resource_id':identity,'name':name,'format':fmt,'sha256':checksum,
        'sha256_scope':sha256_scope,'bytes':bytes,'version':version or checksum,'source_refs':source_refs or [],'updated_at':updated_at or now(),'status':status,
        'verified':False,'integrity_verified':False,'verification_scope':'stored_receipt; bytes_rechecked_on_open','semantic_status':semantic_status,'metric_status':metric_status,'financial_verified':False,'preview_level':_preview_level(fmt),
        'download_url':base+'/download','preview_url':base+'/preview','origin_turn_id':origin,'archived':False,'private':private or {}}
    record['downloads']=downloads or [{'format':fmt,'url':base+'/download'}]
    if downloads:record['downloads']=[{'format':x,'url':base+'/download?format='+x} for x in downloads]
    try:old=get('workspace_artifact',aid)
    except KeyError:old=None
    if old and _gone(old):return None
    put('workspace_artifact',record,aid)
    return _public(record)


def _public(record):return {k:v for k,v in record.items() if k not in ('project_id','private','resource_id')}


def _verified(record):return dict(_public(record),verified=True,integrity_verified=True,verification_scope='complete_selected_version_bytes_checked; no_financial_approval')


def _source_entry(pid,source_id,value):
    revision=value.get('revision',1)
    version=value.get('revision_id') or str(revision)
    # Identical bytes can occur in two revisions, with different names or
    # business context. A later catalog read must not rewrite an older link.
    return _register(pid,'source',source_id+':'+version,value['name'],value['content_hash'],
        bytes=value.get('bytes'),version=version,updated_at=value.get('recorded_at') or value.get('created_at'),
        status='source_snapshot',private={'source_id':source_id,'revision_id':value.get('revision_id'),
                                        'source_revision':revision},
        source_refs=[{'source_id':source_id,'revision':revision}])


def source_artifact(pid,source_id,revision=None,expected_sha256=None):
    """Resolve one exact uploaded version without scanning unrelated artifacts."""
    _live(pid)
    source=get('source',source_id)
    if source.get('project_id')!=pid:raise KeyError(source_id)
    if _gone(source):raise ArtifactGone('来源已归档，不能继续查看或下载')
    if source.get('kind')!='file':raise ArtifactError('只有上传文件支持原件预览')
    if revision is not None and (type(revision) is not int or revision<1):raise ArtifactError('来源版本编号无效')
    if expected_sha256 is not None and (not isinstance(expected_sha256,str) or not re.fullmatch(r'[a-f0-9]{64}',expected_sha256)):
        raise ArtifactError('来源内容摘要无效')
    if expected_sha256 is not None and revision is None:raise ArtifactError('按报告查看原件时必须指定准确来源版本')
    selected=source
    if revision is not None and revision!=source.get('revision',1):
        versions=[value for value in all_items('source_version')
                  if value.get('source_id')==source_id and value.get('revision')==revision]
        if not versions:raise KeyError(source_id)
        if len(versions)!=1:raise ArtifactError('来源版本记录不唯一，未选择替代版本')
        selected=versions[0]
        if selected.get('project_id')!=pid:raise KeyError(source_id)
        if _gone(selected):raise ArtifactGone('所选来源版本已归档')
        selected=dict(selected,revision_id=selected['id'])
    if not selected.get('content_hash'):raise ArtifactError('来源缺少完整性记录，暂不能预览原件')
    if expected_sha256 is not None and selected['content_hash']!=expected_sha256:
        raise ArtifactError('所选来源版本与报告记录的原件摘要不一致，未打开其他版本')
    record=_source_entry(pid,source_id,selected)
    if record is None:raise ArtifactGone('该来源版本已归档')
    _,metadata=file(pid,record['id'])
    return metadata


def _run_project(run):
    # Legacy pipeline queue rows bind ownership through their server-owned parent.
    if run.get('kind')=='pipeline':
        try:return get('pipeline',run['target']).get('project_id')
        except KeyError:return None
    return run.get('project_id')


def _runs(pid):
    from .jobs import jobs,engine,render
    with engine.connect() as connection:
        # Close the read transaction before catalog registrations write receipts.
        # SQLite cannot commit a separate writer while this cursor is active.
        rows=connection.execute(select(jobs).order_by(jobs.c.created.desc())).mappings().all()
    for row in rows:
        run=render(row)
        if _run_project(run)==pid and run.get('validated') and not run.get('artifact_removed') and run['status'] in ('validated','succeeded'):yield run


def catalog(pid,kind='',offset=0,limit=40):
    _live(pid)
    if kind not in ('','all',*KINDS) or type(offset) is not int or offset<0 or type(limit) is not int or not 1<=limit<=100:raise ArtifactError('成果目录筛选或分页无效')
    from . import project_data,business_reports
    items=[];unavailable=[]
    def add(*args,**kwargs):
        item=_register(pid,*args,**kwargs)
        if item:items.append(item)
    for file in workspace.list_files(pid):
        try:
            from .agent_handoff import _atomic_bytes,_read_bytes
            original=workspace.path_for(pid,file['path']);content=_read_bytes(original,256*1024);content.decode('utf-8');digest=_hash(content)
            cached=_root()/(digest+Path(file['path']).suffix.lower());_atomic_bytes(cached,content)
            updated=datetime.fromtimestamp(original.stat().st_mtime,timezone.utc).isoformat()
            add('code',file['path'],file['path'],digest,bytes=len(content),status='draft',updated_at=updated,semantic_status='not_applicable',metric_status='not_applicable',private={'snapshot':str(cached.relative_to(DATA)),'path':file['path']})
        except (ValueError,OSError,KeyError):unavailable.append({'kind':'code','name':file['path'],'reason':'该文件暂不能完整读取'})
    for value in all_items('office_artifact'):
        if value.get('project_id')!=pid or _gone(value):continue
        add('office',value['id'],value['name'],value['sha256'],bytes=value.get('bytes'),updated_at=value.get('created_at'),origin=_origin(value,pid),private={'office_id':value['id']})
    for value in all_items('business_report'):
        if value.get('project_id')!=pid or _gone(value):continue
        add('report',value['id'],value.get('title','业务报告'),value['manifest_sha256'],format='md',version=value.get('recipe_commit') or value['recipe_sha256'],
            source_refs=[{'result_id':value['run_id'],'handoff_sha256':value.get('handoff_sha256')}],updated_at=value.get('created_at'),status='verified_result',
            private={'report_id':value['id']},downloads=['docx','pptx','md','json'],origin=_origin(value,pid),sha256_scope='frozen_report_manifest')
    workflows={v['id']:v for v in all_items('data_workflow') if v.get('project_id')==pid}
    pipelines={v['id']:v for v in all_items('pipeline') if v.get('project_id')==pid}
    for run in _runs(pid):
        if run['kind']=='data_workflow':
            parent=workflows.get(run['target'],{})
            if _gone(parent):continue
            try:
                _,meta=project_data.result_folder(pid,run['id'])
                status='adopted' if parent.get('published_run')==run['id'] else 'candidate' if parent.get('candidate_run')==run['id'] else 'needs_review' if run.get('external_agent_import') else 'historical'
                add('result',run['id'],parent.get('name') or '完整加工结果 '+run['id'][:12],meta['handoff_sha256'],format='data',version=meta['plan_sha256'],
                    source_refs=[{'source_id':sid,'revision':rev} for sid,rev in meta.get('source_revisions',{}).items()],updated_at=meta.get('created_at'),status=status,
                    private={'run_id':run['id']},downloads=['xlsx','csv','handoff'],sha256_scope='complete_result_manifest',semantic_status='requires_review' if meta.get('requires_semantic_review') or run.get('external_agent_import') else 'declared_by_workflow')
            except (ValueError,OSError,KeyError):unavailable.append({'kind':'result','id':run['id'],'reason':'完整结果回执暂不可读取'})
        elif run['kind']=='pipeline':
            parent=pipelines.get(run['target'])
            if not parent or _gone(parent) or run['id'] not in (parent.get('published_run'),parent.get('candidate_run')):continue
            try:
                from .releases import manifest_for
                _,manifest=manifest_for(run)
                from .metric_registry import status as metric_status
                confirmed=metric_status(parent,run,manifest)
                options={'format':'data','version':run['commit'],'source_refs':manifest.get('inputs',[]),'updated_at':run.get('finished_at') or run.get('created_at'),
                    'status':'adopted' if parent.get('published_run')==run['id'] else 'candidate','private':{'run_id':run['id'],'pipeline_id':parent['id']},'downloads':['xlsx','parquet'],'sha256_scope':'complete_parquet','metric_status':confirmed,'semantic_status':'owner_confirmed_version' if confirmed=='confirmed' else 'requires_review'}
                add('application',run['id'],parent.get('name','业务应用')+' · 完整结果',manifest['parquet_hash'],**options)
                application=manifest.get('application') or {}
                if application.get('chart')=='bar' and manifest.get('metric_aggregation')=='sum_decimal_38_12' and manifest.get('spec',{}).get('group_by') and manifest.get('spec',{}).get('amount_field'):
                    add('chart',run['id'],parent.get('name','业务应用')+' · 图表',manifest['parquet_hash'],**options)
            except (ValueError,OSError,KeyError):unavailable.append({'kind':'application','id':run['id'],'reason':'应用完整结果回执暂不可读取'})
    existing={(item['sha256'],Path(item['name']).name) for item in items}
    for value in all_items('agent_task_artifact'):
        if value.get('project_id')!=pid or _gone(value) or (value['sha256'],value['name']) in existing:continue
        add('agent',value['id'],value['name'],value['sha256'],bytes=value['bytes'],updated_at=value.get('created_at'),origin=_origin(value,pid),
            status='needs_review' if value.get('external_agent_import') else 'saved',semantic_status='requires_review' if value.get('requires_semantic_review') else 'not_confirmed',
            private={'agent_artifact_id':value['id'],'turn_id':value['turn_id']},source_refs=[{'task_id':value['turn_id'],'handoff_sha256':value.get('external_handoff_manifest_sha256')}])
    for value in all_items('source'):
        if value.get('project_id')!=pid or value.get('kind')!='file' or _gone(value) or not value.get('content_hash'):continue
        item=_source_entry(pid,value['id'],value)
        if item:items.append(item)
    items.sort(key=lambda x:(x['updated_at'],x['id']),reverse=True)
    if kind not in ('','all'):items=[item for item in items if item['kind']==kind]
    else:
        def priority(item):
            if item['kind']=='agent':
                if item['format'] in ('docx','pptx','pdf','md'):return 0
                if item['format'] in ('xlsx','csv'):return 2
                return 4
            return {'report':0,'office':1,'result':2,'application':2,'chart':3,'code':4,'source':5}[item['kind']]
        items.sort(key=priority)  # Stable: newest first within each business group.
    page=items[offset:offset+limit];end=offset+len(page)
    return {'items':page,'total':len(items),'offset':offset,'limit':limit,'next_offset':end if end<len(items) else None,
            'complete_catalog_available':True,'page_is_complete_catalog':offset==0 and end==len(items),'unavailable':unavailable,'unavailable_count':len(unavailable)}


def _record(pid,aid):
    _live(pid)
    if not isinstance(aid,str) or not re.fullmatch(r'wa_[a-f0-9]{44}',aid):raise KeyError(aid)
    record=get('workspace_artifact',aid)
    if record['project_id']!=pid:raise KeyError(aid)
    if _gone(record):raise ArtifactGone('该成果已归档')
    private=record['private'];kind=record['kind']
    if kind=='source':
        value=get('source',private['source_id'])
        if value.get('project_id')!=pid or _gone(value):raise ArtifactGone('来源已归档，不能继续查看或下载')
    elif kind in ('office','report'):
        value=get('office_artifact' if kind=='office' else 'business_report',record['resource_id'])
        if value.get('project_id')!=pid or _gone(value):raise ArtifactGone('该成果已归档或不可用')
        expected=value['sha256' if kind=='office' else 'manifest_sha256']
        if expected!=record['sha256']:raise ArtifactError('成果版本已变化，请从目录重新选择')
    elif kind in ('result','application','chart'):
        run=get('run',private['run_id'])
        if _run_project(run)!=pid or run.get('artifact_removed') or not run.get('validated') or run['status'] not in ('validated','succeeded'):raise ArtifactGone('该结果版本不可用')
        parent_kind='data_workflow' if kind=='result' else 'pipeline'
        try:parent=get(parent_kind,run['target'])
        except KeyError:parent={}
        if _gone(parent):raise ArtifactGone('所属数据应用已归档')
        if kind=='result':
            from .project_data import result_folder
            _,metadata=result_folder(pid,private['run_id'])
            if metadata['handoff_sha256']!=record['sha256']:raise ArtifactError('结果版本与所选成果摘要不一致')
    elif kind=='agent':
        value=get('agent_task_artifact',private['agent_artifact_id'])
        if value.get('project_id')!=pid or _gone(value):raise ArtifactGone('协作成果已归档')
        if value.get('sha256')!=record['sha256'] or value.get('turn_id')!=private['turn_id']:raise ArtifactError('协作成果版本与所选任务回执不一致')
    return record


def _pipeline(pid,record):
    from .releases import manifest_for,verify_semantics
    run=get('run',record['private']['run_id']);pipeline=get('pipeline',record['private']['pipeline_id'])
    folder,manifest=manifest_for(run);path=folder/'business_objects.parquet'
    _info(path,manifest['parquet_hash']);verify_semantics(pipeline,run,manifest,folder)
    if manifest['parquet_hash']!=record['sha256']:raise ArtifactError('应用结果与已选择版本不一致')
    return path,manifest


def _pipeline_xlsx(pid,record):
    import duckdb
    from .result_delivery import export_xlsx,durable_cached_artifact
    parquet,manifest=_pipeline(pid,record)
    target=_root()/(record['id']+'.xlsx');receipt_id=record['id']+'_xlsx'
    with FileLock(str(DATA/(receipt_id+'.lock')),timeout=30):
        try:receipt=get('workspace_artifact_export',receipt_id)
        except KeyError:receipt={}
        if receipt and (receipt.get('project_id')!=pid or receipt.get('artifact_id')!=record['id'] or receipt.get('source_sha256')!=record['sha256']):
            raise ArtifactError('导出回执与冻结来源版本不一致')
        def build(temporary):
            labels={field['name']:field['label'] for field in manifest['fields']}
            with duckdb.connect() as db:
                db.execute("set memory_limit='256MB'");db.execute('set threads=1')
                relation=db.read_parquet(str(parquet));cursor=relation.execute();names=[item[0] for item in cursor.description]
                keep=[i for i,name in enumerate(names) if not name.startswith('_dlt')]
                schema=[{'name':names[i],'label':labels.get(names[i],names[i]),'type':str(cursor.description[i][1])} for i in keep]
                def rows():
                    while batch:=cursor.fetchmany(1000):
                        for row in batch:yield {names[i]:row[i] for i in keep}
                export_xlsx(rows(),schema,temporary,precision_policy='text_exact',context={'run_id':manifest['run_id'],'commit':manifest['commit']},provenance={'parquet_sha256':manifest['parquet_hash']},controls=[{'kind':'row_count','expected':manifest['records']}])
        def save_receipt(checksum,state):
            put('workspace_artifact_export',{'project_id':pid,'artifact_id':record['id'],'source_sha256':record['sha256'],
                'sha256':checksum if state=='committed' else None,'pending_sha256':checksum if state=='pending' else None,
                'state':state,'run_id':manifest['run_id']},receipt_id)
        durable_cached_artifact(target,expected_sha256=receipt.get('sha256'),pending_sha256=receipt.get('pending_sha256'),build=build,
            save_pending=lambda checksum:save_receipt(checksum,'pending'),save_committed=lambda checksum:save_receipt(checksum,'committed'))
        _info(target,get('workspace_artifact_export',receipt_id)['sha256'])
        return target


def file(pid,aid,format=None):
    record=_record(pid,aid);kind=record['kind'];private=record['private']
    allowed={item['format'] for item in record['downloads']}
    format=format or record['downloads'][0]['format']
    if format not in allowed:raise ArtifactError('该成果没有此下载格式')
    expected=record['sha256'] if record['sha256_scope']=='file' else None
    name=Path(record['name']).name
    if kind=='code':path=DATA/private['snapshot']
    elif kind=='office':
        from .office import artifact
        path,metadata=artifact(pid,private['office_id']);expected=metadata['sha256']
    elif kind=='report':
        from .business_reports import report_file
        selected={'md':'markdown','json':'manifest'}.get(format,format);path=report_file(pid,private['report_id'],selected)
        report=get('business_report',private['report_id'])
        expected=report['markdown_sha256'] if selected=='markdown' else report['manifest_sha256'] if selected=='manifest' else report['office_exports'][selected]
        name=record['name']+'.'+format
    elif kind=='result':
        from .project_data import result_download
        path=result_download(pid,private['run_id'],format,expected_handoff_sha256=record['sha256']);name=record['name']+{'xlsx':'.xlsx','csv':'.csv','handoff':'.zip'}[format]
    elif kind in ('application','chart'):
        if format=='parquet':path,_=_pipeline(pid,record);expected=record['sha256'];name=record['name']+'.parquet'
        else:path=_pipeline_xlsx(pid,record);name=record['name']+'.xlsx'
    elif kind=='source':
        from .project_data import source_path
        source=get('source',private['source_id'])
        if private.get('revision_id'):
            historical=get('source_version',private['revision_id'])
            if (historical.get('source_id')!=source['id'] or historical.get('project_id')!=pid
                    or historical.get('content_hash')!=record['sha256']
                    or private.get('source_revision',historical.get('revision'))!=historical.get('revision')):
                raise ArtifactError('来源版本不一致')
            if _gone(historical):raise ArtifactGone('所选来源版本已归档')
            source=historical
        path=source_path(source)
    else:
        item=get('agent_task_artifact',private['agent_artifact_id']);path=DATA/item['relative_path'];expected=item['sha256']
    if not path.resolve().is_relative_to(DATA.resolve()) and kind!='code':raise ArtifactError('成果路径不在受管目录内')
    path,checksum,size=_snapshot(path,expected,name)
    return path,dict(_verified(record),name=name,download_sha256=checksum,download_bytes=size,download_format=format)


def _table(record,columns,rows,total,offset,limit,**extra):
    schema=columns
    if isinstance(columns,dict):columns=columns['columns']
    columns=[dict(field,label=field.get('label') or field.get('meaning') or field['name']) for field in columns]
    end=offset+len(rows)
    return {'artifact':_verified(record),'render':'table','columns':columns,'schema':schema,'rows':_plain(rows),'offset':offset,'limit':limit,'total':total,
            'next_offset':end if end<total else None,'unit':'rows','page_is_complete_result':offset==0 and end==total,'complete_result_available':True,
            'exact_numeric_representation':'decimal_strings; no_float_rounding','verified':True,**extra}


def _text(record,text,offset,text_limit,**extra):
    page=text[offset:offset+text_limit];end=offset+len(page)
    return {'artifact':_verified(record),'render':'markdown' if record['format']=='md' else 'text','text':page,'offset':offset,'total':len(text),'limit':text_limit,
            'next_offset':end if end<len(text) else None,'unit':'characters','verified':True,'active_content_executed':False,**extra}


def _document_blocks(document):
    blocks=[];tables={}
    for segment in document['segments']:
        text=document['text'][segment['start']:segment['end']]
        if segment.get('table_index') and segment.get('row_index') and segment.get('column_index'):
            key=(segment.get('part'),segment.get('slide_index'),segment.get('shape_index'),segment['table_index'])
            if key not in tables:
                block={'type':'table','columns':[],'rows':[],'cells':{},'locator':segment.get('locator')};tables[key]=block;blocks.append(block)
            block=tables[key];row=segment['row_index'];column=segment['column_index'];block['cells'].setdefault(row,{}).setdefault(column,[]).append(text)
        else:blocks.append({'type':'paragraph','text':text,'locator':segment.get('locator'),'slide_index':segment.get('slide_index'),'outline_level':segment.get('outline_level')})
    for block in tables.values():
        cells=block.pop('cells');maximum=max(column for row in cells.values() for column in row)
        if maximum>200:raise ArtifactError('Office表格列数超过预览预算')
        block['columns']=[{'name':'c_'+str(i),'label':'列'+str(i)} for i in range(1,maximum+1)]
        block['rows']=[{'c_'+str(i):'\n'.join(row.get(i,[])) for i in range(1,maximum+1)} for _,row in sorted(cells.items())]
    return blocks


def _chart(rows,columns,x,y,maximum=100):
    if isinstance(columns,dict):columns=columns['columns']
    fields={field['name']:field for field in columns}
    if x not in fields or y not in fields:raise ArtifactError('图表字段不存在')
    field=fields[y]
    if field.get('aggregation')!='additive' or field.get('value_format','native')!='native':raise ArtifactError('图表汇总必须使用已明确可加的原生数值，不能猜测比例可加性')
    values={}
    with localcontext() as context:
        context.prec=120
        for row in rows:
            if row.get(y) is None:raise ArtifactError('图表数值缺失，未补零')
            number=Decimal(str(row[y]))
            if not number.is_finite():raise ArtifactError('图表数值不是有限数')
            label=str(row.get(x,''));values[label]=values.get(label,Decimal(0))+number
    ordered=sorted(values.items(),key=lambda item:(-item[1],item[0]));shown=ordered[:maximum]
    return {'type':'bar','x':x,'y':y,'unit':field.get('unit'),'points':[{'label':label,'value':str(number)} for label,number in shown],
            'total_categories':len(ordered),'shown_categories':len(shown),'complete':len(shown)==len(ordered),'source_rows':len(rows),'computation':'full_result_exact_decimal_sum',
            'scope_note':'完整结果汇总；仅展示前'+str(maximum)+'类' if len(shown)<len(ordered) else '覆盖完整结果中的全部类别'}


def preview(pid,aid,offset=0,limit=50,text_limit=12000,sheet=None,chart_x=None,chart_y=None):
    if type(offset) is not int or offset<0 or type(limit) is not int or not 1<=limit<=100 or type(text_limit) is not int or not 1<=text_limit<=20000:raise ArtifactError('成果预览分页无效')
    record=_record(pid,aid);kind=record['kind'];private=record['private']
    if kind=='result':
        from .project_data import read_result,result_folder
        page=read_result(pid,private['run_id'],offset,limit)
        if page.get('manifest_sha256')!=record['sha256']:raise ArtifactError('所选结果与报告记录的版本不一致')
        response=_table(record,page['schema'],page['rows'],page['total'],offset,limit,
            context=page['context'],fidelity_note='完整已验证结果的分页视图；本页不等于全量结果')
        if chart_x or chart_y:
            if not chart_x or not chart_y:raise ArtifactError('请同时指定图表的分类与数值字段')
            from .result_delivery import read_all_handoff
            folder,metadata=result_folder(pid,private['run_id'])
            if metadata.get('handoff_sha256')!=record['sha256']:raise ArtifactError('所选结果与报告记录的版本不一致')
            full=read_all_handoff(folder/'handoff',record['sha256'],native_types=True,max_rows=MAX_ROWS)
            response['chart']=_chart(full['rows'],full['schema'],chart_x,chart_y)
        return response
    if kind in ('application','chart'):
        import duckdb
        path,manifest=_pipeline(pid,record)
        with duckdb.connect() as db:
            db.execute("set memory_limit='256MB'");db.execute('set threads=1');relation=db.read_parquet(str(path));total=relation.count('*').fetchone()[0]
            cursor=relation.limit(limit,offset).execute();names=[item[0] for item in cursor.description]
            rows=[dict(zip(names,row)) for row in cursor.fetchall()]
            fields={field['name']:field for field in manifest['fields']};columns=[{'name':name,'label':fields.get(name,{}).get('label',name),'data_type':str(cursor.description[i][1])} for i,name in enumerate(names) if not name.startswith('_dlt')]
            rows=[{field['name']:row[field['name']] for field in columns} for row in rows]
            response=_table(record,columns,rows,total,offset,limit,fidelity_note='直接读取当前固定版本的完整Parquet结果；预览按行分页')
            if kind=='chart':
                spec=manifest['spec'];x=spec['group_by'];y=spec['amount_field']
                if manifest.get('metric_aggregation')!='sum_decimal_38_12':raise ArtifactError('冻结版本没有明确的精确求和规则')
                if x not in names or y not in names:raise ArtifactError('已保存图表字段不存在')
                relation.create_view('artifact_chart_input')
                quote=lambda value:'"'+value.replace('"','""')+'"'
                expression="try_cast(replace(cast("+quote(y)+" as varchar), ',', '') as decimal(38,12))"
                invalid=db.execute('select count(*) from artifact_chart_input where '+quote(y)+' is null or '+expression+' is null').fetchone()[0]
                if invalid:raise ArtifactError('图表数值缺失或不能按冻结规则解析，未补零')
                grouped=db.execute('select cast('+quote(x)+' as varchar),sum('+expression+') from artifact_chart_input group by 1').fetchall()
                allrows=[{x:label,y:number} for label,number in grouped]
                chartfields=[{'name':name,'aggregation':'additive' if name==y else 'unknown','value_format':'native','unit':spec.get('currency','')} for name in names]
                response['chart']=_chart(allrows,chartfields,x,y)
                response['chart'].update(source_rows=total,computation='full_result_frozen_sum_decimal_38_12')
                response['chart']['semantic_note']='按冻结代码明确声明的求和规则计算；指标口径状态单独显示，不代表财务批准'
            return response
    format='md' if kind=='report' else record['format']
    path,download=file(pid,aid,'md' if kind=='report' else None)
    if format in ('md','txt','py','sql','json','yaml','yml','html','htm'):
        if path.stat().st_size>4*1024*1024:return {'artifact':_verified(record),'render':'download_only','fidelity_note':'文本超过安全预览预算，请下载完整文件','verified':True}
        try:text=path.read_text(encoding='utf-8-sig')
        except UnicodeError:return {'artifact':_verified(record),'render':'download_only','fidelity_note':'文本编码不是可安全预览的UTF-8，请下载原件','verified':True}
        if format=='json':
            try:rows=json.loads(text,parse_float=Decimal)
            except (ValueError,RecursionError):rows=None
            if isinstance(rows,list) and rows and all(isinstance(row,dict) for row in rows):
                names=list(dict.fromkeys(key for row in rows for key in row))
                if len(names)>200:raise ArtifactError('JSON表格列数超过预览预算')
                return _table(record,[{'name':name,'label':name} for name in names],rows[offset:offset+limit],len(rows),offset,limit,fidelity_note='JSON对象数组的完整数据分页，数值以精确字符串保留')
        return _text(record,text,offset,text_limit,fidelity_note='安全内容预览；未执行HTML、JavaScript或生成代码')
    if format=='csv':
        saved=[];total=0
        with path.open(encoding='utf-8-sig',newline='') as stream:
            reader=csv.reader(stream);labels=next(reader,None)
            if labels is None:return _table(record,[],[],0,offset,limit)
            width=len(labels)
            for row in reader:
                if offset<=total<offset+limit:saved.append(row)
                total+=1;width=max(width,len(row))
                if total>MAX_ROWS:raise ArtifactError('CSV预览超过完整行数预算，未截断为全量')
                if width>200:raise ArtifactError('CSV表格列数超过预览预算')
        if width>200:raise ArtifactError('CSV表格列数超过预览预算')
        columns=[{'name':'c_'+str(i+1),'label':labels[i] if i<len(labels) else '列'+str(i+1)} for i in range(width)]
        rows=[{field['name']:row[i] if i<len(row) else None for i,field in enumerate(columns)} for row in saved]
        return _table(record,columns,rows,total,offset,limit,fidelity_note='第一行作为展示表头；不猜测字段语义，不运行CSV中的公式样文本')
    if format=='xlsx':
        from .workbook_reader import inspect_workbook,_archive,_raw_rows
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
        metadata=inspect_workbook(path);names=[item['name'] for item in metadata['sheets']]
        selected=sheet or names[0]
        if selected not in names:raise ArtifactError('指定页签不存在')
        selected_metadata=next(item for item in metadata['sheets'] if item['name']==selected)
        # Streaming writers can omit <dimension>; other producers can leave it
        # stale. Use the bounded XML cell scan, never the optional dimension,
        # both for paging and the hard grid-size admission check.
        total=selected_metadata['max_row'];width=selected_metadata['max_column']
        if width>200 or total>MAX_ROWS:raise ArtifactError('Excel原始网格超过完整预览预算，请下载原件或明确数据区域')
        raw_page={}
        with _archive(path) as archive:
            for row_index,cells in _raw_rows(archive,selected_metadata['part']):
                if row_index>offset+limit:break
                if row_index>offset:raw_page.update(cells)
        book=load_workbook(path,read_only=True,data_only=False,keep_links=False)
        try:
            ws=book[selected]
            columns=[{'name':'c_'+str(i),'label':get_column_letter(i)} for i in range(1,width+1)]
            rows=[];formats={};raw_values={}
            if offset<total:
                for index,row in enumerate(ws.iter_rows(min_row=offset+1,max_row=min(total,offset+limit),max_col=width),offset+1):
                    values={}
                    for column,(field,cell) in enumerate(zip(columns,row),1):
                        coordinate=get_column_letter(column)+str(index);raw=raw_page.get(coordinate,{})
                        value=_plain(cell.value)
                        if raw.get('type','n')=='n' and raw.get('value') is not None:
                            raw_values[coordinate]=raw['value']
                            if not raw.get('formula_present') and not isinstance(cell.value,(date,datetime)):
                                value=raw['value']
                        values[field['name']]=value
                    rows.append(values)
                    for column,cell in enumerate(row,1):
                        if cell.value is not None and cell.number_format!='General':formats[get_column_letter(column)+str(index)]=cell.number_format
            return _table(record,columns,rows,total,offset,limit,sheet_names=names,current_sheet=selected,cell_formats=formats,raw_numeric_values=raw_values,
                fidelity_note='原始单元格网格；合并值仅保留锚点，不推断表头。公式显示原文，未重算或保证缓存新鲜度',formula_recalculation=False,business_semantics_approved=False)
        finally:book.close()
    if format in ('docx','pptx'):
        from .document_reader import read_document
        document=read_document(path,max_text_bytes=2*1024*1024,max_blocks=10000);blocks=_document_blocks(document);page=blocks[offset:offset+limit];end=offset+len(page)
        return {'artifact':_verified(record),'render':'document','blocks':page,'offset':offset,'limit':limit,'total':len(blocks),'next_offset':end if end<len(blocks) else None,'unit':'blocks',
                'verified':True,'warnings':document['warnings'],'fidelity_note':'正文及表格内容提取视图，不还原Word分页、PPT排版、字体、图片或图表；请下载原件核对版式','layout_fidelity':False}
    if format=='pdf':
        with path.open('rb') as stream:header=stream.read(5)
        if header!=b'%PDF-':raise ArtifactError('文件不是有效PDF来源')
        return {'artifact':_verified(record),'render':'pdf','url':record['download_url'],'verified':True,'fidelity_note':'原始PDF供下载核对；未执行OCR或证明图形数值准确'}
    if format in ('png','jpg','jpeg','webp','tiff'):
        with path.open('rb') as stream:header=stream.read(16)
        valid=header.startswith(b'\x89PNG\r\n\x1a\n') or header.startswith(b'\xff\xd8\xff') or header.startswith((b'II*\x00',b'MM\x00*')) or header.startswith(b'RIFF') and header[8:12]==b'WEBP'
        if not valid:raise ArtifactError('文件不是受支持的图片来源')
        return {'artifact':_verified(record),'render':'image','url':record['download_url'],'verified':True,'fidelity_note':'原始图片预览；未执行OCR或把图像内容认定为财务事实'}
    return {'artifact':_verified(record),'render':'download_only','verified':True,'fidelity_note':'此格式没有已验证的安全预览器，可下载原件'}


def install(app,auth):
    def checked(function,*args,**kwargs):
        try:return function(*args,**kwargs)
        except ArtifactGone as error:raise HTTPException(410,str(error)) from None
        except KeyError:raise HTTPException(404,'成果不存在或不属于当前项目') from None
        except (ArtifactError,ValueError) as error:raise HTTPException(409,str(error)[:500]) from None
        except OSError:raise HTTPException(409,'成果文件暂不可读取，未返回未经验证内容') from None

    @app.get('/api/projects/{project_id}/artifacts',dependencies=[Depends(auth)])
    def artifacts(project_id:str,kind:str='',offset:int=0,limit:int=40):return checked(catalog,project_id,kind,offset,limit)

    @app.get('/api/projects/{project_id}/sources/{source_id}/artifact',dependencies=[Depends(auth)])
    def uploaded_source_artifact(project_id:str,source_id:str,revision:int|None=None,expected_sha256:str|None=None):
        return checked(source_artifact,project_id,source_id,revision,expected_sha256)

    @app.get('/api/projects/{project_id}/artifacts/{artifact_id}/preview',dependencies=[Depends(auth)])
    def artifact_preview(project_id:str,artifact_id:str,offset:int=0,limit:int=50,text_limit:int=12000,sheet:str|None=None,chart_x:str|None=None,chart_y:str|None=None):
        return checked(preview,project_id,artifact_id,offset,limit,text_limit,sheet,chart_x,chart_y)

    @app.get('/api/projects/{project_id}/artifacts/{artifact_id}/explanation',dependencies=[Depends(auth)])
    def artifact_explanation(project_id:str,artifact_id:str):
        from .business_explanation import explain_artifact
        return checked(explain_artifact,project_id,artifact_id)

    @app.get('/api/projects/{project_id}/artifacts/{artifact_id}/download',dependencies=[Depends(auth)])
    def artifact_download(project_id:str,artifact_id:str,format:str|None=None):
        path,metadata=checked(file,project_id,artifact_id,format)
        media=mimetypes.guess_type(metadata['name'])[0] or 'application/octet-stream'
        if media in ('text/html','image/svg+xml','application/javascript','text/javascript'):media='text/plain'
        return FileResponse(path,filename=metadata['name'],media_type=media,headers={'X-Artifact-SHA256':metadata['download_sha256'],'Cache-Control':'no-store','X-Content-Type-Options':'nosniff'})
