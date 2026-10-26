"""Deterministic editable Office deliverables; templates, never generated code."""
import hashlib,json,re,zipfile
from pathlib import Path
from typing import Literal
from pydantic import Field,ValidationError
from .models import StrictModel
from .config import DATA
from .store import get,put,uid,now
from . import workspace

class OfficeArgs(StrictModel):
    spec_path:str=Field(min_length=1,max_length=200)
    name:str=Field(min_length=1,max_length=91,
        pattern=r'^[A-Za-z0-9\u4e00-\u9fff][A-Za-z0-9\u4e00-\u9fff _-]{0,90}$',
        description='不含路径和扩展名的文件基名，例如净收入明细；扩展名由规格kind决定，不能传report.xlsx')
class Section(StrictModel):
    heading:str=Field(default='',max_length=300)
    body:str=Field(default='',max_length=20000)
    bullets:list[str]=Field(default_factory=list,max_length=30)
class Sheet(StrictModel):
    name:str=Field(min_length=1,max_length=31)
    columns:list[str]=Field(min_length=1,max_length=60)
    rows:list[list[str|int|float|bool|None]]=Field(default_factory=list,max_length=5000)
class Spec(StrictModel):
    kind:Literal['document','spreadsheet','presentation']
    title:str=Field(min_length=1,max_length=300)
    sections:list[Section]=Field(default_factory=list,max_length=30)
    sheets:list[Sheet]=Field(default_factory=list,max_length=10)

def render(turn,spec_path,name):
    from .intents import require
    require(turn,'write')
    try:OfficeArgs.model_validate({'spec_path':spec_path,'name':name})
    except ValidationError:
        raise ValueError('交付文件名须为不含路径和扩展名的基名，仅支持文字、数字、空格、横线，最多91字符') from None
    text=workspace.read_file(turn['project_id'],spec_path)
    try:spec=Spec.model_validate_json(text)
    except ValidationError as exc:
        # Report locations/types, never raw input values or the whole project
        # file. The planner already has the canonical office_spec_schema.
        issues=[{'path':'.'.join(map(str,error['loc'])),'type':error['type']}
                for error in exc.errors(include_input=False,include_context=False)[:8]]
        raise ValueError('Office规格不符合office_spec_schema，请修正原规格文件：'
                         +json.dumps(issues,ensure_ascii=False)) from None
    ext={'document':'docx','spreadsheet':'xlsx','presentation':'pptx'}[spec.kind]
    digest=hashlib.sha256((text+':'+name+':office-renderer-v1').encode()).hexdigest();aid='of_'+hashlib.sha256((turn['project_id']+digest).encode()).hexdigest()[:40]
    folder=DATA/'deliverables'/aid;folder.mkdir(parents=True,exist_ok=True);file=folder/(name+'.'+ext);tmp=folder/('temporary.'+ext)
    if not file.is_file():
        if ext=='docx':
            from docx import Document
            from docx.shared import Pt,Inches
            from docx.oxml.ns import qn
            doc=Document();style=doc.styles['Normal'];style.font.name='Calibri';style.font.size=Pt(11);style.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
            doc.sections[0].top_margin=Inches(.8);doc.sections[0].bottom_margin=Inches(.8);doc.add_heading(spec.title,0)
            for section in spec.sections:
                if section.heading:doc.add_heading(section.heading,1)
                for line in section.body.splitlines():doc.add_paragraph(line)
                for line in section.bullets:doc.add_paragraph(line,style='List Bullet')
            doc.save(tmp)
        elif ext=='xlsx':
            from openpyxl import Workbook
            from openpyxl.styles import Font,PatternFill
            if not spec.sheets:raise ValueError('工作簿至少需要一张表')
            book=Workbook();book.remove(book.active);names=set()
            for sheet in spec.sheets:
                if re.search(r'[\\/*?:\[\]]',sheet.name) or sheet.name.lower() in names:raise ValueError('表名无效或重复')
                names.add(sheet.name.lower());ws=book.create_sheet(sheet.name);ws.append(sheet.columns);ws.freeze_panes='A2'
                for row in sheet.rows:
                    if len(row)!=len(sheet.columns):raise ValueError('行数据与表头列数不一致')
                    ws.append(row)
                for row in ws:
                    for cell in row:
                        # User/imported strings must never become executable formulas.
                        if isinstance(cell.value,str):cell.data_type='s'
                        if cell.row==1:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='214F42')
                ws.auto_filter.ref=ws.dimensions
                for column in ws.columns:ws.column_dimensions[column[0].column_letter].width=min(42,max(14,max(len(str(c.value or '')) for c in column)+2))
            book.save(tmp)
        else:
            from pptx import Presentation
            from pptx.util import Inches,Pt
            if not spec.sections:raise ValueError('演示文稿至少需要一页内容')
            deck=Presentation();deck.slide_width=Inches(13.333);deck.slide_height=Inches(7.5)
            title=deck.slides.add_slide(deck.slide_layouts[0]);title.shapes.title.text=spec.title
            for section in spec.sections:
                slide=deck.slides.add_slide(deck.slide_layouts[1]);slide.shapes.title.text=section.heading or spec.title
                tf=slide.placeholders[1].text_frame;tf.clear()
                lines=section.body.splitlines()+section.bullets
                if len(lines)>12 or any(len(line)>180 for line in lines):raise ValueError('单页内容过多，请拆分为更短的幻灯片')
                for i,line in enumerate(lines):
                    p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.size=Pt(22)
            deck.save(tmp)
        with zipfile.ZipFile(tmp) as z:
            if z.testzip() is not None or '[Content_Types].xml' not in z.namelist():raise ValueError('交付文件格式验证失败')
        tmp.replace(file)
    sha=hashlib.sha256(file.read_bytes()).hexdigest()
    metadata={'project_id':turn['project_id'],'source_spec':spec_path,'source_sha256':hashlib.sha256(text.encode()).hexdigest(),'renderer':'office-v1','name':name+'.'+ext,'sha256':sha,'bytes':file.stat().st_size,'relative_path':str(file.relative_to(DATA)),'created_at':now()}
    put('office_artifact',metadata,aid)
    workspace.write_file(turn['project_id'],'deliverables/'+name+'.manifest.json',json.dumps({k:v for k,v in dict(metadata,id=aid).items() if k!='relative_path'},ensure_ascii=False,indent=2))
    return {'artifact_id':aid,'sha256':sha,'format':ext,'structural_validation':True,'visual_validation':False,'formula_recalculation':False,'download':{'label':name+'.'+ext,'url':'/api/projects/'+turn['project_id']+'/deliverables/'+aid}}

def artifact(pid,aid):
    value=get('office_artifact',aid)
    if value['project_id']!=pid:raise ValueError('交付物不属于该项目')
    file=(DATA/value['relative_path']).resolve()
    if not file.is_relative_to((DATA/'deliverables').resolve()) or hashlib.sha256(file.read_bytes()).hexdigest()!=value['sha256']:raise ValueError('交付物校验失败')
    return file,value
