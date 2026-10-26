"""Complete, verifiable result handoffs and precision-aware Excel delivery.

The handoff is authoritative and lossless. Excel is a constrained presentation
format: formulas are never evaluated and high-precision values require an
explicit text policy. No raw data is sent to another service by this module.
"""
from __future__ import annotations

from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation, localcontext
import hashlib
import csv
import json
import math
import os
from pathlib import Path
import re
import shutil
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from typing import Any, Callable, Iterable
from .platform_files import atomic_replace,assert_plain_path,open_regular_read,fsync_file


class DeliveryError(ValueError):
    pass


HANDOFF_FILES = {'rows.jsonl', 'schema.json', 'context.json', 'provenance.jsonl', 'stages.json'}
EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_COLUMNS = 16_384
EXCEL_MAX_TEXT = 32_767
NUMERIC_TYPES = {'number', 'numeric', 'decimal', 'integer', 'int', 'bigint', 'smallint',
                 'tinyint', 'hugeint', 'float', 'double', 'real', 'ubigint'}


def _json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(',', ':'), allow_nan=False)


def _pack(value: Any) -> Any:
    if isinstance(value, Decimal):
        if not value.is_finite():
            raise DeliveryError('Non-finite Decimal cannot be delivered')
        return {'$atlas': 'decimal', 'value': str(value)}
    if isinstance(value, datetime):
        return {'$atlas': 'datetime', 'value': value.isoformat()}
    if isinstance(value, date):
        return {'$atlas': 'date', 'value': value.isoformat()}
    if isinstance(value, dict):
        if any(not isinstance(k, str) for k in value):
            raise DeliveryError('JSON object keys must be strings')
        # All user objects are wrapped, so they cannot impersonate a type marker.
        return {'$atlas': 'object', 'items': [[k, _pack(v)] for k, v in value.items()]}
    if isinstance(value, (list, tuple)):
        return [_pack(v) for v in value]
    if isinstance(value, float) and not math.isfinite(value):
        raise DeliveryError('Non-finite float cannot be delivered')
    if isinstance(value, int) and not isinstance(value, bool) and abs(value) > 9_007_199_254_740_991:
        return {'$atlas': 'integer', 'value': str(value)}
    if value is None or isinstance(value, (str, bool, int, float)):
        return value
    raise DeliveryError('Unsupported handoff value type: ' + type(value).__name__)


def _unpack(value: Any, native_types: bool = True) -> Any:
    if isinstance(value, list):
        return [_unpack(v, native_types) for v in value]
    if isinstance(value, dict):
        kind = value.get('$atlas')
        if kind == 'object' and set(value) == {'$atlas', 'items'}:
            entries = value['items']
            if not isinstance(entries, list) or any(not isinstance(x, list) or len(x) != 2 or not isinstance(x[0], str) for x in entries):
                raise DeliveryError('Invalid handoff object encoding')
            if len({x[0] for x in entries}) != len(entries):
                raise DeliveryError('Duplicate handoff object keys')
            return {k: _unpack(v, native_types) for k, v in entries}
        if kind in ('decimal', 'integer', 'datetime', 'date') and set(value) == {'$atlas', 'value'}:
            raw = value['value']
            if not isinstance(raw, str):
                raise DeliveryError('Invalid typed handoff value')
            if kind == 'integer':
                if not re.fullmatch(r'-?(?:0|[1-9][0-9]*)', raw):
                    raise DeliveryError('Invalid integer handoff value')
                return int(raw) if native_types else raw
            if kind == 'decimal':
                try:
                    number = Decimal(raw)
                except InvalidOperation as exc:
                    raise DeliveryError('Invalid Decimal handoff value') from exc
                if not number.is_finite():
                    raise DeliveryError('Non-finite handoff value')
                return number if native_types else raw
            try:
                result = datetime.fromisoformat(raw) if kind == 'datetime' else date.fromisoformat(raw)
            except ValueError as exc:
                raise DeliveryError('Invalid date handoff value') from exc
            return result if native_types else raw
        raise DeliveryError('Unknown handoff object encoding')
    return value


def _schema(schema: Any) -> list[dict]:
    if isinstance(schema, dict):
        schema = schema.get('fields', schema.get('columns'))
    if not isinstance(schema, list) or not schema:
        raise DeliveryError('Schema must contain a nonempty list of field definitions')
    fields = []
    names = set()
    for field in schema:
        if isinstance(field, str):
            field = {'name': field, 'label': field}
        if not isinstance(field, dict) or not isinstance(field.get('name'), str) or not field['name']:
            raise DeliveryError('Each schema field needs a name')
        if field['name'] in names:
            raise DeliveryError('Duplicate schema field: ' + field['name'])
        names.add(field['name'])
        fields.append(dict(field))
    _pack(fields)
    return fields


def _row(row: Any, names: set[str]) -> dict:
    if not isinstance(row, dict) or set(row) != names:
        raise DeliveryError('Result row columns must exactly match the declared schema')
    _pack(row)
    return row


def _sha(path: Path) -> str:
    result = hashlib.sha256()
    try:
        with os.fdopen(open_regular_read(path),'rb') as stream:
            for chunk in iter(lambda: stream.read(1024 * 1024), b''):
                result.update(chunk)
    except OSError:
        raise DeliveryError('产物不可读或不是独立的普通文件') from None
    return result.hexdigest()


def _fsync_directory(path: Path) -> None:
    from .platform_files import fsync_directory
    fsync_directory(path)


def durable_cached_artifact(path: str | Path, *, expected_sha256: str | None,
                            pending_sha256: str | None, build: Callable[[Path], Any],
                            save_pending: Callable[[str], Any],
                            save_committed: Callable[[str], Any],
                            before_publish: Callable[[], Any] | None = None) -> Path:
    """Publish under the caller's lock with an independently persisted receipt.

    The pending receipt must be committed to the authoritative metadata store
    before rename. A retry can then complete either side of the rename without
    trusting an orphan file or a receipt kept beside the file itself. ``build``
    must validate its complete output before returning; callbacks bind the hash
    to the frozen source and export options in that metadata store.
    """
    destination = Path(path)
    temporary = destination.with_name('.' + destination.stem + '.pending' + destination.suffix)
    assert_plain_path(destination)
    assert_plain_path(temporary)

    def verified(candidate: Path, expected: str) -> None:
        if candidate.is_symlink() or not candidate.is_file() or _sha(candidate) != expected:
            raise DeliveryError('产物完整性校验失败，未返回或覆盖')

    if destination.is_symlink() or temporary.is_symlink():
        raise DeliveryError('产物路径不允许符号链接')
    if expected_sha256:
        verified(destination, expected_sha256)
        return destination
    if destination.exists():
        if not pending_sha256:
            raise DeliveryError('产物缺少独立完整性回执，未返回或覆盖')
        verified(destination, pending_sha256)
        _fsync_directory(destination.parent)
        save_committed(pending_sha256)
        return destination
    if pending_sha256 and temporary.exists():
        verified(temporary, pending_sha256)
        checksum = pending_sha256
    else:
        # An unpublished staging file without a receipt is disposable. Published
        # files above are never accepted or overwritten without a trusted hash.
        if temporary.exists():
            if not temporary.is_file():
                raise DeliveryError('产物暂存路径无效')
            temporary.unlink()
        destination.parent.mkdir(parents=True, exist_ok=True)
        build(temporary)
        if temporary.is_symlink() or not temporary.is_file():
            raise DeliveryError('产物构建未生成普通文件')
        temporary.chmod(0o600)
        fsync_file(temporary)
        _fsync_directory(destination.parent)
        checksum = _sha(temporary)
        save_pending(checksum)
    if before_publish:
        before_publish()
    # The lock serializes cooperating writers; this check also refuses an
    # unexpected destination created after receipt preparation.
    if destination.exists() or destination.is_symlink():
        raise DeliveryError('产物目标在发布期间出现，未覆盖')
    atomic_replace(temporary, destination)
    _fsync_directory(destination.parent)
    save_committed(checksum)
    return destination


def _write_json(path: Path, value: Any) -> None:
    with path.open('w', encoding='utf-8', newline='\n') as stream:
        stream.write(_json(_pack(value)) + '\n')
        stream.flush()
        os.fsync(stream.fileno())


def _read_json(path: Path, native_types: bool = True) -> Any:
    try:
        with os.fdopen(open_regular_read(path),'r',encoding='utf-8') as stream:
            return _unpack(json.load(stream), native_types)
    except (OSError, json.JSONDecodeError, TypeError, KeyError) as exc:
        raise DeliveryError('Invalid handoff metadata: ' + path.name) from exc


def _provenance_entries(provenance: Any) -> Iterable[Any]:
    if provenance is None:
        return []
    if isinstance(provenance, dict):
        return [provenance]
    if isinstance(provenance, (str, bytes)):
        raise DeliveryError('Provenance must be an object or a sequence of objects')
    return provenance


def write_handoff(folder: str | Path, rows: Iterable[dict], schema: Any,
                  context: Any = None, provenance: Any = None, stages: Any = None) -> dict:
    """Write an immutable, complete result package; never silently truncate rows.

    The destination must be absent or empty. Parent code should store the returned
    manifest_sha256 independently and pass it when verifying/consuming the result.
    Decimal/date types use an unambiguous tagged JSON representation.
    """
    fields = _schema(schema)
    names = {f['name'] for f in fields}
    destination = Path(folder)
    assert_plain_path(destination)
    if destination.is_symlink() or (destination.exists() and not destination.is_dir()):
        raise DeliveryError('Handoff destination is not an ordinary directory')
    if destination.exists() and any(destination.iterdir()):
        raise DeliveryError('Handoff artifacts are immutable; choose a new destination')
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = Path(tempfile.mkdtemp(prefix='.' + destination.name + '-', dir=destination.parent))
    try:
        count = 0
        with (temporary / 'rows.jsonl').open('w', encoding='utf-8', newline='\n') as out:
            for count, value in enumerate(rows, 1):
                _row(value, names)
                out.write(_json({'row_index': count, 'values': _pack(value)}) + '\n')
            out.flush()
            os.fsync(out.fileno())
        _write_json(temporary / 'schema.json', fields)
        _write_json(temporary / 'context.json', context if context is not None else {})
        _write_json(temporary / 'stages.json', stages if stages is not None else [])
        provenance_count = 0
        with (temporary / 'provenance.jsonl').open('w', encoding='utf-8', newline='\n') as out:
            for provenance_count, entry in enumerate(_provenance_entries(provenance), 1):
                out.write(_json({'entry_index': provenance_count, 'value': _pack(entry)}) + '\n')
            out.flush()
            os.fsync(out.fileno())
        manifest = {'version': 1, 'format': 'atlas-complete-result', 'complete': True,
                    'row_count': count, 'field_count': len(fields), 'provenance_count': provenance_count,
                    'created_at': datetime.now(timezone.utc).isoformat(),
                    'encoding': 'atlas-tagged-json-v1; decimals_are_exact; preview_is_not_a_handoff',
                    'files': {name: {'sha256': _sha(temporary / name), 'bytes': (temporary / name).stat().st_size}
                              for name in sorted(HANDOFF_FILES)}}
        with (temporary / 'manifest.json').open('w', encoding='utf-8', newline='\n') as out:
            out.write(_json(manifest) + '\n');out.flush();os.fsync(out.fileno())
        # A concurrent nonempty destination is not overwritten by os.replace.
        _fsync_directory(temporary)
        atomic_replace(temporary, destination)
        _fsync_directory(destination.parent)
        return dict(manifest, manifest_sha256=_sha(destination / 'manifest.json'),
                    path=str(destination))
    except BaseException:
        shutil.rmtree(temporary, ignore_errors=True)
        raise


def verify_handoff(folder: str | Path, expected_manifest_sha256: str | None = None) -> dict:
    root = Path(folder)
    assert_plain_path(root)
    if root.is_symlink() or not root.is_dir():
        raise DeliveryError('Handoff directory is unavailable')
    manifest_file = root / 'manifest.json'
    if manifest_file.is_symlink() or not manifest_file.is_file():
        raise DeliveryError('Handoff manifest is unavailable')
    actual_manifest_sha = _sha(manifest_file)
    if expected_manifest_sha256 is not None and actual_manifest_sha != expected_manifest_sha256:
        raise DeliveryError('Handoff manifest checksum mismatch')
    try:
        manifest = json.loads(manifest_file.read_text(encoding='utf-8'))
    except (OSError, json.JSONDecodeError) as exc:
        raise DeliveryError('Invalid handoff manifest') from exc
    if manifest.get('version') != 1 or manifest.get('format') != 'atlas-complete-result' or manifest.get('complete') is not True:
        raise DeliveryError('Unsupported or incomplete handoff manifest')
    if not isinstance(manifest.get('files'), dict) or set(manifest['files']) != HANDOFF_FILES:
        raise DeliveryError('Invalid handoff file set')
    for key in ('row_count', 'field_count', 'provenance_count'):
        if type(manifest.get(key)) is not int or manifest[key] < 0:
            raise DeliveryError('Invalid handoff count')
    for name, meta in manifest['files'].items():
        if not isinstance(meta, dict):
            raise DeliveryError('Invalid handoff file metadata')
        path = root / name
        if path.is_symlink() or not path.is_file() or path.stat().st_size != meta.get('bytes') or _sha(path) != meta.get('sha256'):
            raise DeliveryError('Handoff artifact checksum mismatch: ' + name)
    fields = _schema(_read_json(root / 'schema.json'))
    names = {f['name'] for f in fields}
    count = 0
    try:
        with (root / 'rows.jsonl').open(encoding='utf-8') as stream:
            for count, line in enumerate(stream, 1):
                value = json.loads(line)
                if type(value.get('row_index')) is not int or value['row_index'] != count:
                    raise DeliveryError('Handoff row order is incomplete')
                _row(_unpack(value.get('values')), names)
        provenance_count = 0
        with (root / 'provenance.jsonl').open(encoding='utf-8') as stream:
            for provenance_count, line in enumerate(stream, 1):
                value = json.loads(line)
                if type(value.get('entry_index')) is not int or value['entry_index'] != provenance_count:
                    raise DeliveryError('Handoff provenance order is incomplete')
                _unpack(value.get('value'))
    except (json.JSONDecodeError, TypeError) as exc:
        raise DeliveryError('Invalid handoff row or provenance encoding') from exc
    if (count != manifest.get('row_count') or len(fields) != manifest.get('field_count')
            or provenance_count != manifest.get('provenance_count')):
        raise DeliveryError('Handoff row/schema/provenance count mismatch')
    _read_json(root / 'context.json');_read_json(root / 'stages.json')
    return dict(manifest, manifest_sha256=actual_manifest_sha, path=str(root), verified=True)


def read_handoff(folder: str | Path, offset: int = 0, limit: int = 100,
                 expected_manifest_sha256: str | None = None, native_types: bool = False) -> dict:
    """Read a page after complete integrity validation; default Decimal values are strings.

    native_types=True returns Decimal/date Python objects for computation. HTTP/LLM
    adapters should use the default to avoid JSON encoders converting Decimal to float.
    """
    if type(offset) is not int or offset < 0 or type(limit) is not int or not 1 <= limit <= 1000:
        raise DeliveryError('Invalid handoff pagination')
    root = Path(folder)
    manifest = verify_handoff(root, expected_manifest_sha256)
    rows = []
    with (root / 'rows.jsonl').open(encoding='utf-8') as stream:
        for index, line in enumerate(stream):
            if index < offset:
                continue
            if index >= offset + limit:
                break
            rows.append(_unpack(json.loads(line)['values'], native_types))
    end = offset + len(rows)
    return {'rows': rows, 'schema': _read_json(root / 'schema.json', native_types),
            'context': _read_json(root / 'context.json', native_types),
            'stages': _read_json(root / 'stages.json', native_types),
            'offset': offset, 'limit': limit, 'total': manifest['row_count'],
            'next_offset': end if end < manifest['row_count'] else None,
            'complete_result_available': True, 'page_is_complete_result': offset == 0 and end == manifest['row_count'],
            'manifest_sha256': manifest['manifest_sha256'],
            'decimal_representation': 'native_decimal' if native_types else 'exact_decimal_strings'}


def _numeric(field: dict) -> bool:
    value = str(field.get('type', field.get('data_type', ''))).lower().split('(')[0]
    return field.get('numeric') is True or value in NUMERIC_TYPES


def _decimal(value: Any) -> Decimal:
    if isinstance(value, bool):
        raise DeliveryError('Boolean is not a numeric measure')
    try:
        result = value if isinstance(value, Decimal) else Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError) as exc:
        raise DeliveryError('Invalid numeric result') from exc
    if not result.is_finite():
        raise DeliveryError('Non-finite numeric result')
    return result


def _significant_digits(value: Decimal) -> int:
    # Decimal.normalize uses the ambient context and may round large inputs.
    digits = list(value.as_tuple().digits)
    while len(digits) > 1 and digits[-1] == 0:
        digits.pop()
    return len(digits)


def _canonical_decimal(value: Decimal) -> str:
    sign, digits, exponent = value.as_tuple()
    if not any(digits):
        return '0'
    digits = list(digits)
    while digits[-1] == 0:
        digits.pop();exponent += 1
    return ('-' if sign else '') + ''.join(map(str, digits)) + 'e' + str(exponent)


def _exact_add(left: Decimal, right: Decimal) -> Decimal:
    a, b = left.as_tuple(), right.as_tuple()
    integer_digits = max(len(a.digits) + a.exponent, len(b.digits) + b.exponent, 1)
    scale = max(-a.exponent, -b.exponent, 0)
    precision = integer_digits + scale + 2
    if precision > 10000:
        raise DeliveryError('Numeric reconciliation exceeds the exact-computation digit budget')
    with localcontext() as ctx:
        ctx.prec = precision
        return left + right


def _verify_excel_numeric_xml(path: Path, expected_sha256: str, expected_count: int) -> None:
    actual = hashlib.sha256();count = 0
    with zipfile.ZipFile(path) as archive:
        with archive.open('xl/worksheets/sheet1.xml') as stream:
            for _, element in ET.iterparse(stream, events=('end',)):
                if element.tag.rsplit('}', 1)[-1] != 'c':
                    continue
                if element.get('t', 'n') == 'n':
                    value = element.find('{*}v')
                    if value is not None and value.text is not None:
                        number = _decimal(value.text)
                        actual.update((element.get('r', '') + ':' + _canonical_decimal(number) + '\n').encode())
                        count += 1
                element.clear()
    if count != expected_count or actual.hexdigest() != expected_sha256:
        raise DeliveryError('Excel serialization changed numeric values; result was not published')


def _excel_number(value: Any, precision_policy: str) -> tuple[Any, str | None]:
    number = _decimal(value)
    absolute = number.copy_abs()
    within_range = number == 0 or Decimal('2.2250738585072E-308') <= absolute <= Decimal('9.99999999999999E307')
    safe = _significant_digits(number) <= 15 and within_range
    if safe:
        converted = int(number) if number == number.to_integral_value() else float(number)
        # openpyxl serializes numeric values using safe_string (currently 16g),
        # which may differ from Python's shortest float representation. Check
        # that exact serializer before deciding that a numeric cell is safe.
        from openpyxl.compat.strings import safe_string
        safe = Decimal(str(converted)) == number and Decimal(safe_string(converted)) == number
    if not safe:
        if precision_policy == 'reject':
            raise DeliveryError('Value exceeds Excel numeric precision/range; explicitly select text_exact or use the exact handoff')
        return str(number), 'exact_numeric_text'
    return converted, 'binary_float_input' if isinstance(value, float) else None


def _excel_text(value: Any) -> str:
    if isinstance(value, (dict, list, tuple)):
        value = _json(_pack(value))
    elif isinstance(value, (date, datetime)):
        value = value.isoformat()
    else:
        value = str(value)
    if len(value) > EXCEL_MAX_TEXT:
        raise DeliveryError('Value exceeds Excel cell text limit; use the complete handoff')
    # Invalid XML control characters must fail instead of being dropped.
    if any(ord(c) < 32 and c not in '\t\n\r' for c in value):
        raise DeliveryError('Value contains characters not representable in Excel XML')
    return value


def _validate_layout(layout: dict | None, fields: list[dict]) -> dict:
    value = dict(layout or {})
    if set(value) - {'sheet_name', 'start_cell', 'columns', 'header_fill', 'header_font', 'column_widths', 'number_formats'}:
        raise DeliveryError('Unsupported Excel layout option')
    value.setdefault('sheet_name', '结果')
    if not isinstance(value['sheet_name'], str) or not 1 <= len(value['sheet_name']) <= 31 or re.search(r'[\\/*?:\[\]]', value['sheet_name']):
        raise DeliveryError('Invalid Excel sheet name')
    value.setdefault('start_cell', 'A1')
    if not isinstance(value['start_cell'], str) or not re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]{0,6}', value['start_cell']):
        raise DeliveryError('Invalid Excel start cell')
    names = [f['name'] for f in fields]
    value.setdefault('columns', names)
    if (not isinstance(value['columns'], list) or len(value['columns']) != len(names)
            or set(value['columns']) != set(names)):
        raise DeliveryError('Excel layout must include every declared column exactly once')
    for key, default in [('header_fill', '214F42'), ('header_font', 'FFFFFF')]:
        value.setdefault(key, default)
        if not isinstance(value[key], str) or not re.fullmatch(r'[A-Fa-f0-9]{6}', value[key]):
            raise DeliveryError('Invalid Excel header color')
    for key in ('column_widths', 'number_formats'):
        value.setdefault(key, {})
        if not isinstance(value[key], dict) or set(value[key]) - set(names):
            raise DeliveryError('Unknown column in Excel formatting')
    for width in value['column_widths'].values():
        if isinstance(width, bool) or not isinstance(width, (int, float)) or not math.isfinite(width) or not 1 <= width <= 255:
            raise DeliveryError('Invalid Excel column width')
    for fmt in value['number_formats'].values():
        if not isinstance(fmt, str) or len(fmt) > 255:
            raise DeliveryError('Invalid Excel number format')
    return value


def _checks(controls: Any, count: int, sums: dict[str, Decimal]) -> list[dict]:
    if controls is None:
        return []
    if not isinstance(controls, list):
        raise DeliveryError('Controls must be a list')
    results = []
    for check in controls:
        if not isinstance(check, dict) or set(check) - {'kind', 'column', 'expected', 'tolerance'}:
            raise DeliveryError('Invalid result control')
        kind = check.get('kind', 'sum')
        if kind == 'row_count':
            actual = Decimal(count)
        elif kind == 'sum' and check.get('column') in sums:
            actual = sums[check['column']]
        else:
            raise DeliveryError('Unsupported or unknown numeric result control')
        if 'expected' not in check:
            raise DeliveryError('Result control needs an expected value')
        expected = _decimal(check['expected']);tolerance = _decimal(check.get('tolerance', '0'))
        if tolerance < 0:
            raise DeliveryError('Control tolerance cannot be negative')
        passed = _exact_add(actual, expected.copy_negate()).copy_abs() <= tolerance
        result = {'kind': kind, 'column': check.get('column'), 'expected': str(expected),
                  'actual': str(actual), 'tolerance': str(tolerance), 'passed': passed}
        results.append(result)
        if not passed:
            raise DeliveryError('Result control failed: ' + kind + ':' + str(check.get('column', 'rows')))
    return results


def export_csv(rows: Iterable[dict], schema: Any, path: str | Path) -> dict:
    """Export the complete table with declared labels and exact scalar values.

    Formula-like text is prefixed with an apostrophe, matching local exports;
    actual numeric cells, including negative Decimals, retain their value.
    The typed handoff preserves the original text and null distinctions.
    """
    fields = _schema(schema)
    names = {field['name'] for field in fields}
    labels = [field.get('label') or field['name'] for field in fields]
    escaped_text_cells = 0
    def literal(value):
        nonlocal escaped_text_cells
        if value.startswith(('\t','\r','\n')) or value.lstrip(' \t\r\n').startswith(('=','+','-','@')):
            escaped_text_cells += 1
            return "'" + value
        return value
    if any(not isinstance(label, str) for label in labels):
        raise DeliveryError('CSV列名必须为明确的业务文字')
    labels = [literal(label) for label in labels]
    if len(set(labels)) != len(labels):
        raise DeliveryError('导出列名重复，请明确各列的不同业务含义后重新导出')
    destination = Path(path)
    if destination.suffix.lower() != '.csv' or destination.exists() or destination.is_symlink():
        raise DeliveryError('CSV destination must be a new .csv file')
    destination.parent.mkdir(parents=True, exist_ok=True)
    def scalar(value, field):
        if value is None:return ''
        # The handoff may encode numbers as exact decimal text. Follow the
        # declared numeric schema before applying the text-only CSV guard.
        if _numeric(field):return format(_decimal(value), 'f')
        if isinstance(value, bool):return 'true' if value else 'false'
        if isinstance(value, Decimal):return format(value, 'f')
        if isinstance(value, (datetime, date)):return value.isoformat()
        if isinstance(value, str):return literal(value)
        if isinstance(value, (int, float)):return str(value)
        raise DeliveryError('CSV只支持标量列，复杂结构请下载完整交接包')
    expected = hashlib.sha256();count = 0
    with destination.open('x', encoding='utf-8-sig', newline='') as stream:
        writer = csv.writer(stream)
        writer.writerow(labels)
        for row in rows:
            _row(row, names)
            values = [scalar(row[field['name']], field) for field in fields]
            writer.writerow(values);count += 1
            expected.update((_json(values) + '\n').encode('utf-8'))
        stream.flush();os.fsync(stream.fileno())
    actual = hashlib.sha256();read_count = 0
    with destination.open('r', encoding='utf-8-sig', newline='') as stream:
        reader = csv.reader(stream)
        if next(reader, None) != labels:raise DeliveryError('CSV列名回读核验失败')
        for values in reader:
            actual.update((_json(values) + '\n').encode('utf-8'));read_count += 1
    if read_count != count or actual.digest() != expected.digest():
        raise DeliveryError('CSV全量回读核验失败，未交付结果')
    return {'records': count, 'columns': labels, 'sha256': _sha(destination), 'complete': True,
            'formula_like_text_escaped': escaped_text_cells, 'original_text_available_in': 'typed_handoff'}


def export_xlsx(rows: Iterable[dict], schema: Any, path: str | Path, *,
                precision_policy: str = 'reject', layout: dict | None = None,
                context: Any = None, provenance: Any = None, controls: Any = None) -> dict:
    """Export all rows with explicit precision and formula-injection protections.

    precision_policy='text_exact' writes only otherwise-unsafe numbers as text.
    It does not silently round them or pretend they can be summed exactly in Excel.
    Layout allows a sheet name, start cell, complete column reordering and styles.
    Arbitrary template fidelity and native formula recalculation are out of scope.
    """
    if precision_policy not in ('reject', 'text_exact'):
        raise DeliveryError('Unknown Excel precision policy')
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter
    fields = _schema(schema);layout = _validate_layout(layout, fields)
    field_map = {f['name']: f for f in fields};names = set(field_map)
    start_row, start_column = coordinate_to_tuple(layout['start_cell'])
    if start_row >= EXCEL_MAX_ROWS or start_column + len(fields) - 1 > EXCEL_MAX_COLUMNS:
        raise DeliveryError('Excel layout exceeds worksheet limits')
    destination = Path(path)
    if destination.suffix.lower() != '.xlsx' or destination.exists() or destination.is_symlink():
        raise DeliveryError('Excel destination must be a new .xlsx file')
    destination.parent.mkdir(parents=True, exist_ok=True)
    temp_dir = Path(tempfile.mkdtemp(prefix='.' + destination.stem + '-', dir=destination.parent))
    temporary = temp_dir / destination.name
    book = Workbook(write_only=True)
    ws = book.create_sheet(layout['sheet_name'])
    ws.freeze_panes = get_column_letter(start_column) + str(start_row + 1)
    for pos, name in enumerate(layout['columns'], start_column):
        ws.column_dimensions[get_column_letter(pos)].width = layout['column_widths'].get(name, 22)
    sum_control_columns = {c.get('column') for c in controls or [] if isinstance(c, dict) and c.get('kind', 'sum') == 'sum'}
    precision = [];sums = {f['name']: Decimal(0) for f in fields if _numeric(f)}
    source_digest = hashlib.sha256();numeric_digest = hashlib.sha256();numeric_count = 0;count = 0
    def cell(sheet, value, force_text=False):
        c = WriteOnlyCell(sheet, value=value)
        if force_text or isinstance(value, str):
            c.data_type = 's'
        return c
    try:
        for _ in range(start_row - 1):
            ws.append([])
        header = [None] * (start_column - 1)
        for name in layout['columns']:
            c = cell(ws, _excel_text((field_map[name].get('label') or name)))
            c.font = Font(bold=True, color=layout['header_font'])
            c.fill = PatternFill('solid', fgColor=layout['header_fill'])
            header.append(c)
        ws.append(header)
        for count, row in enumerate(rows, 1):
            _row(row, names)
            if start_row + count > EXCEL_MAX_ROWS:
                raise DeliveryError('Result exceeds Excel worksheet row limit; use the complete handoff')
            source_digest.update((_json(_pack(row)) + '\n').encode())
            output = [None] * (start_column - 1)
            for pos, name in enumerate(layout['columns'], start_column):
                value = row[name];field = field_map[name];reason = None
                if value is None and name in sum_control_columns:
                    raise DeliveryError('Cannot reconcile a sum with missing numeric values: ' + name)
                numeric = value is not None and (_numeric(field) or (isinstance(value, (Decimal, int, float)) and not isinstance(value, bool)))
                if numeric:
                    number = _decimal(value)
                    if name in sums:
                        sums[name] = _exact_add(sums[name], number)
                    exported, reason = _excel_number(value, precision_policy)
                    if isinstance(exported, str):
                        exported = _excel_text(exported)
                elif value is None or isinstance(value, bool):
                    exported = value
                else:
                    exported = _excel_text(value)
                c = cell(ws, exported)
                if numeric and not isinstance(exported, str):
                    coordinate = get_column_letter(pos) + str(start_row + count)
                    numeric_digest.update((coordinate + ':' + _canonical_decimal(number) + '\n').encode())
                    numeric_count += 1
                    c.number_format = (layout['number_formats'].get(name) or field.get('number_format') or '#,##0.############')
                if reason:
                    precision.append({'cell': get_column_letter(pos) + str(start_row + count),
                                      'column': name, 'reason': reason, 'exact_value': str(number)})
                output.append(c)
            ws.append(output)
        ws.auto_filter.ref = layout['start_cell'] + ':' + get_column_letter(start_column + len(fields) - 1) + str(start_row + count)
        control_results = _checks(controls, count, sums)
        used = {layout['sheet_name'].casefold()}
        def new_sheet(base):
            name = base;counter = 1
            while name.casefold() in used:
                counter += 1;name = base + str(counter)
            used.add(name.casefold());return book.create_sheet(name)
        semantic_sheet = new_sheet('语义说明')
        semantic_sheet.append([cell(semantic_sheet, x) for x in ['field', 'label', 'type', 'unit', 'grain', 'period', 'complete_metadata']])
        for f in fields:
            values = [f['name'], (f.get('label') or f['name']), f.get('type', f.get('data_type', '')),
                      f.get('unit', ''), f.get('grain', ''), f.get('period', ''), _json(_pack(f))]
            semantic_sheet.append([cell(semantic_sheet, _excel_text(v)) for v in values])
        provenance_sheet = new_sheet('来源与校验')
        provenance_sheet.append([cell(provenance_sheet, 'kind'), cell(provenance_sheet, 'details')])
        meta = [('export_policy', {'precision_policy': precision_policy, 'formula_recalculation': False,
                                  'complete_rows': count, 'layout_scope': 'explicit_columns_and_style_only'}),
                ('context', context if context is not None else {}),
                ('controls', control_results)]
        meta.extend(('precision', p) for p in precision)
        meta.extend(('provenance', entry) for entry in _provenance_entries(provenance))
        if len(meta) + 1 > EXCEL_MAX_ROWS:
            raise DeliveryError('Provenance exceeds Excel row limit; use the complete handoff')
        for kind, value in meta:
            provenance_sheet.append([cell(provenance_sheet, kind), cell(provenance_sheet, _excel_text(_json(_pack(value))))])
        book.save(temporary)
        _verify_excel_numeric_xml(temporary, numeric_digest.hexdigest(), numeric_count)
        # Publication is atomic and never exposes a partially written workbook.
        if destination.exists():
            raise DeliveryError('Excel destination appeared during export; refusing to overwrite it')
        fsync_file(temporary)
        atomic_replace(temporary, destination)
        _fsync_directory(destination.parent)
        result = {'version': 1, 'path': str(destination), 'name': destination.name,
                  'sha256': _sha(destination), 'bytes': destination.stat().st_size,
                  'row_count': count, 'field_count': len(fields), 'complete': True,
                  'schema': fields, 'context': context if context is not None else {},
                  'source_rows_sha256': source_digest.hexdigest(), 'layout': layout,
                  'precision_policy': precision_policy, 'precision_notes': precision,
                  'controls': control_results, 'formula_recalculation': False,
                  'visual_validation': False, 'provenance_preserved': True, 'numeric_xml_verified': True,
                  'scope': 'full_values_and_semantics; controlled_layout; not_arbitrary_template_fidelity'}
        return result
    except BaseException:
        # openpyxl write-only XML generators need explicit closure on errors.
        try:
            for sheet in book.worksheets:
                if not sheet.closed:
                    sheet.close()
        except Exception:
            pass
        raise
    finally:
        book.close()
        shutil.rmtree(temp_dir, ignore_errors=True)


def read_all_handoff(folder: str | Path, expected_manifest_sha256: str | None = None,
                     native_types: bool = False, *, max_rows: int = 1_000_000) -> dict:
    """Load a complete bounded result after one integrity validation.

    Internal computation/export should use this instead of repeatedly calling the
    paginated reader, which intentionally verifies the package for each public read.
    """
    if type(max_rows) is not int or max_rows < 0:
        raise DeliveryError('Invalid complete-result row budget')
    root = Path(folder)
    manifest = verify_handoff(root, expected_manifest_sha256)
    if manifest['row_count'] > max_rows:
        raise DeliveryError('Complete result exceeds the consumer row budget; no rows were omitted')
    with (root / 'rows.jsonl').open(encoding='utf-8') as stream:
        rows = [_unpack(json.loads(line)['values'], native_types) for line in stream]
    with (root / 'provenance.jsonl').open(encoding='utf-8') as stream:
        provenance = [_unpack(json.loads(line)['value'], native_types) for line in stream]
    return {'rows': rows, 'schema': _read_json(root / 'schema.json', native_types),
            'context': _read_json(root / 'context.json', native_types),
            'stages': _read_json(root / 'stages.json', native_types),
            'provenance': provenance, 'total': manifest['row_count'],
            'complete_result_available': True, 'page_is_complete_result': True,
            'manifest_sha256': manifest['manifest_sha256'],
            'decimal_representation': 'native_decimal' if native_types else 'exact_decimal_strings'}
