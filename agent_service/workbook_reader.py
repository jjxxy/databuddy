"""Coordinate-preserving XLSX table ingestion with explicit layout decisions.

Only source-declared values and layout are extracted. Formula caches are not
recalculated, indentation is not assumed to be a business ontology, and merged
detail measures are never forward-filled. Read-only openpyxl plus streaming XML
keeps worksheet cell objects out of memory; returned rows remain bounded by the
requested limit. All regular cell coordinates are reconstructable from the
source-row list and each field's Excel column. Exceptional metadata is bounded
and a full read fails rather than silently dropping it.
"""
from bisect import bisect_right
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation, localcontext
from functools import lru_cache
import hashlib
import json
from pathlib import Path, PurePosixPath
import re
import unicodedata
from xml.etree import ElementTree as ET
import zipfile

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.styles.numbers import is_date_format
from openpyxl.utils.cell import column_index_from_string, get_column_letter, range_boundaries
from openpyxl.utils.datetime import from_excel

NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
REL = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
PKG_REL = '{http://schemas.openxmlformats.org/package/2006/relationships}'
MAX_UNCOMPRESSED = 200 * 1024 * 1024
MAX_SCAN_CELLS = 5_000_000
MAX_SHEETS = 100
MAX_MERGES = 20_000
MAX_DISCOVERY_BLOCKS = 30
MAX_DISCOVERY_ROWS = 16
MAX_DISCOVERY_REGIONS = 100
LAYOUT_SELECTION_CODES = frozenset({'HEADER_SELECTION_REQUIRED','SHEET_SELECTION_REQUIRED','WORKBOOK_MERGE_AMBIGUOUS'})


class WorkbookReadError(ValueError):
    def __init__(self, message, code='WORKBOOK_INVALID', details=None):
        super().__init__(message)
        self.code = code
        self.details = details or {}


def _integer(value, label, minimum=1, maximum=1_048_576):
    if type(value) is not int or not minimum <= value <= maximum:
        raise WorkbookReadError(label + '超出支持范围', 'WORKBOOK_SELECTION_INVALID')
    return value


def _archive(path):
    try:
        archive = zipfile.ZipFile(path)
        entries = archive.infolist()
        if len(entries) > 10_000 or sum(i.file_size for i in entries) > MAX_UNCOMPRESSED:
            archive.close()
            raise WorkbookReadError('Excel 解压体积或文件数超过限制', 'WORKBOOK_LIMIT')
        if any(i.flag_bits & 1 for i in entries):
            archive.close()
            raise WorkbookReadError('不支持加密 Excel 压缩条目', 'WORKBOOK_ENCRYPTED')
        return archive
    except (zipfile.BadZipFile, OSError) as exc:
        raise WorkbookReadError('Excel 文件格式损坏或不可读') from exc


def _part(base, target):
    if target.startswith('/'):
        path = PurePosixPath(target.lstrip('/'))
    else:
        path = PurePosixPath(base).parent / target
    parts = []
    for part in path.parts:
        if part == '..':
            if not parts:
                raise WorkbookReadError('Excel 内部引用越界')
            parts.pop()
        elif part != '.':
            parts.append(part)
    return '/'.join(parts)


def _relationships(archive, part):
    p = PurePosixPath(part)
    rel_path = str(p.parent / '_rels' / (p.name + '.rels'))
    try:
        document = ET.fromstring(archive.read(rel_path))
    except KeyError:
        return {}
    return {r.get('Id'): _part(part, r.get('Target', '')) for r in document
            if r.tag == PKG_REL + 'Relationship' and r.get('TargetMode') != 'External'}


class _RegionSketch:
    """Retain bounded structural row profiles, never an entire worksheet grid."""
    def __init__(self):
        self.blocks = []
        self.active = None
        self.truncated = False

    def add(self, row, cells):
        if not cells:
            return
        if self.active is None or row > self.active['end'] + 1:
            if len(self.blocks) >= MAX_DISCOVERY_BLOCKS:
                self.truncated = True
                self.active = None
                return
            self.active = {'start': row, 'end': row, 'rows': 0, 'columns': set(),
                           'column_first': {}, 'column_last': {}, 'profiles': [],
                           'max_width': 0, 'possible_boundaries': []}
            self.blocks.append(self.active)
        b = self.active
        text_count = sum(c['kind'] in ('text', 'shared_text') for c in cells)
        numeric_count = sum(c['kind'] in ('number', 'year') for c in cells)
        profile = {'row': row, 'cells': cells[:32], 'nonempty': len(cells),
                   'text_count': text_count, 'numeric_count': numeric_count}
        first_rows = b['rows'] < 6
        width_growth = len(cells) > b['max_width']
        header_like = text_count >= 2 and numeric_count == 0
        if header_like and b['rows'] >= 2 and b.get('previous_numeric', 0) > 0:
            if len(b['possible_boundaries']) < 10:
                b['possible_boundaries'].append(row)
        if (first_rows or width_growth or header_like) and len(b['profiles']) < MAX_DISCOVERY_ROWS:
            b['profiles'].append(profile)
        # If late header width increases, retain it even after early notes exhausted the sample.
        elif width_growth and b['profiles']:
            b['profiles'][-1] = profile
        b['end'] = row
        b['rows'] += 1
        b['max_width'] = max(b['max_width'], len(cells))
        b['previous_numeric'] = numeric_count
        for cell in cells:
            col = cell['column']
            b['columns'].add(col)
            b['column_first'].setdefault(col, row)
            b['column_last'][col] = row


def _column_groups(columns):
    groups = []
    for column in sorted(columns):
        if not groups or column > groups[-1][-1] + 1:
            groups.append([column])
        else:
            groups[-1].append(column)
    return groups


def _region_proposals(sketch, merges, named_tables):
    candidates = []
    for table in named_tables:
        left, top, right, bottom = range_boundaries(table['range'])
        header_count = table['header_rows']
        candidates.append({'range': table['range'], 'table': table['name'],
                           'header_rows': list(range(top, top + header_count)),
                           'data_start_row': top + header_count,
                           'data_end_row': bottom - table['totals_rows'],
                           'confidence': 'source_declared_table', 'needs_confirmation': True,
                           'reasons': ['Excel 内置命名表提供了明确边界，仍需核对是否为本次目标表'],
                           'header_preview': [], 'context_rows': [],
                           'proposed_options': {'table': table['name']}})
    for block in sketch.blocks:
        # Entirely empty columns are evidence of separation, not proof of two business tables.
        groups = _column_groups(block['columns'])
        usable = [g for g in groups if len(g) >= 2]
        if not usable:
            usable = groups[:1]
        for group in usable:
            left, right = group[0], group[-1]
            profiles = []
            for p in block['profiles']:
                cells = [c for c in p['cells'] if left <= c['column'] <= right]
                if cells:
                    profiles.append(dict(p, cells=cells))
            strong = [p for p in profiles if
                      sum(c['kind'] in ('text', 'shared_text', 'year') for c in p['cells']) >= min(2, len(group))
                      and any(c['kind'] in ('text', 'shared_text') for c in p['cells'])]
            if not strong:
                continue
            header = strong[0]
            header_rows = [header['row']]
            horizontal = False
            for merge in merges:
                a, top, z, bottom = merge['bounds']
                if a <= right and z >= left and top <= header['row'] <= bottom:
                    horizontal |= z > a
                    if bottom - header['row'] < 6:
                        header_rows = list(range(header['row'], max(header_rows[-1], bottom) + 1))
            if horizontal:
                for following in profiles:
                    if following['row'] == header_rows[-1] + 1 and len(header_rows) < 6:
                        if not any(m['bounds'][0] <= right and m['bounds'][2] >= left
                                   and m['bounds'][1] <= header_rows[-1] <= m['bounds'][3]
                                   and m['bounds'][2] > m['bounds'][0] for m in merges):
                            break
                        textish = sum(c['kind'] in ('text', 'shared_text', 'year') for c in following['cells'])
                        if textish >= min(2, len(group)) and not any(c['kind'] == 'number' for c in following['cells']):
                            header_rows.append(following['row'])
            bottom = max(block['column_last'][c] for c in group)
            if bottom <= header_rows[-1]:
                continue
            reference = get_column_letter(left) + str(header_rows[0]) + ':' + get_column_letter(right) + str(bottom)
            # Native tables already provide stronger and exact boundaries.
            if any(reference == t['range'] for t in named_tables):
                continue
            reasons = ['连续非空行与实际有值列形成的布局候选，未推断业务含义']
            if block['start'] < header_rows[0]:
                reasons.append('前置标题/说明行被排除出候选表头，保留其行号供核对')
            if len(usable) > 1:
                reasons.append('全空列分隔出多个并排区域，需确认是否独立表格')
            if len(sketch.blocks) > 1:
                reasons.append('空行分隔出多个纵向区域，未自动拼接')
            if len(header_rows) > 1:
                reasons.append('合并范围及下层文字表头支持多级表头候选')
            boundary_rows = [r for r in block['possible_boundaries'] if header_rows[-1] < r <= bottom]
            if boundary_rows:
                reasons.append('区域内部有再次出现的文字表头候选，可能是分页重复表头或另一张表')
            preview = [{'coordinate': get_column_letter(c['column']) + str(p['row']),
                        **({'text': c['text']} if 'text' in c else {'shared_string_index': c['shared_string_index']})}
                       for p in profiles if p['row'] in header_rows for c in p['cells']
                       if c['kind'] in ('text', 'shared_text')][:12]
            candidates.append({'range': reference, 'table': None, 'header_rows': header_rows,
                               'data_start_row': header_rows[-1] + 1, 'data_end_row': bottom,
                               'confidence': 'layout_ambiguous' if len(usable) > 1 or boundary_rows else 'layout_supported',
                               'needs_confirmation': True, 'reasons': reasons,
                               'header_preview': preview, 'context_rows': [p['row'] for p in profiles if p['row'] < header_rows[0]][:6],
                               'possible_internal_header_rows': boundary_rows,
                               'proposed_options': {'cell_range': reference, 'header_rows': header_rows,
                                                    'data_start_row': header_rows[-1] + 1, 'data_end_row': bottom}})
            if len(candidates) >= MAX_DISCOVERY_BLOCKS:
                return candidates, True
    return candidates, sketch.truncated


def _resolve_region_preview_strings(archive, regions):
    needed = {cell['shared_string_index'] for region in regions for cell in region['header_preview']
              if 'shared_string_index' in cell}
    if not needed or 'xl/sharedStrings.xml' not in archive.namelist():
        return
    found = {}
    with archive.open('xl/sharedStrings.xml') as stream:
        index = 0
        for _, element in ET.iterparse(stream, events=('end',)):
            if element.tag != NS + 'si':
                continue
            if index in needed:
                found[index] = ''.join(item.text or '' for item in element.iter(NS + 't'))[:120]
            index += 1
            element.clear()
            if len(found) == len(needed):
                break
    for region in regions:
        for cell in region['header_preview']:
            if 'shared_string_index' in cell:
                idx = cell.pop('shared_string_index')
                cell['text'] = found.get(idx, '[共享字符串未解析]')


def _sheet_metadata(archive, part):
    merges = []
    hidden_rows = []
    table_ids = []
    max_row = max_col = scanned_cells = 0
    sketch = _RegionSketch()
    row_cells = []
    with archive.open(part) as stream:
        for event, element in ET.iterparse(stream, events=('end',)):
            if element.tag == NS + 'c':
                scanned_cells += 1
                if scanned_cells > MAX_SCAN_CELLS:
                    raise WorkbookReadError('工作表单元格数量超过扫描预算', 'WORKBOOK_LIMIT')
                if any(element.find(NS + t) is not None for t in ('v', 'f', 'is')):
                    coordinate = element.get('r', '')
                    match = re.fullmatch(r'([A-Z]+)([1-9][0-9]*)', coordinate)
                    if match:
                        max_row = max(max_row, int(match[2]))
                        max_col = max(max_col, column_index_from_string(match[1]))
                        typ = element.get('t', 'n')
                        raw = element.find(NS + 'v')
                        payload = raw.text if raw is not None else None
                        inline = element.find(NS + 'is')
                        text = ''.join(t.text or '' for t in inline.iter(NS + 't')) if inline is not None else None
                        is_formula = element.find(NS + 'f') is not None
                        if is_formula or payload is not None or text:
                            kind = 'formula' if is_formula else ('shared_text' if typ == 's' else 'text' if typ in ('str', 'inlineStr') else 'number')
                            cell = {'column': column_index_from_string(match[1]), 'kind': kind}
                            if kind == 'shared_text':
                                try:cell['shared_string_index'] = int(payload)
                                except (TypeError, ValueError):cell.update(kind='text', text='[共享字符串索引无效]')
                            elif kind == 'text':
                                cell['text'] = (text if text is not None else payload or '')[:120]
                            elif kind == 'number' and payload and re.fullmatch(r'(?:19|20|21)[0-9]{2}', payload):
                                cell['kind'] = 'year'
                            row_cells.append(cell)
                element.clear()
            elif element.tag == NS + 'row':
                sketch.add(int(element.get('r', '0')), row_cells)
                row_cells = []
                if element.get('hidden') in ('1', 'true'):
                    hidden_rows.append(int(element.get('r', '0')))
                element.clear()
            elif element.tag == NS + 'mergeCell':
                ref = element.get('ref', '')
                try:
                    bounds = range_boundaries(ref)
                except (ValueError, TypeError) as exc:
                    raise WorkbookReadError('Excel 合并范围无效') from exc
                if len(merges) >= MAX_MERGES:
                    raise WorkbookReadError('Excel 合并区域超过限制', 'WORKBOOK_LIMIT')
                merges.append({'range': ref, 'bounds': bounds,
                               'anchor': get_column_letter(bounds[0]) + str(bounds[1])})
                element.clear()
            elif element.tag == NS + 'tablePart':
                table_ids.append(element.get(REL + 'id'))
                element.clear()
    relationships = _relationships(archive, part)
    tables = []
    for rid in table_ids:
        target = relationships.get(rid)
        if target is None:
            raise WorkbookReadError('Excel 表格引用不完整')
        table = ET.fromstring(archive.read(target))
        tables.append({'name': table.get('name'), 'display_name': table.get('displayName'),
                       'range': table.get('ref'), 'header_rows': int(table.get('headerRowCount', '1')),
                       'totals_rows': int(table.get('totalsRowCount', '0'))})
    regions, regions_truncated = _region_proposals(sketch, merges, tables)
    _resolve_region_preview_strings(archive, regions)
    return {'max_row': max_row, 'max_column': max_col, 'merged_ranges': merges,
            'hidden_rows': hidden_rows, 'tables': tables, 'scanned_cells': scanned_cells,
            'candidate_regions': regions, 'region_discovery_truncated': regions_truncated,
            'discovery_scope': 'bounded_layout_proposals; confirm_header_grain_and_range; no_automatic_selection'}


def inspect_workbook(path):
    """List sheet/table metadata without selecting or combining their contents."""
    path = Path(path)
    try:
        with _archive(path) as archive:
            book = ET.fromstring(archive.read('xl/workbook.xml'))
            relationships = _relationships(archive, 'xl/workbook.xml')
            sheets = book.find(NS + 'sheets')
            if sheets is None or not 1 <= len(sheets) <= MAX_SHEETS:
                raise WorkbookReadError('工作簿页签数量超出限制', 'WORKBOOK_LIMIT')
            output = []
            for sheet in sheets:
                part = relationships.get(sheet.get(REL + 'id'))
                if part is None:
                    raise WorkbookReadError('工作表引用不完整')
                output.append({'name': sheet.get('name'), 'state': sheet.get('state', 'visible'),
                               'part': part, **_sheet_metadata(archive, part)})
            calculation = book.find(NS + 'calcPr')
            proposals = []
            for s in output:
                for candidate in s['candidate_regions']:
                    candidate['proposed_options']['sheet'] = s['name']
                    candidate['sheet'] = s['name']
                    candidate['candidate_id'] = hashlib.sha256(json.dumps(
                        [s['name'], candidate['proposed_options']], sort_keys=True).encode()).hexdigest()
                    if len(proposals) < MAX_DISCOVERY_REGIONS:
                        proposals.append(candidate)
            return {'parser': 'atlas_coordinate_xlsx', 'version': 1, 'sheets': output,
                    'candidate_regions': proposals, 'needs_confirmation': bool(proposals),
                    'region_discovery_truncated': sum(len(s['candidate_regions']) for s in output) > MAX_DISCOVERY_REGIONS or any(s['region_discovery_truncated'] for s in output),
                    'calculation_properties': dict(calculation.attrib) if calculation is not None else {},
                    'formula_cache_freshness': 'unknown; this reader never recalculates formulas'}
    except (ET.ParseError, KeyError, ValueError) as exc:
        if isinstance(exc, WorkbookReadError):
            raise
        raise WorkbookReadError('Excel 内部结构不完整或损坏') from exc


def workbook_sheet_names(path):
    """Cheap inventory for the legacy adapter, without rescanning worksheet cells."""
    try:
        with _archive(Path(path)) as archive:
            document = ET.fromstring(archive.read('xl/workbook.xml'))
            sheets = document.find(NS + 'sheets')
            if sheets is None or not 1 <= len(sheets) <= MAX_SHEETS:
                raise WorkbookReadError('工作簿页签数量超出限制', 'WORKBOOK_LIMIT')
            return [sheet.get('name') for sheet in sheets]
    except (ET.ParseError, KeyError) as exc:
        raise WorkbookReadError('Excel 页签结构不完整或损坏') from exc


def _raw_rows(archive, part):
    with archive.open(part) as stream:
        for _, element in ET.iterparse(stream, events=('end',)):
            if element.tag != NS + 'row':
                continue
            cells = {}
            for cell in element.findall(NS + 'c'):
                value = cell.find(NS + 'v')
                formula = cell.find(NS + 'f')
                cells[cell.get('r')] = {'value': value.text if value is not None else None,
                                       'has_value': value is not None,
                                       'type': cell.get('t', 'n'),
                                       'formula_present': formula is not None}
            yield int(element.get('r', '0')), cells
            element.clear()


class _Merges:
    def __init__(self, ranges, min_col, max_col):
        self.columns = {}
        intervals = 0
        for merge in ranges:
            left, top, right, bottom = merge['bounds']
            for column in range(max(left, min_col), min(right, max_col) + 1):
                intervals += 1
                if intervals > 100_000:
                    raise WorkbookReadError('合并区域索引超过内存预算，请选择更小表格范围', 'WORKBOOK_LIMIT')
                self.columns.setdefault(column, []).append((top, bottom, merge))
        self.starts = {}
        for col, values in self.columns.items():
            values.sort(key=lambda v: v[0])
            for a, b in zip(values, values[1:]):
                if a[1] >= b[0]:
                    raise WorkbookReadError('Excel 合并区域重叠，不能安全读取', 'WORKBOOK_MERGE_AMBIGUOUS')
            self.starts[col] = [value[0] for value in values]

    def at(self, row, column):
        values = self.columns.get(column, [])
        index = bisect_right(self.starts.get(column, []), row) - 1
        return values[index][2] if index >= 0 and row <= values[index][1] else None


def _formula_knowledge(formula):
    references = []
    functions = []
    try:
        for item in Tokenizer(formula).items:
            if item.type == 'OPERAND' and item.subtype == 'RANGE':
                references.append(item.value)
            elif item.type == 'FUNC' and item.subtype == 'OPEN':
                functions.append(item.value.rstrip('(').upper())
    except Exception:
        return {'formula': formula, 'references': [], 'functions': [], 'parse_status': 'unsupported'}
    return {'formula': formula, 'references': list(dict.fromkeys(references)),
            'functions': list(dict.fromkeys(functions)), 'parse_status': 'tokens_only_not_executed',
            'external_reference': any('[' in r for r in references),
            'dynamic_reference': bool(set(functions) & {'INDIRECT', 'OFFSET'})}


@lru_cache(maxsize=2048)
def _number_format_semantics(number_format):
    """Recognize scaling operators, keeping quoted/escaped text literal.

    This is deliberately not an Excel display renderer. Mixed sections and
    unsupported scale operators remain inspectable and fail executable reads.
    """
    sections = [('', '')]
    index = 0
    while index < len(number_format):
        character = number_format[index]
        code, literal = sections[-1]
        if character == '"':
            end = number_format.find('"', index + 1)
            if end < 0:
                return {'percentage': 'unsupported', 'display_scaled': False}
            literal += number_format[index + 1:end]
            index = end + 1
        elif character in ('\\', '_', '*'):
            if character == '\\' and index + 1 < len(number_format):
                literal += number_format[index + 1]
            index += 2
        elif character == '[':
            end = number_format.find(']', index + 1)
            if end < 0:
                return {'percentage': 'unsupported', 'display_scaled': False}
            index = end + 1
        elif character == ';':
            sections.append(('', ''))
            index += 1
            continue
        else:
            code += character
            index += 1
        sections[-1] = (code, literal)
    numeric = [(code, literal) for position, (code, literal) in enumerate(sections[:3])
               if position < 2 or re.search(r'[0#?%]', code) or '%' in literal]
    kinds = set()
    scaled = False
    for code, literal in numeric:
        scaled |= bool(re.search(r'[0#?],+(?![0#?])', code))
        operators, suffixes = code.count('%'), literal.count('%')
        kinds.add('fraction' if operators == 1 and not suffixes else
                  'points' if suffixes == 1 and not operators else
                  'native' if not operators and not suffixes else 'unsupported')
    percentage = ('fraction' if kinds == {'fraction'} else 'points' if kinds == {'points'} else
                  'unsupported' if kinds & {'fraction', 'points', 'unsupported'} else None)
    return {'percentage': percentage, 'display_scaled': scaled}


def _kind(value, number_format):
    if isinstance(value, bool):
        return 'boolean'
    if isinstance(value, (date, datetime, time)):
        return 'date_or_time'
    if isinstance(value, Decimal) or type(value) in (int, float):
        percentage = _number_format_semantics(number_format)['percentage']
        return {'fraction': 'percentage', 'points': 'percentage_points'}.get(percentage, 'number')
    return 'text' if isinstance(value, str) else 'blank'


def _value(cell, raw, epoch, require_formula_cache):
    value = cell.value
    formula = cell.data_type == 'f' or raw.get('formula_present', False)
    details = {}
    if formula:
        details['formula'] = _formula_knowledge(str(value))
        details['cache_freshness'] = 'unknown'
        details['cache_status'] = 'available'
        stored = raw.get('value')
        if stored is None and not (raw.get('type') == 'str' and raw.get('has_value')):
            details['cache_status'] = 'missing'
            if require_formula_cache:
                raise WorkbookReadError('公式缺少保存的计算结果，需在 Excel 等兼容引擎重算后重新保存',
                                        'FORMULA_CACHE_MISSING', {'coordinate': cell.coordinate})
            return None, details
        typ = raw.get('type', 'n')
        if typ == 'e':
            details['cache_status'] = 'error'
            if require_formula_cache:
                raise WorkbookReadError('Excel 公式结果包含错误，不能作为有效数值采集',
                                        'FORMULA_CACHED_ERROR', {'coordinate': cell.coordinate})
            return None, details
        if typ == 'str':
            return stored or '', details
        if typ == 'b':
            return stored == '1', details
        if typ not in ('n', 'd'):
            raise WorkbookReadError('此公式缓存类型未获得验证', 'FORMULA_CACHE_UNSUPPORTED', {'coordinate': cell.coordinate})
        if typ == 'd':
            try:
                datetime.fromisoformat(stored.replace('Z', '+00:00'))
            except (ValueError, AttributeError) as exc:
                raise WorkbookReadError('公式日期缓存格式无效', 'FORMULA_CACHE_UNSUPPORTED', {'coordinate': cell.coordinate}) from exc
            details['source_kind'] = 'date_or_time'
            return stored, details
        value = stored
    elif cell.data_type == 'e':
        raise WorkbookReadError('Excel 单元格包含错误值', 'CELL_VALUE_ERROR', {'coordinate': cell.coordinate})
    if raw.get('type', 'n') == 'n' and raw.get('value') is not None:
        try:
            number = Decimal(raw['value'])
        except InvalidOperation as exc:
            raise WorkbookReadError('Excel 数值编码无效', 'CELL_VALUE_ERROR', {'coordinate': cell.coordinate}) from exc
        if not number.is_finite():
            raise WorkbookReadError('Excel 包含非有限数值', 'CELL_VALUE_ERROR', {'coordinate': cell.coordinate})
        if len(number.as_tuple().digits) > 15 or (number and number.adjusted() >= 15):
            details['numeric_precision_warning'] = 'Excel numeric cells may already have lost original digits beyond 15 significant digits; no recovery inferred'
        if is_date_format(cell.number_format):
            details['raw_numeric_value'] = raw['value']
            value = from_excel(float(number), epoch=epoch)
        else:
            value = number
    if isinstance(value, (date, datetime, time)):
        details['source_kind'] = 'date_or_time'
        value = value.isoformat()
    return value, details


def _unit(header, number_format, *, numeric=True):
    match = re.search(r'[（(]\s*(亿元|千元|万元|元|人民币|CNY|USD|人|个|件|小时|天|%)\s*[）)]', header)
    header_unit = match[1] if match else None
    percentage = _number_format_semantics(number_format)['percentage'] if numeric else None
    quoted = re.findall(r'"([^"]+)"', number_format) if numeric else []
    format_units = {token.strip() for token in quoted
                    if token.strip() in ('亿元', '千元', '万元', '元', 'CNY', 'USD')}
    if percentage in ('fraction', 'points'):
        format_units.add('%')
    if len(format_units | ({header_unit} if header_unit else set())) > 1:
        return {'unit': None, 'unit_basis': 'conflicting_source_declarations', 'unit_conflict': True,
                'conflicting_units': sorted(format_units | ({header_unit} if header_unit else set()))}
    if percentage == 'fraction':
        return {'unit': '%', 'stored_semantics': 'ratio', 'display_multiplier': '100'}
    if percentage == 'points':
        return {'unit': '%', 'stored_semantics': 'percent_points', 'display_multiplier': '1'}
    if header_unit:
        return {'unit': header_unit, 'unit_basis': 'explicit_header_text', 'stored_semantics': 'source_unit; no_implicit_conversion'}
    if format_units:
        return {'unit': next(iter(format_units)), 'unit_basis': 'explicit_number_format_literal',
                'stored_semantics': 'source_format; display_scaling_must_be_reviewed'}
    return {'unit': None, 'unit_basis': 'not_declared'}


def parse_localized_value(value, *, kind, currency=None, unit=None, unit_multiplier='1',
                          date_format=None, missing_policy='error', missing_tokens=()):
    """Explicit text normalization, separate from automatic workbook ingestion.

    The caller declares currency, source unit, multiplier, date format and null
    policy. No inferred exchange rate, unit scale, date order or missing-zero
    conversion. Percent text must contain a percent sign. The original text and
    exact conversion policy are returned in a receipt; no formula is executed.
    """
    if kind not in ('decimal', 'percent', 'date') or missing_policy not in ('error', 'null'):
        raise WorkbookReadError('本地化解析种类或缺失策略无效', 'LOCALIZED_POLICY_INVALID')
    if currency not in (None, 'CNY', 'USD', 'EUR', 'JPY'):
        raise WorkbookReadError('需明确使用受支持的币种代码', 'LOCALIZED_POLICY_INVALID')
    if unit is not None and (not isinstance(unit, str) or not unit or len(unit) > 80):
        raise WorkbookReadError('来源单位声明无效', 'LOCALIZED_POLICY_INVALID')
    if not isinstance(unit_multiplier, str) or not re.fullmatch(r'[0-9]{1,26}(?:\.[0-9]{1,12})?', unit_multiplier):
        raise WorkbookReadError('单位倍率必须为显式正十进制字符串', 'LOCALIZED_POLICY_INVALID')
    multiplier = Decimal(unit_multiplier)
    if multiplier <= 0:
        raise WorkbookReadError('单位倍率必须为正数', 'LOCALIZED_POLICY_INVALID')
    if kind == 'date':
        allowed = {'%Y-%m-%d', '%Y/%m/%d', '%Y年%m月%d日', '%d/%m/%Y', '%m/%d/%Y', '%Y%m%d'}
        if date_format not in allowed or currency is not None or unit is not None or multiplier != 1:
            raise WorkbookReadError('日期必须明确完整年月日格式，不能推断03/04的含义', 'LOCALIZED_POLICY_INVALID')
    elif date_format is not None or (kind == 'percent' and (currency is not None or unit not in (None, '%'))):
        raise WorkbookReadError('数值解析策略含不兼容的日期或货币声明', 'LOCALIZED_POLICY_INVALID')
    if not isinstance(missing_tokens, (list, tuple)) or len(missing_tokens) > 100 or any(not isinstance(t, str) or len(t) > 100 for t in missing_tokens):
        raise WorkbookReadError('缺失标记列表无效', 'LOCALIZED_POLICY_INVALID')
    if value is not None and (not isinstance(value, str) or len(value) > 4000):
        raise WorkbookReadError('本地化规范化只接收原始文本或空值', 'LOCALIZED_VALUE_INVALID')
    receipt = {'raw_text': value, 'kind': kind, 'currency': currency, 'source_unit': unit,
               'unit_multiplier': unit_multiplier, 'date_format': date_format,
               'missing_policy': missing_policy, 'transformations': []}
    text = unicodedata.normalize('NFKC', value).strip() if value is not None else None
    tokens = {unicodedata.normalize('NFKC', t).strip() for t in missing_tokens}
    if text is None or text == '' or text in tokens:
        if missing_policy == 'error':
            raise WorkbookReadError('遇到缺失文本，未将其转换成零', 'LOCALIZED_MISSING_VALUE')
        receipt['transformations'].append('explicit_missing_to_null')
        return {'value': None, 'receipt': receipt}
    if text != value:
        receipt['transformations'].append('unicode_width_and_outer_whitespace_normalized')
    if kind == 'date':
        allowed = {'%Y-%m-%d', '%Y/%m/%d', '%Y年%m月%d日', '%d/%m/%Y', '%m/%d/%Y', '%Y%m%d'}
        if date_format not in allowed or currency is not None or unit is not None or multiplier != 1:
            raise WorkbookReadError('日期必须明确完整年月日格式，不能推断03/04的含义', 'LOCALIZED_POLICY_INVALID')
        try:
            result = datetime.strptime(text, date_format).date().isoformat()
        except ValueError as exc:
            raise WorkbookReadError('文本与声明的日期格式不一致', 'LOCALIZED_VALUE_INVALID') from exc
        receipt['transformations'].append('explicit_date_format_to_iso')
        return {'value': result, 'receipt': receipt}
    if date_format is not None:
        raise WorkbookReadError('数值解析不接受日期格式', 'LOCALIZED_POLICY_INVALID')
    if kind == 'percent' and (currency is not None or unit not in (None, '%')):
        raise WorkbookReadError('百分比解析不能同时声明货币单位', 'LOCALIZED_POLICY_INVALID')
    # Currency tokens are consumed only when their meaning has been declared.
    currency_tokens = {'CNY': ('人民币', 'RMB', 'CNY', '¥'), 'USD': ('USD', '$'),
                       'EUR': ('EUR', '€'), 'JPY': ('JPY', '日元', '¥')}
    known_tokens = ('人民币', 'RMB', 'CNY', 'USD', 'EUR', 'JPY', '日元', '¥', '$', '€')
    def remove_currency(raw):
        if currency is not None:
            for token in currency_tokens[currency]:
                if raw.upper().startswith(token.upper()):
                    receipt['transformations'].append('declared_currency_prefix_removed')
                    return raw[len(token):].strip()
        return raw
    text = remove_currency(text)
    negative = text.startswith('(') and text.endswith(')')
    if negative:
        text = text[1:-1].strip()
        receipt['transformations'].append('accounting_parentheses_to_negative')
    elif '(' in text or ')' in text:
        raise WorkbookReadError('金额括号格式不完整', 'LOCALIZED_VALUE_INVALID')
    text = remove_currency(text)
    if any(token.upper() in text.upper() for token in known_tokens):
        raise WorkbookReadError('金额符号与明示币种不一致或未声明币种', 'LOCALIZED_CURRENCY_REQUIRED')
    if negative and text[:1] in ('+', '-'):
        raise WorkbookReadError('括号负数不能叠加另一个正负号', 'LOCALIZED_VALUE_INVALID')
    if kind == 'percent':
        if not text.endswith('%') or text.count('%') != 1:
            raise WorkbookReadError('百分比文本必须带有明确百分号，不能猜测15还是0.15', 'LOCALIZED_PERCENT_SIGN_REQUIRED')
        text = text[:-1].strip()
        receipt['transformations'].append('percent_sign_to_fraction')
    elif '%' in text:
        raise WorkbookReadError('百分比文本不能直接当普通金额读取', 'LOCALIZED_VALUE_INVALID')
    if unit is not None and kind != 'percent':
        normalized_unit = unicodedata.normalize('NFKC', unit)
        if text.endswith(normalized_unit):
            text = text[:-len(normalized_unit)].strip()
            receipt['transformations'].append('declared_source_unit_suffix_removed')
    elif kind == 'decimal' and currency == 'CNY' and text.endswith('元') and not text.endswith(('万元', '亿元', '千元')):
        text = text[:-1].strip()
        receipt['transformations'].append('declared_base_currency_suffix_removed')
    if not re.fullmatch(r'[+-]?(?:[0-9]+|[0-9]{1,3}(?:,[0-9]{3})+)(?:\.[0-9]+)?', text):
        raise WorkbookReadError('数值文本含非法分组、未声明单位或不能识别的字符', 'LOCALIZED_VALUE_INVALID')
    if ',' in text:
        receipt['transformations'].append('validated_thousands_separators_removed')
    with localcontext() as context:
        context.prec = max(100, len(text) + len(unit_multiplier) + 20)
        result = Decimal(text.replace(',', ''))
        if negative:
            result = -result
        if kind == 'percent':
            result /= 100
        result *= multiplier
        if multiplier != 1:
            receipt['transformations'].append('explicit_unit_multiplier_applied')
        if not result.is_finite() or abs(result) >= Decimal('1e26') or result != result.quantize(Decimal('0.000000000001')):
            raise WorkbookReadError('规范化数值不能以DECIMAL(38,12)精确表示，未静默舍入', 'LOCALIZED_PRECISION')
    return {'value': result, 'receipt': receipt}


def _normalize_percentage_cell(value, number_format, target):
    """Convert only explicit percentage representations, retaining the raw value.

    A bare number is deliberately not a percentage declaration. Column titles
    and the chosen target representation cannot authorize guessing whether 13
    means 13% or 1300%. Numeric cells need an Excel percent operator or literal
    percent suffix; text cells need a percent sign.
    """
    if value is None:
        return value, None
    original_kind = _kind(value, number_format)
    if isinstance(value, str):
        parsed = parse_localized_value(value, kind='percent', unit='%')
        fraction = parsed['value']
        transformations = parsed['receipt']['transformations']
    elif isinstance(value, Decimal) or type(value) in (int, float):
        semantics = _number_format_semantics(number_format)
        if semantics['display_scaled'] or semantics['percentage'] not in ('fraction', 'points'):
            raise WorkbookReadError('数值单元格没有明确百分比表示，不能猜测13与0.13的含义',
                                    'PERCENTAGE_NUMERIC_UNDECLARED')
        with localcontext() as context:
            context.prec = 100
            fraction = Decimal(str(value)) / (100 if semantics['percentage'] == 'points' else 1)
        transformations = ['excel_percentage_fraction_preserved' if semantics['percentage'] == 'fraction'
                           else 'literal_percentage_points_to_fraction']
    else:
        raise WorkbookReadError('百分比列包含非数值、非百分比文本的单元格', 'LOCALIZED_VALUE_INVALID')
    with localcontext() as context:
        context.prec = 100
        result = fraction * (100 if target == 'points' else 1)
        if not result.is_finite() or abs(result) >= Decimal('1e26') or result != result.quantize(Decimal('0.000000000001')):
            raise WorkbookReadError('规范化百分比不能以DECIMAL(38,12)精确表示，未静默舍入', 'LOCALIZED_PRECISION')
    receipt = {'raw_value': str(value), 'raw_kind': original_kind, 'number_format': number_format,
               'output_representation': target, 'normalized_value': str(result),
               'transformations': transformations + (['fraction_to_percentage_points'] if target == 'points' else [])}
    return result, receipt


def read_table(path, *, sheet=None, table=None, header_rows=None, cell_range=None,
               data_start_row=None, data_end_row=None, merged_label_columns=(),
               limit=100_000, provenance_limit=20_000, require_formula_cache=True,
               hidden_rows='include', subtotal_policy='preserve', percentage_columns=None):
    """Extract one explicitly selected rectangular table, retaining its meaning.

    Multi-sheet workbooks require ``sheet`` or an unambiguous named ``table``.
    ``header_rows`` is an ordered consecutive list, e.g. [3,4]. Merged headers
    without an explicit header boundary are refused. ``merged_label_columns``
    permits only vertical merged text labels in the listed c_N/Excel columns;
    numeric detail merges are refused even when listed. No formula evaluation.
    ``percentage_columns={'c_2':'fraction'}`` explicitly normalizes signed
    percentage text and Excel percent-formatted numbers to a common ratio;
    ``points`` chooses percentage points. Bare numeric values remain refused.
    ``require_formula_cache=False`` is for a preview, never a verified full run.
    """
    _integer(limit, '行数限制', maximum=1_000_000)
    _integer(provenance_limit, '特殊单元格元数据预算', minimum=0, maximum=1_000_000)
    if type(require_formula_cache) is not bool or hidden_rows not in ('include', 'exclude', 'error') or subtotal_policy not in ('preserve', 'error'):
        raise WorkbookReadError('Excel 读取策略无效', 'WORKBOOK_SELECTION_INVALID')
    path = Path(path)
    document = inspect_workbook(path)
    matches = [s for s in document['sheets'] if sheet is None or s['name'] == sheet]
    if table is not None:
        matches = [s for s in matches if any(table in (t['name'], t['display_name']) for t in s['tables'])]
    if len(matches) != 1:
        raise WorkbookReadError('请明确选择工作表或唯一命名表格，不会自动拼接多个页签',
                                'SHEET_SELECTION_REQUIRED', {'sheets': [s['name'] for s in document['sheets']]})
    selected = matches[0]
    named = next((t for t in selected['tables'] if table in (t['name'], t['display_name'])), None) if table else None
    if cell_range is not None and named is not None:
        raise WorkbookReadError('命名表格与手动单元格范围不能同时指定', 'WORKBOOK_SELECTION_INVALID')
    try:
        bounds = range_boundaries(cell_range or (named or {}).get('range') or
                                  ('A1:' + get_column_letter(max(1, selected['max_column'])) + str(max(1, selected['max_row']))))
    except (ValueError, TypeError) as exc:
        raise WorkbookReadError('Excel 单元格范围无效', 'WORKBOOK_SELECTION_INVALID') from exc
    min_col, start, max_col, end = bounds
    if any(v is None for v in bounds) or max_col - min_col + 1 > 256:
        raise WorkbookReadError('需明确选择最多256列的有界表格范围', 'WORKBOOK_LIMIT')
    percentage_columns = {} if percentage_columns is None else percentage_columns
    if not isinstance(percentage_columns, dict) or len(percentage_columns) > 50 or any(
            not isinstance(key, str) or not re.fullmatch(r'c_[1-9][0-9]{0,2}', key)
            or not 1 <= int(key[2:]) <= max_col - min_col + 1
            or value not in ('fraction', 'points') for key, value in percentage_columns.items()):
        raise WorkbookReadError('百分比规范化须明确所选范围内的c_N列及fraction或points', 'WORKBOOK_SELECTION_INVALID')
    for value, maximum in ((min_col, 16_384), (max_col, 16_384), (start, 1_048_576), (end, 1_048_576)):
        _integer(value, '单元格范围', maximum=maximum)
    explicit_headers = header_rows is not None
    if header_rows is None:
        if named and named['header_rows'] != 1:
            raise WorkbookReadError('无表头或多表头命名表格需要明确 header_rows', 'HEADER_SELECTION_REQUIRED')
        header_rows = [start]
    if not isinstance(header_rows, list) or not 1 <= len(header_rows) <= 20:
        raise WorkbookReadError('header_rows 必须为1到20个连续行号', 'WORKBOOK_SELECTION_INVALID')
    for row in header_rows:
        _integer(row, '表头行')
    if header_rows != list(range(header_rows[0], header_rows[-1] + 1)) or header_rows[0] < start or header_rows[-1] > end:
        raise WorkbookReadError('表头行必须在所选范围内连续排列', 'WORKBOOK_SELECTION_INVALID')
    merges = _Merges(selected['merged_ranges'], min_col, max_col)
    if not explicit_headers and not named and any(
            m['bounds'][1] <= start <= m['bounds'][3] and m['bounds'][0] <= max_col and m['bounds'][2] >= min_col
            for m in selected['merged_ranges']):
        raise WorkbookReadError('多级或合并表头需明确 header_rows，避免将标题或数据误当字段', 'HEADER_SELECTION_REQUIRED',
                                {'merged_ranges': [m['range'] for m in selected['merged_ranges'][:30]]})
    first_data = data_start_row if data_start_row is not None else header_rows[-1] + 1
    last_data = data_end_row if data_end_row is not None else end - ((named or {}).get('totals_rows', 0))
    _integer(first_data, '数据起始行')
    _integer(last_data, '数据结束行', minimum=0)
    if first_data <= header_rows[-1] or first_data < start or last_data > end:
        raise WorkbookReadError('数据范围与表头边界冲突', 'WORKBOOK_SELECTION_INVALID')
    label_columns = set()
    for col in merged_label_columns:
        if type(col) is int:
            _integer(col, '合并标签列', maximum=16_384)
            resolved = col
        elif isinstance(col, str) and re.fullmatch(r'c_[1-9][0-9]*', col):
            resolved = min_col + int(col[2:]) - 1
        elif isinstance(col, str) and re.fullmatch(r'[A-Z]{1,3}', col):
            resolved = column_index_from_string(col)
        else:
            raise WorkbookReadError('合并标签列须使用 c_N、Excel列字母或列号', 'WORKBOOK_SELECTION_INVALID')
        if not min_col <= resolved <= max_col:
            raise WorkbookReadError('合并标签列不在选择范围内', 'WORKBOOK_SELECTION_INVALID')
        label_columns.add(resolved)
    warnings = []
    special = {}
    rows = []
    source_rows = []
    row_roles = {}
    fields = []
    formats = {col: [] for col in range(min_col, max_col + 1)}
    numeric_formats = {col: set() for col in range(min_col, max_col + 1)}
    kinds = {col: set() for col in range(min_col, max_col + 1)}
    anchors = {}
    headers = {}
    header_cells = {}
    stacks = {}
    hidden = set(selected['hidden_rows'])
    truncated = False
    missing_caches = 0
    precision_warnings = 0
    metadata_complete = True
    metadata_count = 0
    representation_conflicts = []
    header_sources = set()
    for header_row in header_rows:
        for column in range(min_col, max_col + 1):
            merge = merges.at(header_row, column)
            header_sources.add(merge['anchor'] if merge else get_column_letter(column) + str(header_row))
    reader = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    row_iterator = raw_iterator = None
    try:
        ws = reader[selected['name']]
        # Anchor columns can precede a selected subrange; read only enough to resolve them.
        read_min = min([min_col] + [m['bounds'][0] for m in selected['merged_ranges']
                       if m['bounds'][0] <= max_col and m['bounds'][2] >= min_col and m['bounds'][1] <= last_data])
        with _archive(path) as archive:
            raw_iterator = iter(_raw_rows(archive, selected['part']))
            raw_row, raw_cells = next(raw_iterator, (None, {}))
            scan_end = max(header_rows[-1], min(last_data, selected['max_row']))
            row_iterator = ws.iter_rows(min_row=1, max_row=scan_end,
                                        min_col=read_min, max_col=max_col)
            for row_number, cells in enumerate(row_iterator, 1):
                while raw_row is not None and raw_row < row_number:
                    raw_row, raw_cells = next(raw_iterator, (None, {}))
                raw = raw_cells if raw_row == row_number else {}
                values = {}
                details = {}
                for column, cell in enumerate(cells, read_min):
                    coordinate = get_column_letter(column) + str(row_number)
                    needed = coordinate in header_sources or (first_data <= row_number <= last_data and min_col <= column <= max_col)
                    if not needed:
                        values[column] = None
                        continue
                    v, extra = _value(cell, raw.get(coordinate, {}), reader.epoch, require_formula_cache)
                    values[column] = v
                    if extra:
                        details[column] = extra
                    if row_number <= header_rows[-1]:
                        headers[coordinate] = v
                        if extra:
                            if extra.get('cache_status') in ('missing', 'error'):
                                missing_caches += 1
                            if extra.get('numeric_precision_warning'):
                                precision_warnings += 1
                            if metadata_count >= provenance_limit:
                                metadata_complete = False
                                if require_formula_cache:
                                    raise WorkbookReadError('表头特殊单元格元数据超过预算', 'WORKBOOK_METADATA_LIMIT')
                            else:
                                header_cells[coordinate] = extra
                                metadata_count += 1
                    merge = merges.at(row_number, column)
                    if merge and coordinate == merge['anchor']:
                        anchors[coordinate] = v
                if row_number <= header_rows[-1] or row_number < first_data or row_number > last_data:
                    continue
                if not any(values.get(c) is not None for c in range(min_col, max_col + 1)) and not details and not any(merges.at(row_number, c) for c in range(min_col, max_col + 1)):
                    continue
                if row_number in hidden:
                    if hidden_rows == 'error':
                        raise WorkbookReadError('数据包含隐藏行，需明确包含或排除策略', 'HIDDEN_ROWS_PRESENT', {'row': row_number})
                    if hidden_rows == 'exclude':
                        continue
                if len(rows) >= limit:
                    truncated = True
                    break
                record = {}
                row_details = {}
                roles = []
                for column, cell in enumerate(cells, read_min):
                    if column < min_col:
                        continue
                    key = 'c_' + str(column - min_col + 1)
                    coordinate = get_column_letter(column) + str(row_number)
                    v = values.get(column)
                    meta = {k: value for k, value in details.get(column, {}).items()
                            if k != 'source_kind' and (k != 'raw_numeric_value' or 'formula' in details.get(column, {}))}
                    merge = merges.at(row_number, column)
                    if merge:
                        left, top, right, bottom = merge['bounds']
                        if top < first_data:
                            raise WorkbookReadError('合并区域跨越表头与数据，需重新明确布局', 'WORKBOOK_MERGE_AMBIGUOUS', {'range': merge['range']})
                        anchor = anchors.get(merge['anchor'])
                        if left != right or column not in label_columns:
                            raise WorkbookReadError('明细区有未确认含义的合并单元格，不能自动补齐', 'WORKBOOK_MERGE_AMBIGUOUS', {'range': merge['range'], 'column': key})
                        if not isinstance(anchor, str) or not anchor.strip():
                            raise WorkbookReadError('数值或空白合并明细不能作为标签填充，避免重复金额', 'MERGED_MEASURE_UNSAFE', {'range': merge['range']})
                        v = anchor
                        meta.update(merged_range=merge['range'], source_coordinate=merge['anchor'],
                                    propagation='explicit_vertical_text_label')
                    number_format = cell.number_format or 'General'
                    if isinstance(v, Decimal) or type(v) in (int, float):
                        numeric_formats[column].add(number_format)
                        format_semantics = _number_format_semantics(number_format)
                        format_error = ('NUMERIC_DISPLAY_SCALE_UNVERIFIED' if format_semantics['display_scaled'] else
                                        'PERCENTAGE_FORMAT_UNSUPPORTED' if format_semantics['percentage'] == 'unsupported' else None)
                        if format_error:
                            if require_formula_cache:
                                raise WorkbookReadError('数字格式包含未核对的显示缩放或百分比表示，保留原值供检查，不能猜测存储单位',
                                                        format_error, {'coordinate': coordinate, 'number_format': number_format})
                            if key not in representation_conflicts:
                                representation_conflicts.append(key)
                            meta['number_format_review_required'] = format_error
                    percentage_target = percentage_columns.get(key)
                    if percentage_target:
                        v, receipt = _normalize_percentage_cell(v, number_format, percentage_target)
                        if receipt:
                            meta['percentage_normalization'] = receipt
                    record[key] = v
                    if (v is not None or meta) and number_format not in formats[column]:
                        formats[column].append(number_format)
                        if len(formats[column]) > 20:
                            raise WorkbookReadError('单列数值格式变化超过预算，请选择更小范围', 'WORKBOOK_LIMIT')
                    kind = (('percentage' if percentage_target == 'fraction' else 'percentage_points')
                            if percentage_target and v is not None else
                            details.get(column, {}).get('source_kind') or _kind(v, number_format))
                    if kind != 'blank':
                        kinds[column].add(kind)
                    if formats[column] and number_format != formats[column][0]:
                        meta['number_format'] = number_format
                    if meta.get('cache_status') in ('missing', 'error'):
                        missing_caches += 1
                    if meta.get('numeric_precision_warning'):
                        precision_warnings += 1
                    if isinstance(v, str):
                        leading = len(v) - len(v.lstrip(' \t\u3000'))
                        indent = float(getattr(cell.alignment, 'indent', 0) or 0)
                        stack = stacks.setdefault(column, [])
                        score = indent if indent else leading
                        while stack and stack[-1][0] >= score:
                            stack.pop()
                        if leading or indent:
                            meta.update(style_indent=indent, leading_whitespace=leading,
                                        hierarchy_path=[item[1] for item in stack] + [v.strip()],
                                        hierarchy_semantics='layout_candidate_not_business_mapping')
                        if v.strip():
                            stack.append((score, v.strip()))
                        if re.fullmatch(r'(?:小计|合计|总计|累计|subtotal|total)', v.strip(), re.I):
                            roles.append('subtotal_label_candidate')
                    functions = meta.get('formula', {}).get('functions', [])
                    if 'SUBTOTAL' in functions:
                        roles.append('subtotal_formula')
                    if meta:
                        row_details[key] = meta
                if roles:
                    if subtotal_policy == 'error':
                        raise WorkbookReadError('数据区包含汇总候选行，需明确行粒度后再加工', 'SUBTOTAL_ROWS_PRESENT', {'row': row_number})
                    row_roles[str(row_number)] = list(dict.fromkeys(roles))
                if row_number in hidden:
                    row_roles.setdefault(str(row_number), []).append('hidden_in_source')
                if row_details:
                    if metadata_count + len(row_details) > provenance_limit:
                        metadata_complete = False
                        if require_formula_cache:
                            raise WorkbookReadError('特殊单元格元数据超过预算，不能丢弃证据后继续全量执行', 'WORKBOOK_METADATA_LIMIT')
                    else:
                        special[str(row_number)] = row_details
                        metadata_count += len(row_details)
                rows.append(record)
                source_rows.append(row_number)
        labels = set()
        for column in range(min_col, max_col + 1):
            path_parts = []
            header_coordinates = []
            for row_number in header_rows:
                coordinate = get_column_letter(column) + str(row_number)
                merge = merges.at(row_number, column)
                source_coordinate = merge['anchor'] if merge else coordinate
                value = headers.get(source_coordinate)
                header_coordinates.append({'coordinate': coordinate, 'source_coordinate': source_coordinate,
                                           'merged_range': merge['range'] if merge else None})
                if value is not None and str(value).strip() and (not path_parts or path_parts[-1] != str(value).strip()):
                    path_parts.append(str(value).strip())
            label = ' / '.join(path_parts) or ('未命名列' + get_column_letter(column))
            review = not path_parts
            if label in labels:
                label += ' [' + get_column_letter(column) + ']'
                warnings.append({'code': 'DUPLICATE_HEADER_DISAMBIGUATED', 'column': get_column_letter(column)})
                review = True
            labels.add(label)
            if not path_parts:
                warnings.append({'code': 'HEADER_MISSING', 'column': get_column_letter(column)})
            seen_formats = formats[column] or ['General']
            numeric_column = bool(kinds[column] & {'number', 'percentage', 'percentage_points'})
            unit_details = [_unit(label, fmt, numeric=numeric_column) for fmt in seen_formats]
            percentage_target = percentage_columns.get('c_' + str(column - min_col + 1))
            if percentage_target:
                # The normalized field no longer has each cell's source storage
                # representation. Preserve source formats in receipts and check
                # their currency/unit declarations before describing the target.
                original_units = [_unit(label, fmt) for fmt in numeric_formats[column]]
                target_format = '0%' if percentage_target == 'fraction' else '0"%"'
                unit_details = [_unit(label, target_format)]
                if any(detail.get('unit_conflict') for detail in original_units):
                    unit_details.append({'unit': None, 'unit_conflict': True})
            unit_conflict = (any(detail.get('unit_conflict') for detail in unit_details) or
                             len({detail.get('unit') for detail in unit_details}) > 1)
            percentage_kinds = kinds[column] & {'percentage', 'percentage_points'}
            mixed_percentage = bool(percentage_kinds) and len(kinds[column]) > 1
            if mixed_percentage:
                if require_formula_cache:
                    raise WorkbookReadError('同一列混合百分比与其他表示法，需明确逐类转换规则后再执行，不能统一当作比例',
                                            'MIXED_PERCENTAGE_REPRESENTATION', {'column': get_column_letter(column), 'formats': seen_formats})
                representation_conflicts.append('c_' + str(column - min_col + 1))
                review = True
            if unit_conflict:
                if require_formula_cache:
                    raise WorkbookReadError('表头或同列数字格式声明了不一致的单位，不能以第一种单位解释整列',
                                            'UNIT_FORMAT_CONFLICT', {'column': get_column_letter(column), 'formats': seen_formats})
                if 'c_' + str(column - min_col + 1) not in representation_conflicts:
                    representation_conflicts.append('c_' + str(column - min_col + 1))
                review = True
            field = {'name': 'c_' + str(column - min_col + 1), 'label': label,
                     'header_path': path_parts, 'header_coordinates': header_coordinates,
                     'excel_column': get_column_letter(column), 'excel_column_index': column,
                     'number_formats': seen_formats, 'observed_kinds': sorted(kinds[column]),
                     'numeric': bool(kinds[column]) and kinds[column] <= {'number', 'percentage', 'percentage_points'},
                     'semantic_status': 'requires_review' if review else 'source_declared_header',
                     **unit_details[0]}
            if mixed_percentage or field['name'] in representation_conflicts:
                field.update(unit=None, stored_semantics=('ambiguous_percentage_representation' if mixed_percentage else 'ambiguous_numeric_representation'),
                             representation_conflict=True, numeric=False)
            if percentage_target:
                field['normalization'] = {'kind': 'explicit_percentage', 'output_representation': percentage_target,
                                          'receipts': 'special_cells[source_row][field].percentage_normalization'}
            if len(seen_formats) > 1:
                warnings.append({'code': 'MIXED_NUMBER_FORMATS', 'field': field['name'], 'formats': seen_formats})
            fields.append(field)
    finally:
        # Read-only worksheet streams survive Workbook.close() when iteration
        # stops early. Close both generators so failed reads release Windows
        # file handles as promptly as successful reads.
        if row_iterator is not None:
            row_iterator.close()
        if raw_iterator is not None:
            raw_iterator.close()
        reader.close()
    if missing_caches:
        warnings.append({'code': 'FORMULA_VALUES_UNAVAILABLE', 'count': missing_caches})
    if precision_warnings:
        warnings.append({'code': 'EXCEL_PRECISION_LIMIT', 'count': precision_warnings,
                         'note': '原文件以数值保存的长编号可能已损失位数，不能据此恢复原编号'})
    if row_roles:
        warnings.append({'code': 'ROW_ROLE_REVIEW_REQUIRED', 'note': '隐藏/汇总候选行已保留并标记，汇总前需确认粒度，避免重复计数'})
    if not metadata_complete:
        warnings.append({'code': 'METADATA_PREVIEW_TRUNCATED', 'note': '预览不能用于完整语义传递或正式执行'})
    if representation_conflicts:
        warnings.append({'code': 'NUMERIC_REPRESENTATION_UNVERIFIED', 'fields': list(dict.fromkeys(representation_conflicts)),
                         'note': '源值与格式已保留；显示缩放、单位或百分比表示存在冲突，须明确转换规则'})
    hasher = hashlib.sha256()
    with path.open('rb') as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b''):
            hasher.update(chunk)
    source_hash = hasher.hexdigest()
    table_identity = {'file_sha256': source_hash, 'sheet': selected['name'], 'table': (named or {}).get('name'),
                      'range': [min_col, start, max_col, end], 'header_rows': header_rows,
                      'data_rows': [first_data, last_data], 'merged_label_columns': sorted(label_columns),
                      'hidden_rows_policy': hidden_rows, 'subtotal_policy': subtotal_policy}
    if percentage_columns:
        table_identity['percentage_columns'] = dict(sorted(percentage_columns.items()))
    table_id = hashlib.sha256(json.dumps(table_identity, sort_keys=True).encode()).hexdigest()
    provenance = {'version': 1, 'file_sha256': source_hash, 'table_id': table_id, 'sheet': selected['name'],
                  'source_rows': source_rows,
                  'columns': {f['name']: {'excel_column': f['excel_column'], 'header_path': f['header_path'],
                                        'number_formats': f['number_formats']} for f in fields},
                  'special_cells': special, 'complete': metadata_complete and not truncated and missing_caches == 0 and not representation_conflicts,
                  'header_cells': header_cells,
                  'coordinate_rule': 'columns[field].excel_column + source_rows[row_index]; special_cells may override source_coordinate'}
    return {'version': 1, 'parser': 'atlas_coordinate_xlsx', 'table_id': table_id, 'selection': table_identity,
            'rows': rows, 'fields': fields, 'source_rows': source_rows, 'special_cells': special,
            'header_cells': header_cells,
            'provenance': provenance,
            'row_roles': row_roles, 'warnings': warnings, 'truncated': truncated,
            'metadata_complete': metadata_complete, 'formula_cache_freshness': 'unknown',
            'executable_values_complete': missing_caches == 0 and metadata_complete and not truncated and not representation_conflicts,
            'coordinate_rule': 'field.excel_column + source_rows[row_index]; special_cells override source_coordinate for explicit merged labels',
            'scope': 'source_values_and_layout_only; no_formula_recalculation_or_implicit_semantic_conversion'}
