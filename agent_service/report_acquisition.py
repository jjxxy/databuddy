"""Bounded, repeatable acquisition of already-authorized report exports.

Only GET is supported. Query parameters are frozen requests, not evidence that a
remote report applied them. Published receipts distinguish requested filters,
export-content assertions, and declared/echoed semantic context.
"""
from __future__ import annotations

import csv
import hashlib
import json
import math
import os
import re
import tempfile
import time
import zipfile
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Literal
from urllib.parse import parse_qsl, urlsplit

import httpx
from filelock import FileLock
from pydantic import Field, field_validator, model_validator

from .connectors import SourceError, json_path, safe_api_url
from .models import StrictModel


class ReportError(SourceError):
    pass


def _digest(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True,
                                     separators=(',', ':'), default=str).encode()).hexdigest()


def _secret_key(key):
    normalized=re.sub(r'[^a-z0-9]', '', key.lower())
    return normalized in {'authorization','cookie','setcookie','sig','signature','key'} or normalized.endswith(
        ('password','passcode','token','apikey','secret','credential','signature','accesskey'))


class ReportRequest(StrictModel):
    scope_id: str = Field(pattern=r'^[A-Za-z0-9_-]{1,64}$')
    request_id: str = Field(pattern=r'^[A-Za-z0-9_-]{1,80}$')
    url: str = Field(min_length=8, max_length=8192)
    format: Literal['xlsx','csv','json'] = 'xlsx'
    filters: dict = Field(default_factory=dict, max_length=40)
    context: dict[str,str] = Field(default_factory=dict, max_length=30)
    expected_columns: list[str] = Field(default_factory=list, max_length=1000)
    expected_rows: int | None = Field(default=None, ge=0, le=10000000)
    expected_sha256: str | None = Field(default=None, pattern=r'^[a-f0-9]{64}$')
    filter_columns: dict[str,str] = Field(default_factory=dict, max_length=40)
    filter_echo_paths: dict[str,str] = Field(default_factory=dict, max_length=40)
    context_echo_paths: dict[str,str] = Field(default_factory=dict, max_length=30)
    json_rows_path: str = Field(default='', max_length=300)
    csv_encoding: Literal['utf-8-sig','utf-8','gb18030'] = 'utf-8-sig'
    xlsx_sheet: str | None = Field(default=None, max_length=100)
    xlsx_header_row: int = Field(default=1, ge=1, le=1000)
    allow_empty: bool = False
    max_bytes: int = Field(default=25*1024*1024, ge=1, le=500*1024*1024)
    max_rows: int = Field(default=100000, ge=1, le=10000000)
    max_columns: int = Field(default=1000, ge=1, le=16384)
    timeout_seconds: float = Field(default=30, ge=1, le=120)
    total_timeout_seconds: float = Field(default=120, ge=1, le=900)
    max_attempts: int = Field(default=3, ge=1, le=5)

    @field_validator('url')
    @classmethod
    def valid_url(cls,value):
        p=urlsplit(value)
        if p.scheme not in ('http','https') or not p.hostname or p.username or p.password or p.fragment:
            raise ValueError('报告导出地址必须是无内嵌凭证和片段的 HTTP(S) 地址')
        return value

    @model_validator(mode='after')
    def validate_contract(self):
        if any(not isinstance(k,str) or not k or len(k)>100 or _secret_key(k) for k in self.filters):
            raise ValueError('筛选参数无效；认证凭证必须单独放在请求头')
        for value in self.filters.values():
            values=value if isinstance(value,list) else [value]
            if not values or len(values)>100:
                raise ValueError('筛选值数量无效')
            if any(type(v) not in (str,int,float,bool) or (isinstance(v,float) and not math.isfinite(v)) or len(str(v))>2000 for v in values):
                raise ValueError('筛选值必须是有界的有限标量或标量数组')
        if any(not isinstance(k,str) or not k or _secret_key(k) or len(k)>100 or len(v)>2000 for k,v in self.context.items()):
            raise ValueError('报告上下文无效；不得包含认证信息')
        if len(set(self.expected_columns))!=len(self.expected_columns) or any(not x or len(x)>300 for x in self.expected_columns):
            raise ValueError('预期字段为空、重复或过长')
        if self.expected_rows is not None and self.expected_rows>self.max_rows:
            raise ValueError('预期行数超出读取预算')
        if (set(self.filter_columns)|set(self.filter_echo_paths))-set(self.filters):
            raise ValueError('筛选核对规则引用了不存在的请求参数')
        if set(self.context_echo_paths)-set(self.context):
            raise ValueError('上下文核对规则引用了不存在的声明')
        if self.format!='json' and (self.filter_echo_paths or self.context_echo_paths or self.json_rows_path):
            raise ValueError('元数据回显路径只适用于 JSON 报告')
        for value in list(self.filter_columns.values())+list(self.filter_echo_paths.values())+list(self.context_echo_paths.values()):
            if not value or len(value)>300:raise ValueError('核对规则字段或路径无效')
        return self


def _safe_endpoint(url):
    p=urlsplit(url)
    return {'scheme':p.scheme, 'host':p.hostname, 'port':p.port,
            'path':p.path or '/',
            # Signed download tokens are transport credentials. Never persist or hash them.
            'query':[(k,v) for k,v in parse_qsl(p.query,keep_blank_values=True) if not _secret_key(k)]}


def request_digest(request):
    value=request.model_dump()
    value['url']=_safe_endpoint(request.url)
    return _digest(value)


def _value(value):
    if isinstance(value,(date,datetime)):return value.isoformat()
    if type(value) is bool:return 'true' if value else 'false'
    if value is None:return None
    return str(value)


def _expected_values(value):
    return {_value(v) for v in (value if isinstance(value,list) else [value])}


def _echo_equals(observed,expected):
    if isinstance(expected,list):
        return isinstance(observed,list) and len(observed)==len(expected) and all(_echo_equals(a,b) for a,b in zip(observed,expected))
    if type(expected) is bool:return type(observed) is bool and observed==expected
    if type(expected) in (int,float):
        return type(observed) in (int,float,Decimal) and Decimal(str(observed))==Decimal(str(expected))
    return type(observed) is type(expected) and observed==expected


def _invalid_json_constant(value):
    raise ValueError('Non-finite JSON value')


def _headers(values, request_url):
    if not isinstance(values,dict) or len(values)>40:raise ReportError('报告认证请求头无效')
    result={}
    for key,value in values.items():
        if not isinstance(key,str) or not isinstance(value,str) or '\n' in key+value or '\r' in key+value:
            raise ReportError('报告认证请求头无效')
        if key.lower() in ('host','accept-encoding','content-length','range','if-range'):continue
        result[key]=value
    result['Host']=request_url.netloc.decode()
    result['Accept-Encoding']='identity'
    return result


def _file_hash(path):
    value=hashlib.sha256()
    with path.open('rb') as f:
        for chunk in iter(lambda:f.read(1024*1024),b''):value.update(chunk)
    return value.hexdigest()


def _check_deadline(deadline):
    if time.monotonic()>deadline:raise ReportError('报告获取超过总时间预算','REPORT_TIMEOUT',True)


def validate_download(path, request):
    """Check a complete local file; never execute formulas or return source rows."""
    request=ReportRequest.model_validate(request)
    path=Path(path)
    if not path.is_file() or path.stat().st_size==0:
        raise ReportError('报告文件为空或未完成','REPORT_INCOMPLETE')
    if path.stat().st_size>request.max_bytes:raise ReportError('报告文件超过体积预算','REPORT_TOO_LARGE')
    deadline=time.monotonic()+request.timeout_seconds
    with path.open('rb') as f:prefix=f.read(4096)
    text_prefix=prefix.lstrip(b'\xef\xbb\xbf \t\r\n').lower()
    if re.match(rb'<(?:!doctype\b|html\b|head\b|body\b|form\b|script\b|meta\b|title\b|div\b|section\b|table\b)',text_prefix):
        raise ReportError('导出返回登录页或 HTML 页面，未取得报告','REPORT_NOT_DATA')
    hash_value=_file_hash(path)
    if request.expected_sha256 and hash_value!=request.expected_sha256:
        raise ReportError('导出文件与预期摘要不一致','REPORT_HASH_MISMATCH')
    checks={key:{'method':'export_column_exact_membership','column':column,'rows_checked':0}
            for key,column in request.filter_columns.items()}
    count=0; columns=[]; all_columns=set(); document=None; selection={}

    def accept_row(row):
        nonlocal count
        count+=1
        if count>request.max_rows:raise ReportError('报告超过行数预算，未截断交付','REPORT_TOO_MANY_ROWS')
        if len(row)>request.max_columns:raise ReportError('报告超过字段数量预算','REPORT_TOO_MANY_COLUMNS')
        if set(request.expected_columns)-set(row):raise ReportError('报告记录缺少预期字段','REPORT_SCHEMA_MISMATCH')
        all_columns.update(row)
        for key,check in checks.items():
            if check['column'] not in row or _value(row[check['column']]) not in _expected_values(request.filters[key]):
                raise ReportError('报告数据与要求的筛选条件不一致','REPORT_FILTER_MISMATCH')
            check['rows_checked']+=1
        if count%100==0:_check_deadline(deadline)

    if request.format=='csv':
        if text_prefix.startswith((b'{',b'[')):
            try:
                with path.open(encoding='utf-8-sig') as candidate:json.load(candidate,parse_constant=_invalid_json_constant)
            except (ValueError,UnicodeDecodeError):pass
            else:raise ReportError('导出实际返回 JSON，未按 CSV 报告处理','REPORT_NOT_DATA')
        if b'\x00' in prefix or prefix.startswith(b'PK'):
            raise ReportError('导出内容不是文本 CSV','REPORT_NOT_DATA')
        selection={'encoding':request.csv_encoding,'table':'entire_csv'}
        try:
            with path.open(encoding=request.csv_encoding,newline='') as f:
                reader=csv.reader(f)
                columns=[v.strip() for v in next(reader,[])]
                if not columns or any(not v for v in columns) or len(set(columns))!=len(columns):
                    raise ReportError('CSV 表头为空或重复','REPORT_SCHEMA_MISMATCH')
                if len(columns)>request.max_columns:raise ReportError('报告超过字段数量预算','REPORT_TOO_MANY_COLUMNS')
                if set(request.expected_columns)-set(columns):raise ReportError('CSV 缺少预期字段','REPORT_SCHEMA_MISMATCH')
                for values in reader:
                    if not values or not any(v.strip() for v in values):continue
                    if len(values)!=len(columns):raise ReportError('CSV 行列数量不一致','REPORT_SCHEMA_MISMATCH')
                    accept_row(dict(zip(columns,values)))
        except (UnicodeDecodeError,csv.Error) as exc:
            raise ReportError('CSV 编码或格式不符合已声明的读取规则','REPORT_NOT_DATA') from exc
    elif request.format=='json':
        try:
            with path.open(encoding='utf-8-sig') as f:document=json.load(f,parse_float=Decimal,parse_constant=_invalid_json_constant)
        except (ValueError,UnicodeDecodeError) as exc:
            raise ReportError('导出内容不是有效 JSON','REPORT_NOT_DATA') from exc
        selection={'records_path':request.json_rows_path}
        rows=json_path(document,request.json_rows_path)
        if not isinstance(rows,list) or any(not isinstance(v,dict) for v in rows):
            raise ReportError('JSON 记录路径不是对象数组，可能为接口错误页','REPORT_SCHEMA_MISMATCH')
        for row in rows:accept_row(row)
        columns=sorted(all_columns)
    else:
        try:
            with zipfile.ZipFile(path) as z:
                names=set(z.namelist())
                if '[Content_Types].xml' not in names or 'xl/workbook.xml' not in names:
                    raise ReportError('导出内容不是 XLSX 工作簿','REPORT_NOT_DATA')
                if sum(x.file_size for x in z.infolist())>200*1024*1024:
                    raise ReportError('Excel 解压后体积超过预算','REPORT_TOO_LARGE')
                if any('vbaproject' in name.lower() for name in names):
                    raise ReportError('XLSX 导出包含宏内容，拒绝按普通表格读取','REPORT_NOT_DATA')
                if z.testzip() is not None:raise ReportError('Excel 压缩内容不完整','REPORT_INCOMPLETE')
            from openpyxl import load_workbook
            source=path.open('rb')
            try:
                book=load_workbook(source,read_only=True,data_only=True,keep_links=False)
            except BaseException:
                source.close();raise
            try:
                if request.xlsx_sheet and request.xlsx_sheet not in book.sheetnames:
                    raise ReportError('导出工作簿缺少指定页签','REPORT_SCHEMA_MISMATCH')
                sheet=book[request.xlsx_sheet or book.sheetnames[0]]
                selection={'sheet':sheet.title,'header_row':request.xlsx_header_row,'available_sheets':book.sheetnames}
                # Stored dimensions may lie or include far-away formatting; scan bounded rows.
                sheet.reset_dimensions()
                for physical,values in enumerate(sheet.iter_rows(values_only=True),1):
                    if physical%1000==0:_check_deadline(deadline)
                    if physical>request.max_rows+request.xlsx_header_row:
                        raise ReportError('Excel 扫描范围超过预算，请明确页签和表格范围','REPORT_TOO_MANY_ROWS')
                    if physical<request.xlsx_header_row:continue
                    if physical==request.xlsx_header_row:
                        columns=[str(v).strip() if v is not None else '' for v in values]
                        if not columns or any(not v for v in columns) or len(set(columns))!=len(columns):
                            raise ReportError('Excel 表头为空或重复，请指定正确表头行或使用底稿读取配置','REPORT_SCHEMA_MISMATCH')
                        if len(columns)>request.max_columns:raise ReportError('报告超过字段数量预算','REPORT_TOO_MANY_COLUMNS')
                        if set(request.expected_columns)-set(columns):raise ReportError('Excel 缺少预期字段','REPORT_SCHEMA_MISMATCH')
                        continue
                    if not any(v is not None for v in values):continue
                    if len(values)>len(columns) and any(v is not None for v in values[len(columns):]):
                        raise ReportError('Excel 数据超出声明表头范围','REPORT_SCHEMA_MISMATCH')
                    accept_row(dict(zip(columns,list(values)+[None]*max(0,len(columns)-len(values)))))
            finally:book.close();source.close()
        except (zipfile.BadZipFile,KeyError,OSError,ValueError) as exc:
            if isinstance(exc,ReportError):raise
            raise ReportError('Excel 内容损坏或格式无法读取','REPORT_NOT_DATA') from exc
    if request.format in ('xlsx','csv') and not columns:raise ReportError('报告没有可确认的表头','REPORT_SCHEMA_MISMATCH')
    if set(request.filter_columns.values())-set(columns):raise ReportError('报告缺少筛选核对字段','REPORT_SCHEMA_MISMATCH')
    if not count and not request.allow_empty:raise ReportError('报告没有可用记录，未按成功处理','REPORT_EMPTY')
    if request.expected_rows is not None and count!=request.expected_rows:
        raise ReportError('报告行数与要求不一致，可能存在漏导出','REPORT_ROW_COUNT_MISMATCH')
    if set(request.expected_columns)-set(columns):raise ReportError('报告缺少预期字段','REPORT_SCHEMA_MISMATCH')
    filter_evidence=[]
    for key,check in checks.items():
        if count:filter_evidence.append({'filter':key,**check})
    for key,path_name in request.filter_echo_paths.items():
        observed=json_path(document,path_name)
        if not _echo_equals(observed,request.filters[key]):
            raise ReportError('报告筛选元数据回显不一致','REPORT_FILTER_MISMATCH')
        filter_evidence.append({'filter':key,'method':'json_metadata_echo','path':path_name})
    context_evidence=[]
    for key,path_name in request.context_echo_paths.items():
        if json_path(document,path_name)!=request.context[key]:
            raise ReportError('报告语义上下文回显不一致','REPORT_CONTEXT_MISMATCH')
        context_evidence.append({'context':key,'method':'json_metadata_echo','path':path_name})
    _check_deadline(deadline)
    return {'format':request.format,'selection':selection,'rows':count,'columns':columns,'content_sha256':hash_value,
            'bytes':path.stat().st_size,'completion_verified':True,'filters_requested':request.filters,
            'filter_verification':filter_evidence,
            'filters_verified':bool(request.filters) and {v['filter'] for v in filter_evidence}==set(request.filters),
            'context_declared':request.context,'context_verification':context_evidence,
            'context_verified':bool(request.context) and {v['context'] for v in context_evidence}==set(request.context),
            'scope':'file_integrity_and_configured_checks_only; no_formula_recalculation; no_automatic_business_meaning_inference'}


def _atomic_json(path,value):
    from .platform_files import chmod_descriptor,fsync_directory,atomic_replace
    fd,name=tempfile.mkstemp(prefix='.report-receipt-',suffix='.tmp',dir=path.parent)
    temp=Path(name)
    try:
        chmod_descriptor(fd)
        with os.fdopen(fd,'w') as f:
            json.dump(value,f,ensure_ascii=False,indent=2);f.flush();os.fsync(f.fileno())
        atomic_replace(temp,path)
        fsync_directory(path.parent)
    finally:temp.unlink(missing_ok=True)


def acquire_report(request,destination_dir,*,headers=None,client_factory=None):
    """Acquire one safe GET report and atomically promote only validated bytes.

    destination_dir must be selected by the server, scoped to the correct project
    and connector principal. request_id identifies one logical acquisition;
    refreshes must use a new request_id. Credentials never enter the receipt.
    """
    request=ReportRequest.model_validate(request)
    destination=Path(destination_dir).resolve();destination.mkdir(parents=True,exist_ok=True)
    identity=_digest([request.scope_id,request.request_id])[:32]
    final=destination/(identity+'.'+request.format)
    receipt_path=destination/(identity+'.receipt.json')
    fingerprint=request_digest(request)
    with FileLock(str(destination/(identity+'.lock')),timeout=5):
        if receipt_path.is_symlink() or final.is_symlink():raise ReportError('报告路径不能为符号链接','REPORT_PATH_INVALID')
        if receipt_path.is_file():
            try:previous=json.loads(receipt_path.read_text())
            except (OSError,ValueError):raise ReportError('报告完成回执损坏，未复用文件','REPORT_RECEIPT_INVALID') from None
            if previous.get('request_sha256')!=fingerprint:
                raise ReportError('同一获取编号的筛选或读取规则已变化，请使用新编号','REPORT_IDEMPOTENCY_CONFLICT')
            if not final.is_file() or _file_hash(final)!=previous.get('content_sha256'):
                raise ReportError('已完成报告缺失或内容已变化，未复用','REPORT_ARTIFACT_CHANGED')
            return dict(previous,reused=True)
        deadline=time.monotonic()+request.total_timeout_seconds
        factory=client_factory or httpx.Client
        for attempt in range(1,request.max_attempts+1):
            _check_deadline(deadline)
            fd,name=tempfile.mkstemp(prefix='.'+identity+'-',suffix='.'+request.format,dir=destination)
            from .platform_files import chmod_descriptor
            temporary=Path(name);chmod_descriptor(fd);os.close(fd)
            try:
                original=httpx.URL(request.url)
                # Configured filters override same-named URL query parameters once.
                original=original.copy_merge_params(request.filters)
                ip=safe_api_url(str(original))
                remaining=min(request.timeout_seconds,max(.1,deadline-time.monotonic()))
                with factory(timeout=remaining,follow_redirects=False,trust_env=False) as client:
                    with client.stream('GET',original.copy_with(host=ip),headers=_headers(headers or {},original),
                                       extensions={'sni_hostname':original.host}) as response:
                        if response.status_code>=300:
                            raise ReportError('报告导出请求失败：HTTP '+str(response.status_code),'REPORT_HTTP_FAILED',
                                              response.status_code==429 or response.status_code>=500)
                        if response.status_code!=200:
                            raise ReportError('报告导出未返回完整文件','REPORT_INCOMPLETE')
                        content_type=response.headers.get('content-type','').split(';',1)[0].strip().lower()
                        if content_type in ('text/html','application/xhtml+xml'):
                            raise ReportError('导出返回登录或错误页面，未取得报告','REPORT_NOT_DATA')
                        if response.headers.get('content-encoding','identity').lower() not in ('identity',''):
                            raise ReportError('报告返回了未约定的压缩传输，无法核对完整字节数','REPORT_ENCODING_UNSUPPORTED')
                        length=response.headers.get('content-length')
                        if length is not None:
                            try:length=int(length)
                            except (ValueError,TypeError):raise ReportError('报告文件长度响应头无效','REPORT_INCOMPLETE') from None
                            if length<0 or length>request.max_bytes:raise ReportError('报告文件超过体积预算','REPORT_TOO_LARGE')
                        size=0
                        with temporary.open('wb') as f:
                            for chunk in response.iter_bytes():
                                _check_deadline(deadline)
                                size+=len(chunk)
                                if size>request.max_bytes:raise ReportError('报告文件超过体积预算','REPORT_TOO_LARGE')
                                f.write(chunk)
                            f.flush();os.fsync(f.fileno())
                        if length is not None and size!=length:
                            raise ReportError('报告下载未完成，未发布部分文件','REPORT_INCOMPLETE',True)
                        etag=response.headers.get('etag')
                _check_deadline(deadline)
                evidence=validate_download(temporary,request)
                _check_deadline(deadline)
                receipt={**evidence,'version':1,'acquisition_id':identity,'scope_id':request.scope_id,
                         'request_id':request.request_id,'request_sha256':fingerprint,
                         'endpoint':{'scheme':original.scheme,'host':original.host,'path':original.path},
                         'file':final.name,'attempts':attempt,'reused':False,
                         'etag_sha256':_digest(etag) if etag else None,
                         'retry_policy':'bounded_GET_only; interrupted_downloads_restart; no_POST_or_browser_click_replay'}
                from .platform_files import atomic_replace
                atomic_replace(temporary,final)
                _atomic_json(receipt_path,receipt)
                return receipt
            except (httpx.TransportError,httpx.TimeoutException) as exc:
                error=ReportError('报告连接或下载中断，请检查网络和导出端点','REPORT_NETWORK_FAILED',True)
                if attempt==request.max_attempts:raise error from None
            except SourceError as exc:
                if not exc.retryable or attempt==request.max_attempts:raise
            finally:
                temporary.unlink(missing_ok=True)
            _check_deadline(deadline)
            pause=min(2**(attempt-1)*.25,2,max(0,deadline-time.monotonic()))
            if pause:time.sleep(pause)
    raise ReportError('报告未取得完整文件','REPORT_INCOMPLETE')
