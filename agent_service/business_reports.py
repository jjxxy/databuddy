"""Reusable business briefs whose numerical assertions come from frozen results.

Recipes are versioned workspace code. Numerical output stays in private runtime
artifacts; a report contains observations and declared assumptions, never an
invented causal explanation. This intentionally does not emulate arbitrary PPT
templates or make OCR output authoritative.
"""
from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import tempfile
from decimal import Decimal, ROUND_HALF_EVEN, localcontext
from pathlib import Path
from typing import Literal

from pydantic import Field, model_validator

from . import workspace
from .config import DATA, MAX_ROWS
from .data_workflows import NAME, TableSchema, decimal_value, digest
from .models import StrictModel
from .store import all_items, get, now, put


class BriefMetric(StrictModel):
    id: str = Field(pattern=NAME)
    label: str = Field(min_length=1, max_length=120)
    kind: Literal['sum', 'value', 'ratio', 'variance', 'row_count']
    column: str | None = Field(default=None, pattern=NAME)
    baseline_column: str | None = Field(default=None, pattern=NAME)
    reduction: Literal['sum', 'value'] = 'sum'
    zero_policy: Literal['error', 'null'] = 'error'
    rounding_scale: int | None = Field(default=None, ge=0, le=12)

    @model_validator(mode='after')
    def operands(self):
        if self.kind != 'row_count' and self.column is None:
            raise ValueError('指标必须指定数值字段')
        if self.kind in ('ratio', 'variance') and self.baseline_column is None:
            raise ValueError('比率和差异必须指定明确的比较基数字段')
        if self.kind not in ('ratio', 'variance') and self.baseline_column is not None:
            raise ValueError('该指标不接受比较基数字段')
        return self


class ReportRecipe(StrictModel):
    version: Literal[1] = 1
    title: str = Field(min_length=1, max_length=160)
    period: str = Field(pattern=r'^\d{4}-(?:0[1-9]|1[0-2])$')
    expected_schema_sha256: str = Field(pattern=r'^[a-f0-9]{64}$')
    period_column: str | None = Field(default=None, pattern=NAME)
    filters: dict[str, str] = Field(default_factory=dict, max_length=20)
    metrics: list[BriefMetric] = Field(min_length=1, max_length=30)
    assumptions: list[str] = Field(default_factory=list, max_length=20)

    @model_validator(mode='after')
    def safe_text(self):
        if len({m.id for m in self.metrics}) != len(self.metrics):
            raise ValueError('报告指标编号不能重复')
        if any(not x.strip() or len(x)>1000 for x in self.assumptions):
            raise ValueError('假设必须为1至1000字的明确文字')
        if any(not re.fullmatch(NAME,k) for k in self.filters):
            raise ValueError('筛选字段编号无效')
        if self.period_column in self.filters:
            raise ValueError('期间由period指定，不可在筛选中重复覆盖')
        return self


def schema_fingerprint(schema):
    """Periods may advance for recurring briefs; all other semantics must match."""
    schema=TableSchema.model_validate(schema).model_dump()
    for column in schema['columns']:
        column.pop('period',None)
    schema.pop('source_ref',None)
    return digest(schema)


def _read(pid,run_id):
    from .project_data import result_folder
    from .result_delivery import read_all_handoff
    folder,meta=result_folder(pid,run_id)
    adopted=get('run',run_id).get('status')=='succeeded'
    if meta.get('requires_semantic_review') and not adopted:
        raise ValueError('该结果仍有待核对的语义或SQL口径，不能生成已核实数值简报')
    table=read_all_handoff(folder/'handoff',expected_manifest_sha256=meta['handoff_sha256'],max_rows=MAX_ROWS)
    schema=TableSchema.model_validate(table['context'].get('output_semantics') or table['schema'])
    return table,dict(meta,server_adopted_result=adopted),schema


def inspect_result(pid,run_id):
    table,meta,schema=_read(pid,run_id)
    return {'run_id':run_id,'schema_sha256':schema_fingerprint(schema),'schema':schema.model_dump(),
            'rows':table['total'],'handoff_sha256':meta['handoff_sha256'],'complete':True,
            'requires_semantic_review':bool(meta.get('requires_semantic_review')),
            'server_adopted_result':meta['server_adopted_result'],
            'supported_metrics':['sum','value','ratio','variance','row_count'],
            'period_required':True,'arbitrary_causal_explanation_verified':False}


def _text(value):
    if value is None:return '不适用（比较基数为0）'
    return format(value,'f') if isinstance(value,Decimal) else str(value)


def _escape(value):
    # Markdown is a downloadable artifact; escape markup and table delimiters.
    return re.sub(r'([\\`*_{}\[\]<>|#!])',r'\\\1',str(value)).replace('\r',' ').replace('\n',' ')


def _divide(left,right,metric):
    if right==0:
        if metric.zero_policy=='null':return None
        raise ValueError('比较基数为0，无法计算；请明确选择保留不适用或修改业务规则')
    with localcontext() as context:
        context.prec=120
        value=left/right
        if metric.rounding_scale is None:
            # A finite decimal is exact iff multiplication recovers numerator.
            if value*right!=left or value.as_tuple().exponent < -12:
                raise ValueError('比率需要明确小数保留位数，未静默舍入')
        else:value=value.quantize(Decimal(1).scaleb(-metric.rounding_scale),rounding=ROUND_HALF_EVEN)
    return value


def evaluate(recipe,table):
    """Pure deterministic numerical assertions, useful to test without a worker."""
    recipe=ReportRecipe.model_validate(recipe)
    schema=TableSchema.model_validate(table.get('context',{}).get('output_semantics') or table['schema'])
    if schema_fingerprint(schema)!=recipe.expected_schema_sha256:
        raise ValueError('结果字段或业务口径已变化，请核对后更新报告方案，不能直接套用上期逻辑')
    fields={c.name:c for c in schema.columns};rows=table['rows']
    if table.get('total')!=len(rows) or not table.get('complete_result_available'):
        raise ValueError('报告必须读取完整且核验过的结果，不能对预览样本生成汇总')
    if not schema.grain:
        if len(rows)>1:raise ValueError('多行结果缺少唯一业务粒度，无法排除重复汇总')
    else:
        keys=[tuple(row.get(k) for k in schema.grain) for row in rows]
        if any(any(v is None for v in key) for key in keys) or len(set(keys))!=len(keys):
            raise ValueError('结果业务粒度重复或为空，不能生成可能重复计数的报告')
    if not set(recipe.filters)<=set(fields):raise ValueError('报告筛选字段不存在')
    if any(m.kind=='row_count' for m in recipe.metrics) and not recipe.period_column:
        numeric=[c for c in schema.columns if c.data_type in ('decimal','integer')]
        if not numeric or any(c.period!=recipe.period for c in numeric):
            raise ValueError('行数指标缺少明确的目标期间范围，请指定期间字段或核对所有数值列期间')
    if recipe.period_column:
        if recipe.period_column not in fields:raise ValueError('报告期间字段不存在')
        rows=[r for r in rows if r[recipe.period_column]==recipe.period]
    rows=[r for r in rows if all(r.get(k)==v for k,v in recipe.filters.items())]
    if not rows:raise ValueError('报告期间和筛选范围没有数据，未把缺失数据当作0')

    def number(column,reduction):
        if column not in fields:raise ValueError('报告数值字段不存在')
        f=fields[column]
        if f.data_type not in ('decimal','integer'):raise ValueError('报告数值字段类型不正确')
        if f.period and f.period!=recipe.period:raise ValueError('报告字段期间与目标期间不一致')
        if not f.period and not recipe.period_column:raise ValueError('字段未声明期间，必须明确期间字段后再生成报告')
        if reduction=='sum' and (f.aggregation!='additive' or f.value_format!='native'):
            raise ValueError('只能汇总明确可加的原生数值；比率和百分比不能直接相加')
        if reduction=='value' and len(rows)!=1:raise ValueError('直接取值需要恰好一行，不能随意取第一行')
        if any(r.get(column) is None for r in rows):raise ValueError('数值缺失，未把空值自动当作0')
        with localcontext() as context:
            context.prec=120
            values=[decimal_value(r[column]) for r in rows]
            return sum(values,Decimal(0)) if reduction=='sum' else values[0]

    assertions=[]
    with localcontext() as context:
        context.prec=120
        for metric in recipe.metrics:
            proof={'id':metric.id,'label':metric.label,'kind':metric.kind,'rows':len(rows),
                   'period':recipe.period,'filters':recipe.filters,'rounding_scale':metric.rounding_scale,
                   'rounding_mode':'half_even' if metric.rounding_scale is not None else 'exact_only',
                   'zero_policy':metric.zero_policy,'columns':[]}
            if metric.kind=='row_count':proof.update(value=str(len(rows)),unit='行')
            elif metric.kind in ('sum','value'):
                value=number(metric.column,metric.kind);f=fields[metric.column]
                proof.update(value=_text(value),unit=f.unit,value_format=f.value_format,columns=[metric.column])
                if f.value_format in ('percent_fraction','percent_points'):
                    proof.update(display_value=_text(value*100 if f.value_format=='percent_fraction' else value),display_unit='%')
            else:
                value=number(metric.column,metric.reduction);baseline=number(metric.baseline_column,metric.reduction)
                a,b=fields[metric.column],fields[metric.baseline_column]
                if (a.unit,a.value_format,a.period)!=(b.unit,b.value_format,b.period):
                    raise ValueError('比较字段的单位、百分比表示或期间不一致')
                proof.update(actual=_text(value),baseline=_text(baseline),unit=a.unit,
                             columns=[metric.column,metric.baseline_column],reduction=metric.reduction)
                if metric.kind=='ratio':
                    proof.update(value=_text(_divide(value,baseline,metric)),unit='1',value_format='ratio')
                else:
                    delta=value-baseline
                    # Explicit fraction-vs-points representation is never guessed.
                    delta_unit=a.unit
                    if a.value_format=='percent_fraction':delta*=100;delta_unit='百分点'
                    elif a.value_format=='percent_points':delta_unit='百分点'
                    rate=_divide(value-baseline,baseline,metric)
                    proof.update(value=_text(delta),unit=delta_unit,relative_change=_text(rate),relative_change_unit='1',
                                 value_format='difference',relative_change_formula='(actual-baseline)/baseline')
            assertions.append(proof)
    return {'period':recipe.period,'input_rows':table['total'],'selected_rows':len(rows),
            'schema_sha256':schema_fingerprint(schema),'assertions':assertions,
            'no_causal_inference':True,'assumptions':recipe.assumptions}


def render(pid,run_id,recipe,*,recipe_path='reports/month-end.json'):
    recipe=ReportRecipe.model_validate(recipe)
    # Validate path before performing any work or writing artifacts.
    if Path(recipe_path).suffix.lower()!='.json':raise ValueError('可复用报告方案必须使用JSON文件')
    workspace.path_for(pid,recipe_path)
    table,meta,_=_read(pid,run_id);result=evaluate(recipe,table)
    payload={'version':1,'project_id':pid,'run_id':run_id,'plan_sha256':meta['plan_sha256'],
             'handoff_sha256':meta['handoff_sha256'],'recipe':recipe.model_dump(),'evaluation':result,
             'output_semantics':table.get('context',{}).get('output_semantics') or table['schema'],
             'requires_semantic_review':bool(meta.get('requires_semantic_review')),
             'server_adopted_result':meta['server_adopted_result'],
             'source_reference':f'/api/projects/{pid}/data-results/{run_id}',
             'complete_handoff_reference':f'/api/projects/{pid}/data-results/{run_id}/download?format=handoff'}
    rid='brief_'+digest(payload)[:40]
    lines=['# '+_escape(recipe.title),'','报告期间：'+recipe.period,'',
           '以下为完整结果计算出的观察值；业务原因和后续判断需要另附依据。','',
           '| 指标 | 数值 | 单位 | 覆盖行数 |','| --- | ---: | --- | ---: |']
    for assertion in result['assertions']:
        lines.append('| '+_escape(assertion['label'])+' | '+_escape(assertion.get('display_value',assertion['value']))+' | '+_escape(assertion.get('display_unit',assertion['unit']))+' | '+str(assertion['rows'])+' 行 |')
        if 'relative_change' in assertion:
            lines.append('| '+_escape(assertion['label']+'：相对基数变化')+' | '+_escape(assertion['relative_change'])+' | 比值 | '+str(assertion['rows'])+' 行 |')
    if recipe.assumptions:
        lines.extend(['','## 用户声明的假设（未作为事实核实）',''])
        lines.extend('- '+_escape(x) for x in recipe.assumptions)
    lines.extend(['','## 数值依据','',f'- 完整结果：[{run_id}]({payload["source_reference"]})',
                  '- 交接摘要：`'+meta['handoff_sha256']+'`',
                  '- 加工方案摘要：`'+meta['plan_sha256']+'`',
                  '- 总行数：'+str(result['input_rows'])+'；本次筛选行数：'+str(result['selected_rows']),
                  '- 精确字段、比较公式、筛选和舍入规则见同目录 manifest.json。',''])
    markdown='\n'.join(lines);markdown_bytes=markdown.encode()
    manifest={**payload,'report_id':rid,'markdown_sha256':hashlib.sha256(markdown_bytes).hexdigest()}
    manifest_bytes=json.dumps(manifest,ensure_ascii=False,sort_keys=True,indent=2).encode()
    root=DATA/'business-reports'/pid;target=root/rid
    with workspace.lock(pid):
        root.mkdir(parents=True,exist_ok=True)
        if target.exists():
            if (target/'report.md').read_bytes()!=markdown_bytes or (target/'manifest.json').read_bytes()!=manifest_bytes:
                raise ValueError('已有报告产物完整性不匹配，未覆盖')
        else:
            from .result_delivery import _fsync_directory
            from .platform_files import fsync_file
            temporary=Path(tempfile.mkdtemp(prefix='.brief-',dir=root))
            try:
                temporary.chmod(0o700)
                (temporary/'report.md').write_bytes(markdown_bytes);(temporary/'manifest.json').write_bytes(manifest_bytes)
                for path in temporary.iterdir():
                    path.chmod(0o600)
                    fsync_file(path)
                _fsync_directory(temporary)
                temporary.rename(target)
                _fsync_directory(root)
            finally:
                if temporary.exists():shutil.rmtree(temporary)
        # Recipes contain logic and declared assumptions; full numerical output
        # and raw rows remain outside the project code repository.
        workspace.write_file(pid,recipe_path,json.dumps(recipe.model_dump(),ensure_ascii=False,indent=2))
        checkpoint=workspace._checkpoint(pid,None,'保存可复用业务简报方案')
        try:existing=get('business_report',rid);created_at=existing['created_at']
        except KeyError:existing={};created_at=now()
        record=put('business_report',{'project_id':pid,'run_id':run_id,'recipe_path':recipe_path,
            'recipe_sha256':digest(recipe.model_dump()),'recipe_commit':checkpoint['commit'],
            'artifact_dir':target.relative_to(DATA).as_posix(),'handoff_sha256':meta['handoff_sha256'],
            'manifest_sha256':hashlib.sha256(manifest_bytes).hexdigest(),'markdown_sha256':manifest['markdown_sha256'],
            'title':recipe.title,'period':recipe.period,'created_at':created_at,
            'office_exports':existing.get('office_exports',{}),
            'office_export_pending':existing.get('office_export_pending',{}),'complete':True},rid)
    return dict(record,assertions=result['assertions'],download_url=f'/api/projects/{pid}/business-reports/{rid}/download')


def render_saved(pid,run_id,recipe_path):
    return render(pid,run_id,json.loads(workspace.read_file(pid,recipe_path)),recipe_path=recipe_path)


def list_reports(pid):
    get('project',pid)
    return sorted((r for r in all_items('business_report') if r.get('project_id')==pid),key=lambda r:r['created_at'],reverse=True)


def report_file(pid,report_id,format='markdown'):
    record=get('business_report',report_id)
    if record.get('project_id')!=pid:raise ValueError('报告不属于当前工作空间')
    if format not in ('markdown','manifest','docx','pptx'):raise ValueError('报告格式无效')
    root=(DATA/record['artifact_dir']).resolve()
    if not root.is_relative_to((DATA/'business-reports'/pid).resolve()):raise ValueError('报告路径无效')
    if format in ('docx','pptx'):
        return _office_file(pid,report_id,root,format)
    path=root/('report.md' if format=='markdown' else 'manifest.json')
    expected=record['markdown_sha256' if format=='markdown' else 'manifest_sha256']
    if hashlib.sha256(path.read_bytes()).hexdigest()!=expected:raise ValueError('报告产物完整性校验失败')
    return path


def _assertion_rows(manifest):
    rows=[]
    for assertion in manifest['evaluation']['assertions']:
        rows.append([assertion['label'],assertion.get('display_value',assertion['value']),
                     assertion.get('display_unit',assertion['unit']),str(assertion['rows'])])
        if 'relative_change' in assertion:
            rows.append([assertion['label']+'：相对基数变化',assertion['relative_change'],'比值',str(assertion['rows'])])
    return rows


def _office_file(pid,report_id,root,format):
    """Export frozen strings; the model cannot rewrite the document's numbers."""
    from .result_delivery import durable_cached_artifact
    from .locks import assert_owned
    with workspace.lock(pid):
        record=get('business_report',report_id)
        manifest_path=report_file(pid,report_id,'manifest')
        manifest=json.loads(manifest_path.read_text(encoding='utf-8'))
        path=root/('report.'+format)
        expected=record.get('office_exports',{}).get(format)
        pending=record.get('office_export_pending',{}).get(format)
        if pending and pending.get('manifest_sha256')!=record['manifest_sha256']:
            raise ValueError('报告Office待发布回执与冻结报告不一致')
        def build(temporary):
            if format=='docx':_write_docx(manifest,temporary)
            else:_write_pptx(manifest,temporary)
            _verify_office_values(manifest,temporary,format)
        def save_pending(checksum):
            pending_exports=dict(record.get('office_export_pending',{}))
            pending_exports[format]={'sha256':checksum,'manifest_sha256':record['manifest_sha256']}
            record.update(put('business_report',dict(record,office_export_pending=pending_exports),report_id))
        def save_committed(checksum):
            exports=dict(record.get('office_exports',{}));exports[format]=checksum
            pending_exports=dict(record.get('office_export_pending',{}));pending_exports.pop(format,None)
            record.update(put('business_report',dict(record,office_exports=exports,office_export_pending=pending_exports),report_id))
        return durable_cached_artifact(path,expected_sha256=expected,
            pending_sha256=pending.get('sha256') if pending else None,build=build,
            save_pending=save_pending,save_committed=save_committed,before_publish=assert_owned)


def _write_docx(manifest,path):
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches,Pt,RGBColor
    document=Document();section=document.sections[0]
    section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.top_margin=section.bottom_margin=Inches(.8)
    section.left_margin=section.right_margin=Inches(.85)
    for name in ('Normal','Title','Heading 1','Heading 2'):
        font=document.styles[name].font;font.name='Arial';font.color.rgb=RGBColor(0,0,0)
        document.styles[name].element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    document.styles['Normal'].font.size=Pt(11)
    document.styles['Normal'].paragraph_format.space_after=Pt(7)
    document.add_paragraph(manifest['recipe']['title'],'Title')
    document.add_paragraph('报告期间：'+manifest['recipe']['period'])
    document.add_paragraph('以下数值来自完整加工结果。业务原因和后续判断需要另附依据。')
    table=document.add_table(rows=1,cols=4);table.autofit=False
    for cell,width in zip(table.rows[0].cells,(2.55,2.55,.8,.9)):cell.width=Inches(width)
    headers=['指标','数值','单位','范围行数']
    for cell,text in zip(table.rows[0].cells,headers):cell.text=text
    repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in _assertion_rows(manifest):
        row=table.add_row()
        for cell,text,width in zip(row.cells,values,(2.55,2.55,.8,.9)):
            cell.text=text;cell.width=Inches(width)
    for index,row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            properties=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
            for edge in ('top','left','bottom','right'):
                border=OxmlElement('w:'+edge);border.set(qn('w:val'),'single');border.set(qn('w:sz'),'4');border.set(qn('w:color'),'D9D9D9');borders.append(border)
            properties.append(borders)
            if index==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EAF0F6');properties.append(shade)
                for run in cell.paragraphs[0].runs:run.bold=True
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before=Pt(5);paragraph.paragraph_format.space_after=Pt(5)
    assumptions=manifest['recipe']['assumptions']
    if assumptions:
        document.add_heading('用户声明的假设',level=1)
        document.add_paragraph('以下假设未作为事实核实。')
        for value in assumptions:document.add_paragraph(value)
    document.add_heading('数值依据',level=1)
    for text in ('完整结果编号：'+manifest['run_id'],'交接摘要：'+manifest['handoff_sha256'],
                 '加工方案摘要：'+manifest['plan_sha256'],'完整数据行数：'+str(manifest['evaluation']['input_rows']),
                 '本次筛选行数：'+str(manifest['evaluation']['selected_rows'])):
        document.add_paragraph(text)
    document.save(path)


def _write_pptx(manifest,path):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches,Pt
    presentation=Presentation();presentation.slide_width=Inches(13.333);presentation.slide_height=Inches(7.5)
    def text(slide,value,x,y,w,h,size=20,bold=False):
        box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));frame=box.text_frame
        frame.word_wrap=True;frame.margin_left=frame.margin_right=Inches(.03)
        for i,line in enumerate(value.split('\n')):
            paragraph=frame.paragraphs[0] if i==0 else frame.add_paragraph()
            paragraph.text=line;paragraph.font.name='Microsoft YaHei';paragraph.font.size=Pt(size)
            paragraph.font.bold=bold;paragraph.font.color.rgb=RGBColor(20,31,43)
        return box
    def slide(title):
        current=presentation.slides.add_slide(presentation.slide_layouts[6])
        text(current,title,.7,.45,11.9,1.0,30,True)
        return current
    cover=presentation.slides.add_slide(presentation.slide_layouts[6])
    text(cover,manifest['recipe']['title'],.7,.45,11.9,2.8,30,True)
    text(cover,'报告期间 '+manifest['recipe']['period'],.75,3.7,11.7,.65,28)
    text(cover,'以下数值来自完整加工结果。\n业务原因和后续判断需要另附依据。',.75,4.8,11.7,1.5,22)
    rows=_assertion_rows(manifest);offset=0;page_number=1
    while offset<len(rows):
        # Long business labels consume more vertical space. Never truncate them.
        chunk=[];height=.55
        while offset+len(chunk)<len(rows) and len(chunk)<6:
            values=rows[offset+len(chunk)]
            row_height=max(.68,.29*((len(values[0])+17)//18),.29*((len(values[1])+36)//37))
            if chunk and height+row_height>4.9:break
            chunk.append((values,row_height));height+=row_height
        current=slide('经营指标 '+str(page_number))
        native=current.shapes.add_table(len(chunk)+1,4,Inches(.75),Inches(1.7),Inches(11.8),Inches(height)).table
        for column,width in zip(native.columns,(3.5,5.3,1.3,1.7)):column.width=Inches(width)
        for i,values in enumerate([['指标','数值','单位','范围行数']]+[values for values,_ in chunk]):
            native.rows[i].height=Inches(.55 if i==0 else chunk[i-1][1])
            for j,value in enumerate(values):
                cell=native.cell(i,j);cell.text=value;cell.margin_left=cell.margin_right=Inches(.1)
                cell.margin_top=cell.margin_bottom=Inches(.06)
                cell.fill.solid();cell.fill.fore_color.rgb=RGBColor(234,240,246) if i==0 else RGBColor(255,255,255)
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.font.name='Microsoft YaHei';paragraph.font.size=Pt(17)
                    paragraph.font.bold=i==0;paragraph.font.color.rgb=RGBColor(20,31,43)
        text(current,'期间 '+manifest['recipe']['period']+'    数值单位与原始口径一致',.75,6.85,11.7,.35,13)
        offset+=len(chunk);page_number+=1
    for assumption in manifest['recipe']['assumptions']:
        for offset in range(0,len(assumption),400):
            current=slide('用户声明的假设')
            text(current,'以下假设未作为事实核实。',.75,1.65,11.7,.55,20)
            text(current,assumption[offset:offset+400],.75,2.4,11.7,4.3,20)
    current=slide('数值依据')
    for index,value in enumerate(['完整结果编号：'+manifest['run_id'],'交接摘要：'+manifest['handoff_sha256'],
        '加工方案摘要：'+manifest['plan_sha256'],'完整数据行数：'+str(manifest['evaluation']['input_rows']),
        '本次筛选行数：'+str(manifest['evaluation']['selected_rows'])]):
        text(current,value,.75,1.65+index*.9,11.8,.7,18)
    presentation.save(path)


def _verify_office_values(manifest,path,format):
    if format=='docx':
        from docx import Document
        document=Document(path)
        actual=[[c.text for c in r.cells] for table in document.tables for r in list(table.rows)[1:]]
    else:
        from pptx import Presentation
        actual=[[c.text for c in row.cells] for slide in Presentation(path).slides
                for shape in slide.shapes if shape.has_table for row in list(shape.table.rows)[1:]]
    if actual!=_assertion_rows(manifest):raise ValueError('Office报告数值读回与冻结断言不一致')
