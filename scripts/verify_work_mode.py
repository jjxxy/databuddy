"""Exercise an actual deployed runner; no mock backend or model is used.

Set ATLAS_WORK_RUNNER_URL and ATLAS_WORK_RUNNER_TOKEN. This creates a fresh
verification workspace, runs commands, and retains its checkpoints for review.
"""
import argparse
import base64
import hashlib
import io
import json
import os
from pathlib import Path
import secrets
import sys
import time

import httpx
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from work_runner.protocol import Command, InputFile, digest


def verify():
    url = os.environ['ATLAS_WORK_RUNNER_URL'].rstrip('/')
    token = os.environ['ATLAS_WORK_RUNNER_TOKEN']
    wid = digest(('atlas-work-verification-' + secrets.token_hex(16)).encode())
    with httpx.Client(base_url=url, headers={'Authorization':'Bearer ' + token}, timeout=20,
                      trust_env=False, follow_redirects=False) as client:
        def call(method, path, **kwargs):
            response = client.request(method, '/v1/' + path, **kwargs)
            response.raise_for_status()
            return response.json()
        health = call('GET', 'health')
        assert health['ready'] and health['runtime'] == 'runsc' and health['network'] == 'none'

        def submit(code, timeout=30):
            eid = digest(secrets.token_bytes(32))
            data = code.encode()
            request = Command(execution_id=eid, workspace_id=wid, command='python verify.py', timeout_seconds=timeout,
                files=[InputFile(path='verify.py', content_base64=base64.b64encode(data).decode(), sha256=digest(data))])
            first = call('POST', 'executions', json=request.model_dump())
            repeated = call('POST', 'executions', json=request.model_dump())
            assert first['execution_id'] == repeated['execution_id'] == eid
            return request

        def wait(request):
            deadline = time.monotonic() + request.timeout_seconds + 60
            while time.monotonic() < deadline:
                result = call('GET', 'executions/' + request.execution_id, params={'workspace_id':wid})
                if result['state'] in ('succeeded','failed','unknown','timed_out','cancelled'):
                    return result
                time.sleep(.5)
            raise TimeoutError('Runner did not return a terminal receipt')

        first = submit('''import os, socket, json
from pathlib import Path
from openpyxl import Workbook
assert os.getuid() == 10001
assert not any(k in os.environ for k in ('APP_TOKEN','LLM_API_KEY','DATABASE_URL','ATLAS_WORK_RUNNER_TOKEN'))
assert not Path('/var/run/docker.sock').exists()
try:
    Path('/etc/atlas-write-probe').write_text('must be blocked')
except OSError:
    pass
else:
    raise AssertionError('Root filesystem is writable')
try:
    socket.create_connection(('1.1.1.1', 443), timeout=1)
except OSError:
    pass
else:
    raise AssertionError('Unapproved egress is possible')
Path('outputs').mkdir(exist_ok=True)
Path('outputs/exact.json').write_text(json.dumps({'amount':'9007199254740993.00000001'}))
book=Workbook();book.active.append(['amount']);book.active.append(['9007199254740993.00000001']);book.save('outputs/exact.xlsx')
print('nonroot, rootfs, network, no-secrets and editable workbook probes passed')
''')
        result = wait(first)
        assert result['state'] == 'succeeded' and result['exit_code'] == 0, result.get('stderr') or result
        file = call('GET', 'executions/' + first.execution_id + '/file', params={'workspace_id':wid,'path':'outputs/exact.xlsx'})
        raw = base64.b64decode(file['content_base64'], validate=True)
        assert digest(raw) == file['sha256']
        book = load_workbook(io.BytesIO(raw), read_only=True)
        assert book.active['A2'].value == '9007199254740993.00000001'; book.close()
        second = submit("from pathlib import Path\nimport json\nassert json.loads(Path('outputs/exact.json').read_text())['amount']=='9007199254740993.00000001'\nPath('outputs/resumed.txt').write_text('continued from checkpoint')\n")
        assert wait(second)['state'] == 'succeeded'
        long = submit('import time\ntime.sleep(120)\n', timeout=150)
        call('POST', 'executions/' + long.execution_id + '/cancel', params={'workspace_id':wid}, content=b'')
        assert wait(long)['state'] == 'cancelled'
        return {'ok':True,'real_runner_exercised':True,'simulated_model':False,'model_evaluated':False,
                'workspace_id':wid,'executions':[first.execution_id,second.execution_id,long.execution_id],
                'checks':['gvisor_readiness','nonroot','no_application_secrets','no_docker_socket',
                          'readonly_rootfs','no_network','exact_xlsx_text','immutable_file_hash','checkpoint_resume','idempotency','cancellation'],
                'security_penetration_test':False,'production_ha_evaluated':False}


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path)
    options = parser.parse_args()
    result = verify()
    encoded = json.dumps(result, ensure_ascii=False, indent=2) + '\n'
    if options.output:
        options.output.write_text(encoded)
    print(encoded)
