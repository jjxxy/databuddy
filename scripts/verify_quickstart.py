"""Exercise the quickstart cloud contract with a local simulated model.

Starts the real cloud supervisor, API, worker, and agent in isolated storage.
With --opencode-binary it also executes the actual pinned OpenCode session loop.
The OpenAI-compatible fixture only returns predetermined synthetic responses;
this is integration evidence, never a real-model or Docker qualification.
"""
from __future__ import annotations

import argparse
from contextlib import contextmanager
from decimal import Decimal
import hashlib
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
import io
import json
import os
from pathlib import Path
import socket
import subprocess
import sys
import tempfile
import threading
import time
import traceback
import zipfile

import httpx

ROOT = Path(__file__).resolve().parents[1]
FIXTURE_KEY = 'quickstart-fixture-key-not-a-real-credential'
FIXTURE_MODEL = 'quickstart-simulated-model'
FIXTURE_MODEL_ALTERNATE = 'quickstart-simulated-model-alternate'
FIXTURE_OPERATOR_TOKEN = 'quickstart-fixture-operator-token-not-a-real-credential'
EXACT_TOTAL = '300.000000000002'


class ModelFixture:
    """A deliberately narrow OpenAI-compatible protocol fixture, not an LLM."""

    def __init__(self):
        self.workflow = None
        self.recipe = None
        self.calls = {'connection': 0, 'source_planning': 0, 'agent': 0}
        self.tool_probes = 0
        self.models_seen = []
        self.agent_models = []
        self.native_loop_calls = 0
        self.native_compaction_calls = 0
        self.scenarios = {}
        self.observed_contexts = {}

    @staticmethod
    def context(messages):
        """Locate the Atlas context in either direct or OpenCode messages."""
        decoder = json.JSONDecoder()
        found = None
        for message in messages:
            content = message.get('content', '')
            if isinstance(content, list):
                content = '\n'.join(part.get('text', '') for part in content if isinstance(part, dict))
            if not isinstance(content, str):continue
            for start, character in enumerate(content):
                if character != '{':continue
                try:value, _ = decoder.raw_decode(content[start:])
                except json.JSONDecodeError:continue
                if isinstance(value, dict) and ({'request','mode','completed_steps'} <= value.keys()
                                               or {'baseline','fields'} <= value.keys()):found=value
        if found is not None:return found
        raise ValueError('No synthetic Atlas context')

    def reply(self, body):
        messages = body['messages']
        if len(messages) == 1 and not any(tool.get('function',{}).get('name')=='atlas_step' for tool in body.get('tools',[])):
            self.calls['connection'] += 1
            return 'OK'
        context = self.context(messages)
        if 'baseline' in context and 'fields' in context:
            self.calls['source_planning'] += 1
            return json.dumps({'proposed_spec': context['baseline'], 'suggestions': []})
        self.calls['agent'] += 1
        self.agent_models.append(body['model'])
        request = context['request']
        self.observed_contexts[request] = context
        scenario = next((actions for marker,actions in self.scenarios.items() if marker in request),None)
        if scenario is not None:
            actions=scenario
        elif context['mode'] == 'chat' and 'quickstart-fixture-check' in request:
            return json.dumps({'kind': 'final', 'answer': '本地模拟模型接线检查完成。'})
        elif 'quickstart-fixture-plan' in request:
            path, value = 'workflows/month.json', self.workflow
            actions = None
        elif 'quickstart-fixture-recipe' in request:
            path, value = 'reports/month.json', self.recipe
            actions = None
        else:
            raise ValueError('Unexpected request to the simulated model')
        if actions is None:
            if value is None:raise ValueError('The synthetic fixture document has not been prepared')
            actions = [
            {'kind': 'tool', 'tool': 'project_context'},
            {'kind': 'tool', 'tool': 'write_file', 'arguments': {
                'path': path, 'content': json.dumps(value, ensure_ascii=False, indent=2)}},
            {'kind': 'tool', 'tool': 'read_file', 'arguments': {'path': path}},
            {'kind': 'final', 'answer': '说明文件已保存并读取核对。数据加工由后续明确操作执行。'},
            ]
        # Direct Atlas calls rebuild completed_steps; OpenCode native sessions
        # append their tool calls to the original immutable prompt context.
        native_steps = sum(1 for message in messages for call in message.get('tool_calls', [])
                           if call.get('function', {}).get('name') == 'atlas_step')
        return json.dumps(actions[min(max(len(context['completed_steps']), native_steps), len(actions)-1)], ensure_ascii=False)

    def message(self, body):
        self.models_seen.append(body.get('model'))
        if any('Summarize ongoing work without inventing facts' in str(message.get('content',''))
               for message in body.get('messages',[]) if message.get('role')=='system'):
            self.native_compaction_calls+=1
            self.agent_models.append(body.get('model'))
            return {'role':'assistant','content':'## Objective\n继续当前合成业务验收。\n## Constraints\n不要将期末余额求和。金额单位为元。\n## Next Move\n沿用完整来源和已验证回执；压缩摘要不是数值证据。'},'stop'
        names = {tool.get('function', {}).get('name') for tool in body.get('tools', [])}
        if 'atlas_step' in names:
            if names != {'atlas_step'}:raise ValueError('OpenCode exposed an unexpected native tool')
            self.native_loop_calls += 1
        if 'connection_check' in names:
            self.tool_probes += 1
            name, arguments = 'connection_check', {'ok':True}
        else:
            decision = self.reply(body)
            if not names:return {'role':'assistant','content':decision}, 'stop'
            value = json.loads(decision)
            if 'atlas_step' in names:
                name, arguments = 'atlas_step', value
            elif value['kind'] == 'final':
                name = 'finish_task'
                arguments = {key:item for key,item in value.items() if key != 'kind'}
            else:
                name, arguments = value['tool'], value.get('arguments', {})
            if name not in names:raise ValueError('Synthetic action is not an advertised tool')
        return {'role':'assistant','content':None,'tool_calls':[{'id':'call_fixture_' + str(len(self.models_seen)),
            'type':'function','function':{'name':name,'arguments':json.dumps(arguments,ensure_ascii=False)}}]}, 'tool_calls'


@contextmanager
def local_model():
    fixture = ModelFixture()

    class Handler(BaseHTTPRequestHandler):
        def do_POST(self):
            if self.path != '/v1/chat/completions' or self.headers.get('Authorization') != 'Bearer ' + FIXTURE_KEY:
                self.send_response(401); self.end_headers(); return
            try:
                body = json.loads(self.rfile.read(int(self.headers['Content-Length'])))
                if body.get('model') not in (FIXTURE_MODEL, FIXTURE_MODEL_ALTERNATE):
                    raise ValueError('Unexpected simulated model identifier')
                message, finish_reason = fixture.message(body)
                payload = {'id': 'quickstart-simulated-only', 'object': 'chat.completion',
                    'model': body['model'], 'choices': [{'index': 0,
                    'message': message, 'finish_reason': finish_reason}],
                    'usage': {'prompt_tokens': 10, 'completion_tokens': 10, 'total_tokens': 20}}
                if body.get('stream'):
                    chunks=[]
                    delta={'role':'assistant'}
                    if message.get('tool_calls'):
                        delta['tool_calls']=[dict(call,index=index) for index,call in enumerate(message['tool_calls'])]
                    else:delta['content']=message['content']
                    for value in [
                        {'id':payload['id'],'object':'chat.completion.chunk','created':int(time.time()),
                         'model':body['model'],'choices':[{'index':0,'delta':delta,'finish_reason':None}]},
                        {'id':payload['id'],'object':'chat.completion.chunk','created':int(time.time()),
                         'model':body['model'],'choices':[{'index':0,'delta':{},'finish_reason':finish_reason}],
                         'usage':payload['usage']},
                    ]:
                        chunks.append('data: '+json.dumps(value,ensure_ascii=False)+'\n\n')
                    encoded=(''.join(chunks)+'data: [DONE]\n\n').encode()
                    content_type='text/event-stream'
                else:
                    encoded = json.dumps(payload, ensure_ascii=False).encode()
                    content_type='application/json'
                self.send_response(200); self.send_header('Content-Type', content_type)
                self.send_header('Content-Length', str(len(encoded))); self.end_headers(); self.wfile.write(encoded)
            except (ValueError, KeyError, TypeError):
                self.send_response(400); self.end_headers()

        def log_message(self, *args):
            pass

    server = ThreadingHTTPServer(('127.0.0.1', 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True); thread.start()
    try:
        yield fixture, f'http://127.0.0.1:{server.server_port}/v1'
    finally:
        server.shutdown(); server.server_close(); thread.join(timeout=2)


def isolated_environment(folder: Path, port: int, token: str) -> dict:
    """Avoid real account, backup, migration, or runner settings from the shell."""
    env = {key: value for key, value in os.environ.items()
           if not key.startswith(('ATLAS_', 'LLM_', 'RAILWAY_', 'RESTIC_', 'AWS_', 'OPENCODE_',
                                  'OPENAI_', 'ANTHROPIC_', 'GEMINI_', 'GOOGLE_', 'AZURE_', 'DEEPSEEK_'))
           and key not in ('APP_TOKEN', 'DATABASE_URL', 'DATA_DIR', 'PORT', 'API_ALLOWED_HOSTS')}
    data = folder / 'data'
    env.update(DATA_DIR=str(data), DATABASE_URL='sqlite:///' + str(data / 'metadata.sqlite'),
        APP_TOKEN=token, PORT=str(port), PYTHONPATH=str(ROOT),
        ATLAS_LLM_CONFIG_FILE=str(folder / 'llm.json'),
        ATLAS_CLOUD_VERIFY='0', ATLAS_DROP_PRIVILEGES='0', ATLAS_MIGRATE_SQLITE='0',
        ATLAS_OPENCODE_RUNTIME='0',
        ATLAS_DEPLOYMENT_PROFILE='single', ATLAS_AUTH_MODE='token',
        ATLAS_OPERATOR_TOKEN=FIXTURE_OPERATOR_TOKEN,
        APP_TIMEZONE='Asia/Shanghai',
        DO_NOT_TRACK='1', DLT_TELEMETRY='false')
    return env


def prepare_quickstart(folder: Path, configuration: dict) -> str:
    """Exercise the shipped stdlib setup command, without a Docker dependency."""
    path = folder / 'llm.json'
    original = json.dumps(configuration, ensure_ascii=False, indent=2) + '\n'
    path.write_text(original)
    token = None
    for _ in range(2):
        setup = subprocess.run([sys.executable, str(ROOT / 'quickstart.py'), 'setup',
            '--directory', str(folder)], cwd=ROOT, capture_output=True, text=True, timeout=15)
        if setup.returncode:
            raise RuntimeError('Quickstart setup failed')
        if FIXTURE_KEY in setup.stdout + setup.stderr:
            raise AssertionError('Quickstart setup reflected a model credential')
        assert path.read_text() == original
        values = {}
        for line in (folder / '.env').read_text().splitlines():
            if line.strip() and not line.lstrip().startswith('#'):
                key, separator, value = line.partition('=')
                if separator: values[key.strip()] = value.strip().strip('"').strip("'")
        current = values.get('APP_TOKEN', '')
        if len(current) < 16 or (token is not None and token != current):
            raise AssertionError('Quickstart did not preserve its generated application token')
        token = current
    return token


def workflow_document(source_id, budget_source_id=None):
    columns = [{'name': 'office', 'data_type': 'text', 'meaning': '办事处'}] + [
        {'name': field, 'data_type': 'decimal', 'meaning': meaning, 'unit': 'CNY',
         'period': '2026-09', 'aggregation': 'additive'}
        for field, meaning in [('actual', '本月净收入'), ('budget', '本月净收入预算')]]
    schema = {'columns': columns, 'grain': ['office'], 'meaning': '办事处月末实际和预算'}
    result={'version': 1, 'plan': {'version': 1, 'name': '快速启动完整月报验证',
        'inputs': {'detail': schema}, 'stages': [{'id': 'complete_rows', 'kind': 'select',
        'description': '保留每一行已声明实际和预算', 'input': 'detail',
        'columns': ['office', 'actual', 'budget']}], 'output': 'complete_rows'},
        'bindings': {'detail': {'source_id': source_id,
        'column_map': {'c_1': 'office', 'c_2': 'actual', 'c_3': 'budget'}}}}
    if budget_source_id:
        result['plan']['inputs']['detail']['columns']=columns[:2]
        result['plan']['inputs']['budget']={'columns':[columns[0],columns[2]],'grain':['office'],
                                          'meaning':'办事处本月净收入预算'}
        result['plan']['stages']=[{'id':'complete_rows','kind':'lookup','description':'逐办事处关联完整实际和预算',
            'input':'detail','right':'budget','left_key':'office','right_key':'office',
            'outputs':{'budget':columns[2]}}]
        result['bindings']['detail']['column_map']={'c_1':'office','c_2':'actual'}
        result['bindings']['budget']={'source_id':budget_source_id,'column_map':{'c_1':'office','c_2':'budget'}}
    return result


def verify(folder: Path, *, opencode_binary: Path | None = None, workspace_features=False) -> dict:
    with socket.socket() as socket_:
        socket_.bind(('127.0.0.1', 0)); port = socket_.getsockname()[1]
    with local_model() as (fixture, endpoint):
        configuration = {'default_model_id':'finance','models':[
            {'id':'finance','name':'财经分析','base_url':endpoint,'model':FIXTURE_MODEL,
             'api_key':FIXTURE_KEY,'share_data':True},
            {'id':'fast','name':'快速处理','base_url':endpoint,'model':FIXTURE_MODEL_ALTERNATE,
             'api_key':FIXTURE_KEY,'share_data':True}]}
        token = prepare_quickstart(folder, configuration)
        assert not any(fixture.calls.values()), 'Setup must not contact a model endpoint'
        original_configuration = (folder / 'llm.json').read_bytes()
        env = isolated_environment(folder, port, token)
        selected_folder=folder/'业务资料'
        if workspace_features:
            selected_folder.mkdir()
            data=selected_folder/'.atlas'
            env.update(ATLAS_LOCAL_WORKSPACE=str(selected_folder),ATLAS_LOCAL_MODE='1',
                       DATA_DIR=str(data),DATABASE_URL='sqlite:///'+str(data/'metadata.sqlite'))
        if opencode_binary:
            binary=opencode_binary.resolve(strict=True)
            checked=subprocess.run([str(binary),'--version'],capture_output=True,text=True,timeout=15)
            if checked.returncode or checked.stdout.strip() not in ('2.1.2','opencode v2.1.2'):
                raise ValueError('The optional OpenCode integration requires verified version 2.1.2')
            env.update(ATLAS_OPENCODE_RUNTIME='1',ATLAS_OPENCODE_BINARY=str(binary))
        process = None
        log_path = folder / 'service.log'
        with log_path.open('w') as log, httpx.Client(base_url=f'http://127.0.0.1:{port}',
                headers={'Authorization': 'Bearer ' + env['APP_TOKEN']}, trust_env=False, timeout=20) as client:
            def response(method, path, **kwargs):
                value = client.request(method, path, **kwargs); value.raise_for_status(); return value

            def call(method, path, **kwargs):
                return response(method, path, **kwargs).json()

            def wait(check, label, seconds=90):
                deadline = time.monotonic() + seconds
                while time.monotonic() < deadline:
                    if process.poll() is not None:
                        raise RuntimeError('Cloud supervisor stopped during ' + label)
                    try:
                        value = check()
                        if value: return value
                    except (httpx.ConnectError, httpx.ReadError):
                        pass
                    time.sleep(.2)
                raise TimeoutError('Timed out waiting for ' + label)

            def start(previous=None):
                nonlocal process
                process = subprocess.Popen([sys.executable, '-m', 'agent_service.cloud'],
                    cwd=ROOT, env=env, stdout=log, stderr=log)
                wait(lambda: call('GET', '/health'), 'API health')
                wait(lambda: (lambda state: state['worker']['healthy'] and state['agent']['healthy'] and
                    (previous is None or all(state[role].get('pid') != previous[role].get('pid')
                                             for role in ('worker', 'agent'))))(
                    call('GET', '/api/status')), 'worker and agent heartbeats')
                health = subprocess.run([sys.executable, '-m', 'agent_service.quickstart_health'],
                    cwd=ROOT, env=env, capture_output=True, timeout=10)
                assert health.returncode == 0, 'The actual container health command did not pass'

            def stop():
                nonlocal process
                if process and process.poll() is None:
                    process.terminate()
                    try: process.wait(timeout=15)
                    except subprocess.TimeoutExpired:
                        process.kill(); process.wait(timeout=5)
                        raise RuntimeError('Cloud supervisor did not stop its services cleanly')
                if process and process.returncode != 0:
                    raise RuntimeError('Cloud supervisor exited unsuccessfully')
                process = None

            def say(root, marker, key, mode='work',model_id='fast',source_ids=None,expected=None):
                payload={'text': marker,'request_id': key, 'interaction_mode': mode,'model_id':model_id}
                if source_ids is not None:payload['source_ids']=source_ids
                queued = call('POST', root + '/messages', json=payload)
                def completed():
                    turn = next(value for value in call('GET', root + '/conversation')['turns']
                                if value['id'] == queued['id'])
                    if turn['state'] == 'failed': raise AssertionError('Synthetic agent task failed: '+key)
                    return turn if turn['state'] == 'done' and (mode == 'chat' or turn.get('work_version')) else None
                turn = wait(completed, 'synthetic agent task')
                if turn.get('outcome') != (expected or ('answered' if mode == 'chat' else 'verified_complete')):
                    raise AssertionError('Agent task did not satisfy its server verification: '+key)
                if turn.get('model_usage', {}).get('calls', 0) < 1:
                    raise AssertionError('Agent task did not call the local simulated model')
                assert turn['model_id']==model_id
                return turn

            def scripted(root,marker,actions,*,text=None,mode='work',source_ids=None,expected='answered'):
                fixture.scenarios[marker]=actions
                return say(root,text or marker,marker,mode,source_ids=source_ids,expected=expected)

            def final_step(answer='本地合成场景已根据真实工具回执核对。'):
                return {'kind':'final','answer':answer}

            def action(name,**arguments):return {'kind':'tool','tool':name,'arguments':arguments}

            try:
                start()
                catalog = call('GET','/api/models')
                assert catalog == {'configured':True,'default_model_id':'finance','models':[
                    {'id':'finance','name':'财经分析'},{'id':'fast','name':'快速处理'}]}
                assert client.get('/api/model').status_code == 403
                operator_headers={'X-Atlas-Operator-Token':FIXTURE_OPERATOR_TOKEN}
                status = call('GET', '/api/model',headers=operator_headers)
                assert status['configured'] and status['share_data'] and status['model'] == FIXTURE_MODEL
                assert 'api_key' not in status
                assert (folder / 'llm.json').read_bytes() == original_configuration
                assert client.get('/api/model', headers={'Authorization': ''}).status_code == 401
                project = call('POST', '/api/projects', json={
                    'name': '快速启动隔离验证', 'request_id': 'quickstart-fixture-project'})
                root = '/api/projects/' + project['id']
                assert client.get(root+'/company-search').status_code==403
                assert client.put('/api/model',json={'base_url':'https://not-allowed.example/v1',
                    'model':'unapproved','api_key':'unapproved'}).status_code==403
                raw = b'office,actual,budget\nEast,100.000000000001,100\nWest,200.000000000001,200\n'
                input_sources=None;feature_checks={}
                if workspace_features:
                    actual=b'office,actual\nEast,100.000000000001\nWest,200.000000000001\n'
                    budget=b'office,budget\nWest,200\nEast,100\n'
                    (selected_folder/'实际.csv').write_bytes(actual)
                    (selected_folder/'预算.csv').write_bytes(budget)
                    (selected_folder/'llm.json').write_text('{"private":"fixture-only"}')
                    local=call('GET','/api/local-workspace')
                    assert local['enabled'] and local['imports_are_snapshots'] and not local['originals_overwritten']
                    listed=call('GET','/api/local-workspace/files')
                    assert {item['name'] for item in listed['entries']}=={'实际.csv','预算.csv'}
                    imported=call('POST',root+'/local-import',json={'paths':['预算.csv','实际.csv'],'request_id':'fixture-local-import'})
                    assert not imported['failures'] and not imported['partial']
                    assert [item['name'] for item in imported['sources']]==['预算.csv','实际.csv']
                    assert imported==call('POST',root+'/local-import',json={'paths':['预算.csv','实际.csv'],'request_id':'fixture-local-import'})
                    input_sources=imported['sources'];source=input_sources[1]
                    assert all(item['revision_id'] and item['local_origin']['sha256']==item['content_hash'] for item in input_sources)
                    feature_checks.update(local_folder_import=True,local_import_retry_same_snapshots=True,
                                          local_private_files_excluded=True)
                else:
                    source = call('POST', root + '/attachments',
                        files={'file': ('month.csv', raw, 'text/csv')}, data={'request_id': 'quickstart-fixture-source'})
                def source_ready():
                    all_sources={item['id']:item for item in call('GET', '/api/sources')}
                    currents=[all_sources[item['id']] for item in (input_sources or [source])]
                    if any(item['status']=='failed' for item in currents):raise AssertionError('Synthetic CSV source probe failed')
                    return all_sources[source['id']] if all(item['status']=='ready' for item in currents) else None
                probed = wait(source_ready, 'CSV source planning')
                assert probed['planner_mode'] == 'llm' and fixture.calls['source_planning'] >= 1
                fixture.workflow = workflow_document(source['id'],input_sources[0]['id'] if input_sources else None)
                plan_turn=say(root, '请保存并读取核对JSON说明文件 workflows/month.json。quickstart-fixture-plan', 'quickstart-plan',
                              source_ids=[item['id'] for item in input_sources] if input_sources else None)
                if workspace_features:
                    assert plan_turn['source_ids']==[item['id'] for item in input_sources]
                    attached=plan_turn['attached_sources']
                    assert [item['content_hash'] for item in attached]==[item['content_hash'] for item in input_sources]
                    observed=next(value for request,value in fixture.observed_contexts.items() if 'quickstart-fixture-plan' in request)
                    assert observed['attached_source_ids']==plan_turn['source_ids']
                    assert [item['revision_id'] for item in observed['attached_sources']]==[item['revision_id'] for item in input_sources]
                    assert (selected_folder/'Atlas工作区'/project['id']/'files/workflows/month.json').is_file()
                    feature_checks.update(multifile_input_order_and_versions=True,agent_authored_code_visible_locally=True,
                                          complete_cross_file_lookup=True)
                assert json.loads(response('GET', root + '/file', params={'path': 'workflows/month.json'}).text) == fixture.workflow
                workflow = call('POST', root + '/data-workflows', json={'path': 'workflows/month.json'})
                wid = workflow['workflow_id']; job = call('POST', root + '/data-workflows/' + wid + '/run')['id']
                def result_ready():
                    current = next(item for item in call('GET', '/api/runs', params={'target': wid}) if item['id'] == job)
                    if current['status'] in ('failed', 'cancelled'): raise AssertionError('Synthetic complete workflow failed')
                    return current if current['status'] in ('validated', 'succeeded') else None
                result = wait(result_ready, 'complete workflow')
                page = call('GET', root + '/data-results/' + job)
                assert page['total'] == 2 and len(page['rows']) == 2 and page['next_offset'] is None
                assert sum((Decimal(row['actual']) for row in page['rows']), Decimal(0)) == Decimal(EXACT_TOTAL)
                if workspace_features:
                    offices={row['office']:row for row in page['rows']}
                    assert set(offices)=={'East','West'}
                    for name,actual_value,budget_value in [('East','100.000000000001','100'),('West','200.000000000001','200')]:
                        assert Decimal(offices[name]['actual'])==Decimal(actual_value)
                        assert Decimal(offices[name]['budget'])==Decimal(budget_value)
                call('POST', root + '/data-results/' + job + '/publish', json={'expected_plan_sha256': result['commit']})
                excel = response('GET', root + '/data-results/' + job + '/download', params={'format': 'xlsx'}).content
                from openpyxl import load_workbook
                book = load_workbook(io.BytesIO(excel), read_only=True, data_only=False)
                try:
                    rows = list(book.worksheets[0].values)
                    assert len(rows) == 3 and all(row[1].data_type != 'f' for row in book.worksheets[0].iter_rows(min_row=2))
                    assert sum((Decimal(str(row[1])) for row in rows[1:]), Decimal(0)) == Decimal(EXACT_TOTAL)
                finally: book.close()
                info = call('GET', root + '/business-report-inputs/' + job)
                fixture.recipe = {'version': 1, 'title': '快速启动经营月报', 'period': '2026-09',
                    'expected_schema_sha256': info['schema_sha256'], 'metrics': [
                    {'id': 'actual', 'label': '净收入', 'kind': 'sum', 'column': 'actual'},
                    {'id': 'variance', 'label': '预算差异', 'kind': 'variance', 'column': 'actual',
                     'baseline_column': 'budget', 'rounding_scale': 12}]}
                say(root, '请保存并读取核对JSON说明文件 reports/month.json。quickstart-fixture-recipe', 'quickstart-recipe')
                brief = call('POST', root + '/business-reports', json={'run_id': job, 'recipe_path': 'reports/month.json'})
                assert brief['assertions'][0]['value'] == EXACT_TOTAL
                assert brief['assertions'][1]['value'] == '0.000000000002'
                downloads = {'xlsx': hashlib.sha256(excel).hexdigest()}
                for format_ in ('docx', 'pptx'):
                    content = response('GET', brief['download_url'], params={'format': format_}).content
                    with zipfile.ZipFile(io.BytesIO(content)) as package: assert package.testzip() is None
                    downloads[format_] = hashlib.sha256(content).hexdigest()
                if workspace_features:
                    # Install an actual versioned SKILL package and make the
                    # real agent tool loop read its reference before writing.
                    skill_package=io.BytesIO()
                    skill_text='---\nname: fixture-close\ndescription: 本地合成月结验收技能\nversion: 1\n---\n先读取 references/check.md，再按本轮明确目标记录核验说明。只使用已授权工具，不执行任意脚本。\n'
                    reference='FIXTURE-SKILL-REFERENCE：本案例金额使用十进制完整结果核对，摘要不能代替原始数据。\n'
                    with zipfile.ZipFile(skill_package,'w') as package:
                        package.writestr('fixture-close/SKILL.md',skill_text)
                        package.writestr('fixture-close/references/check.md',reference)
                    skill=call('POST','/api/extensions/skills',files={'file':('fixture-close.zip',skill_package.getvalue(),'application/zip')})
                    marker='fixture-v18-skill'
                    proof='技能参考已核阅：FIXTURE-SKILL-REFERENCE。完整金额沿用数据运行回执。\n'
                    skill_turn=scripted(root,marker,[action('discover_skills',query='fixture-close'),
                        action('load_skill',skill_id=skill['id']),action('read_skill_file',skill_id=skill['id'],path='references/check.md'),
                        action('write_file',path='notes/skill-proof.md',content=proof),final_step()],
                        text='请读取已安装技能及参考文件，生成 notes/skill-proof.md。'+marker,expected='verified_complete')
                    assert all(step['status']=='done' for step in skill_turn['agent_trace'])
                    assert (selected_folder/'Atlas工作区'/project['id']/'files/notes/skill-proof.md').read_text()==proof
                    reference_step=next(step for step in skill_turn['agent_trace'] if step['tool']=='read_skill_file')
                    assert 'FIXTURE-SKILL-REFERENCE' in json.dumps(reference_step['observation'],ensure_ascii=False)
                    feature_checks['installed_skill_loaded_by_agent']=True

                    # A real stdio MCP subprocess, explicitly approved for the
                    # synthetic models; no simulated MCP response is substituted.
                    mcp=call('POST',root+'/extensions/mcp',json={'name':'本地时钟验收','transport':'stdio',
                        'command':'atlas-local-clock','read_tools':['get_local_time'],'write_tools':[],
                        'storage_approved':True,'model_approved':True,'approved_model_ids':['finance','fast'],
                        'command_approved':True})
                    marker='fixture-v18-mcp'
                    mcp_turn=scripted(root,marker,[action('list_work_connectors'),action('list_work_tools',connector_id=mcp['id']),
                        action('read_work_tool',connector_id=mcp['id'],tool='get_local_time',arguments={'timezone':'UTC'}),final_step()],
                        text='通过已连接MCP读取本机UTC时钟用于检查连接。'+marker)
                    clock=next(step for step in mcp_turn['agent_trace'] if step['tool']=='read_work_tool')
                    assert clock['status']=='done' and 'local_system_clock' in json.dumps(clock['observation'])
                    feature_checks['actual_stdio_mcp_called_by_agent']=True

                    quote='我偏好先看核对表格，再看简明结论。'
                    marker='fixture-v18-memory-save'
                    memory_turn=scripted(root,marker,[action('remember_memory',title='验收展示偏好',quote=quote,scope='personal',category='preference'),final_step('已保存本轮明确指定的个人偏好。')],
                        text='请记住：'+quote+marker)
                    memory=next(item for item in call('GET',root+'/memories')['items'] if item['title']=='验收展示偏好')
                    assert memory['content']==quote and memory['provenance']['source_turn_id']==memory_turn['id']
                    assert memory['trusted_financial_rule'] is False
                    marker='fixture-v18-memory-recall'
                    scripted(root,marker,[action('recall_memories',query='验收展示偏好'),final_step()],
                        text='请回查我的记忆偏好。'+marker,mode='chat')
                    recalled=next(value for request,value in fixture.observed_contexts.items() if marker in request)
                    assert any(item['content']==quote for item in recalled['memory_context']['items'])
                    for index in range(9):
                        marker='fixture-v18-history-'+str(index)
                        constraint='不要将期末余额求和。金额单位为元。' if index==0 else '本轮只记录合成验收进度。'
                        scripted(root,marker,[final_step('第'+str(index)+'轮合成验收进度已记录。')],
                            text=constraint+'上下文连续验收第'+str(index)+'轮。'+marker,mode='chat')
                    compact=call('GET',root+'/context')
                    assert compact['summary_id'] and compact['compression']['active']
                    assert compact['compression']['method']=='extractive_work_index_v1'
                    assert compact['compression']['covered_older_turns']>=3 and compact['compression']['originals_retained']
                    assert '不要将期末余额求和' in json.dumps(compact['critical_constraints'],ensure_ascii=False)
                    marker='fixture-v18-checkpoint'
                    checkpoint_turn=scripted(root,marker,[action('read_context_checkpoint',checkpoint_id=compact['summary_id'],offset=0,limit=20),
                        action('read_context_original',turn_id=plan_turn['id'],offset=0,limit=8000),final_step()],
                        text='请回查上下文压缩检查点与原始工作记录。'+marker,mode='chat')
                    assert all(step['status']=='done' for step in checkpoint_turn['agent_trace'])
                    feature_checks.update(explicit_personal_memory_saved=True,memory_injected_into_model=True,
                                          context_checkpoint_and_original_recovered=True,critical_constraints_retained=True)
                    # Snapshot export checks the bytes actually visible in the
                    # selected folder, including full result and Office output.
                    marker='fixture-v18-local-export'
                    exported=scripted(root,marker,[action('local_workspace_status'),action('export_local_workspace'),final_step()],
                        text='导出本地成果。'+marker)
                    export=next(step['observation'] for step in exported['agent_trace'] if step['tool']=='export_local_workspace')
                    assert export['complete'] and not export['failures']
                    snapshot=selected_folder/export['relative_path']
                    manifest=json.loads((snapshot/'manifest.json').read_text())
                    assert manifest['complete'] and any(item['path'].startswith('data/') and item['path'].endswith('.xlsx') for item in manifest['files'])
                    for item in manifest['files']:
                        assert hashlib.sha256((snapshot/item['path']).read_bytes()).hexdigest()==item['sha256']
                    assert (selected_folder/'实际.csv').read_bytes()==actual and (selected_folder/'预算.csv').read_bytes()==budget
                    feature_checks.update(local_export_complete_hashes=True,local_originals_preserved=True)
                assert call('GET', '/api/audit/verify')['ok']
                first = call('GET', '/api/status')
                stop()
                # Catalogs are owner-authoritative: a backend-only name update
                # creates a revision; historical artifacts retain their bytes.
                configuration['models'][1]['name']='快速处理（已更新）'
                (folder / 'llm.json').write_text(json.dumps(configuration))
                start(first)
                restored = call('GET', '/api/model',headers=operator_headers)
                assert restored['configured'] and restored['share_data'] and restored['model'] == FIXTURE_MODEL
                assert call('GET','/api/models')['models'][1]['name']=='快速处理（已更新）'
                current = call('GET', '/api/status')
                assert current['worker']['pid'] != first['worker']['pid'] and current['agent']['pid'] != first['agent']['pid']
                assert call('GET', root + '/data-workflows/' + wid)['published_run'] == job
                assert call('GET', root + '/data-results/' + job)['rows'] == page['rows']
                for format_, expected in downloads.items():
                    url = root + '/data-results/' + job + '/download' if format_ == 'xlsx' else brief['download_url']
                    assert hashlib.sha256(response('GET', url, params={'format': format_}).content).hexdigest() == expected
                final=say(root, 'quickstart-fixture-check：请回复本地模拟接线检查完成。', 'quickstart-restart-check', 'chat')
                assert final['model_name']=='快速处理（已更新）'
                if workspace_features:
                    assert any(item['id']==skill['id'] for item in call('GET','/api/extensions',params={'project_id':project['id']})['skills'])
                    assert any(item['id']==memory['id'] for item in call('GET',root+'/memories')['items'])
                    assert call('GET',root+'/context')['compression']['active']
                    recalled=next(value for request,value in fixture.observed_contexts.items() if 'quickstart-fixture-check' in request)
                    assert any(item['content']==quote for item in recalled['memory_context']['items'])
                    marker='fixture-v18-memory-forget'
                    scripted(root,marker,[action('forget_memory',memory_id=memory['id'],expected_version=memory['version']),final_step('已撤回指定个人偏好。')],
                        text='请忘记“验收展示偏好”这条记忆。'+marker)
                    assert not any(item['id']==memory['id'] for item in call('GET',root+'/memories')['items'])
                    assert quote not in json.dumps(call('GET',root+'/context'),ensure_ascii=False)
                    current=call('GET','/api/status');stop();start(current)
                    marker='fixture-v18-forgotten-restart'
                    scripted(root,marker,[action('recall_memories',query='验收展示偏好'),final_step('已回查当前有效记忆。')],
                        text='请回查当前有效记忆，遵循撤回状态。'+marker,mode='chat')
                    forgotten=next(value for request,value in fixture.observed_contexts.items() if marker in request)
                    assert quote not in json.dumps(forgotten,ensure_ascii=False)
                    assert not any(item['id']==memory['id'] for item in call('GET',root+'/memories')['items'])
                    assert all(hashlib.sha256((snapshot/item['path']).read_bytes()).hexdigest()==item['sha256'] for item in manifest['files'])
                    current=call('GET','/api/status')
                    feature_checks.update(memory_survives_restart=True,forgotten_memory_not_reinjected_after_restart=True,
                                          installed_skill_survives_restart=True,context_checkpoint_persists_restart=True)
                assert fixture.calls['connection'] == 4 and fixture.tool_probes==4
                assert FIXTURE_MODEL_ALTERNATE in fixture.models_seen
                assert set(fixture.agent_models)=={FIXTURE_MODEL_ALTERNATE}
                if opencode_binary:
                    assert fixture.native_loop_calls >= 9, 'OpenCode did not own the native multi-step tool loop'
                else:
                    assert fixture.native_loop_calls==0
                assert call('GET', '/api/audit/verify')['ok']
                evidence={'ok': True, 'version': current['version'], 'simulated_model': True,
                    'real_model_evaluated': False, 'docker_executed': False,
                    'opencode_runtime_executed':bool(opencode_binary),
                    'opencode_version':'2.1.2' if opencode_binary else None,
                    'opencode_native_tool_loop_calls':fixture.native_loop_calls,
                    'opencode_native_compaction_calls':fixture.native_compaction_calls,
                    'quickstart_cli_executed': True, 'quickstart_setup_idempotent': True,
                    'cloud_supervisor': True, 'container_health_command_verified': True,
                    'api_worker_agent_healthy': True, 'model_configured_share_data': True,
                    'managed_model_selection':True,'business_cannot_configure_models':True,
                    'selected_model_routed_to_upstream':True,
                    'company_configuration_operator_only':True,'catalog_owner_restart_update':True,
                    'agent_authored_versioned_files': True, 'complete_csv_workflow': True,
                    'full_rows': 2, 'exact_total': EXACT_TOTAL, 'complete_excel': True,
                    'docx_pptx_exact_assertions': True, 'restart_persistence': True,
                    'model_credentials_survive_restart': True, 'model_calls': dict(fixture.calls),
                    'native_tool_capability_probes':fixture.tool_probes}
                return dict(evidence,workspace_features_executed=workspace_features,workspace_features=feature_checks,
                            offline_evaluated=False,network_mode='connected')
            finally:
                stop()


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--evidence-dir', type=Path, help='New empty directory for synthetic evidence and service logs')
    parser.add_argument('--opencode-binary',type=Path,
        help='Optional verified OpenCode 2.1.2 executable; runs actual native session tool loop with the simulated model')
    parser.add_argument('--workspace-features',action='store_true',
        help='Also verify multi-file local folder processing, installed Skills, stdio MCP, memory and context checkpoints')
    options = parser.parse_args(argv)
    def execute(folder):
        try:
            result = verify(folder,opencode_binary=options.opencode_binary,workspace_features=options.workspace_features)
            (folder / 'verification.json').write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
            print(json.dumps(result, ensure_ascii=False)); return 0
        except Exception as error:
            # Never echo process environment, connection JSON, or request bodies.
            frames=traceback.extract_tb(error.__traceback__)
            diagnostic={'ok':False,'simulated_model':True,'error_type':type(error).__name__,
                        'location':{'file':Path(frames[-1].filename).name,'line':frames[-1].lineno,
                                    'function':frames[-1].name} if frames else None}
            (folder/'verification.failure.json').write_text(json.dumps(diagnostic,ensure_ascii=False,indent=2)+'\n')
            print('Quickstart simulated-model verification failed: ' + type(error).__name__, file=sys.stderr)
            return 1
    if options.evidence_dir:
        folder = options.evidence_dir.resolve()
        if folder.exists() and any(folder.iterdir()): parser.error('--evidence-dir must be empty')
        folder.mkdir(parents=True, exist_ok=True); folder.chmod(0o700)
        return execute(folder)
    with tempfile.TemporaryDirectory(prefix='atlas-quickstart-check-') as temporary:
        return execute(Path(temporary))


if __name__ == '__main__':
    raise SystemExit(main())
